"""WP-047 staging + paired delivery; READY is independent of decision assurance."""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
import uuid
import time
from collections import Counter
from common import json_text, sha256_file
from db import connect
from report_model import build_preliminary, matrix_values
from scenario_formulas import literal, add_scenarios, validate_formulas
from pdf_renderer import render_pdf, label


def _new_stage(root, prefix):
    # Windows tempfile.mkdtemp uses a private DACL. A delivered revision must
    # inherit its destination's existing access, including the human user.
    # This creates a new unique directory; no existing permissions are edited.
    for _ in range(10):
        stage=Path(root)/(prefix+uuid.uuid4().hex)
        try:
            stage.mkdir(mode=0o777 if os.name=='nt' else 0o700)
            return stage
        except FileExistsError:
            continue
    raise FileExistsError('Could not allocate a fresh report stage')


def _publish_directory(stage, target):
    # Windows may briefly deny a closed file while an indexer scans it.
    # Retry the same atomic rename; never replace an existing revision.
    for attempt in range(4):
        if target.exists():
            raise FileExistsError("Report revision already exists")
        try:
            stage.rename(target)
            return
        except PermissionError:
            if attempt == 3:
                raise
            time.sleep(0.15 * (attempt + 1))

def _validate_pair(directory, contract, model):
    from openpyxl import load_workbook
    import pymupdf as fitz
    xlsx=directory/"Teklif_Degerlendirme.xlsx"; pdf=directory/"Teklif_Degerlendirme.pdf"
    result=validate_formulas(xlsx,contract)
    if result["status"]!="PASS": raise ValueError("; ".join(result["errors"]))
    cache_check={'status':'NOT_CALCULATED'}
    if model.get('recalculation')=='LOCAL_RECALCULATED':
        from calculation_cache import validate_caches
        evidence=model.get('calculation_evidence') or {}
        if (evidence.get('available') is not True or evidence.get('status')!='LOCAL_RECALCULATED'
                or evidence.get('engine')!=model.get('calculation_engine')
                or evidence.get('output_sha256')!=sha256_file(xlsx)):
            raise ValueError('Yerel hesap kanıtı Excel dosyasıyla uyuşmuyor')
        cache_check=validate_caches(xlsx,contract)
    wb=load_workbook(xlsx,data_only=False)
    if len(wb.sheetnames)<2: raise ValueError("Çok sekmeli XLSX yok")
    # Independently count rendered assessment rows, excluding headers.
    actual=Counter(row[4].value for row in list(wb["Sartname Matrisi"].rows)[1:])
    expected=Counter(label(r["status"]) for r in model.get("matrix",[]))
    if actual!=expected: raise ValueError("Şartname matris satır sayımı uyuşmuyor")
    wb.close()
    with fitz.open(pdf) as doc:
        if len(doc)<2 or any(not p.get_text().strip() for p in doc):
            raise ValueError("PDF iki anlamlı sayfa değil")
        pages=len(doc)
        text="\n".join(p.get_text() for p in doc)
    for phrase in (label(model["decision_assurance"]),label(model["comparison_status"])):
        if phrase not in text: raise ValueError("PDF rapor statüsü kayboldu")
    return {"status":"PASS","formula_count":result["formula_count"],"pdf_pages":pages,"calculation_cache":cache_check,
            "recalculation":model.get('recalculation','PENDING_LOCAL_RECALC'),
            "calculation_engine":model.get('calculation_engine'),
            "native_recalculation":"NOT_RUN" if model.get('recalculation')=='LOCAL_RECALCULATED' else model.get('recalculation','PENDING_LOCAL_RECALC'),
            "formula_check_recalculation":"NOT_EVALUATED_BY_STRUCTURAL_CHECK"}

def _workbook(model,scenarios,scoring=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb=Workbook(); wb.remove(wb.active)
    def table(name,headers,rows):
        ws=wb.create_sheet(name); literal(ws,headers)
        for row in rows: literal(ws,row)
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.protection.sheet=True
        for cell in ws[1]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="DFE8F0")
        for row in ws:
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical="top")
                ws.column_dimensions[c.column_letter].width=35
        return ws
    table("Ozet",["Alan","Değer"],[
        ["Rapor içeriği",label(model["content_status"])],["Karar güvencesi",label(model["decision_assurance"])],
        ["Karşılaştırma",label(model["comparison_status"])],["Profil",model["profile"]],
        ["Kesin öneri","Yok"],["Kaynak özeti",model.get("source_digest","Bütünlük doğrulanamadı")],
        ["Sınırlamalar","\n".join(model["limitations"]+model.get("integrity_findings",[]))],
        ["Excel formül hesabı","Hesap durumu eşlik eden PDF ve rapor doğrulama kaydında gösterilir."]])
    table("Envanter",["Belge","Durum","Kaynak hash"],[[d["relative_path"],d["parse_status"],d["sha256"]]
                                                    for d in model["documents"]])
    names={s["supplier_id"]:s["display_name"] for s in model["suppliers"]}
    table("Teklifler",["Firma","Teklif","Durum"],[[names.get(o["supplier_id"],o["supplier_id"]),o["offer_id"],o.get("status")]
                                               for o in model["offers"]])
    table("Kaynak Olgulari",["Firma","Teklif","Alan","Kaynak metni","Kaynak sayı","Kanıt durumu","Kanıt"],
          [[names.get(f.get("supplier_id"),f.get("supplier_id")),f.get("offer_id"),f.get("field_key"),
            f.get("value_text"),f.get("value_number"),f.get("evidence_status"),f.get("evidence_id")] for f in model["facts"]])
    table("Bilgi Talepleri",["Önem","Soru","Durum"],[[i.get("severity"),i.get("question"),i.get("status")] for i in model["open_issues"]])
    table("Kanitlar",["Kanıt","Belge","Konum","Alıntı","Hash"],[[e["evidence_id"],e["document_id"],e.get("locator_json"),e.get("quote"),e.get("source_sha256")] for e in model["evidence"]])
    if model.get('buyer_declarations'):
        from display_values import label as business_label
        table('Alici Beyanlari',['Firma','Teklif','Alan','Kullanıcı değeri','Birim','Hedef','Yanıt','Soru','Kaynak hash','Durum'],
              [[names.get(next(o['supplier_id'] for o in model['offers'] if o['offer_id']==r['offer_id'])),r['offer_id'],
                business_label(r['kind']),r['value'],r['unit'],r['target'],r['answer_text'],r['issue_id'],
                json_text(r['source_hashes']),'Kullanıcı beyanı; karar/hesap teyidi ayrı'] for r in model['buyer_declarations']])
    if model.get('buyer_quantity_totals'):
        table('Alici Miktar Esasi',['Teklif','Kaynak toplamı','Miktar esasıyla ara tutar','Para birimi','Fark','Sınır'],
              [[r['offer_id'],r['source_total'],r['quantity_adjusted_total'],r['currency'],r['delta'],r['limitation']]
               for r in model['buyer_quantity_totals']])
    from source_scenarios import from_sources
    from scenario_scoring import add_scoring
    contract=add_scenarios(wb,from_sources(model) if scenarios is None else scenarios)
    from matrix_workbook import add_matrix, seal_source_cells
    add_matrix(wb, model, contract)
    add_scoring(wb,model,contract,scoring)
    contract['version']='teklif-formula-contract/v3'
    seal_source_cells(wb, contract)
    return wb,contract

def _register(run_root,revision,manifest):
    con=connect(run_root)
    try:
        con.execute("BEGIN TRANSACTION")
        con.execute("INSERT OR REPLACE INTO run_meta(key,value) VALUES (?,?)",
                    ["report_delivery_"+revision,json_text(manifest)])
        con.execute("INSERT OR REPLACE INTO run_meta(key,value) VALUES ('preliminary_report_delivery',?)",[json_text(manifest)])
        con.commit()
    except BaseException:
        try: con.rollback()
        except Exception: pass
        raise
    finally: con.close()

def export_preliminary(run_root: Path,reason: str,scenarios=None,*,scoring=None,_directory="reports",_allow_fallback=True):
    """No analytical state/snapshot is promoted; never overwrite a completed run."""
    from question_waits import assert_operation_allowed
    assert_operation_allowed(run_root, 'report_delivery.export_preliminary')
    model=build_preliminary(run_root,reason)
    model["scenario_reference"]={}
    if _directory not in {"reports", "reports_fallback"}:
        raise ValueError("Unknown safe report directory")
    reports=run_root/_directory
    phase="prepare"
    if not reports.resolve().is_relative_to(run_root.resolve()):
        return {"delivery_status":"FAILED","deliverable_artifacts":False,"error":"Report path leaves run"}
    try:
        # Closed runs are read-only under the existing contract.
        if model["content_status"]!="SAFETY_NOTICE":
            con=connect(run_root)
            try:
                if con.execute("SELECT 1 FROM workflow_state WHERE phase='COMPLETE' AND status IN ('PASS','PARTIAL')").fetchone():
                    raise ValueError("Completed run preserved; use a new run")
            finally: con.close()
        reports.mkdir(parents=True,exist_ok=True)
        wb,contract=_workbook(model,scenarios,scoring)
        model["scenario_reference"]=contract["reference"]
        model["scenario_assumptions"]=contract.get("assumptions",{})
        model['scenario_scoring']=(contract.get('scoring') or {}).get('configuration')
        model['scenario_scoring_results']=contract.get('scoring_results',{})
        identity=hashlib.sha256(json_text({"model":model,"scenarios":scenarios,"formula_version":contract["version"]}).encode()).hexdigest()
        revision=identity[:24]; target=reports/revision
        if not target.resolve().is_relative_to(reports.resolve()):
            raise ValueError("Report revision path leaves report directory")
        if target.exists():
            manifest=json.loads((target/"manifest.json").read_text(encoding="utf-8"))
            if manifest.get("model_digest") != identity or manifest.get("revision") != revision:
                raise ValueError("Previous report model identity changed")
            required = {"Teklif_Degerlendirme.xlsx", "Teklif_Degerlendirme.pdf", "report_model.json", "formula_contract.json"}
            if set(manifest.get("files", {})) != required:
                raise ValueError("Previous report pair manifest changed")
            for name,sha in manifest["files"].items():
                if sha256_file(target/name)!=sha: raise ValueError("Previous report revision hash changed")
            saved_contract=json.loads((target/"formula_contract.json").read_text(encoding="utf-8"))
            if saved_contract != contract:
                raise ValueError("Previous formula contract changed")
            saved_model=json.loads((target/"report_model.json").read_text(encoding="utf-8"))
            source_projection={k:v for k,v in saved_model.items() if k not in ('recalculation','calculation_engine','calculation_evidence')}
            if source_projection != model:
                raise ValueError("Previous report projection changed")
            _validate_pair(target,contract,saved_model)
            phase="register"
            if model["content_status"]!="SAFETY_NOTICE": _register(run_root,revision,manifest)
            wb.close()
            return {**manifest,"delivery_status":"READY","deliverable_artifacts":True,"reused":True,
                    "xlsx":str(target/"Teklif_Degerlendirme.xlsx"),"pdf":str(target/"Teklif_Degerlendirme.pdf")}
        stage=_new_stage(reports,".report-stage-")
        xlsx=stage/"Teklif_Degerlendirme.xlsx"; pdf=stage/"Teklif_Degerlendirme.pdf"
        wb.save(xlsx); wb.close()
        from local_calculation import recalculate
        recalc=recalculate(xlsx)
        model['recalculation']='LOCAL_RECALCULATED' if recalc.get('available') else 'PENDING_LOCAL_RECALC'
        model['calculation_engine']=recalc.get('engine')
        model['calculation_evidence']={k:v for k,v in recalc.items() if k not in ('path','process')}
        (stage/"report_model.json").write_text(json.dumps(model,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
        (stage/"formula_contract.json").write_text(json.dumps(contract,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
        try:
            pdf_result=render_pdf(model,pdf)
        except Exception as primary:
            # Bounded alternative renderer, same model. No success with a missing PDF.
            from pdf_fallback import render_safety_pdf
            pdf_result=render_safety_pdf(model,pdf,str(primary))
        qa=_validate_pair(stage,contract,model)
        manifest={"schema":"teklif-paired-delivery/v1","revision":revision,"model_digest":identity,
                  "relative_directory":str(target.relative_to(run_root)),
                  "decision_assurance":model["decision_assurance"],"comparison_status":model["comparison_status"],
                  "content_status":model["content_status"],"pair_validation":"PASS","qa":qa,"profile":model["profile"],
                  "metadata_mode":"SAFETY_FILESYSTEM_ONLY" if model["content_status"]=="SAFETY_NOTICE" else "DATABASE",
                  "files":{name:sha256_file(stage/name) for name in
                           ["Teklif_Degerlendirme.xlsx","Teklif_Degerlendirme.pdf","report_model.json","formula_contract.json"]}}
        (stage/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
        # Same filesystem directory rename; an incomplete pair is never exposed as a revision.
        _publish_directory(stage, target)
        phase="register"
        if model["content_status"]!="SAFETY_NOTICE": _register(run_root,revision,manifest)
        return {**manifest,"delivery_status":"READY","deliverable_artifacts":True,
                "xlsx":str(target/"Teklif_Degerlendirme.xlsx"),"pdf":str(target/"Teklif_Degerlendirme.pdf"),
                "renderer":pdf_result}
    except Exception as exc:
        if isinstance(exc, OSError) and phase != "register" and _allow_fallback:
            result=export_preliminary(run_root,reason,scenarios,scoring=scoring,_directory="reports_fallback",_allow_fallback=False)
            return {**result,"alternate_report_path_attempted":True,"primary_path_error":str(exc)}
        return {"delivery_status":"FAILED","deliverable_artifacts":False,
                "decision_assurance":model["decision_assurance"],"error":str(exc)}


def inspect_preliminary(run_root: Path):
    """Read-only delivery status from the committed pointer, with fresh pair checks."""
    from db import get_meta
    manifest = get_meta(run_root, "preliminary_report_delivery", None, parse_json=True)
    if not manifest:
        return {"delivery_status": "NOT_CREATED", "deliverable_artifacts": False}
    try:
        relative = Path(manifest["relative_directory"])
        if len(relative.parts) != 2 or relative.parts[0] not in {"reports", "reports_fallback"} or relative.parts[1] != manifest["revision"]:
            raise ValueError("Invalid committed report path")
        directory = run_root / relative
        if not directory.resolve().is_relative_to(run_root.resolve()):
            raise ValueError("Committed report path leaves run")
        required = {"Teklif_Degerlendirme.xlsx", "Teklif_Degerlendirme.pdf", "report_model.json", "formula_contract.json"}
        if set(manifest.get("files", {})) != required:
            raise ValueError("Committed pair manifest incomplete")
        for name, sha in manifest["files"].items():
            if sha256_file(directory / name) != sha:
                raise ValueError("Committed report file changed: " + name)
        model = json.loads((directory / "report_model.json").read_text(encoding="utf-8"))
        contract = json.loads((directory / "formula_contract.json").read_text(encoding="utf-8"))
        current = build_preliminary(run_root, model["limitations"][0])
        current["scenario_reference"] = model["scenario_reference"]
        current["scenario_assumptions"] = model.get("scenario_assumptions",{})
        current['scenario_scoring']=model.get('scenario_scoring')
        current['scenario_scoring_results']=model.get('scenario_scoring_results',{})
        for key in ('recalculation','calculation_engine','calculation_evidence'):
            if key in model:current[key]=model[key]
        if current != model:
            raise ValueError("Report source state changed; create a new revision")
        qa = _validate_pair(directory, contract, model)
        return {"delivery_status": "READY", "deliverable_artifacts": True,
                "decision_assurance": model["decision_assurance"], "revision": manifest["revision"],
                "xlsx": str(directory / "Teklif_Degerlendirme.xlsx"),
                "pdf": str(directory / "Teklif_Degerlendirme.pdf"), "qa": qa}
    except Exception as exc:
        return {"delivery_status": "FAILED", "deliverable_artifacts": False, "error": str(exc)}


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Create or inspect a validated preliminary Excel/PDF pair; never approves a purchase")
    parser.add_argument("--run", required=True)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--reason", help="Actual reason for the preliminary or safety report")
    action.add_argument("--inspect", action="store_true", help="Inspect the current committed preliminary pair")
    args = parser.parse_args()
    run = Path(args.run)
    from run_compatibility import assert_compatible
    assert_compatible(run, Path(__file__).resolve().parents[1])
    if args.inspect:
        result = inspect_preliminary(run)
    else:
        if not args.reason.strip():
            parser.error("A non-empty report reason is required")
        result = export_preliminary(run, args.reason.strip())
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 0 if result.get("delivery_status") == "READY" and result.get("deliverable_artifacts") else 3


if __name__ == "__main__":
    from question_waits import run_guarded_main
    raise SystemExit(run_guarded_main(main))
