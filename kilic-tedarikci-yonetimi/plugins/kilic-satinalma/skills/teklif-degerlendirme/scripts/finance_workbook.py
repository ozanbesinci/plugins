"""Editable payment/FX assumptions and guarded financial TCO formulas."""
from openpyxl.styles import Protection, PatternFill, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from scenario_finance import number, policy_valid, finance_reference, lifecycle_reference, scenario_date, FINANCIAL_FIELDS, COST_MODES, MAX_COST_YEARS

FINANCIAL_SHEETS = ('Finansal Varsayimlar','Odeme Plani','Finansal TCO','Finansal Karsilastirma','Yillik Yasam Giderleri')


def excel_date_serial(ref):
    return f'DATE(VALUE(LEFT({ref},4)),VALUE(MID({ref},6,2)),VALUE(RIGHT({ref},2)))'


def excel_date_guard(ref):
    serial=excel_date_serial(ref)
    return (f'AND(ISTEXT({ref}),LEN({ref})=10,MID({ref},5,1)="-",MID({ref},8,1)="-",'
            f'VALUE(LEFT({ref},4))>=1900,VALUE(LEFT({ref},4))<=9999,{ref}<>"1900-02-29",'
            f'EXACT(LEFT({ref},4),RIGHT("0000"&YEAR({serial}),4)),'
            f'EXACT(MID({ref},6,2),RIGHT("00"&MONTH({serial}),2)),'
            f'EXACT(RIGHT({ref},2),RIGHT("00"&DAY({serial}),2)))')

def validate_input(kind, value):
    if kind.startswith('ledger_'):
        from dated_cashflows import validate_ledger_input
        validate_ledger_input(kind,value)
        return
    if value is None: return
    if kind in {"currency", "iso_date", "text", "direction", "policy_kind", "cost_mode"}:
        if not isinstance(value, str) or not value.strip(): raise ValueError("Financial policy text required")
        if kind == "currency" and (len(value)!=3 or not value.isascii() or not value.isalpha() or value!=value.upper()): raise ValueError("Invalid financial currency")
        if kind == "iso_date": scenario_date(value)
        if kind == "text" and len(value)>200: raise ValueError("Financial source label too long")
        if kind == "direction" and value!="TARGET_PER_SOURCE": raise ValueError("Invalid financial rate direction")
        if kind == "policy_kind" and value!="FIXED_SCENARIO": raise ValueError("Invalid financial policy kind")
        if kind == "cost_mode" and value not in COST_MODES: raise ValueError("Invalid lifecycle cost mode")
        return
    limits={"count":12,"days":36500,"share":1,"money":None}
    if kind not in limits: raise ValueError("Unknown financial input contract")
    result=number(value,kind,whole=kind in {"count","days"},upper=limits[kind])
    if kind=="count" and result==0: raise ValueError("Positive payment count required")

def add_finance(wb, scenarios, contract):
    from scenario_formulas import INPUTS, literal
    vars,pay,tco,comparison,annual=(wb.create_sheet(name) for name in FINANCIAL_SHEETS)
    literal(vars,["Senaryo","Ödeme adedi","KDV geri kazanım günü","Aynı kapsam/kur/dönem NBD bazı","Hedef para","Kaynak para","Kur tarihi","Kur dayanağı","Kur yönü","Kur politikası","Yaşam gideri modeli"])
    literal(pay,["Senaryo","Sıra","Ödeme payı","Analiz tarihinden gün","Nominal ödeme","Ödeme NBD","Girdi durumu"])
    literal(tco,["Senaryo","Ödeme NBD","Geri kazanılan KDV NBD","Yaşam giderleri NBD","Kalıntı NBD","Finansal TCO NBD","Durum / kapsam","Yaşam giderleri nominal","Gider kapsamı"])
    literal(annual,["Senaryo","Yıl (analiz tarihinden)","Enerji: hedef para","Bakım: hedef para","Diğer: hedef para","Yıl sonu gider toplamı","Yıl sonu gider NBD","Durum / kullanıcı varsayımı"])
    literal(comparison,["Senaryo","Kullanıcı NBD bazı","Senaryo TCO NBD","Fark","Oransal fark","Durum"])
    contract["version"]="teklif-formula-contract/v2"
    contract.setdefault("extra_inputs",{})
    def input(ws,cell,value,kind):
        validate_input(kind,value);ws[cell]=value
        if isinstance(value,str):ws[cell].data_type="s"
        ws[cell].protection=Protection(locked=False);ws[cell].fill=PatternFill("solid",fgColor="DDEBFF")
        contract["extra_inputs"][ws.title+"!"+cell]=kind
        if kind in {"count","days","share","money"}:
            upper={"count":12,"days":36500,"share":1,"money":None}[kind]
            dv=DataValidation(type="whole" if kind in {"count","days"} else "decimal",operator="between" if upper is not None else "greaterThanOrEqual",formula1="1" if kind=="count" else "0",formula2=str(upper) if upper is not None else None,allow_blank=True)
        elif kind in {"direction","policy_kind","cost_mode"}:
            choices={"direction":"TARGET_PER_SOURCE","policy_kind":"FIXED_SCENARIO","cost_mode":"CONSTANT_ANNUAL,YEARLY_SCHEDULE"}
            dv=DataValidation(type="list",formula1='"'+choices[kind]+'"',allow_blank=True)
        else:
            dv=DataValidation(type="textLength",operator="between",formula1="1",formula2="200",allow_blank=True)
        dv.showErrorMessage=True;dv.errorStyle="stop";dv.error="Geçerli varsayım girin; bilinmiyorsa boş bırakın."
        ws.add_data_validation(dv);dv.add(cell)
    def formula(ws,cell,expression,guard=None):
        text="="+expression if guard is None else '=IF(IFERROR('+guard+',FALSE),'+expression+',"")'
        ws[cell]=text;contract["formulas"][ws.title+"!"+cell]=text
    for row,scenario in enumerate(scenarios,start=2):
        financial=scenario.get("financial");financial={} if financial is None else financial;values=scenario.get("inputs",{})
        if not isinstance(financial,dict) or set(financial)-FINANCIAL_FIELDS:raise ValueError("Unknown financial scenario field")
        lifecycle=lifecycle_reference(values,financial)
        payments=financial.get("payments") or []
        if not isinstance(payments,list) or len(payments)>12:raise ValueError("At most twelve payments per scenario")
        policy=financial.get("policy") or {}
        if policy:policy_valid(policy,values.get("fx_rate"))
        literal(vars,[scenario["id"]])
        specs=[(len(payments) or None,"count"),(financial.get("vat_refund_days"),"days"),(financial.get("base_present_value"),"money")]
        specs += [(policy.get(key),kind) for key,kind in (("base_currency","currency"),("source_currency","currency"),("rate_date","iso_date"),("rate_source","text"),("direction","direction"),("kind","policy_kind"))]
        specs.append((lifecycle['mode'],"cost_mode"))
        for col,(value,kind) in enumerate(specs,start=2):input(vars,get_column_letter(col)+str(row),value,kind)
        refs={key:"'Senaryo Girdileri'!"+get_column_letter(i)+str(row) for i,key in enumerate(INPUTS,start=2)}
        v=lambda col:"'Finansal Varsayimlar'!"+col+str(row)
        count=v("B");refund=v("C");currency=v("E");source=v("F");rate=refs["annual_rate"];years=refs["years"]
        cash="'Kur Vergi Ithalat'!E"+str(row);cost="'Kur Vergi Ithalat'!F"+str(row)
        codes=lambda ref:",".join(f"CODE(MID({ref},{i},1))>=65,CODE(MID({ref},{i},1))<=90" for i in (1,2,3))
        policy_guard=f'AND(LEN({currency})=3,LEN({source})=3,{codes(currency)},{codes(source)},{excel_date_guard(v("G"))},LEN({v("H")})>0,{v("I")}="TARGET_PER_SOURCE",{v("J")}="FIXED_SCENARIO",ISNUMBER({refs["fx_rate"]}),{refs["fx_rate"]}>0,OR({currency}<>{source},{refs["fx_rate"]}=1))'
        rate_guard=f"AND(ISNUMBER({rate}),{rate}>=0,{rate}<=10)"
        horizon=f"AND(ISNUMBER({years}),{years}>=1,{years}<=100,{years}=INT({years}))"
        mode=v("K")
        first_year=2+(row-2)*MAX_COST_YEARS;last_year=first_year+MAX_COST_YEARS-1
        cost_rows={item['year']:item for item in financial.get('annual_costs') or []}
        for year,yr in enumerate(range(first_year,last_year+1),start=1):
            item=cost_rows.get(year,{})
            literal(annual,[scenario['id'],year])
            for col,key in (("C","energy"),("D","maintenance"),("E","other")):
                input(annual,col+str(yr),item.get(key),"money")
            known=f"AND(ISNUMBER(C{yr}),C{yr}>=0,ISNUMBER(D{yr}),D{yr}>=0,ISNUMBER(E{yr}),E{yr}>=0)"
            formula(annual,"F"+str(yr),f"SUM(C{yr}:E{yr})",f'AND({mode}="YEARLY_SCHEDULE",{policy_guard},{horizon},{year}<={years},{known})')
            formula(annual,"G"+str(yr),f"F{yr}/(1+{rate})^{year}",f"AND(ISNUMBER(F{yr}),{rate_guard})")
            formula(annual,"H"+str(yr),f'IF(ISNUMBER(F{yr}),"Yıl sonu, hedef para; yalnız kullanıcı senaryosu","Hesaplanamadı veya model/dönem dışında")')
        annual_range=lambda col:f"'Yillik Yasam Giderleri'!{col}{first_year}:{col}{last_year}"
        complete_years=f"AND({policy_guard},{horizon},COUNT({annual_range('F')})={years},COUNTA({annual_range('C')})={years},COUNTA({annual_range('D')})={years},COUNTA({annual_range('E')})={years})"
        count_guard=f"AND(ISNUMBER({count}),{count}>=1,{count}<=12,{count}=INT({count}))"
        start=2+(row-2)*12;end=start+11
        for slot,payrow in enumerate(range(start,end+1),start=1):
            item=payments[slot-1] if slot<=len(payments) else {}
            if not isinstance(item,dict) or set(item)-{"percent","days"}:raise ValueError("Unknown payment field")
            literal(pay,[scenario["id"],slot])
            input(pay,"C"+str(payrow),item.get("percent"),"share");input(pay,"D"+str(payrow),item.get("days"),"days")
            share="C"+str(payrow);days="D"+str(payrow)
            valid=f"AND(ISNUMBER({share}),{share}>=0,{share}<=1,ISNUMBER({days}),{days}>=0,{days}<=36500,{days}=INT({days}))"
            active=f"AND({count_guard},{slot}<={count})"
            formula(pay,"E"+str(payrow),cash+"*"+share,f"AND({policy_guard},{active},{valid},ISNUMBER({cash}))")
            formula(pay,"F"+str(payrow),f"E{payrow}/(1+{rate})^({days}/365)",f"AND(ISNUMBER(E{payrow}),{rate_guard})")
            formula(pay,"G"+str(payrow),f'IF(ISNUMBER(F{payrow}),"Kullanıcı ödeme varsayımı; 365 gün bazı","Hesaplanamadı veya kullanılmayan satır")')
        literal(tco,[scenario["id"]]);literal(comparison,[scenario["id"]])
        pref=lambda col:f"'Odeme Plani'!{col}{start}:{col}{end}"
        # Qualify per-row predicates because this guard is evaluated on another sheet.
        qualified=[]
        for slot,payrow in enumerate(range(start,end+1),start=1):
            p=f"'Odeme Plani'!C{payrow}";d=f"'Odeme Plani'!D{payrow}"
            qualified.append(f'IF({slot}<={count},AND(ISNUMBER({p}),{p}>=0,{p}<=1,ISNUMBER({d}),{d}>=0,{d}<=36500,{d}=INT({d})),AND({p}="",{d}=""))')
        pay_guard=f"AND({policy_guard},{rate_guard},{count_guard},ISNUMBER({cash}),{','.join(qualified)},ABS(SUM({pref('C')})-1)<=0.00000001)"
        formula(tco,"B"+str(row),"SUM("+pref("F")+")",pay_guard)
        tax_guard=f"AND({policy_guard},ISNUMBER({cash}),ISNUMBER({cost}),{cash}>={cost})"
        refund_guard=f"AND({rate_guard},ISNUMBER({refund}),{refund}>=0,{refund}<=36500,{refund}=INT({refund}))"
        formula(tco,"C"+str(row),f'IF({cash}={cost},0,IF(IFERROR({refund_guard},FALSE),({cash}-{cost})/(1+{rate})^({refund}/365),""))',tax_guard)
        energy=refs["energy_annual"];maint=refs["maintenance_annual"];residual=refs["residual"]
        constant_known=f"AND({horizon},ISNUMBER({energy}),{energy}>=0,ISNUMBER({maint}),{maint}>=0)"
        constant_pv=f"({energy}+{maint})*IF({rate}=0,{years},(1-(1+{rate})^(-{years}))/{rate})"
        yearly_pv=f"AND({complete_years},{rate_guard},COUNT({annual_range('G')})={years})"
        unused_costs=f"COUNTA({annual_range('C')},{annual_range('D')},{annual_range('E')})=0"
        constant_guard=f"AND({policy_guard},{rate_guard},{constant_known},{unused_costs})"
        formula(tco,"D"+str(row),f'IF({mode}="YEARLY_SCHEDULE",IF(IFERROR({yearly_pv},FALSE),SUM({annual_range("G")}),""),IF({mode}="CONSTANT_ANNUAL",IF(IFERROR({constant_guard},FALSE),{constant_pv},""),""))')
        formula(tco,"H"+str(row),f'IF({mode}="YEARLY_SCHEDULE",IF(IFERROR({complete_years},FALSE),SUM({annual_range("F")}),""),IF({mode}="CONSTANT_ANNUAL",IF(IFERROR(AND({constant_known},{unused_costs}),FALSE),({energy}+{maint})*{years},""),""))')
        formula(tco,"I"+str(row),f'IF({mode}="YEARLY_SCHEDULE","Her yıl enerji/bakım/diğer; hedef para, yıl sonu; kullanıcı kapsamı",IF({mode}="CONSTANT_ANNUAL","Sabit enerji/bakım; diğer yaşam giderleri kapsam dışında","Yaşam gideri modeli bilinmiyor; toplam hesaplanamadı"))')
        formula(tco,"E"+str(row),f"{residual}/(1+{rate})^{years}",f"AND({policy_guard},{rate_guard},{horizon},ISNUMBER({residual}),{residual}>=0)")
        formula(tco,"F"+str(row),f"B{row}-C{row}+D{row}-E{row}",f"COUNT(B{row}:E{row})=4")
        formula(tco,"G"+str(row),f'IF(ISNUMBER(F{row}),"Kullanıcı senaryosu; I sütunundaki gider kapsamı; resmî sıralama değildir","Hesaplanamadı: ödeme/KDV/kur politikası/yıllık gider/dönem/kalıntı teyidi")')
        # The nominal scenario tabs must follow the selected lifecycle model too.
        nominal=wb['TCO Yasam Dongusu'];old=nominal['B'+str(row)].value[1:]
        yearly_nominal_guard=f"AND(ISNUMBER({cost}),ISNUMBER(H{row}),ISNUMBER({residual}),{residual}>=0)"
        # H is qualified because the formula is written on the nominal TCO sheet.
        yearly_nominal_guard=yearly_nominal_guard.replace(f'H{row}',f"'Finansal TCO'!H{row}")
        formula(nominal,'B'+str(row),f'IF({mode}="YEARLY_SCHEDULE",IF(IFERROR({yearly_nominal_guard},FALSE),{cost}+\'Finansal TCO\'!H{row}-{residual},""),IF({mode}="CONSTANT_ANNUAL",IF(IFERROR({unused_costs},FALSE),{old},""),""))')
        formula(nominal,'C'+str(row),f'IF(ISNUMBER(B{row}),"Kullanıcı nominal TCO; seçili gider modeli; resmî sıralama yok","Hesaplanamadı: seçili gider modeli/kapsam/dönem/kalıntı")')
        total=f"'Finansal TCO'!F{row}";baseline=v("D")
        formula(comparison,"B"+str(row),baseline,f"AND({policy_guard},ISNUMBER({baseline}))")
        formula(comparison,"C"+str(row),total,f"ISNUMBER({total})")
        formula(comparison,"D"+str(row),f"C{row}-B{row}",f"COUNT(B{row}:C{row})=2")
        formula(comparison,"E"+str(row),f"D{row}/B{row}",f"AND(ISNUMBER(D{row}),ISNUMBER(B{row}),B{row}<>0)")
        formula(comparison,"F"+str(row),f'IF(ISNUMBER(D{row}),"Aynı kapsam/kur/dönem NBD bazı kullanıcı varsayımı; sıralama yok","NBD karşılaştırması hesaplanamadı")')
        ref=contract["reference"][str(row)]
        result=finance_reference(values,financial,cash=ref.get("Kur Vergi Ithalat!E"),cost=ref.get("Kur Vergi Ithalat!F"))
        if lifecycle['mode']!='CONSTANT_ANNUAL':
            for key in ('TCO Yasam Dongusu!B','Senaryo Karsilastirma!C','Senaryo Karsilastirma!D','Senaryo Karsilastirma!E','Senaryo Kontrol!D'):
                ref.pop(key,None)
            if lifecycle['mode']=='YEARLY_SCHEDULE' and policy_valid(policy,values.get('fx_rate')) and lifecycle['nominal'] is not None and ref.get('Kur Vergi Ithalat!F') is not None and values.get('residual') is not None:
                nominal_total=ref['Kur Vergi Ithalat!F']+lifecycle['nominal']-values['residual']
                ref.update({'TCO Yasam Dongusu!B':nominal_total,'Senaryo Karsilastirma!C':nominal_total,'Senaryo Kontrol!D':nominal_total})
                if values.get('base_total') is not None:
                    base=values['base_total'];ref['Senaryo Karsilastirma!B']=base;ref['Senaryo Karsilastirma!D']=nominal_total-base
                    if base!=0:ref['Senaryo Karsilastirma!E']=(nominal_total-base)/base
        if lifecycle['nominal'] is not None and (lifecycle['mode']=='CONSTANT_ANNUAL' or policy_valid(policy,values.get('fx_rate'))):
            ref['Finansal TCO!H']=lifecycle['nominal']
        for key,col in (("payment_present_value","B"),("recoverable_vat_present_value","C"),("opex_present_value","D"),("residual_present_value","E"),("total_present_value","F")):
            if result[key] is not None:ref["Finansal TCO!"+col]=result[key]
        for key,col in (("baseline_present_value","B"),("total_present_value","C"),("difference","D"),("relative_difference","E")):
            if result[key] is not None:ref["Finansal Karsilastirma!"+col]=result[key]
    for ws in (vars,pay,tco,comparison,annual):
        ws.freeze_panes="C2";ws.auto_filter.ref=ws.dimensions;ws.protection.sheet=True
        for cell in ws[1]:cell.font=Font(bold=True)
        for row in ws:
            for cell in row:ws.column_dimensions[cell.column_letter].width=28

def read_lifecycle_inputs(wb, row):
    if 'Yillik Yasam Giderleri' not in wb:
        return {}
    result={'cost_mode':wb['Finansal Varsayimlar'].cell(row,11).value}
    start=2+(row-2)*MAX_COST_YEARS;costs=[]
    for year,yr in enumerate(range(start,start+MAX_COST_YEARS),start=1):
        item={key:wb['Yillik Yasam Giderleri'].cell(yr,col).value for col,key in ((3,'energy'),(4,'maintenance'),(5,'other'))}
        if any(value is not None for value in item.values()):
            costs.append({'year':year,**item})
    if costs:result['annual_costs']=costs
    return result


def validate_extra_inputs(wb, contract):
    errors=[]
    from dated_workbook import validate_ledgers
    errors.extend(validate_ledgers(wb))
    for locator,kind in contract.get("extra_inputs",{}).items():
        try:
            sheet,cell=locator.rsplit("!",1);validate_input(kind,wb[sheet][cell].value)
        except (ValueError,KeyError,TypeError) as exc:errors.append(str(exc))
    if "Finansal Varsayimlar" not in wb:return errors
    vars=wb["Finansal Varsayimlar"];pay=wb["Odeme Plani"]
    for r in range(2,vars.max_row+1):
        try:
            from scenario_formulas import INPUTS
            values={key:wb['Senaryo Girdileri'].cell(r,col).value for col,key in enumerate(INPUTS,start=2)}
            lifecycle_reference(values,read_lifecycle_inputs(wb,r))
        except (ValueError,KeyError,TypeError) as exc:errors.append(str(exc))
        count=vars.cell(r,2).value
        if not isinstance(count,(int,float)) or isinstance(count,bool) or count!=int(count) or not 1<=count<=12:continue
        shares=[]
        for slot in range(1,13):
            row=2+(r-2)*12+slot-1;p=pay.cell(row,3).value;d=pay.cell(row,4).value
            if slot>count and (p is not None or d is not None):errors.append("Payment data outside declared count")
            elif slot<=count:shares.append(p)
        if all(isinstance(v,(int,float)) and not isinstance(v,bool) for v in shares) and abs(sum(shares)-1)>0.00000001:errors.append("Payment shares do not sum to one")
    return errors
