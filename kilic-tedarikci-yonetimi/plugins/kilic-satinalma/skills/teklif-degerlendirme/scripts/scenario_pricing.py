"""Editable price units, VAT-included prices and any number of discount stages."""
from decimal import Decimal as D
import math
from openpyxl.styles import Protection,PatternFill,Font,Alignment
from openpyxl.worksheet.datavalidation import DataValidation

SHEET='Kalem Fiyat Varsayimlari'
DISCOUNTS='Ardisik Iskontolar'

def validate(kind,value):
    if value is None:return
    if kind=='pricing_flag':
        if type(value) is not int or value not in (0,1):raise ValueError('KDV dahil seçimi 0 veya 1 olmalı')
        return
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):raise ValueError('Fiyat varsayımı sonlu sayı olmalı')
    if kind=='pricing_unit' and value<=0:raise ValueError('Fiyat birimi pozitif olmalı')
    if kind=='pricing_discount' and not 0<=value<1:raise ValueError('İskonto 0 dahil, 1 hariç aralığında olmalı')
    if kind not in ('pricing_unit','pricing_discount'):raise ValueError('Bilinmeyen fiyat varsayımı')

def net_price(values,pricing=None):
    pricing={} if pricing is None else pricing
    if not isinstance(pricing,dict) or set(pricing)-{'price_unit','price_includes_vat','discounts'}:raise ValueError('Bilinmeyen fiyat varsayımı alanı')
    unit=pricing.get('price_unit',1);included=pricing.get('price_includes_vat',0)
    extra=pricing.get('discounts',[])
    if not isinstance(extra,list):raise ValueError('Ardışık iskontolar liste olmalı')
    validate('pricing_unit',unit);validate('pricing_flag',included)
    for rate in extra:validate('pricing_discount',rate)
    if any(values.get(k) is None for k in ('quantity','unit_price','discount')) or unit is None or included is None:return None
    if included and values.get('vat_rate') is None:return None
    result=D(str(values['quantity']))*D(str(values['unit_price']))/D(str(unit))*(1-D(str(values['discount'])))
    for rate in extra:
        if rate is not None:result*=1-D(str(rate))
    if included:result/=1+D(str(values['vat_rate']))
    return result

def add_pricing(wb,scenarios,contract):
    from scenario_formulas import INPUTS,literal
    from openpyxl.utils import get_column_letter
    ws=wb.create_sheet(SHEET);ds=wb.create_sheet(DISCOUNTS)
    literal(ws,['Senaryo','Fiyat birimi','Fiyat KDV dahil: 1 / hariç: 0','Hesap ilkesi'])
    literal(ds,['Senaryo','Ek iskonto sırası','İskonto oranı'])
    contract.setdefault('extra_inputs',{});contract['pricing_rows']={}
    def input(sheet,cell,value,kind):
        validate(kind,value);sheet[cell]=value;sheet[cell].protection=Protection(locked=False)
        sheet[cell].fill=PatternFill('solid',fgColor='DDEBFF');contract['extra_inputs'][sheet.title+'!'+cell]=kind
        if kind=='pricing_flag':dv=DataValidation(type='whole',operator='between',formula1='0',formula2='1',allow_blank=True)
        elif kind=='pricing_unit':dv=DataValidation(type='decimal',operator='greaterThan',formula1='0',allow_blank=True)
        else:dv=DataValidation(type='decimal',operator='between',formula1='0',formula2='0.999999999999',allow_blank=True)
        dv.showErrorMessage=True;dv.errorStyle='stop';sheet.add_data_validation(dv);dv.add(cell)
    for row,scenario in enumerate(scenarios,start=2):
        pricing=scenario.get('pricing') or {};values=scenario.get('inputs',{})
        net_price(values,pricing)
        literal(ws,[scenario['id']]);input(ws,'B'+str(row),pricing.get('price_unit',1),'pricing_unit')
        input(ws,'C'+str(row),pricing.get('price_includes_vat',0),'pricing_flag')
        ws.cell(row,4,'Ara aşamada yuvarlama yok; para görünümü iki ondalık. Boş ek iskonto uygulanmaz.')
        first=ds.max_row+1;rates=pricing.get('discounts',[])
        for index in range(max(3,len(rates)+2)):
            n=ds.max_row+1;literal(ds,[scenario['id'],index+2])
            input(ds,'C'+str(n),rates[index] if index<len(rates) else None,'pricing_discount')
        last=ds.max_row
        contract['pricing_rows'][str(row)]={'first':first,'last':last}
        refs={key:f"'Senaryo Girdileri'!{get_column_letter(i)}{row}" for i,key in enumerate(INPUTS,start=2)}
        unit=f"'{SHEET}'!B{row}";flag=f"'{SHEET}'!C{row}"
        q,p,d,vat=(refs[k] for k in ('quantity','unit_price','discount','vat_rate'))
        extras='*'.join(f'IF(ISNUMBER(\'{DISCOUNTS}\'!C{n}),1-\'{DISCOUNTS}\'!C{n},1)' for n in range(first,last+1))
        checks=','.join(f'OR(\'{DISCOUNTS}\'!C{n}="",AND(ISNUMBER(\'{DISCOUNTS}\'!C{n}),\'{DISCOUNTS}\'!C{n}>=0,\'{DISCOUNTS}\'!C{n}<1))' for n in range(first,last+1))
        guard=f'AND(ISNUMBER({q}),{q}>0,ISNUMBER({p}),{p}>=0,ISNUMBER({d}),{d}>=0,{d}<1,ISNUMBER({unit}),{unit}>0,ISNUMBER({flag}),OR({flag}=0,AND({flag}=1,ISNUMBER({vat}),{vat}>=0)),{checks})'
        expression=f'{q}*{p}/{unit}*(1-{d})*{extras}/IF({flag}=1,1+{vat},1)'
        target=wb['Miktar ve Kapsam'];target['B'+str(row)]=f'=IF(IFERROR({guard},FALSE),{expression},"")'
        target['C'+str(row)]=f'=IF(ISNUMBER(B{row}),"SİMÜLASYON HESAPLANDI","HESAPLANAMADI — fiyat birimi/KDV/miktar/fiyat/iskonto")'
        for col in ('B','C'):contract['formulas'][target.title+'!'+col+str(row)]=target[col+str(row)].value
        target['B'+str(row)].number_format='#,##0.00'
    for sheet in (ws,ds):
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.protection.sheet=True
        sheet.row_dimensions[1].height=45
        for cell in sheet[1]:cell.font=Font(bold=True)
        for cells in sheet:
            for cell in cells:cell.alignment=Alignment(wrap_text=True,vertical='top')
        for col in ('A','B','C','D'):sheet.column_dimensions[col].width=32
    ws.column_dimensions['D'].width=70

def read_pricing(wb,row,contract):
    rows=contract['pricing_rows'][str(row)]
    rates=[wb[DISCOUNTS].cell(n,3).value for n in range(rows['first'],rows['last']+1)]
    while rates and rates[-1] is None:rates.pop()
    return {'price_unit':wb[SHEET].cell(row,2).value,'price_includes_vat':wb[SHEET].cell(row,3).value,
            'discounts':rates}

