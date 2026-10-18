"""Location-sealed dated ledger formulas and editable user assumptions."""
from openpyxl.styles import Font,PatternFill,Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from dated_cashflows import MAX_FLOWS,CATEGORIES,METHOD,SCOPE,ledger_reference,validate_ledger_input
from finance_workbook import excel_date_serial,excel_date_guard

SUMMARY='Tarihli Nakit Ozeti'
DETAIL='Tarihli Nakit Akislari'

def read_ledger(wb,row):
    ws=wb[SUMMARY];count=ws.cell(row,5).value
    result=dict(zip(('valuation_date','base_currency','annual_rate','baseline','scope','method'),
                    (ws.cell(row,col).value for col in (2,3,4,6,12,13))))
    start=2+(row-2)*MAX_FLOWS;flows=[]
    for slot,yr in enumerate(range(start,start+MAX_FLOWS),start=1):
        item=dict(zip(('cost_id','category','amount','currency','cash_date','rate_date','fx_rate','rate_source'),
                      (wb[DETAIL].cell(yr,col).value for col in range(3,11))))
        if isinstance(count,(int,float)) and not isinstance(count,bool) and count==int(count) and 1<=count<=MAX_FLOWS and slot<=count:
            flows.append(item)
        elif any(value is not None for value in item.values()):
            raise ValueError('Cash-flow data outside declared count')
    result['flows']=flows
    return result

def validate_ledgers(wb):
    if SUMMARY not in wb:return []
    errors=[]
    for row in range(2,wb[SUMMARY].max_row+1):
        try:ledger_reference(read_ledger(wb,row))
        except (ValueError,KeyError,TypeError) as exc:errors.append(str(exc))
    return errors

def add_dated_ledgers(wb,scenarios,contract):
    from scenario_formulas import literal
    summary=wb.create_sheet(SUMMARY);detail=wb.create_sheet(DETAIL)
    literal(summary,['Senaryo','NBD baz tarihi','Hedef para','Hedef para yıllık efektif oran','Kayıt adedi','Aynı kapsam NBD bazı',
                     'Tam kapsam nominal net nakit','Tam kapsam NBD','Baza göre NBD farkı','Oransal fark','Durum / kapsam','Kapsam beyanı','Kur/NBD yöntemi'])
    literal(detail,['Senaryo','Sıra','Tekil maliyet kimliği','Maliyet türü','Kaynak para tutarı','Kaynak para',
                    'Nakit tarihi','Projeksiyon kur tarihi','Hedef/kaynak kur','Kur varsayımı dayanağı','Baz tarihinden gün','İşaretli hedef nakit','NBD','Durum'])
    def inp(ws,cell,value,kind):
        validate_ledger_input(kind,value);ws[cell]=value
        if isinstance(value,str):ws[cell].data_type='s'
        ws[cell].protection=Protection(locked=False);ws[cell].fill=PatternFill('solid',fgColor='DDEBFF')
        contract['extra_inputs'][ws.title+'!'+cell]=kind
        if kind in ('ledger_amount','ledger_baseline','ledger_rate','ledger_fx','ledger_count'):
            upper=MAX_FLOWS if kind=='ledger_count' else 10 if kind=='ledger_rate' else None
            dv=DataValidation(type='whole' if kind=='ledger_count' else 'decimal',operator='between' if upper else 'greaterThan' if kind=='ledger_fx' else 'greaterThanOrEqual',formula1='1' if kind=='ledger_count' else '0',formula2=str(upper) if upper else None,allow_blank=True)
        elif kind=='ledger_category':dv=DataValidation(type='list',formula1='"'+','.join(CATEGORIES)+'"',allow_blank=True)
        else:dv=DataValidation(type='textLength',operator='between',formula1='1',formula2='200',allow_blank=True)
        dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(cell)
    def formula(ws,cell,expression,guard=None):
        text='='+expression if guard is None else '=IF(IFERROR('+guard+',FALSE),'+expression+',"")'
        ws[cell]=text;contract['formulas'][ws.title+'!'+cell]=text
    for row,scenario in enumerate(scenarios,start=2):
        ledger=scenario.get('dated_costs');ledger={} if ledger is None else ledger
        reference=ledger_reference(ledger);flows=ledger.get('flows') or []
        literal(summary,[scenario['id']])
        for col,key,kind in ((2,'valuation_date','date'),(3,'base_currency','currency'),(4,'annual_rate','rate'),(5,None,'count'),(6,'baseline','baseline'),(12,'scope','scope'),(13,'method','method')):
            inp(summary,get_column_letter(col)+str(row),len(flows) or None if key is None else ledger.get(key),'ledger_'+kind)
        v=lambda col:f"'{SUMMARY}'!{col}{row}"
        first=2+(row-2)*MAX_FLOWS;last=first+MAX_FLOWS-1
        codes=lambda ref:','.join(f'CODE(MID({ref},{i},1))>=65,CODE(MID({ref},{i},1))<=90' for i in (1,2,3))
        policy=f'AND({excel_date_guard(v("B"))},LEN({v("C")})=3,{codes(v("C"))},{v("M")}="{METHOD}")'
        count=f'AND(ISNUMBER({v("E")}),{v("E")}>=1,{v("E")}<=48,{v("E")}=INT({v("E")}))'
        rate=f'AND(ISNUMBER({v("D")}),{v("D")}>=0,{v("D")}<=10)'
        detail_range=lambda col:f"'{DETAIL}'!{col}{first}:{col}{last}"
        for slot,yr in enumerate(range(first,last+1),start=1):
            item=flows[slot-1] if slot<=len(flows) else {};literal(detail,[scenario['id'],slot])
            for col,key,kind in ((3,'cost_id','id'),(4,'category','category'),(5,'amount','amount'),(6,'currency','currency'),(7,'cash_date','date'),(8,'rate_date','date'),(9,'fx_rate','fx'),(10,'rate_source','text')):
                inp(detail,get_column_letter(col)+str(yr),item.get(key),'ledger_'+kind)
            active=f'AND({count},{slot}<={v("E")})'
            dates=f'AND({policy},{active},{excel_date_guard("G"+str(yr))},{excel_date_guard("H"+str(yr))},G{yr}=H{yr})'
            origin=excel_date_serial(v('B'));when=excel_date_serial('G'+str(yr))
            # Excel includes a fictitious 1900-02-29; LibreOffice does not.
            # Derive the engine's extra day from two real dates, then remove it
            # only when the interval crosses the boundary. Both engines must
            # equal Gregorian date subtraction, including pre-March 1900.
            extra_day='(DATE(1900,3,1)-DATE(1900,1,1)-59)'
            difference=f'({when}-{origin}-IF(AND({origin}<DATE(1900,3,1),{when}>=DATE(1900,3,1)),{extra_day},0))'
            formula(detail,'K'+str(yr),difference,f'AND({dates},{difference}>=0,{difference}<=36500)')
            known=f'AND({active},ISNUMBER(K{yr}),LEN(TRIM(C{yr}))>0,LEN(C{yr})<=100,SUMPRODUCT(--EXACT({detail_range("C")},C{yr}))=1,OR('+','.join(f'D{yr}="{kind}"' for kind in CATEGORIES)+f'),ISNUMBER(E{yr}),E{yr}>=0,LEN(F{yr})=3,{codes("F"+str(yr))},ISNUMBER(I{yr}),I{yr}>0,OR(F{yr}<>{v("C")},I{yr}=1),LEN(TRIM(J{yr}))>0,LEN(J{yr})<=200)'
            formula(detail,'L'+str(yr),f'E{yr}*I{yr}*IF(OR(D{yr}="VAT_REFUND",D{yr}="RESIDUAL"),-1,1)',known)
            formula(detail,'M'+str(yr),f'L{yr}/(1+{v("D")})^(K{yr}/365)',f'AND(ISNUMBER(L{yr}),{rate})')
            formula(detail,'N'+str(yr),f'IF(ISNUMBER(M{yr}),"Kullanıcı nakit tarihi kur projeksiyonu; hedef para NBD","Hesaplanamadı veya kullanılmayan kayıt")')
        coverage=','.join(f'COUNTIF({detail_range("D")},"{kind}")>=1' for kind in CATEGORIES)
        filled=','.join(f'COUNTA({detail_range(col)})={v("E")}' for col in 'CDEFGHIJ')
        complete=f'AND({count},{v("L")}="{SCOPE}",{coverage},{filled},COUNT({detail_range("L")})={v("E")})'
        formula(summary,'G'+str(row),'SUM('+detail_range('L')+')',complete)
        formula(summary,'H'+str(row),'SUM('+detail_range('M')+')',f'AND({complete},COUNT({detail_range("M")})={v("E")},{rate})')
        formula(summary,'I'+str(row),f'H{row}-F{row}',f'AND(ISNUMBER(H{row}),ISNUMBER(F{row}),F{row}>=0)')
        formula(summary,'J'+str(row),f'I{row}/F{row}',f'AND(ISNUMBER(I{row}),ISNUMBER(F{row}),F{row}<>0)')
        formula(summary,'K'+str(row),f'IF(ISNUMBER(H{row}),"Beyan edilen tam kullanıcı nakit senaryosu; diğer sayfa toplamları eklenmez; resmî sıralama yok","Hesaplanamadı: eksik maliyet türü/tutar/tarih/kur/dayanak/kapsam")')
        ref=contract['reference'][str(row)]
        for key,col in (('nominal','G'),('present_value','H'),('difference','I'),('relative_difference','J')):
            if reference[key] is not None:ref[SUMMARY+'!'+col]=reference[key]
    for ws in (summary,detail):
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions;ws.protection.sheet=True
        for cell in ws[1]:cell.font=Font(bold=True)
        for cells in ws:
            for cell in cells:ws.column_dimensions[cell.column_letter].width=28
