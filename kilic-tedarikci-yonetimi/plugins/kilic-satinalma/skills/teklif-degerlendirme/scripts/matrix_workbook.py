"""Shared offer-by-requirement matrix and independently checkable count formulas."""
from openpyxl.styles import Alignment, Font, PatternFill
from report_model import MATRIX_HEADERS, matrix_values
from pdf_renderer import label
from scenario_formulas import literal, source_cell_value


def add_matrix(wb, model, contract):
    ws = wb.create_sheet("Sartname Matrisi")
    literal(ws, MATRIX_HEADERS)
    for assessment in model.get("matrix", []):
        values = matrix_values(assessment)
        for index in (4, 5, 6):
            values[index] = label(values[index])
        literal(ws, values)
    counts = wb.create_sheet("Sartname Sayim")
    literal(counts, ["Firma", "Teklif", "Uygunsuz", "Belirtilmemiş", "Teyit bekliyor", "Kritik uygunsuz"])
    names = {s["supplier_id"]: s["display_name"] for s in model.get("suppliers", [])}
    last = max(2, ws.max_row)
    for rn, offer in enumerate(model.get("offers", []), start=2):
        literal(counts, [names.get(offer["supplier_id"], offer["supplier_id"]), offer["offer_id"]])
        for column, state in (("C", "Uygunsuz"), ("D", "Belirtilmemiş"),
                              ("E", "Teyit bekliyor"), ("F", "Uygunsuz")):
            value = (f'=COUNTIFS(\'Sartname Matrisi\'!$B$2:$B${last},B{rn},'
                     f'\'Sartname Matrisi\'!$E$2:$E${last},"{state}"')
            if column == "F":
                value += f',\'Sartname Matrisi\'!$F$2:$F${last},"Kritik"'
            value += ")"
            counts[f"{column}{rn}"] = value
            contract["formulas"][counts.title + "!" + f"{column}{rn}"] = value
    for sheet in (ws, counts):
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.protection.sheet = True
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DFE8F0")
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                sheet.column_dimensions[cell.column_letter].width = 32


def seal_source_cells(wb, contract):
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    continue
                if sheet.title == "Senaryo Girdileri" and cell.coordinate in contract["inputs"]:
                    continue
                if sheet.title + "!" + cell.coordinate in contract.get("extra_inputs",{}):
                    continue
                contract["source_cells"][sheet.title + "!" + cell.coordinate] = source_cell_value(cell.value)
