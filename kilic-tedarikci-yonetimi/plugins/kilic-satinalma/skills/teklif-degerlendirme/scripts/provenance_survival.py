from __future__ import annotations
import argparse, json
from pathlib import Path
from typing import Any
from common import json_text, sha256_file, stable_id, utc_now
from db import connect, rows_as_dict
from active_records import current_artifacts
from display_values import display_cell
from reviewed_snapshot import load_latest

def _j(v,default=None):
    try: return v if isinstance(v,(dict,list)) else json.loads(v) if v is not None else default
    except Exception: return default

def check(run_root:Path)->dict[str,Any]:
    from openpyxl import load_workbook
    snap,srow=load_latest(run_root); sha=srow['sha256']; con=connect(run_root); findings=[]
    try:
        facts={x['fact_id']:x for x in snap.get('facts',[])}; calcs={x['calculation_id']:x for x in snap.get('calculations',[])}; reqs={x['requirement_id']:x for x in snap.get('requirements',[])}; claims={x['claim_id']:x for x in snap.get('claims',[])}
        latest=current_artifacts(con)
        for kind in ['xlsx', 'decision_snapshot', 'pdf']:
            a=latest.get(kind)
            if not a: findings.append({'severity':'critical','kind':kind,'message':'required artifact missing'}); continue
            path=run_root/a['relative_path']
            if not path.exists() or sha256_file(path)!=a['sha256']: findings.append({'severity':'critical','kind':kind,'message':'current artifact hash mismatch'})
            parents=_j(a.get('parent_hashes_json'),[]) or []
            if sha not in parents: findings.append({'severity':'critical','kind':kind,'message':'reviewed snapshot parent hash missing'})
            binds=rows_as_dict(con,"SELECT * FROM artifact_bindings WHERE artifact_id=?",[a['artifact_id']])
            if kind in {'xlsx','decision_snapshot'} and not binds: findings.append({'severity':'critical','kind':kind,'message':'artifact bindings missing'})
            for b in binds:
                if b.get('snapshot_sha256')!=sha: findings.append({'severity':'critical','kind':kind,'component':b.get('component_id'),'message':'binding snapshot hash mismatch'})
                for field,universe in [('fact_ids_json',facts),('calculation_ids_json',calcs),('requirement_ids_json',reqs),('claim_ids_json',claims)]:
                    for ref in _j(b.get(field),[]) or []:
                        if ref not in universe: findings.append({'severity':'critical','kind':kind,'component':b.get('component_id'),'message':f'unknown bound reference {ref}'})
        x=latest.get('xlsx')
        if x:
            path=run_root/x['relative_path']; wb=load_workbook(path,data_only=False,read_only=False)
            if 'Kanitlar' not in wb.sheetnames: findings.append({'severity':'critical','kind':'xlsx','message':'Kanitlar sheet missing'})
            else:
                ws=wb['Kanitlar']; obs={ws.cell(r,1).value:(ws.cell(r,4).value,ws.cell(r,5).value,ws.cell(r,6).value,ws.cell(r,10).value) for r in range(2,ws.max_row+1)}
                evid={x['evidence_id']:x for x in snap.get('evidence',[])}; docs={x['document_id']:x for x in snap.get('documents',[])}
                for fid,f in facts.items():
                    if not f.get('evidence_id'): continue
                    row=obs.get(fid)
                    if not row: findings.append({'severity':'critical','kind':'xlsx','fact_id':fid,'message':'evidence row missing'}); continue
                    ev=evid.get(f['evidence_id'],{}); source=docs.get(ev.get('document_id'),{}).get('relative_path')
                    expected=(display_cell('Kanitlar',4,f.get('evidence_status')),source,display_cell('Kanitlar',6,ev.get('locator_json')),sha)
                    if tuple(row)!=expected: findings.append({'severity':'critical','kind':'xlsx','fact_id':fid,'message':'status/source/locator/snapshot provenance did not survive export'})
            if 'Iddia Kaydi' not in wb.sheetnames: findings.append({'severity':'critical','kind':'xlsx','message':'claim registry sheet missing'})
            else:
                ws=wb['Iddia Kaydi']; obs={ws.cell(r,1).value:(ws.cell(r,2).value,ws.cell(r,4).value,ws.cell(r,5).value,ws.cell(r,6).value,ws.cell(r,7).value,ws.cell(r,8).value,ws.cell(r,9).value,ws.cell(r,10).value) for r in range(2,ws.max_row+1)}
                for cid,c in claims.items():
                    expected=tuple(display_cell('Iddia Kaydi',col,val) for col,val in zip([2,4,5,6,7,8,9,10],[c.get('claim_type'),c.get('support_status'),c.get('fact_ids_json'),c.get('calculation_ids_json'),c.get('requirement_ids_json'),c.get('assertions_json'),c.get('caveat'),sha]))
                    if obs.get(cid)!=expected: findings.append({'severity':'critical','kind':'xlsx','claim_id':cid,'message':'claim type/support references/assertions/caveat/snapshot did not survive export'})
            # Bound fact cells must carry inspectable provenance comments.
            binds=rows_as_dict(con,"SELECT component_id,locator_json,fact_ids_json FROM artifact_bindings WHERE artifact_id=?",[x['artifact_id']])
            for b in binds:
                fids=_j(b.get('fact_ids_json'),[]) or []
                if not fids: continue
                loc=_j(b.get('locator_json'),{}) or {}; sheet=loc.get('sheet'); cell=loc.get('cell')
                if sheet in wb.sheetnames and cell:
                    comment=wb[sheet][cell].comment.text if wb[sheet][cell].comment else ''
                    if any(fid not in comment for fid in fids) or sha not in comment:
                        # Evidence ledger bindings themselves do not need comments; their row is provenance.
                        if sheet not in {'Kanitlar','Iddia Kaydi'}: findings.append({'severity':'high','kind':'xlsx','component':b['component_id'],'message':'bound fact cell lacks fact/snapshot provenance comment'})
        d=latest.get('decision_snapshot')
        if d:
            obj=json.loads((run_root/d['relative_path']).read_text(encoding='utf-8'))
            if obj.get('reviewed_snapshot_sha256')!=sha: findings.append({'severity':'critical','kind':'decision_snapshot','message':'snapshot hash not preserved'})
            if {c.get('claim_id') for c in obj.get('claims',[])} != set(claims): findings.append({'severity':'critical','kind':'decision_snapshot','message':'claim set differs from reviewed snapshot'})
        status='PASS' if not any(f.get('severity') in {'critical','high'} for f in findings) else 'FAIL'
        out={'status':status,'reviewed_snapshot_sha256':sha,'findings':findings,'checked_at':utc_now()}
        p=run_root/'output'/'provenance_survival.json'; p.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8'); ps=sha256_file(p); aid=stable_id('art','provenance_survival',ps,length=22)
        con.execute("INSERT OR REPLACE INTO artifacts(artifact_id,artifact_type,relative_path,sha256,parent_hashes_json,status,created_at) VALUES (?,?,?,?,?,?,?)",[aid,'provenance_survival',str(p.relative_to(run_root)),ps,json_text([sha]),'ready',utc_now()])
        con.execute("DELETE FROM qa_findings WHERE finding_type='provenance_survival'")
        for f in findings:
            fid=stable_id('qa','provenance',json_text(f),length=22); sev=f.get('severity','high')
            con.execute("INSERT OR REPLACE INTO qa_findings(finding_id,artifact_id,finding_type,locator_json,severity,status,message,finding_hash,created_at) VALUES (?,?,?,?,?,?,?,?,?)",[fid,(x or {}).get('artifact_id'),'provenance_survival',json_text({'component':f.get('component'),'fact_id':f.get('fact_id'),'claim_id':f.get('claim_id')}),sev,'open',f.get('message','provenance failure'),stable_id('h',json_text(f),length=20),utc_now()])
        if hasattr(con,'commit'): con.commit()
        return out
    finally: con.close()

def main():
    ap=argparse.ArgumentParser(description='Verify that source status, lineage and caveats survive export'); ap.add_argument('--run',required=True); a=ap.parse_args(); print(json.dumps(check(Path(a.run)),ensure_ascii=False,indent=2))
if __name__=='__main__': main()
