from __future__ import annotations
from artifact_values import comparison_fact_path, comparison_fact_value, evidence_label, offer_display_status, pending_offer_fields, ranking_visible
from effective_report_values import buyer_term, weighted_score_display

import argparse
import json
import platform
from pathlib import Path
from typing import Any, Iterable

from common import json_text, sha256_file, stable_id, utc_now
from db import connect, get_meta, rows_as_dict


def _json(v: Any, default: Any = None) -> Any:
    if isinstance(v, (dict, list)):
        return v
    if v is None:
        return default
    try:
        return json.loads(v)
    except Exception:
        return default


def _recommendation_text(recommendation: Any, suppliers: dict[str, str], offers: list[dict[str, Any]]) -> str:
    """Render the recorded decision action with a canonical supplier name, if bound."""
    if recommendation is None:
        return "Kesin oneri yok"
    if isinstance(recommendation, str):
        return recommendation.strip() or "Kesin oneri yok"
    if not isinstance(recommendation, dict):
        return "Kesin oneri yok"
    action = recommendation.get("text") or recommendation.get("action")
    if not isinstance(action, str) or not action.strip():
        return "Kesin oneri yok"
    supplier_id = recommendation.get("supplier_id")
    if not supplier_id and recommendation.get("offer_id"):
        supplier_id = next((o.get("supplier_id") for o in offers if o.get("offer_id") == recommendation["offer_id"]), None)
    name = suppliers.get(supplier_id) if supplier_id else None
    return f"{name}: {action.strip()}" if name else action.strip()


def _safe_sheet(name: str) -> str:
    return name[:31]


def _excel_safe(value: Any) -> Any:
    """Prevent untrusted source text from becoming an Excel formula.

    Numbers/dates remain native values. Text beginning with =,+,-,@ is written
    as a literal string. Intentional workbook formulas must be created by code
    outside this helper (v5.0.4 currently avoids them for decision-critical
    values so cached-value validation is deterministic).
    """
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _append_safe(ws, row: Iterable[Any]) -> None:
    ws.append([_excel_safe(v) for v in row])


def _style_workbook(wb):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from display_values import present_workbook
    present_workbook(wb)
    fills = {
        "header": PatternFill("solid", fgColor="1F4E78"),
        "section": PatternFill("solid", fgColor="D9EAF7"),
        "input": PatternFill("solid", fgColor="E2F0D9"),
        "warning": PatternFill("solid", fgColor="FCE4D6"),
        "error": PatternFill("solid", fgColor="F4CCCC"),
        "control": PatternFill("solid", fgColor="E4DFEC"),
    }
    thin_gray = Side(style="thin", color="D9D9D9")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fills["header"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        for col in ws.columns:
            values = [str(c.value) if c.value is not None else "" for c in list(col)[:200]]
            max_len = min(60, max([len(x) for x in values] + [8]))
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = max(10, min(max_len + 2, 32))
        if ws.title == "Hesap Detayi":
            from calculation_detail import COLUMN_WIDTHS
            for letter, width in COLUMN_WIDTHS.items():
                ws.column_dimensions[letter].width = width
        # Preserve readable type on wide comparison printouts. Excel repeats
        # supplier/offer identity on every horizontal page instead of scaling
        # all 22 columns into a single A3 page.
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape" if ws.max_column > 2 else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws.max_column > 8 else ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        ws.print_area = ws.dimensions
        if ws.title == "Teklif Karsilastirma":
            ws.sheet_properties.pageSetUpPr.fitToPage = False
            ws.page_setup.fitToWidth = 0
            ws.page_setup.fitToHeight = 0
            ws.page_setup.scale = 100
            ws.print_title_cols = "A:B"
            # Let native Excel paginate at readable 100% scale. Manual column
            # breaks can consume the repeated-title band and omit supplier
            # identities on intervening horizontal pages (native regression).
            # PDF has its own explicit bands in build_print_workbook.
        for row in ws.iter_rows(min_row=2):
            lines = max((sum(max(1, (len(part)+int(ws.column_dimensions[c.column_letter].width)-1)//int(ws.column_dimensions[c.column_letter].width)) for part in str(c.value or "").splitlines()) for c in row), default=1)
            ws.row_dimensions[row[0].row].height = min(400, max(18, 15 * lines))
    return fills, thin_gray


def _write_table(ws, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    _append_safe(ws, headers)
    found=False
    for row in rows:
        found=True
        _append_safe(ws, row)
    if not found:
        _append_safe(ws, ["Kayıt bulunmuyor."])


def _artifact(con, path: Path, run_root: Path, artifact_type: str, parent_hashes: list[str] | None = None) -> str:
    sha = sha256_file(path)
    aid = stable_id("art", artifact_type, sha, length=22)
    rel = str(path.resolve().relative_to(run_root.resolve()))
    con.execute(
        "INSERT OR REPLACE INTO artifacts(artifact_id,artifact_type,relative_path,sha256,parent_hashes_json,status,created_at) VALUES (?,?,?,?,?,?,?)",
        [aid, artifact_type, rel, sha, json_text(parent_hashes or []), "ready", utc_now()],
    )
    return aid


def _recalc_with_excel(path: Path) -> dict[str, Any]:
    if platform.system() != "Windows":
        return {"available": False, "reason": "not_windows"}
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        return {"available": False, "reason": f"pywin32_unavailable:{exc}"}
    excel = None
    wb = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        prior_security = excel.AutomationSecurity
        prior_events = excel.EnableEvents
        excel.AutomationSecurity = 3  # msoAutomationSecurityForceDisable
        excel.EnableEvents = False
        wb = excel.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False, AddToMru=False, IgnoreReadOnlyRecommended=True)
        excel.CalculateFullRebuild()
        wb.Save()
        wb.Close(SaveChanges=True)
        wb = None
        return {"available": True, "status": "recalculated"}
    except Exception as exc:
        return {"available": False, "reason": f"excel_recalc_error:{exc}"}
    finally:
        try:
            if wb is not None:
                wb.Close(SaveChanges=False)
        except Exception:
            pass
        try:
            if excel is not None:
                try:
                    excel.AutomationSecurity = prior_security
                    excel.EnableEvents = prior_events
                except Exception:
                    pass
                excel.Quit()
        except Exception:
            pass


def _export_pdf_with_excel(xlsx: Path, pdf: Path) -> dict[str, Any]:
    if platform.system() != "Windows":
        return {"available": False, "reason": "not_windows"}
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        return {"available": False, "reason": f"pywin32_unavailable:{exc}"}
    excel = None
    wb = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        prior_security = excel.AutomationSecurity
        prior_events = excel.EnableEvents
        excel.AutomationSecurity = 3  # msoAutomationSecurityForceDisable
        excel.EnableEvents = False
        from print_layout import build_print_workbook
        print_xlsx = pdf.with_name(pdf.stem + "_Baski.xlsx")
        print_info = build_print_workbook(xlsx, print_xlsx)
        wb = excel.Workbooks.Open(str(print_xlsx.resolve()), UpdateLinks=0, ReadOnly=True, AddToMru=False, IgnoreReadOnlyRecommended=True)
        wb.ExportAsFixedFormat(0, str(pdf.resolve()))
        wb.Close(False)
        wb = None
        return {"available": True, "path": str(pdf), "print_workbook": print_info}
    except Exception as exc:
        return {"available": False, "reason": f"excel_pdf_error:{exc}"}
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if excel is not None:
                try:
                    excel.AutomationSecurity = prior_security
                    excel.EnableEvents = prior_events
                except Exception:
                    pass
                excel.Quit()
        except Exception:
            pass


def _insert_bindings(con, artifact_id: str, snapshot_sha: str, bindings: list[dict[str, Any]]) -> None:
    con.execute("DELETE FROM artifact_bindings WHERE artifact_id=?", [artifact_id])
    for item in bindings:
        locator = item.get("locator") or {}
        bid = stable_id("bind", artifact_id, item["component_id"], json_text(locator), snapshot_sha, length=24)
        con.execute(
            "INSERT INTO artifact_bindings(binding_id,artifact_id,component_id,component_type,locator_json,fact_ids_json,calculation_ids_json,requirement_ids_json,claim_ids_json,snapshot_sha256,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            [bid, artifact_id, item["component_id"], item.get("component_type", "cell"), json_text(locator),
             json_text(item.get("fact_ids") or []), json_text(item.get("calculation_ids") or []),
             json_text(item.get("requirement_ids") or []), json_text(item.get("claim_ids") or []), snapshot_sha, utc_now()],
        )


def export_report(run_root: Path) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from reviewed_snapshot import load_latest

    snap, snap_row = load_latest(run_root)
    snapshot_sha = snap_row["sha256"]
    meta = snap.get("meta", {})
    profile = str(meta.get("profile") or "standart")
    summary = meta.get("decision_summary") or ""
    recommendation = _json(meta.get("decision_recommendation"), None)
    decision_status = meta.get("decision_status") or "UNKNOWN"

    inventory = snap.get("documents", [])
    suppliers = {r["supplier_id"]: r["display_name"] for r in snap.get("suppliers", [])}
    offers = snap.get("offers", [])
    facts = snap.get("facts", [])
    evidence = {r["evidence_id"]: r for r in snap.get("evidence", [])}
    requirements = snap.get("requirements", [])
    scope = snap.get("scope_items", [])
    issues = snap.get("open_issues", [])
    calcs = snap.get("calculations", [])
    exclusions = snap.get("exclusions", [])
    events = snap.get("workflow_events", [])
    docs = {r["document_id"]: r for r in inventory}
    authority = {r["document_id"]: r for r in snap.get("source_authority", [])}
    data_quality = snap.get("data_quality_findings", [])
    comparison = snap.get("comparability_checks", [])
    independent = snap.get("independent_calculations", [])
    bridges = snap.get("cost_bridges", [])
    claims = snap.get("claims", [])
    ledger = snap.get("scope_ledger", [])

    from display_values import label as display_label, structured as display_structured

    def source_comment(fact: dict[str, Any] | None) -> str | None:
        if not fact:
            return None
        if not fact.get("evidence_id") or fact["evidence_id"] not in evidence:
            return (f"Veri kimliği: {fact.get('fact_id')} | Kaynak kanıtı yok; "
                    f"kaynak durum: {display_label(fact.get('evidence_status'))}. "
                    f"Alıcı hesabı ayrı gösterilir. | İncelenen sürüm özeti: {snapshot_sha}")
        ev = evidence[fact["evidence_id"]]
        doc = docs.get(ev.get("document_id"), {})
        auth = authority.get(ev.get("document_id"), {})
        return (f"Veri kimliği: {fact.get('fact_id')} | Kaynak: {doc.get('relative_path')} | "
                f"Kaynak konumu: {display_structured(ev.get('locator_json'))} | Durum: {display_label(fact.get('evidence_status'))} | "
                f"Otorite: {display_label(auth.get('authority_status') or 'unresolved')} | Alinti: {ev.get('quote') or ''} | İncelenen sürüm özeti: {snapshot_sha}")

    con = connect(run_root)
    try:
        wb = Workbook(); wb.remove(wb.active)
        bindings: list[dict[str, Any]] = []
        def bind(sheet: str, cell: str, *, fact_ids=None, calculation_ids=None, requirement_ids=None, claim_ids=None, component_type="cell", value_path=None):
            bindings.append({"component_id":f"xlsx:{sheet}:{cell}","component_type":component_type,"locator":{"sheet":sheet,"cell":cell,**({"value_path":value_path} if value_path else {})},
                             "fact_ids":fact_ids or [],"calculation_ids":calculation_ids or [],"requirement_ids":requirement_ids or [],"claim_ids":claim_ids or []})

        ws=wb.create_sheet(_safe_sheet("Ozet")); _append_safe(ws,["Alan","Deger"])
        summary_claim_ids=[c["claim_id"] for c in claims]
        snapshot_status = {"READY":"İnceleme hazır", "PRELIMINARY":"Ön değerlendirme"}.get(snap.get("status"),snap.get("status") or "Bilinmiyor")
        for label,value in [
            ("Karar İçin Bilgi Yeterliliği",decision_status),("İncelenen Kayıt Durumu",snapshot_status),
            ("Profil",profile),("Yonetici Ozeti",summary),
            ("Oneri",_recommendation_text(recommendation,suppliers,offers)),
            ("Aktif Teklif Sayisi",sum(1 for x in offers if x.get("status")=="active")),
            ("Acik RFI",sum(1 for x in issues if x.get("status")=="open")),
            ("Reviewed Snapshot SHA256",snapshot_sha),
            ("Not","Nihai secim yetkili karar merciine aittir.")]: _append_safe(ws,[label,value])
        bind(ws.title,"B5",claim_ids=summary_claim_ids); bind(ws.title,"B6",claim_ids=[c['claim_id'] for c in claims if c.get('claim_type')=='recommendation'])

        ws=wb.create_sheet(_safe_sheet("Envanter")); _write_table(ws,["Dosya","Format","Rol","Tedarikci","Teklif","Revizyon","Okuma","Kaynak Otoritesi","Not","SHA256"],[
            [x.get("relative_path"),x.get("format_family"),x.get("document_role"),suppliers.get(x.get("supplier_id"),x.get("supplier_id")),x.get("offer_id"),x.get("revision_label"),x.get("parse_status"),authority.get(x.get('document_id'),{}).get("authority_status"),x.get("parse_notes"),x.get("sha256")] for x in inventory])

        active_offers=[x for x in offers if x.get("status")=="active"]
        fact_map={(x.get("offer_id"),x["field_key"]):x for x in facts if x.get("offer_id")}
        cost_map={}; cost_calc_id={}; price_score_map={}; weighted_score_map={}; scoring_calc_id=None; scoring_result={}
        for c in calcs:
            obj=_json(c.get("result_json"),{}) or {}
            if c.get("calculation_type")=="cost_summary" and c.get("offer_id"):
                cost_map[c["offer_id"]]=obj; cost_calc_id[c["offer_id"]]=c["calculation_id"]
            if c.get("calculation_type")=="scoring" and not c.get("offer_id"):
                scoring_result=obj
                scoring_calc_id=c["calculation_id"]
                for item in obj.get("price_scores",[]): price_score_map[item["offer_id"]]=item.get("price_score_0_10")
                for item in obj.get("weighted_scores",[]): weighted_score_map[item["offer_id"]]=item.get("total_score_0_10")

        if not ranking_visible(meta):
            price_score_map.clear(); weighted_score_map.clear(); scoring_calc_id=None
        ws=wb.create_sheet(_safe_sheet("Teklif Karsilastirma")); headers=["Tedarikci","Teklif","Teklif Toplam","Para Birimi","Kapsam Duzeltmesi","Bilinen Ara Toplam","Odeme NBD","TCO NBD","Karsilastirilabilir Maliyet","Baz Para","Baz Tutar","Fiyat Puani (0-10)","Toplam Puan (0-10)","Teslim","Odeme","Gecerlilik","Durum","Teklif Kanıt Durumu","Teyit Bekleyen Alanlar","Hesap Para Birimi","Para Birimi Esası","Alıcı Miktar Düzeltmesi"]
        _append_safe(ws,headers)
        for idx,off in enumerate(active_offers,start=2):
            oid=off["offer_id"]; total=fact_map.get((oid,"total"),{}); currency=fact_map.get((oid,"currency"),{}); delivery=fact_map.get((oid,"delivery"),{}); payment=fact_map.get((oid,"payment"),{}); cost=cost_map.get(oid,{})
            row=[suppliers.get(off.get("supplier_id"),off.get("display_name") or off.get("supplier_id")),oid,comparison_fact_value(total),comparison_fact_value(currency),float(cost["known_scope_adjustment"]) if cost.get("known_scope_adjustment") is not None else None,float(cost["acquisition_cost_basis"]) if cost.get("acquisition_cost_basis") is not None else None,float(cost["acquisition_present_value"]) if cost.get("acquisition_present_value") is not None else None,float(cost["tco_present_value"]) if cost.get("tco_present_value") is not None else None,float(cost["comparable_total"]) if cost.get("comparable_total") is not None else None,cost.get("base_currency"),float(cost["base_amount"]) if cost.get("base_amount") is not None else None,float(price_score_map[oid]) if price_score_map.get(oid) is not None else None,float(weighted_score_map[oid]) if weighted_score_map.get(oid) is not None else None,comparison_fact_value(delivery),comparison_fact_value(payment),off.get("validity_end"),offer_display_status(oid,facts,exclusions,meta),evidence_label(total),", ".join(pending_offer_fields(oid,facts))]
            row.extend([cost.get("currency") if cost.get("currency_basis") else None,
                        cost.get("currency_basis"),
                        float(cost["buyer_quantity_adjustment"]) if cost.get("buyer_quantity_adjustment") is not None else None])
            terms={field:buyer_term(cost,field) for field in ('payment','validity')}
            for field,col in [('payment',15),('validity',16)]:
                if terms[field]:row[col-1]=terms[field]['value']
            row[12]=weighted_score_display(scoring_result,weighted_score_map.get(oid),ranking_visible(meta))
            _append_safe(ws,row)
            for field,col in [("total",3),("currency",4),("delivery",14),("payment",15)]:
                f=fact_map.get((oid,field)); comment=source_comment(f)
                if comment: ws.cell(idx,col).comment=Comment(comment,"OpenAI")
                if f and not terms.get(field): bind(ws.title,ws.cell(idx,col).coordinate,fact_ids=[f['fact_id']],value_path=comparison_fact_path(f))
            for field,col in [('payment',15),('validity',16)]:
                term=terms[field]
                if term:
                    prior=ws.cell(idx,col).comment.text if ws.cell(idx,col).comment else ''
                    ws.cell(idx,col).comment=Comment('Geçerli değerlendirme esası: uygulanmış kullanıcı kanıtı. Özgün belge bilgisi Ticari Kosullar ve Kanitlar sayfalarında korunur.\n'+prior,'OpenAI')
                    bind(ws.title,ws.cell(idx,col).coordinate,calculation_ids=[cost_calc_id[oid]],value_path=term['path'])
            if oid in cost_calc_id:
                for col in (5,6,7,8,9,10,11): bind(ws.title,ws.cell(idx,col).coordinate,calculation_ids=[cost_calc_id[oid]])
                for col,path in ((20,"currency"),(21,"currency_basis"),(22,"buyer_quantity_adjustment")):
                    if row[col-1] is not None:
                        bind(ws.title,ws.cell(idx,col).coordinate,calculation_ids=[cost_calc_id[oid]],value_path=path)
            if scoring_calc_id:
                for col in (12,13): bind(ws.title,ws.cell(idx,col).coordinate,calculation_ids=[scoring_calc_id])
        # These columns deliberately describe the retained supplier source,
        # while payment and validity show the current reviewed evaluation basis.
        for col,title in [(15,'Ödeme (değerlendirme)'),(16,'Geçerlilik (değerlendirme)'),(17,'Kaynak Belge Durumu'),(18,'Kaynak Toplam Kanıtı'),(19,'Kaynak Teyit Alanları')]:
            ws.cell(1,col,title)
        for r in range(2,ws.max_row+1):
            for c in (3,5,6,7,8,9,11): ws.cell(r,c).number_format='#,##0.00;[Red](#,##0.00);-'
            ws.cell(r,22).number_format='#,##0.00;[Red](#,##0.00);0.00'
            for c in (12,13): ws.cell(r,c).number_format='0.00'
            if ws.cell(r,18).value!="Doğrulandı":
                label=ws.cell(r,18).value
                ws.cell(r,3).number_format=f'#,##0.00" ({label})";[Red](#,##0.00)" ({label})";0" ({label})"'

        buyer_rows=[]
        for off in active_offers:
            oid=off["offer_id"]
            cost=cost_map.get(oid,{})
            for kind in ("currency","reference_quantity","scope_cost","payment_terms","warranty_terms","validity_date","quote_date"):
                for index,item in enumerate((cost.get("buyer_inputs") or {}).get(kind) or []):
                    target=item.get("quantity_item_id") if kind=="reference_quantity" else item.get("scope_id") if kind=="scope_cost" else {"currency":"Para birimi","payment_terms":"Ödeme koşulları","warranty_terms":"Garanti koşulları","validity_date":"Son geçerlilik tarihi","quote_date":"Teklif tarihi"}.get(kind)
                    raw=item.get("quantity") if kind=="reference_quantity" else item.get("amount") if kind=="scope_cost" else item.get("value") if kind=="currency" else None
                    value=(float(raw) if kind in {"reference_quantity","scope_cost"} and raw is not None else raw)
                    if kind=="payment_terms":
                        steps=item.get("milestones") or []
                        value="; ".join(f"{display_label(x.get('trigger'))}: %{x.get('percent')}" for x in steps if isinstance(x,dict))
                    elif kind=="warranty_terms":
                        value=f"{item.get('duration_years')} yıl, işin tamamı" if item.get('scope')=='whole_work' else None
                    elif kind in {"validity_date","quote_date"}:
                        value=item.get('value') if item.get('date_status')=='known' else 'Bilinmiyor (kullanıcı yanıtı)'
                    buyer_rows.append((oid,kind,index,[oid,kind,target,value,item.get("unit") if kind=="reference_quantity" else item.get("currency") if kind=="scope_cost" else None,
                                                       item.get("answer_text"),item.get("issue_id"),cost_calc_id.get(oid)]))
        if buyer_rows:
            ws=wb.create_sheet(_safe_sheet("Alici Beyanlari"))
            _write_table(ws,["Teklif","Beyan Türü","Bağlı Kalem veya Kapsam","Beyan Değeri","Birim veya Para Birimi","Alıcı Yanıtı","Soru Kimliği","Hesap Kimliği"],
                         [row for _,_,_,row in buyer_rows])
            for rn,(oid,kind,index,row) in enumerate(buyer_rows,start=2):
                calc_id=cost_calc_id.get(oid)
                for col,key in ((2,"kind"),(3,"quantity_item_id" if kind=="reference_quantity" else "scope_id" if kind=="scope_cost" else "kind"),
                                (4,"quantity" if kind=="reference_quantity" else "amount" if kind=="scope_cost" else "milestones" if kind=="payment_terms" else "duration_years" if kind=="warranty_terms" else "date_status" if kind in {"validity_date","quote_date"} and row[3]=="Bilinmiyor (kullanıcı yanıtı)" else "value"),
                                (5,"unit" if kind=="reference_quantity" else "currency" if kind=="scope_cost" else "value"),
                                (6,"answer_text"),(7,"issue_id")):
                    bind(ws.title,ws.cell(rn,col).coordinate,calculation_ids=[calc_id] if calc_id else [],
                         value_path=f"buyer_inputs.{kind}.{index}.{key}")
                if calc_id: bind(ws.title,ws.cell(rn,8).coordinate,calculation_ids=[calc_id])
                if kind in {"reference_quantity","scope_cost"}: ws.cell(rn,4).number_format='#,##0.00;[Red](#,##0.00);0.00'

        ws=wb.create_sheet(_safe_sheet("Ticari Kosullar")); fields=["quote_date","validity","delivery","payment","incoterms","warranty","vat"]
        commercial=[f for f in facts if f.get("offer_id") and f.get("field_key") in fields]
        _write_table(ws,["Tedarikci","Teklif","Alan","Deger","Kanit Durumu","Kaynak"],[[suppliers.get(next((o.get("supplier_id") for o in active_offers if o.get("offer_id")==f.get("offer_id")),""),""),f.get("offer_id"),f.get("field_key"),f.get("value_text") or f.get("value_date") or f.get("value_number"),f.get("evidence_status"),docs.get(evidence.get(f.get("evidence_id"),{}).get("document_id"),{}).get("relative_path") if f.get("evidence_id") else ""] for f in commercial])
        for row_no,f in enumerate(commercial,start=2):
            comment=source_comment(f)
            if comment: ws.cell(row_no,4).comment=Comment(comment,"OpenAI")
            bind(ws.title,ws.cell(row_no,4).coordinate,fact_ids=[f['fact_id']])

        if requirements:
            ws=wb.create_sheet(_safe_sheet("Sartname Uygunluk")); _append_safe(ws,["Madde","Kategori","Kritik","Gereksinim"]+[suppliers.get(x.get("supplier_id"),x.get("supplier_id")) for x in active_offers])
            by_req_offer={(f.get("requirement_id"),f.get("offer_id")):f for f in facts if f.get("requirement_id")}
            for req in requirements:
                row=[req["requirement_id"],req.get("category"),"Evet" if req.get("critical") else "Hayir",req.get("text")]; rn=ws.max_row+1
                for off in active_offers:
                    f=by_req_offer.get((req["requirement_id"],off["offer_id"]),{}); row.append(f.get("value_text") or f.get("evidence_status") or "not_stated")
                _append_safe(ws,row)
                for col,off in enumerate(active_offers,start=5):
                    f=by_req_offer.get((req["requirement_id"],off["offer_id"]),{}); comment=source_comment(f) if f else None
                    if comment: ws.cell(rn,col).comment=Comment(comment,"OpenAI")
                    bind(ws.title,ws.cell(rn,col).coordinate,fact_ids=[f['fact_id']] if f else [],requirement_ids=[req['requirement_id']])

        ws=wb.create_sheet(_safe_sheet("Kalem Detayi")); _append_safe(ws,["Teklif","Alan","Kalem Kimligi","Aciklama","Miktar","Birim","Tutar Turu","Tutar","Para Birimi","Kanit Durumu","Kanit Referanslari"])
        for fact in facts:
            normalized = _json(fact.get('normalized_json'), {}) or {}
            for index,item in enumerate(normalized.get('items', [])):
                rn = ws.max_row+1
                _append_safe(ws,[fact.get('offer_id'),fact.get('field_key'),item.get('item_id'),item.get('description'),item.get('quantity'),item.get('unit'),item.get('amount_basis'),item.get('amount'),item.get('currency'),item.get('evidence_status') or fact.get('evidence_status'),json_text(item.get('evidence_refs',[]))])
                for col,key in [(3,'item_id'),(4,'description'),(5,'quantity'),(6,'unit'),(7,'amount_basis'),(8,'amount'),(9,'currency')]:
                    cell = ws.cell(rn,col)
                    bind(ws.title,cell.coordinate,fact_ids=[fact['fact_id']],value_path=f'normalized.items.{index}.{key}')
                    refs = item.get('evidence_refs', [])
                    cell.comment = Comment((source_comment(fact) or f"Veri kimliği: {fact['fact_id']} | İncelenen sürüm özeti: {snapshot_sha}")+f" | Kalem: {item.get('item_id')} | Kanıtlar: {json_text(refs)}",'OpenAI')
        if ws.max_row == 1:
            _append_safe(ws,['Kalem ayrintisi kaydi bulunmuyor.'])

        ws=wb.create_sheet(_safe_sheet('Kalem Hesabi')); _append_safe(ws,['Teklif','Alan','Hesap Turu','Hesap Durumu','Hesaplanan Alt Toplam','Para Birimi','Acik Kalemler','Beyan Genel Toplami Mi'])
        for fact in facts:
            calculation=(_json(fact.get('normalized_json'),{}) or {}).get('item_calculation')
            if not calculation:
                continue
            rn=ws.max_row+1
            _append_safe(ws,[fact.get('offer_id'),fact.get('field_key'),calculation.get('basis'),calculation.get('status'),calculation.get('amount'),calculation.get('currency'),json_text(calculation.get('unresolved_items',[])),'Hayir'])
            for col,key in [(3,'basis'),(4,'status'),(5,'amount'),(6,'currency')]:
                cell=ws.cell(rn,col)
                bind(ws.title,cell.coordinate,fact_ids=[fact['fact_id']],value_path=f'normalized.item_calculation.{key}')
                cell.comment=Comment((source_comment(fact) or f"Veri kimliği: {fact['fact_id']} | İncelenen sürüm özeti: {snapshot_sha}")+' | Hesaplanan kalem alt toplamıdır; beyan genel toplamı değildir.','OpenAI')
        if ws.max_row==1:
            _append_safe(ws,['Kalem hesabi kaydi bulunmuyor.'])

        ws=wb.create_sheet(_safe_sheet("Kapsam ve Riskler")); _write_table(ws,["Teklif","Sinif","Aciklama","Durum","Tutar","Para Birimi","Kritik","Maliyet Rolu"],[[x.get("offer_id"),x.get("classification"),x.get("description"),x.get("status"),x.get("amount"),x.get("currency"),x.get("critical"),x.get('cost_role') or 'unresolved_adjustment'] for x in scope])
        st=ws.max_row+3; ws.cell(st,1,"Elemeler"); ws.cell(st,1).font=Font(bold=True); _append_safe(ws,["Teklif","Madde","Karar","Gerekce","Durum"])
        for x in exclusions: _append_safe(ws,[x.get("offer_id"),x.get("requirement_id"),x.get("decision"),x.get("reason"),x.get("status")])
        if not exclusions: _append_safe(ws,["Kayıt bulunmuyor."])

        def issue_owner(item):
            raw=item.get('owner') or ''
            if raw.startswith('FIRMA:'):
                sid=raw.split(':',1)[1]
                return 'Firma: '+suppliers.get(sid,sid)+' ('+sid+')'
            return raw
        def issue_question(item):
            text=item.get('question') or ''
            # Generated authority questions use internal supplier IDs.
            for sid,name in suppliers.items():
                if text.startswith('Confirm the controlling offer/revision for supplier '+sid+';'):
                    text=text.replace('supplier '+sid+';', 'supplier '+name+' ('+sid+');',1)
            return text
        ws=wb.create_sheet(_safe_sheet("Bilgi Talepleri")); _write_table(ws,["Muhatap","Onem","Tur","Soru","Durum","Fact IDs"],[[issue_owner(x),x.get("severity"),x.get("issue_type"),issue_question(x),x.get("status"),x.get("fact_ids_json")] for x in issues])
        ws=wb.create_sheet(_safe_sheet("Kanitlar")); ev_rows=[]
        for f in facts:
            ev=evidence.get(f.get("evidence_id")) if f.get("evidence_id") else None
            if not ev: continue
            doc=docs.get(ev.get("document_id"),{}); ev_rows.append([f["fact_id"],f.get("offer_id"),f.get("field_key"),f.get("evidence_status"),doc.get("relative_path"),ev.get("locator_json"),ev.get("quote"),ev.get("extraction_path"),ev.get("confidence"),snapshot_sha])
        _write_table(ws,["Fact ID","Teklif","Alan","Durum","Kaynak","Locator","Alinti","Yol","Guven","Snapshot SHA"],ev_rows)
        for rn,f in enumerate([f for f in facts if f.get('evidence_id') and f.get('evidence_id') in evidence],start=2): bind(ws.title,f"A{rn}",fact_ids=[f['fact_id']])

        ws=wb.create_sheet(_safe_sheet("Karsilastirma Sozlesmesi")); _write_table(ws,["Kontrol","Boyut","Teklif","Durum","Kritik","Normalizasyon","Referans","Gozlenen"],[[x.get('check_id'),x.get('dimension'),x.get('offer_id'),x.get('status'),x.get('critical'),x.get('normalization'),x.get('reference_value'),x.get('observed_value')] for x in comparison])
        ws=wb.create_sheet(_safe_sheet("Maliyet Koprusu")); _write_table(ws,["Teklif","Para Birimi","Beyan Toplam","Kopru Toplam","Karsilastirilabilir","Residual","Tolerans","Durum","Bilesenler"],[[x.get('offer_id'),x.get('currency'),x.get('stated_total'),x.get('bridge_total'),x.get('comparable_total'),x.get('residual'),x.get('tolerance'),x.get('status'),x.get('components_json')] for x in bridges])
        ws=wb.create_sheet(_safe_sheet("Bagimsiz Hesap Kontrolu")); _write_table(ws,["Teklif","Hesap","Yontem","Durum","Tolerans","Detay"],[[x.get('offer_id'),x.get('calculation_type'),x.get('method'),x.get('status'),x.get('tolerance'),x.get('detail_json')] for x in independent])
        ws=wb.create_sheet(_safe_sheet("Iddia Kaydi")); _write_table(ws,["Claim ID","Tur","Iddia","Destek","Fact IDs","Calculation IDs","Requirement IDs","Assertions","Caveat","Snapshot SHA"],[[x.get('claim_id'),x.get('claim_type'),x.get('claim_text'),x.get('support_status'),x.get('fact_ids_json'),x.get('calculation_ids_json'),x.get('requirement_ids_json'),x.get('assertions_json'),x.get('caveat'),snapshot_sha] for x in claims])
        for rn,c in enumerate(claims,start=2): bind(ws.title,f"C{rn}",fact_ids=_json(c.get('fact_ids_json'),[]) or [],calculation_ids=_json(c.get('calculation_ids_json'),[]) or [],requirement_ids=_json(c.get('requirement_ids_json'),[]) or [],claim_ids=[c['claim_id']])
        ws=wb.create_sheet(_safe_sheet("Kapsam Izleme")); _write_table(ws,["Tur","Varlik","Aciklama","Durum","Kritik","Beklenen","Gozlenen"],[[x.get('scope_type'),x.get('entity_id'),x.get('label'),x.get('status'),x.get('critical'),x.get('expected_json'),x.get('observed_json')] for x in ledger])

        if profile=="yuksek_guvence":
            from calculation_detail import HEADERS as detail_headers, calculation_rows
            offer_names={o["offer_id"]:suppliers.get(o.get("supplier_id"),o.get("display_name") or o["offer_id"]) for o in offers}
            offer_currency={oid:(c.get("base_currency") or c.get("currency")) for oid,c in cost_map.items()}
            detail=calculation_rows(calcs,offer_names,offer_currency)
            ws=wb.create_sheet(_safe_sheet("Hesap Detayi")); _write_table(ws,detail_headers,[x["values"] for x in detail])
            for rn,x in enumerate(detail,start=2):
                if x["number_format"]: ws.cell(rn,5).number_format=x["number_format"]
                if x.get("calculation_id"): bind(ws.title,f"E{rn}",calculation_ids=[x["calculation_id"]])
            ws=wb.create_sheet(_safe_sheet("Uzman Teyidi")); _write_table(ws,["Muhatap","Onem","Tur","Soru","Durum"],[[issue_owner(x),x.get("severity"),x.get("issue_type"),issue_question(x),x.get("status")] for x in issues if x.get("severity") in {"critical","high"}])

        ws=wb.create_sheet(_safe_sheet("Kaynak Otoritesi")); _write_table(ws,["Kaynak","Kapsam","Sinif","Otorite Durumu","Oncelik","Gerekce"],[[docs.get(did,{}).get("relative_path"),a.get("authority_scope"),a.get("authority_class"),a.get("authority_status"),a.get("precedence_rank"),a.get("rationale")] for did,a in authority.items()])
        ws=wb.create_sheet(_safe_sheet("Veri Kalitesi")); _write_table(ws,["Bulgu","Kontrol","Varlik Turu","Varlik","Onem","Guven","Durum","Risk","Duzeltme"],[[x.get("finding_id"),x.get("check_id"),x.get("entity_type"),x.get("entity_id"),x.get("severity"),x.get("confidence"),x.get("status"),x.get("risk"),x.get("remediation")] for x in data_quality])
        ws=wb.create_sheet(_safe_sheet("Degisim Kaydi")); _write_table(ws,["Event ID","Asama","Olay","Detay"],[[x.get("event_id"),x.get("phase"),x.get("event_type"),x.get("payload_json")] for x in events])

        _style_workbook(wb); out_dir=run_root/"output"; out_dir.mkdir(parents=True,exist_ok=True); xlsx=out_dir/"Teklif_Degerlendirme.xlsx"; wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.save(xlsx)
        recalc=_recalc_with_excel(xlsx); xlsx_id=_artifact(con,xlsx,run_root,"xlsx",[snapshot_sha]); _insert_bindings(con,xlsx_id,snapshot_sha,bindings)

        pdf_info=None
        if profile=="yuksek_guvence":
            pdf=out_dir/"Teklif_Degerlendirme.pdf"; pdf_info=_export_pdf_with_excel(xlsx,pdf)
            if pdf_info.get("available") and pdf.exists():
                pdf_id=_artifact(con,pdf,run_root,"pdf",[snapshot_sha,sha256_file(xlsx)])
                pdf_bindings=[{**b,"component_id":b["component_id"].replace("xlsx:","pdf-derived:"),"component_type":"derived_from_xlsx","locator":{"derived_from":b["component_id"]}} for b in bindings]
                _insert_bindings(con,pdf_id,snapshot_sha,pdf_bindings)

        decision_out=out_dir/"decision_snapshot.json"
        decision_payload={"reviewed_snapshot_sha256":snapshot_sha,"decision_status":decision_status,"summary":summary,"recommendation":recommendation,"claims":claims,"preliminary_reasons":snap.get('preliminary_reasons',[])}
        decision_out.write_text(json.dumps(decision_payload,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
        decision_id=_artifact(con,decision_out,run_root,"decision_snapshot",[snapshot_sha]); _insert_bindings(con,decision_id,snapshot_sha,[{"component_id":"decision_snapshot:claims","component_type":"json","locator":{"path":"claims"},"claim_ids":summary_claim_ids}])
        if hasattr(con,"commit"): con.commit()
        return {"xlsx":str(xlsx),"xlsx_artifact_id":xlsx_id,"recalc":recalc,"pdf":pdf_info,"decision_snapshot":str(decision_out),"reviewed_snapshot":snap_row["relative_path"],"reviewed_snapshot_sha256":snapshot_sha,"bindings":len(bindings)}
    finally:
        con.close()

def main() -> None:
    ap = argparse.ArgumentParser(description="Export source-traceable procurement workbook and optional PDF")
    ap.add_argument("--run", required=True)
    args = ap.parse_args()
    print(json.dumps(export_report(Path(args.run)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
