from __future__ import annotations

from question_waits import guard_run_operation

import argparse
import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from common import clean_text, json_text, stable_id, utc_now
from db import connect, rows_as_dict, set_state
from normalize import vat_basis

CRITICAL_FIELDS = [
    "total",
    "currency",
    "vat",
    "quantity",
    "quote_date",
    "validity",
    "delivery",
    "payment",
    "incoterms",
    "warranty",
    "scope_signature",
]

CURRENCY_ALIASES = {
    "TRY": "TRY",
    "TL": "TRY",
    "TRL": "TRY",
    "USD": "USD",
    "$": "USD",
    "EUR": "EUR",
    "EURO": "EUR",
    "GBP": "GBP",
    "STERLIN": "GBP",
    "STERLING": "GBP",
    "€": "EUR",
    "₺": "TRY",
    "£": "GBP",
}

INCOTERMS = ["EXW", "FCA", "FAS", "FOB", "CFR", "CIF", "CPT", "CIP", "DAP", "DPU", "DDP"]


def _fold(text: str) -> str:
    table = str.maketrans({"ı": "i", "İ": "i", "ş": "s", "Ş": "s", "ğ": "g", "Ğ": "g", "ü": "u", "Ü": "u", "ö": "o", "Ö": "o", "ç": "c", "Ç": "c"})
    return text.translate(table).casefold()


def parse_decimal(raw: str) -> Decimal | None:
    s = clean_text(raw).replace(" ", "")
    # Formula text is never a numeric source. XLSX formulas must have a cached/
    # externally recalculated value before they can become canonical facts.
    if s.startswith("="):
        return None
    s = re.sub(r"[^0-9,.-]", "", s)
    if not s or not re.search(r"\d", s):
        return None
    neg = s.startswith("-")
    s = s.lstrip("-")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        tail = s.rsplit(",", 1)[-1]
        s = s.replace(".", "")
        s = s.replace(",", "." if len(tail) <= 2 else "")
    elif "." in s:
        parts = s.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[-1]) == 3 and len(parts[0]) >= 1):
            s = "".join(parts)
    try:
        val = Decimal(s)
        return -val if neg else val
    except InvalidOperation:
        return None


def parse_date(raw: str) -> str | None:
    raw = clean_text(raw)
    patterns = [
        (r"\b(\d{1,2})[./-](\d{1,2})[./-](20\d{2})\b", "%d-%m-%Y"),
        (r"\b(20\d{2})[./-](\d{1,2})[./-](\d{1,2})\b", "%Y-%m-%d"),
    ]
    for pat, fmt in patterns:
        m = re.search(pat, raw)
        if not m:
            continue
        token = "-".join(m.groups())
        try:
            return datetime.strptime(token, fmt).date().isoformat()
        except ValueError:
            pass
    return None


STRONG_TOTAL_LABELS = ["genel toplam", "son toplam", "toplam teklif", "teklif bedeli", "grand total", "total price", "toplam fiyat", "toplam tutar", "toplam bedel"]
TERMS_CLAUSE_WORDS = ["avans", "iptal", "odeme", "kambiyo", "fatura", "tcmb", "advance", "cancel", "payment", "invoice", "exchange rate"]


def _currency_codes(line: str) -> list[str]:
    codes: list[str] = []
    for alias, code in CURRENCY_ALIASES.items():
        hit = ("$" in line) if alias == "$" else re.search(rf"(?i)(?<![A-Z]){re.escape(alias)}(?![A-Z])", line) is not None
        if hit and code not in codes:
            codes.append(code)
    return codes


def _money_amounts(line: str, currency_on_line: bool) -> list[Decimal]:
    """Numbers that can be an amount: never a percentage; without a currency on the
    line only values written with a thousands separator or of at least 1000 qualify."""
    out: list[Decimal] = []
    for m in re.finditer(r"[-+]?\d[\d., ]*\d|\d", line):
        token = m.group(0).strip()
        before = line[:m.start()].rstrip()
        after = line[m.end():].lstrip()
        if before.endswith("%") or after.startswith("%") or after.startswith("'"):
            continue
        v = parse_decimal(token)
        if v is None:
            continue
        grouped = re.search(r"\d[.,]\d{3}(?:\D|$)", token) is not None
        if currency_on_line or grouped or abs(v) >= 1000:
            out.append(v)
    return out


def _candidate(score: float, unit: dict[str, Any], quote: str, **values: Any) -> dict[str, Any]:
    return {"score": score, "source_unit": unit, "quote": clean_text(quote)[:1000], **values}


def _text_candidates(unit: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    text = unit.get("text_content") or ""
    out: dict[str, list[dict[str, Any]]] = {k: [] for k in CRITICAL_FIELDS}
    for raw_line in text.splitlines():
        line = clean_text(raw_line)
        if not line:
            continue
        f = _fold(line)
        currency_codes = _currency_codes(line)
        amounts = _money_amounts(line, bool(currency_codes))
        terms_clause = any(k in f for k in TERMS_CLAUSE_WORDS)
        # Total. A generic "toplam" inside a payment/cancellation clause ("%20 of the total")
        # is not a total: percentages are never amounts, and only an explicit total label
        # with a money-like amount may become verified without review.
        if any(k in f for k in ["genel toplam", "toplam teklif", "teklif bedeli", "grand total", "total price", "toplam"]):
            if amounts:
                v = amounts[-1]
                strong = any(k in f for k in STRONG_TOTAL_LABELS) and not terms_clause
                out["total"].append(_candidate(0.82 if strong else 0.6, unit, line, value_number=float(v), normalized_decimal=str(v)))
        # Currency: verified only on an explicit total line with an amount. Any other price
        # mention (a day rate, an option, a clause) stays below review threshold, so the
        # semantic task resolves the offer currency from the actual total.
        explicit_total = any(k in f for k in STRONG_TOTAL_LABELS)
        for code in currency_codes:
            if amounts and not terms_clause:
                score = 0.9 if explicit_total else 0.7
            else:
                score = 0.6
            out["currency"].append(_candidate(score, unit, line, value_text=code))
        # VAT
        if "kdv" in f or "vat" in f:
            value = vat_basis(line)
            if value is None:
                m = re.search(r"%\s*(\d+(?:[.,]\d+)?)", line)
                value = f"rate:{m.group(1)}" if m else line
            out["vat"].append(_candidate(0.84, unit, line, value_text=value))
        # Quote date
        if any(k in f for k in ["teklif tarihi", "tarih", "date", "quotation date", "offer date"]):
            d = parse_date(line)
            if d:
                out["quote_date"].append(_candidate(0.8, unit, line, value_date=d, value_text=d))
        # Validity
        if any(k in f for k in ["gecerli", "gecerlilik", "validity", "valid for"]):
            m = re.search(r"(?i)(\d+)\s*(gun|day|days|hafta|week|weeks|ay|month|months)", _fold(line))
            if m:
                out["validity"].append(_candidate(0.82, unit, line, value_text=f"{m.group(1)} {m.group(2)}"))
            else:
                d = parse_date(line)
                out["validity"].append(_candidate(0.8 if d else 0.6, unit, line, value_text=d or line, value_date=d))
        # Delivery
        if any(k in f for k in ["teslim", "termin", "delivery", "lead time"]):
            out["delivery"].append(_candidate(0.78, unit, line, value_text=line))
        # Payment
        if any(k in f for k in ["odeme", "vade", "payment", "advance", "pesin", "cek", "letter of credit", "lc"]):
            out["payment"].append(_candidate(0.78, unit, line, value_text=line))
        # Incoterms
        for term in INCOTERMS:
            if re.search(rf"(?i)\b{term}\b", line):
                out["incoterms"].append(_candidate(0.9, unit, line, value_text=term))
                break
        # Warranty
        if "garanti" in f or "warranty" in f:
            out["warranty"].append(_candidate(0.8, unit, line, value_text=line))
        # Offer-level quantity only if explicitly labelled
        if re.search(r"(?i)\b(miktar|quantity|qty)\b", f):
            nums = re.findall(r"\b\d+(?:[.,]\d+)?\b", line)
            if nums:
                v = parse_decimal(nums[-1])
                if v is not None:
                    unit_match = re.search(
                        r"(?i)\b(kg|g|ton|t|adet|ad|pcs?|piece|pieces|m2|m²|m3|m³|m|cm|mm|lt|l|litre|liter|set|takim|paket|package)\b",
                        line,
                    )
                    fact_unit = unit_match.group(1).lower() if unit_match else None
                    out["quantity"].append(
                        _candidate(0.55, unit, line, value_number=float(v), value_text=str(v), fact_unit=fact_unit)
                    )
        # Product / scope signature. Use only explicit labels; do not infer scope
        # from a supplier name or filename.
        if re.search(r"(?i)\b(urun|ürün|malzeme|product|item|scope|kapsam|konu)\b", f):
            parts = re.split(r"[:;\-]", line, maxsplit=1)
            if len(parts) == 2 and clean_text(parts[1]):
                out["scope_signature"].append(_candidate(0.72, unit, line, value_text=clean_text(parts[1])))
    return out


def _sheet_candidates(unit: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = {k: [] for k in CRITICAL_FIELDS}
    try:
        structured = json.loads(unit.get("structured_json") or "{}")
    except Exception:
        return out
    cells = structured.get("cells", [])
    by_cell = {c.get("cell"): c for c in cells if c.get("cell")}
    try:
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
    except Exception:
        return out
    labels = {
        "total": ["genel toplam", "toplam teklif", "teklif bedeli", "grand total", "total price"],
        "currency": ["para birimi", "doviz", "currency"],
        "vat": ["kdv", "vat"],
        "quote_date": ["teklif tarihi", "quotation date", "offer date"],
        "validity": ["gecerlilik", "validity"],
        "delivery": ["teslim", "termin", "delivery", "lead time"],
        "payment": ["odeme", "vade", "payment"],
        "incoterms": ["incoterm", "teslim sekli"],
        "warranty": ["garanti", "warranty"],
        "quantity": ["miktar", "quantity", "qty"],
        "scope_signature": ["urun", "malzeme", "product", "item", "scope", "kapsam", "konu"],
    }
    label_qualifiers = {"suresi", "tarihi", "sekli", "kosulu", "durumu", "adi",
                        "tanimi", "aciklamasi", "bilgisi", "bedeli", "tutari",
                        "birimi", "no", "numarasi", "gun", "days"}
    for c in cells:
        # Label cells can be formulas only when a trusted cached value exists.
        value = c.get("calculated_value") if c.get("calculated_value") is not None else c.get("formula_or_value")
        if c.get("is_formula") and not c.get("calculated_available"):
            continue
        if not isinstance(value, str):
            continue
        # A worksheet line description can mention "teslim" or "malzeme" and
        # have adjacent category/UOM cells. It is not a field label/value pair.
        fv = ' '.join(_fold(value).split()).strip(' :;-()')
        if len(value) > 48 or '\n' in value or '\r' in value:
            continue
        def is_label(term: str) -> bool:
            if fv == term:
                return True
            if not fv.startswith(term + ' '):
                return False
            remainder = fv[len(term):].strip(' :;-()')
            words = remainder.split()
            return 0 < len(words) <= 2 and all(word.strip(' :;-()') in label_qualifiers for word in words)
        field = next((k for k, terms in labels.items() if any(is_label(t) for t in terms)), None)
        if not field:
            continue
        coord = c.get("cell")
        col_letters, row = coordinate_from_string(coord)
        col = column_index_from_string(col_letters)
        # Prefer same-row label/value layouts. Falling down to the next row
        # when a horizontal value exists can mistake the next field label
        # (for example KDV, Odeme) for the current field value.
        neighbors = []
        for dc, dr in [(1, 0), (2, 0)]:
            nc = f"{get_column_letter(col + dc)}{row + dr}"
            if nc in by_cell:
                neighbors.append(by_cell[nc])
        for n in neighbors:
            if n.get("is_formula") and not n.get("calculated_available"):
                # Never fall back to the formula string itself. That can turn
                # '=100000+5000' into 1000005000 after punctuation stripping.
                continue
            nval = n.get("calculated_value") if n.get("calculated_value") is not None else n.get("formula_or_value")
            if nval is None:
                continue
            locator = {"sheet": structured.get("title"), "range": f"{coord}:{n.get('cell')}"}
            q = f"{coord}={value}; {n.get('cell')}={nval}"
            u2 = dict(unit)
            u2["locator_json"] = json_text(locator)
            if field == "total":
                dec = parse_decimal(str(nval))
                if dec is not None:
                    out[field].append(_candidate(0.96, u2, q, value_number=float(dec), normalized_decimal=str(dec)))
            elif field == "quote_date":
                if hasattr(nval, "date"):
                    d = nval.date().isoformat() if hasattr(nval, "date") else None
                else:
                    d = parse_date(str(nval))
                if d:
                    out[field].append(_candidate(0.95, u2, q, value_date=d, value_text=d))
            elif field == "currency":
                code = CURRENCY_ALIASES.get(str(nval).upper(), str(nval).upper())
                out[field].append(_candidate(0.95, u2, q, value_text=code))
            elif field == "quantity":
                dec = parse_decimal(str(nval))
                if dec is not None:
                    # Look one additional cell to the right for a separately
                    # labelled/entered UOM (common: Quantity | Unit).
                    fact_unit = None
                    uom_cell = by_cell.get(f"{get_column_letter(col + 2)}{row}")
                    if uom_cell and not (uom_cell.get("is_formula") and not uom_cell.get("calculated_available")):
                        uom_val = uom_cell.get("calculated_value") if uom_cell.get("calculated_value") is not None else uom_cell.get("formula_or_value")
                        if uom_val not in (None, ""):
                            fact_unit = clean_text(str(uom_val)).lower()
                    if fact_unit is None:
                        unit_match = re.search(
                            r"(?i)\b(kg|g|ton|t|adet|ad|pcs?|piece|pieces|m2|m²|m3|m³|m|cm|mm|lt|l|litre|liter|set|takim|paket|package)\b",
                            q,
                        )
                        fact_unit = unit_match.group(1).lower() if unit_match else None
                    if fact_unit is None:
                        # Offer forms often keep the unit on a separate
                        # ``Birim | adet`` row. Accept it only when exactly one
                        # explicit unit label/value pair exists on the sheet.
                        uom_values = []
                        for uc in cells:
                            uv = uc.get("calculated_value") if uc.get("calculated_value") is not None else uc.get("formula_or_value")
                            if not isinstance(uv, str) or _fold(uv).strip() not in {"birim", "unit", "uom", "olcu birimi"}:
                                continue
                            uletters, urow = coordinate_from_string(uc.get("cell"))
                            ucol = column_index_from_string(uletters)
                            rc = by_cell.get(f"{get_column_letter(ucol + 1)}{urow}")
                            if rc and not (rc.get("is_formula") and not rc.get("calculated_available")):
                                rv = rc.get("calculated_value") if rc.get("calculated_value") is not None else rc.get("formula_or_value")
                                if rv not in (None, ""):
                                    uom_values.append(clean_text(str(rv)).lower())
                        if len(set(uom_values)) == 1:
                            fact_unit = uom_values[0]
                    out[field].append(_candidate(0.94, u2, q, value_number=float(dec), value_text=str(dec), fact_unit=fact_unit))
            else:
                out[field].append(_candidate(0.94, u2, q, value_text=str(nval)))
    return out



def _tabular_candidates(unit: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    """Conservative CSV/TSV label/value pairs; column headers stay semantic."""
    out: dict[str, list[dict[str, Any]]] = {k: [] for k in CRITICAL_FIELDS}
    try:
        structured = json.loads(unit.get("structured_json") or "{}")
    except Exception:
        return out
    rows = structured.get("rows")
    if not isinstance(rows, list):
        return out
    labels = {
        "total": ["genel toplam", "toplam teklif", "teklif bedeli", "grand total", "total price"],
        "currency": ["para birimi", "doviz", "currency"],
        "vat": ["kdv", "vat"],
        "quote_date": ["teklif tarihi", "quotation date", "offer date"],
        "validity": ["gecerlilik", "validity"],
        "delivery": ["teslim", "termin", "delivery", "lead time"],
        "payment": ["odeme", "vade", "payment"],
        "incoterms": ["incoterm", "teslim sekli"],
        "warranty": ["garanti", "warranty"],
        "quantity": ["miktar", "quantity", "qty"],
        "scope_signature": ["urun", "malzeme", "product", "item", "scope", "kapsam", "konu"],
    }
    qualifiers = {"suresi", "tarihi", "sekli", "kosulu", "durumu", "adi", "tanimi",
                  "aciklamasi", "bilgisi", "bedeli", "tutari", "birimi", "no",
                  "numarasi", "gun", "days", "ay", "sonu", "baslangici"}

    def field_label(value):
        if not isinstance(value, str) or len(value) > 48 or "\n" in value or "\r" in value:
            return None
        folded = " ".join(_fold(value).replace("_", " ").split()).strip(" :;-()")
        for key, terms in labels.items():
            for term in terms:
                if folded == term:
                    return key
                if folded.startswith(term + " "):
                    tail = folded[len(term):].strip(" :;-()").split()
                    if 0 < len(tail) <= 2 and all(word in qualifiers for word in tail):
                        return key
        return None

    for row_index, row in enumerate(rows, start=1):
        if not isinstance(row, list) or len(row) < 2:
            continue
        # In a horizontal table, ``odeme,garanti_ay,nakliye`` are column
        # names, not payment=garanti_ay and warranty=nakliye. Multi-label
        # rows are deliberately left to the bounded semantic task, including
        # ambiguous mixed layouts; never invent deterministic verified values.
        if sum(field_label(value) is not None for value in row) >= 2:
            continue
        for col_index, raw_label in enumerate(row[:-1], start=1):
            if raw_label is None:
                continue
            field = field_label(raw_label)
            if not field:
                continue
            # Prefer the next non-empty cell in the same row.
            nval = None
            ncol = None
            for j in range(col_index, len(row)):
                if row[j] not in (None, ""):
                    nval = row[j]
                    ncol = j + 1
                    break
            if nval is None:
                continue
            locator = {"row": row_index, "columns": [col_index, ncol]}
            u2 = dict(unit)
            u2["locator_json"] = json_text(locator)
            q = f"{raw_label};{nval}"
            if field == "total":
                dec = parse_decimal(str(nval))
                if dec is not None:
                    out[field].append(_candidate(0.96, u2, q, value_number=float(dec), normalized_decimal=str(dec)))
            elif field == "quote_date":
                d = parse_date(str(nval))
                if d:
                    out[field].append(_candidate(0.95, u2, q, value_date=d, value_text=d))
            elif field == "currency":
                raw = str(nval).strip().upper()
                out[field].append(_candidate(0.95, u2, q, value_text=CURRENCY_ALIASES.get(raw, raw)))
            elif field == "quantity":
                dec = parse_decimal(str(nval))
                if dec is not None:
                    fact_unit = None
                    if ncol is not None and ncol < len(row) and row[ncol] not in (None, ""):
                        fact_unit = clean_text(str(row[ncol])).lower()
                    out[field].append(_candidate(0.94, u2, q, value_number=float(dec), value_text=str(dec), fact_unit=fact_unit))
            elif field == "vat":
                value = vat_basis(str(nval)) or str(nval)
                out[field].append(_candidate(0.94, u2, q, value_text=value))
            elif field == "incoterms":
                raw = str(nval).strip().upper()
                value = next((term for term in INCOTERMS if term == raw or re.search(rf"\\b{term}\\b", raw)), raw)
                out[field].append(_candidate(0.95, u2, q, value_text=value))
            else:
                out[field].append(_candidate(0.94, u2, q, value_text=str(nval)))
    return out

def _merge_candidates(target: dict[str, list[dict[str, Any]]], source: dict[str, list[dict[str, Any]]]) -> None:
    for k, vals in source.items():
        target[k].extend(vals)


def _existing_priority(path: str | None) -> int:
    return {
        "adjudicated": 100,
        "semantic_llm": 80,
        "vision": 75,
        "duckdb_pdf_secondary": 70,
        "python_primary": 60,
        "deterministic": 50,
    }.get(path or "", 0)


def _upsert_fact(con, offer_id: str, supplier_id: str, field: str, candidate: dict[str, Any] | None) -> None:
    fact_id = f"{offer_id}/field/{field}"
    critical = field != "scope_signature"
    existing = con.execute("SELECT extraction_path,evidence_status FROM facts WHERE fact_id=?", [fact_id]).fetchone()
    if existing and _existing_priority(existing[0]) > _existing_priority("deterministic"):
        return
    if candidate is None:
        con.execute(
            """
            INSERT OR REPLACE INTO facts(fact_id,offer_id,supplier_id,field_key,evidence_status,critical,extraction_path,normalized_json,updated_at)
            VALUES (?,?,?,?,?,?,?,?,?)
            """,
            [fact_id, offer_id, supplier_id, field, "missing", critical, "deterministic", "{}", utc_now()],
        )
        return
    source_unit = candidate["source_unit"]
    document_id = source_unit["document_id"]
    locator = json.loads(source_unit.get("locator_json") or "{}")
    quote = candidate.get("quote", "")
    ev_id = stable_id("ev", document_id, json_text(locator), quote, field, length=22)
    confidence = float(candidate.get("score", 0.5))
    status = "verified" if confidence >= 0.8 else "requires_confirmation"
    source_sha = con.execute("SELECT sha256 FROM documents WHERE document_id=?", [document_id]).fetchone()
    con.execute(
        """
        INSERT OR REPLACE INTO evidence(evidence_id,document_id,unit_id,locator_json,quote,evidence_status,extraction_path,confidence,source_sha256)
        VALUES (?,?,?,?,?,?,?,?,?)
        """,
        [ev_id, document_id, source_unit.get("unit_id"), json_text(locator), quote, status, "deterministic", confidence, source_sha[0] if source_sha else None],
    )
    normalized = {}
    if candidate.get("normalized_decimal"):
        normalized["decimal"] = candidate["normalized_decimal"]
    con.execute(
        """
        INSERT OR REPLACE INTO facts(
          fact_id,offer_id,supplier_id,field_key,value_text,value_number,value_date,unit,currency,
            evidence_status,critical,evidence_id,extraction_path,normalized_json,updated_at
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        [
            fact_id, offer_id, supplier_id, field, candidate.get("value_text"), candidate.get("value_number"), candidate.get("value_date"),
            candidate.get("fact_unit"), candidate.get("currency"), status, critical, ev_id, "deterministic", json_text(normalized), utc_now(),
        ],
    )


def _has_material_conflict(field: str, candidates: list[dict[str, Any]]) -> bool:
    """Return True when multiple high-confidence candidates disagree materially.

    This is intentionally conservative: a visible and hidden grand total that
    disagree should not be silently resolved by whichever cell scored first.
    """
    best_score = max((float(c.get("score", 0.0)) for c in candidates), default=0.0)
    # Compare only peer-quality candidates. Lower-confidence text heuristics
    # must not conflict with a stronger structured cell extraction merely
    # because a cell address contains digits or a label line contains text.
    cutoff = max(0.8, best_score - 0.05)
    high = [c for c in candidates if float(c.get("score", 0.0)) >= cutoff]
    if len(high) < 2:
        return False
    if field in {"total", "quantity"}:
        vals = [Decimal(str(c["value_number"])) for c in high if c.get("value_number") is not None]
        if len(vals) < 2:
            return False
        lo, hi = min(vals), max(vals)
        tol = max(Decimal("0.01"), abs(hi) * Decimal("0.001"))
        return hi - lo > tol
    if field == "vat":
        # Normalize only exact, known aliases. Unknown or negated wording stays
        # source-distinct and cannot be promoted to a VAT basis.
        vals = {
            vat_basis(str(c.get("value_text"))) or _fold(clean_text(str(c.get("value_text"))))
            for c in high if c.get("value_text") is not None
        }
    else:
        vals = {
            clean_text(str(c.get("value_text") or c.get("value_date") or "")).casefold()
            for c in high
            if c.get("value_text") is not None or c.get("value_date") is not None
        }
    return len(vals) > 1


def _mark_conflicting(con, offer_id: str, field: str, candidates: list[dict[str, Any]]) -> None:
    fact_id = f"{offer_id}/field/{field}"
    previous = con.execute("SELECT normalized_json FROM facts WHERE fact_id=?", [fact_id]).fetchone()
    payload = json.loads(previous[0] or "{}") if previous else {}
    payload["candidates"] = [
            {
                "score": c.get("score"),
                "value_number": c.get("value_number"),
                "value_text": c.get("value_text"),
                "value_date": c.get("value_date"),
                "quote": c.get("quote"),
                "document_id": c.get("source_unit", {}).get("document_id"),
                "locator": json.loads(c.get("source_unit", {}).get("locator_json") or "{}"),
            }
            for c in candidates[:10]
        ]
    con.execute(
        "UPDATE facts SET evidence_status='conflicting', normalized_json=?, updated_at=? WHERE fact_id=?",
        [json_text(payload), utc_now(), fact_id],
    )


@guard_run_operation
def extract(run_root: Path) -> dict[str, Any]:
    con = connect(run_root)
    offer_count = 0
    found = 0
    missing = 0
    try:
        offers = rows_as_dict(con, "SELECT offer_id,supplier_id FROM offers WHERE status <> 'superseded' ORDER BY offer_id")
        for offer in offers:
            offer_count += 1
            docs = rows_as_dict(
                con,
                """
                SELECT d.document_id,d.format_family FROM documents d
                WHERE d.offer_id=? AND d.active=true AND d.parse_status <> 'unreadable'
                ORDER BY d.relative_path
                """,
                [offer["offer_id"]],
            )
            all_candidates: dict[str, list[dict[str, Any]]] = {k: [] for k in CRITICAL_FIELDS}
            for doc in docs:
                units = rows_as_dict(
                    con,
                    "SELECT unit_id,document_id,unit_type,unit_index,locator_json,text_content,structured_json FROM document_units WHERE document_id=? ORDER BY unit_index",
                    [doc["document_id"]],
                )
                for unit in units:
                    # Structured formats must stay on their structured parser
                    # path. Re-parsing synthetic XLSX cell text such as
                    # ``A4=Genel Toplam`` can turn the cell address itself into
                    # a fake numeric candidate when a formula has no cached
                    # value.
                    if unit["unit_type"] == "xlsx_sheet":
                        _merge_candidates(all_candidates, _sheet_candidates(unit))
                    elif unit["unit_type"] == "csv_chunk":
                        _merge_candidates(all_candidates, _tabular_candidates(unit))
                    else:
                        _merge_candidates(all_candidates, _text_candidates(unit))
            for field in CRITICAL_FIELDS:
                candidates = sorted(all_candidates[field], key=lambda x: x.get("score", 0.0), reverse=True)
                best = candidates[0] if candidates else None
                _upsert_fact(con, offer["offer_id"], offer["supplier_id"], field, best)
                if best and _has_material_conflict(field, candidates):
                    _mark_conflicting(con, offer["offer_id"], field, candidates)
                if best:
                    found += 1
                else:
                    missing += 1
        if hasattr(con, "commit"):
            con.commit()
    finally:
        con.close()
    detail = {"offers": offer_count, "fields_found": found, "fields_missing": missing}
    set_state(run_root, "EXTRACT", "PASS" if offer_count else "PARTIAL", detail)
    return detail


def main() -> None:
    ap = argparse.ArgumentParser(description="Deterministic critical-field extraction")
    ap.add_argument("--run", required=True)
    args = ap.parse_args()
    print(json.dumps(extract(Path(args.run)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    from question_waits import run_guarded_main
    raise SystemExit(run_guarded_main(main))
