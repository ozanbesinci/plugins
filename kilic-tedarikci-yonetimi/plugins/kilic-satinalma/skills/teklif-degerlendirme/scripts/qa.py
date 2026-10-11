from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from common import json_text, sha256_bytes, sha256_file, stable_id, utc_now
from db import connect, get_meta, rows_as_dict

ERROR_TOKENS = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def _num_tokens(text: str) -> list[Decimal]:
    out=[]
    number_pattern = r"[-+]?(?:\d{1,3}(?:[ .]\d{3})+|\d+)(?:[,.]\d+)?"
    for tok in re.findall(number_pattern, text or ""):
        s=tok.replace(" ","")
        if "," in s and "." in s:
            if s.rfind(",")>s.rfind("."): s=s.replace(".","").replace(",",".")
            else: s=s.replace(",","")
        elif "," in s:
            tail=s.rsplit(",",1)[-1]; s=s.replace(",","." if len(tail)<=2 else "")
        elif s.count(".")==1 and len(s.rsplit(".",1)[-1])==3:
            s=s.replace(".","")
        try: out.append(Decimal(s))
        except (InvalidOperation,ValueError): pass
    return out


def _pdf_content_checks(pdf_path: Path, run_root: Path) -> list[str]:
    """Verify critical snapshot values survive the XLSX->PDF conversion."""
    from pypdf import PdfReader
    from reviewed_snapshot import load_latest
    errors=[]
    try:
        reader=PdfReader(str(pdf_path))
        text="\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception as exc:
        return [f"PDF critical-content extraction failed: {exc}"]
    if not text.strip():
        return ["PDF has no extractable text; critical values cannot be mechanically verified"]
    folded=" ".join(text.casefold().split())
    snap,_=load_latest(run_root)
    suppliers={r.get('supplier_id'):str(r.get('display_name') or '') for r in snap.get('suppliers',[])}
    for off in [x for x in snap.get('offers',[]) if x.get('status')=='active']:
        name=suppliers.get(off.get('supplier_id'),"")
        if name and " ".join(name.casefold().split()) not in folded:
            errors.append(f"PDF does not contain active supplier label: {name}")
    from pdf_bound_values import check_bound_values
    errors.extend(check_bound_values(text, snap))
    from display_values import label
    status=str(label((snap.get('meta') or {}).get('decision_status') or '')).casefold()
    if status and status not in folded:
        errors.append(f"PDF does not preserve decision status label/value: {status}")
    return errors


def _finding(con, finding_type: str, severity: str, message: str, *, artifact_id: str | None = None, locator: dict[str, Any] | None = None, status: str = "open") -> str:
    fid = stable_id("qa", finding_type, artifact_id or "", json_text(locator or {}), message, length=22)
    con.execute(
        "INSERT OR REPLACE INTO qa_findings(finding_id,artifact_id,finding_type,locator_json,severity,status,message,finding_hash,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
        [fid, artifact_id, finding_type, json_text(locator or {}), severity, status, message, stable_id("h", message, length=20), utc_now()],
    )
    return fid


def _latest_artifact(con, artifact_type: str) -> dict[str, Any] | None:
    rows = rows_as_dict(con, "SELECT * FROM artifacts WHERE artifact_type=? AND status='ready' ORDER BY created_at DESC", [artifact_type])
    return rows[0] if rows else None


def _mechanical_db_checks(con, recommendation_present: bool) -> list[str]:
    errors: list[str] = []
    orphan_fact = con.execute("SELECT count(*) FROM facts f LEFT JOIN evidence e ON f.evidence_id=e.evidence_id WHERE f.evidence_id IS NOT NULL AND e.evidence_id IS NULL").fetchone()[0]
    if orphan_fact:
        errors.append(f"{orphan_fact} facts reference missing evidence")
    orphan_ev = con.execute("SELECT count(*) FROM evidence e LEFT JOIN documents d ON e.document_id=d.document_id WHERE d.document_id IS NULL").fetchone()[0]
    if orphan_ev:
        errors.append(f"{orphan_ev} evidence rows reference missing documents")
    open_critical = con.execute("SELECT count(*) FROM open_issues WHERE status='open' AND severity='critical'").fetchone()[0]
    if open_critical and recommendation_present:
        errors.append("Final recommendation exists while critical issues remain open")
    unresolved_diff = con.execute("SELECT count(*) FROM differences WHERE status='open'").fetchone()[0]
    if unresolved_diff and recommendation_present:
        errors.append("Final recommendation exists while adjudication differences remain open")
    invalid_status = con.execute("SELECT count(*) FROM facts WHERE evidence_status NOT IN ('verified','missing','excluded','unreadable','conflicting','assumption','requires_confirmation')").fetchone()[0]
    if invalid_status:
        errors.append(f"{invalid_status} facts have invalid evidence status")
    return errors


def _workbook_checks(path: Path, profile: str, formula_contract=None) -> tuple[list[str], list[str]]:
    from openpyxl import load_workbook

    errors: list[str] = []
    warnings: list[str] = []
    if not path.exists():
        return [f"Workbook missing: {path}"], warnings
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        return [f"Workbook cannot be opened: {exc}"], warnings
    required = {"Ozet", "Envanter", "Teklif Karsilastirma", "Ticari Kosullar", "Kapsam ve Riskler", "Bilgi Talepleri", "Kanitlar", "Kaynak Otoritesi", "Veri Kalitesi", "Degisim Kaydi", "Karsilastirma Sozlesmesi", "Maliyet Koprusu", "Bagimsiz Hesap Kontrolu", "Iddia Kaydi", "Kapsam Izleme"}
    if profile == "yuksek_guvence":
        required |= {"Hesap Detayi", "Uzman Teyidi"}
    missing = required - set(wb.sheetnames)
    if missing:
        errors.append(f"Missing required sheets: {sorted(missing)}")
    if formula_contract is not None:
        from scenario_formulas import validate_formulas
        checked = validate_formulas(path, formula_contract)
        errors.extend(checked["errors"])
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and any(tok in v for tok in ERROR_TOKENS):
                    errors.append(f"Formula/error token in {ws.title}!{cell.coordinate}: {v}")
                if cell.data_type == "f" and (formula_contract is None or
                        formula_contract["formulas"].get(ws.title + "!" + cell.coordinate) != v):
                    errors.append(f"Unverified formula remains in final workbook {ws.title}!{cell.coordinate}: {v}")
        if ws.max_column > 60:
            warnings.append(f"Very wide sheet: {ws.title} has {ws.max_column} columns")
    wb.close()
    return errors, warnings


def _render_pdf(pdf_path: Path, target: Path) -> list[dict[str, Any]]:
    import pypdfium2 as pdfium

    target.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(pdf_path))
    out = []
    try:
        for idx in range(len(pdf)):
            page = pdf[idx]
            try:
                bitmap = page.render(scale=1.35)
                try:
                    image = bitmap.to_pil()
                    try:
                        path = target / f"page-{idx + 1:04d}.png"
                        image.save(path)
                    finally:
                        image.close()
                finally:
                    bitmap.close()
            finally:
                page.close()
            data = path.read_bytes()
            out.append({"page": idx + 1, "path": str(path), "sha256": sha256_bytes(data)})
    finally:
        pdf.close()
    return out


def _visual_tasks(con, run_root: Path, artifact: dict[str, Any], renders: list[dict[str, Any]], full_first: bool, chunk_size: int = 3) -> int:
    previous = rows_as_dict(con, "SELECT artifact_id FROM artifacts WHERE artifact_type='pdf' AND status='ready' AND artifact_id<>? ORDER BY created_at DESC", [artifact["artifact_id"]])
    previous_id = previous[0]["artifact_id"] if previous else None
    failed_keys = set()
    for finding in rows_as_dict(con, "SELECT locator_json FROM qa_findings WHERE artifact_id=? AND finding_type='visual_layout' AND status='open'", [artifact["artifact_id"]]):
        failed_keys.add(json.loads(finding["locator_json"] or "{}").get("render_key"))
    for task in rows_as_dict(con, "SELECT task_id,payload_json FROM task_queue WHERE task_kind='artifact_visual_review' AND status='pending'"):
        payload = json.loads(task.get('payload_json') or '{}')
        if payload.get('artifact_id') and payload['artifact_id'] != artifact['artifact_id']:
            con.execute("UPDATE task_queue SET status='superseded' WHERE task_id=?", [task['task_id']])

    required = []
    for r in renders:
        key = f"page:{r['page']}"
        current = con.execute("SELECT qa_status FROM render_hashes WHERE artifact_id=? AND render_key=?", [artifact["artifact_id"], key]).fetchone()
        if key in failed_keys or (current and current[0] == "fail"):
            continue
        old = None
        if previous_id:
            old = con.execute("SELECT render_sha256,qa_status FROM render_hashes WHERE artifact_id=? AND render_key=?", [previous_id, f"page:{r['page']}"]).fetchone()
        if old and old[0] == r["sha256"] and old[1] == "pass":
            con.execute(
                "INSERT OR REPLACE INTO render_hashes(artifact_id,render_key,render_sha256,qa_status,qa_note) VALUES (?,?,?,?,?)",
                [artifact["artifact_id"], f"page:{r['page']}", r["sha256"], "pass", "Inherited from identical prior render hash."],
            )
            continue
        if full_first or previous_id or not old:
            required.append(r)
    for r in renders:
        if r not in required:
            continue
        existing = con.execute("SELECT qa_status FROM render_hashes WHERE artifact_id=? AND render_key=?", [artifact["artifact_id"], f"page:{r['page']}"]).fetchone()
        if f"page:{r['page']}" in failed_keys or (existing and existing[0] in {"pass", "fail"}):
            continue
        con.execute(
            "INSERT OR REPLACE INTO render_hashes(artifact_id,render_key,render_sha256,qa_status,qa_note) VALUES (?,?,?,?,?)",
            [artifact["artifact_id"], f"page:{r['page']}", r["sha256"], "pending", "Visual review pending."],
        )
    pending = [r for r in required if f"page:{r['page']}" not in failed_keys and (con.execute("SELECT qa_status FROM render_hashes WHERE artifact_id=? AND render_key=?", [artifact["artifact_id"], f"page:{r['page']}"]).fetchone() or [None])[0] == "pending"]
    created = 0
    for i in range(0, len(pending), chunk_size):
        chunk = pending[i : i + chunk_size]
        payload = {
            "artifact_id": artifact["artifact_id"],
            "artifact_path": str((run_root / artifact["relative_path"]).resolve()),
            "pages": [{"page": r["page"], "render_key": f"page:{r['page']}", "image_path": r["path"], "sha256": r["sha256"]} for r in chunk],
            "goal": "Check every supplied page image for clipping, overflow, unreadable tables/text, broken glyphs, blank/duplicate pages, and materially misleading layout. Return pass/fail per page with concise observations.",
        }
        task_id = stable_id("task", "QA", artifact["artifact_id"], json_text([r["sha256"] for r in chunk]), length=24)
        if con.execute("SELECT 1 FROM task_queue WHERE task_id=?", [task_id]).fetchone():
            continue
        con.execute(
            """
            INSERT OR IGNORE INTO task_queue(task_id,phase,role,task_kind,payload_json,priority,critical,status,attempts,created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?)
            """,
            [task_id, "QA", "visual_exception", "artifact_visual_review", json_text(payload), 80, True, "pending", 0, utc_now()],
        )
        created += 1
    return created


def _supersede_visual_findings(con, run_root: Path, artifact: dict[str, Any], renders: list[dict[str, Any]]) -> int:
    """Retire only older PDF layout findings after current complete acceptance."""
    current = _latest_artifact(con, "pdf")
    if not current or current["artifact_id"] != artifact["artifact_id"] or not renders:
        return 0
    path = run_root / artifact["relative_path"]
    if not path.is_file() or sha256_file(path) != artifact["sha256"]:
        return 0
    expected = {f"page:{r['page']}": r["sha256"] for r in renders}
    if len(expected) != len(renders):
        return 0
    reviewed = rows_as_dict(con, "SELECT render_key,render_sha256,qa_status FROM render_hashes WHERE artifact_id=?", [artifact["artifact_id"]])
    if len(reviewed) != len(expected) or any(r["qa_status"] != "pass" or expected.get(r["render_key"]) != r["render_sha256"] for r in reviewed):
        return 0
    if con.execute("SELECT count(*) FROM qa_findings WHERE artifact_id=? AND status='open' AND severity='critical'", [artifact["artifact_id"]]).fetchone()[0]:
        return 0
    candidates = rows_as_dict(con, """
        SELECT q.finding_id,q.artifact_id FROM qa_findings q JOIN artifacts a ON q.artifact_id=a.artifact_id
        WHERE q.finding_type='visual_layout' AND q.status='open' AND a.artifact_type='pdf'
          AND a.status='ready' AND a.artifact_id<>? AND a.created_at<?
    """, [artifact["artifact_id"], artifact["created_at"]])
    con.execute("""CREATE TABLE IF NOT EXISTS qa_finding_supersessions(
        finding_id VARCHAR PRIMARY KEY, previous_artifact_id VARCHAR NOT NULL,
        replacement_artifact_id VARCHAR NOT NULL, replacement_sha256 VARCHAR NOT NULL,
        render_hashes_json VARCHAR NOT NULL, superseded_at VARCHAR NOT NULL)""")
    for row in candidates:
        con.execute("INSERT OR IGNORE INTO qa_finding_supersessions VALUES (?,?,?,?,?,?)",
                    [row["finding_id"],row["artifact_id"],artifact["artifact_id"],artifact["sha256"],json_text(expected),utc_now()])
        con.execute("UPDATE qa_findings SET status='superseded' WHERE finding_id=? AND status='open'", [row["finding_id"]])
    return len(candidates)


def run_qa(run_root: Path) -> dict[str, Any]:
    profile = str(get_meta(run_root, "profile", "standart"))
    from cross_output_consistency import check as cross_output_check
    from provenance_survival import check as provenance_check
    cross_output = cross_output_check(run_root)
    provenance = provenance_check(run_root)
    from official_delivery import inspect_official
    official_delivery = inspect_official(run_root)
    recommendation_raw = get_meta(run_root, "decision_recommendation", "null")
    try:
        recommendation_present = json.loads(recommendation_raw) is not None
    except Exception:
        recommendation_present = bool(recommendation_raw and recommendation_raw != "null")
    con = connect(run_root)
    try:
        # Clear only regenerated mechanical findings, preserving semantic/visual history.
        con.execute("DELETE FROM qa_findings WHERE finding_type LIKE 'mechanical_%'")
        con.execute("DELETE FROM qa_findings WHERE finding_type LIKE 'formula_audit_%'")
        errors = _mechanical_db_checks(con, recommendation_present)
        if official_delivery["delivery_status"] == "FAILED":
            errors.append("Official paired delivery failed: " + official_delivery["error"])
        if cross_output.get("status") != "PASS":
            errors.append(f"Cross-output reconciliation failed with {len(cross_output.get('findings', []))} mismatch(es)")
        if provenance.get("status") != "PASS":
            errors.append(f"Provenance survival failed with {len(provenance.get('findings', []))} finding(s)")
        xlsx_art = _latest_artifact(con, "xlsx")
        warnings: list[str] = []
        if not xlsx_art:
            errors.append("XLSX artifact is missing")
        else:
            xlsx_path = run_root / xlsx_art["relative_path"]
            if sha256_file(xlsx_path) != xlsx_art["sha256"]:
                errors.append("XLSX artifact hash mismatch")
            formula_contract = get_meta(run_root, "formula_contract_" + xlsx_art["artifact_id"], None, parse_json=True)
            wb_errors, wb_warnings = _workbook_checks(xlsx_path, profile, formula_contract)
            errors.extend(wb_errors)
            warnings.extend(wb_warnings)
            from xlsx_formula_audit import audit as formula_audit
            formula_result = formula_audit(xlsx_path)
            for finding in formula_result.get("findings", []):
                severity = "critical" if finding.get("severity") == "critical" else "medium"
                message = f"{finding.get('sheet')}!{finding.get('locator')}: {finding.get('message')}"
                fid = stable_id("qa", "formula", finding.get("sheet",""), finding.get("locator",""), finding.get("category",""), length=22)
                con.execute(
                    "INSERT OR REPLACE INTO qa_findings(finding_id,artifact_id,finding_type,locator_json,severity,status,message,finding_hash,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                    [fid, xlsx_art["artifact_id"], f"formula_audit_{finding.get('category')}", json_text({"sheet":finding.get("sheet"),"locator":finding.get("locator")}), severity, "open", message, stable_id("h", message, length=20), utc_now()],
                )
                if severity == "critical": errors.append(message)
                else: warnings.append(message)
        for msg in errors:
            _finding(con, "mechanical_error", "critical", msg, artifact_id=xlsx_art["artifact_id"] if xlsx_art else None)
        for msg in warnings:
            _finding(con, "mechanical_warning", "medium", msg, artifact_id=xlsx_art["artifact_id"] if xlsx_art else None)

        visual_tasks = 0
        visual_pending = 0
        visual_status = "NOT_APPLICABLE" if profile != "yuksek_guvence" else "NOT_RUN"
        pdf_art = _latest_artifact(con, "pdf")
        if profile in {"hizli", "standart", "yuksek_guvence"}:
            if not pdf_art:
                _finding(con, "mechanical_missing_pdf", "critical", "Every profile requires final PDF export.")
                errors.append("Mandatory PDF missing")
            else:
                pdf_path = run_root / pdf_art["relative_path"]
                if sha256_file(pdf_path) != pdf_art["sha256"]:
                    errors.append("PDF artifact hash mismatch")
                    _finding(con, "mechanical_error", "critical", "PDF artifact hash mismatch", artifact_id=pdf_art["artifact_id"])
                else:
                    for msg in _pdf_content_checks(pdf_path, run_root):
                        errors.append(msg)
                        _finding(con, "mechanical_pdf_content", "critical", msg, artifact_id=pdf_art["artifact_id"])
                    # Validate mechanical content/hash first. A failed or
                    # preliminary artifact cannot begin final visual acceptance.
                    decision_status = str(get_meta(run_root, "decision_status", "UNKNOWN"))
                    if profile == "yuksek_guvence" and not errors and decision_status == "SUFFICIENT":
                        renders = _render_pdf(pdf_path, run_root / "qa" / pdf_art["artifact_id"])
                        visual_tasks = _visual_tasks(con, run_root, pdf_art, renders, full_first=True)
                        visual_status = "PENDING" if renders else "NOT_RUN"
                        visual_pending = con.execute("SELECT count(*) FROM render_hashes WHERE artifact_id=? AND qa_status<>'pass'", [pdf_art["artifact_id"]]).fetchone()[0]
        if hasattr(con, "commit"):
            con.commit()
        if visual_status == "PENDING" and visual_pending == 0:
            visual_status = "PASS"
        superseded_findings = 0
        if not errors and visual_status == "PASS" and pdf_art:
            superseded_findings = _supersede_visual_findings(con, run_root, pdf_art, renders)
            con.commit()
        status = "FAIL" if errors else ("NEEDS_VISUAL_QA" if visual_pending else "PASS")
        return {"visual_findings_superseded": superseded_findings, "status": status, "errors": errors, "warnings": warnings, "cross_output": cross_output, "provenance_survival": provenance, "visual_tasks_created": visual_tasks, "visual_pending_pages": int(visual_pending), "visual_status": visual_status}
    finally:
        con.close()


def main() -> None:
    ap = argparse.ArgumentParser(description="Mechanical and delta visual QA")
    ap.add_argument("--run", required=True)
    args = ap.parse_args()
    print(json.dumps(run_qa(Path(args.run)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
