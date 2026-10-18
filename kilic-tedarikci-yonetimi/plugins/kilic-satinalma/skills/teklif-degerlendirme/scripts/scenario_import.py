"""Import only editable scenario cells into a new preliminary report revision."""

from question_waits import guard_run_operation
import argparse
import json
from pathlib import Path
from common import sha256_file
from db import connect,get_meta
from report_delivery import inspect_preliminary,export_preliminary
from scenario_formulas import INPUTS,validate_formulas
from finance_workbook import read_lifecycle_inputs

@guard_run_operation
def import_scenarios(run_root,workbook):
    con=connect(run_root)
    try:
        if con.execute("SELECT 1 FROM workflow_state WHERE phase='COMPLETE' AND status IN ('PASS','PARTIAL')").fetchone():
            raise ValueError("Completed run preserved; create a new run")
    finally:con.close()
    source=inspect_preliminary(run_root)
    if source['delivery_status']!='READY':raise ValueError("Current preliminary source revision is not trusted: "+str(source.get('error')))
    manifest=get_meta(run_root,'preliminary_report_delivery',None,parse_json=True)
    root=Path(source['xlsx']).parent;contract_path=root/'formula_contract.json'
    if sha256_file(contract_path)!=manifest['files']['formula_contract.json']:raise ValueError("Stored formula contract changed")
    contract=json.loads(contract_path.read_text(encoding='utf-8'))
    audit=validate_formulas(workbook,contract)
    if audit['status']!='PASS':raise ValueError("Scenario workbook rejected: "+'; '.join(audit['errors']))
    from openpyxl import load_workbook
    wb=load_workbook(workbook,data_only=False)
    scenarios=[]
    try:
        for r,assumption in sorted(contract['assumptions'].items(),key=lambda pair:int(pair[0])):
            row=int(r);inputs={key:wb['Senaryo Girdileri'].cell(row,i).value for i,key in enumerate(INPUTS,start=2)}
            scenario={**assumption,'inputs':inputs,'basis':assumption.get('basis') or 'Kullanıcının Excel senaryo girdileri; kaynak ve resmî karar bazı değişmedi.'}
            scenario.pop('financial',None)
            scenario.pop('dated_costs',None)
            if contract.get('pricing_rows'):
                from scenario_pricing import read_pricing
                scenario['pricing']=read_pricing(wb,row,contract)
            from dated_workbook import SUMMARY,read_ledger
            if SUMMARY in wb:
                ledger=read_ledger(wb,row)
                if ledger['flows'] or any(ledger.get(key) is not None for key in ('valuation_date','base_currency','annual_rate','baseline','scope','method')):
                    scenario['dated_costs']=ledger
            if 'Finansal Varsayimlar' in wb:
                ws=wb['Finansal Varsayimlar'];count=ws.cell(row,2).value
                financial={'vat_refund_days':ws.cell(row,3).value,'base_present_value':ws.cell(row,4).value,
                           'policy':{key:ws.cell(row,i).value for i,key in enumerate(('base_currency','source_currency','rate_date','rate_source','direction','kind'),start=5)}}
                lifecycle=read_lifecycle_inputs(wb,row)
                financial.update(lifecycle)
                if count is not None:
                    start=2+(row-2)*12
                    financial['payments']=[{'percent':wb['Odeme Plani'].cell(k,3).value,'days':wb['Odeme Plani'].cell(k,4).value} for k in range(start,start+int(count))]
                has_lifecycle=bool(lifecycle) and (lifecycle.get('cost_mode')!='CONSTANT_ANNUAL' or bool(lifecycle.get('annual_costs')))
                if has_lifecycle or count is not None or financial['vat_refund_days'] is not None or financial['base_present_value'] is not None or any(v is not None for v in financial['policy'].values()):
                    scenario['financial']=financial
            scenarios.append(scenario)
        from scenario_scoring import read_scoring
        scoring=read_scoring(wb,contract)
    finally:wb.close()
    report=export_preliminary(run_root,'Kullanıcı Excel senaryosu; kaynak olguları ve resmî karar bazı korunmuştur.',scenarios,scoring=scoring)
    return {**report,'scenario_import':'IMPORTED' if report['delivery_status']=='READY' else 'FAILED',
            'imported_workbook_sha256':sha256_file(workbook),'decision_assurance':report.get('decision_assurance','NO_DECISION')}

def main():
    ap=argparse.ArgumentParser(description='Import user Excel scenarios into a new preliminary paired report; source facts remain unchanged')
    ap.add_argument('--run',required=True);ap.add_argument('--workbook',required=True);args=ap.parse_args()
    try:result=import_scenarios(Path(args.run),Path(args.workbook))
    except (ValueError,OSError) as exc:result={'scenario_import':'FAILED','deliverable_artifacts':False,'error':str(exc)}
    print(json.dumps(result,ensure_ascii=False,indent=2));raise SystemExit(0 if result.get('scenario_import')=='IMPORTED' else 2)
if __name__ == "__main__":
    from question_waits import run_guarded_main
    raise SystemExit(run_guarded_main(main))
