"""Multi-criterion Excel scenarios; source membership is fixed, decision is untouched."""
from copy import deepcopy
import math
from openpyxl.styles import Font,PatternFill,Protection,Alignment
from openpyxl.utils import get_column_letter as letter
from openpyxl.worksheet.datavalidation import DataValidation

SHEET='Cok Kriterli Puanlama'
SENSITIVITY='Agirlik Duyarliligi'

def validate(kind,value):
    if value is None:return
    choices={'score_direction':('direct','lower','higher'),'score_method':('proportional','minmax')}
    if kind in choices:
        if value not in choices[kind]:raise ValueError('Puanlama yönü/yöntemi geçersiz')
        return
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0:raise ValueError('Puanlama girdisi negatif olmayan sonlu sayı olmalı')
    if kind=='score_weight' and value>100:raise ValueError('Ağırlık 0–100 aralığında olmalı')
    if kind not in ('score_weight','score_value'):raise ValueError('Bilinmeyen puanlama girdisi')

def variants(criteria):
    base=[c.get('weight') for c in criteria]
    if len(base)<2 or any(v is None for v in base) or abs(sum(base)-100)>1e-6:return []
    donor=next(i for i,w in enumerate(base) if w>0);receiver=next(i for i in range(len(base)) if i!=donor)
    delta=min(5,base[donor]/3, (100-base[receiver])/3)
    if delta<=0:return []
    result=[]
    for multiplier in (1,2):
        weights=base.copy();weights[donor]-=delta*multiplier;weights[receiver]+=delta*multiplier
        result.append({'id':f'Ornek {multiplier}','weights':{c['id']:w for c,w in zip(criteria,weights)}})
    return result

def configuration(model):
    from source_scenarios import decode,numeric
    from report_model import current_excluded_offers
    params=decode(model.get('meta',{}).get('scoring_parameters'),{}) or {}
    if not params.get('criteria'):return None
    excluded=current_excluded_offers({'facts':[],'exclusions':[],**model,'requirements':model.get('requirements',[])})
    offers=[o for o in model.get('offers',[]) if o.get('status')=='active' and o['offer_id'] not in excluded]
    if len({o['supplier_id'] for o in offers})<2:return None
    criteria=[]
    for i,c in enumerate(params.get('criteria',[])):
        if c.get('monetized_in_cost'):continue
        scores=c.get('values',c.get('scores',{})) or {}
        criteria.append({'id':c.get('id') or c.get('criterion_id') or f'K{i+1}',
                         'weight':numeric(c.get('weight')),'direction':c.get('direction','direct'),
                         'method':c.get('method','proportional'),'values':{o['offer_id']:numeric(scores.get(o['offer_id'])) for o in offers}})
    if not criteria:return None
    sensitivity=params.get('sensitivity') or variants(criteria)
    sensitivity=[{**case,'weights':{key:numeric(value) for key,value in case['weights'].items()}} for case in sensitivity]
    return {'offers':[o['offer_id'] for o in offers],'criteria':criteria,'sensitivity':sensitivity,
            'basis':'Kaydedilmiş kriterlerden kullanıcı senaryosu; örnek alternatifler şirket politikası veya resmî satınalma sıralaması değildir.'}

def recompute(config):
    offers=config['offers'];criteria=config['criteria'];normalized={o:[] for o in offers}
    for c in criteria:
        validate('score_direction',c.get('direction'));validate('score_method',c.get('method'));validate('score_weight',c.get('weight'))
        values=[c['values'].get(o) for o in offers]
        for v in values:validate('score_value',v)
        if c.get('direction') is None or c.get('method') is None or any(v is None for v in values):return {}
        low,high=min(values),max(values)
        for oid,value in zip(offers,values):
            if c['direction']=='direct':
                if value>10:raise ValueError('Doğrudan puan 0–10 aralığında olmalı')
                score=value
            elif c['method']=='minmax':score=10 if high==low else 10*((high-value) if c['direction']=='lower' else (value-low))/(high-low)
            elif c['direction']=='lower':score=(10 if value==0 else 0) if low==0 else 10*low/value
            else:score=10 if high==0 else 10*value/high
            normalized[oid].append(score)
    results={}
    cases=[{'id':'Baz','weights':{c['id']:c.get('weight') for c in criteria}},*config.get('sensitivity',[])]
    if len({c['id'] for c in cases})!=len(cases):raise ValueError('Ağırlık senaryosu kimlikleri benzersiz olmalı')
    complete=[]
    for case in cases:
        weights=[case['weights'].get(c['id']) for c in criteria]
        if all(isinstance(w,(int,float)) and not isinstance(w,bool) for w in weights) and abs(sum(weights)-100)<1e-6:
            signature=tuple(weights)
            if signature in complete:raise ValueError('Baz ve alternatif ağırlık setleri birbirinden farklı olmalı')
            complete.append(signature)
    for case in cases:
        weights=[case['weights'].get(c['id']) for c in criteria]
        for weight in weights:validate('score_weight',weight)
        if any(w is None for w in weights) or abs(sum(weights)-100)>1e-6:continue
        scores={oid:sum(w*s for w,s in zip(weights,normalized[oid]))/100 for oid in offers}
        ratio=sum(w for w,c in zip(weights,criteria) if len(set(c['values'][o] for o in offers))>1)/100
        results[case['id']]={'scores':scores,'ranks':{o:1+sum(v>scores[o] for v in scores.values()) for o in offers},'distinguishing_weight_ratio':ratio}
    return results

def add_scoring(wb,model,contract,config=None):
    from scenario_formulas import literal
    config=deepcopy(config if config is not None else configuration(model))
    ws=wb.create_sheet(SHEET);ss=wb.create_sheet(SENSITIVITY)
    contract['scoring']=None
    if not config:
        literal(ws,['Puanlama durumu']);literal(ws,['En az iki uygun bağımsız firma ve kayıtlı kriter/ağırlık gerekir; değer uydurulmadı.'])
        literal(ss,['Duyarlılık durumu']);literal(ss,['Çok kriterli puanlama bazı hazır değil.'])
        for sheet in (ws,ss):sheet.column_dimensions['A'].width=90;sheet.protection.sheet=True
        return
    from report_model import current_excluded_offers
    excluded=current_excluded_offers({'facts':[],'exclusions':[],**model,'requirements':model.get('requirements',[])})
    members={o['offer_id']:o for o in model.get('offers',[]) if o.get('status')=='active' and o['offer_id'] not in excluded}
    if not set(config['offers'])<=set(members) or len({members[o]['supplier_id'] for o in config['offers']})<2:
        raise ValueError('Puan senaryosu en az iki uygun bağımsız kaynak firmasına bağlı olmalı')
    if len(set(config['offers']))!=len(config['offers']) or len({c['id'] for c in config['criteria']})!=len(config['criteria']):raise ValueError('Teklif/kriter kimlikleri benzersiz olmalı')
    criteria=config['criteria'];offers=config['offers'];n=len(criteria)
    if n<1:raise ValueError('En az bir kriter gerekli')
    if n>1 and len(config.get('sensitivity',[]))<2:config['sensitivity']=variants(criteria)
    results=recompute(config)
    contract.setdefault('extra_inputs',{})
    def input(sheet,cell,value,kind):
        validate(kind,value);sheet[cell]=value;sheet[cell].protection=Protection(locked=False);sheet[cell].fill=PatternFill('solid',fgColor='DDEBFF')
        contract['extra_inputs'][sheet.title+'!'+cell]=kind
        if kind in ('score_direction','score_method'):
            values='direct,lower,higher' if kind=='score_direction' else 'proportional,minmax'
            dv=DataValidation(type='list',formula1='"'+values+'"',allow_blank=True)
        else:dv=DataValidation(type='decimal',operator='between' if kind=='score_weight' else 'greaterThanOrEqual',formula1='0',formula2='100' if kind=='score_weight' else None,allow_blank=True)
        dv.showErrorMessage=True;dv.errorStyle='stop';sheet.add_data_validation(dv);dv.add(cell)
    def formula(sheet,cell,value):sheet[cell]=value;contract['formulas'][sheet.title+'!'+cell]=value
    names={s['supplier_id']:s['display_name'] for s in model.get('suppliers',[])}
    labels={o['offer_id']:names.get(o['supplier_id'],o['supplier_id'])+' / '+o['offer_id'] for o in model.get('offers',[])}
    literal(ws,['Kriter','Ağırlık: toplam 100','Yön: direct=puan; lower=düşük iyi; higher=yüksek iyi','Yöntem: proportional=oransal; minmax=aralık',*[labels.get(o,o) for o in offers]])
    for r,c in enumerate(criteria,start=2):
        literal(ws,[c['id']]);input(ws,'B'+str(r),c.get('weight'),'score_weight');input(ws,'C'+str(r),c.get('direction'),'score_direction');input(ws,'D'+str(r),c.get('method'),'score_method')
        for col,oid in enumerate(offers,start=5):input(ws,letter(col)+str(r),c['values'].get(oid),'score_value')
    normalized_first=n+4
    literal(ws,[]);literal(ws,['Normalleştirilmiş puanlar — 0–10'])
    for i,c in enumerate(criteria):
        source=i+2;r=normalized_first+i;ws.cell(r,1,c['id']).data_type='s'
        refs=f'E{source}:{letter(4+len(offers))}{source}';direction=f'C{source}';method=f'D{source}'
        for col,oid in enumerate(offers,start=5):
            raw=f'{letter(col)}{source}';low=f'MIN({refs})';high=f'MAX({refs})'
            direct=f'IF({raw}<=10,{raw},"")'
            minmax=f'IF({high}={low},10,10*IF({direction}="lower",{high}-{raw},{raw}-{low})/({high}-{low}))'
            proportional=f'IF({direction}="lower",IF({low}=0,IF({raw}=0,10,0),10*{low}/{raw}),IF({high}=0,10,10*{raw}/{high}))'
            guard=f'AND(COUNT({refs})={len(offers)},COUNTIF({refs},"<0")=0,OR({direction}="direct",{direction}="lower",{direction}="higher"),OR({method}="proportional",{method}="minmax"))'
            formula(ws,letter(col)+str(r),f'=IF(IFERROR({guard},FALSE),IF({direction}="direct",{direct},IF({method}="minmax",{minmax},{proportional})),"")')
    totalrow=normalized_first+n+1;rankrow=totalrow+1;ws.cell(totalrow,1,'Kullanıcı senaryosu toplam puanı');ws.cell(rankrow,1,'Senaryo sırası; satınalma önerisi değildir')
    weight=f'$B$2:$B${n+1}';weightguard=f'AND(COUNT({weight})={n},COUNTIF({weight},"<0")=0,COUNTIF({weight},">100")=0,ABS(SUM({weight})-100)<0.000001)'
    for col,oid in enumerate(offers,start=5):
        let=letter(col);scores=f'{let}{normalized_first}:{let}{normalized_first+n-1}'
        formula(ws,f'{let}{totalrow}',f'=IF(AND({weightguard},COUNT({scores})={n}),SUMPRODUCT({weight},{scores})/100,"")')
        totals=f'$E${totalrow}:${letter(4+len(offers))}${totalrow}'
        formula(ws,f'{let}{rankrow}',f'=IF(COUNT({totals})={len(offers)},1+COUNTIF({totals},">"&{let}{totalrow}),"")')
        if 'Baz' in results:
            contract['reference'].setdefault(str(totalrow),{})[SHEET+'!'+let]=results['Baz']['scores'][oid]
            contract['reference'].setdefault(str(rankrow),{})[SHEET+'!'+let]=results['Baz']['ranks'][oid]
    cases=config.get('sensitivity',[])
    literal(ss,['Kriter','Baz ağırlık',*[case['id'] for case in cases]])
    for r,c in enumerate(criteria,start=2):
        literal(ss,[c['id']]);formula(ss,f'B{r}',f"='{SHEET}'!B{r}")
        for col,case in enumerate(cases,start=3):input(ss,letter(col)+str(r),case['weights'].get(c['id']),'score_weight')
    sumrow=n+3;ss.cell(sumrow,1,'Ağırlık toplamı — 100 olmalı')
    for col in range(2,len(cases)+3):formula(ss,letter(col)+str(sumrow),f'=SUM({letter(col)}2:{letter(col)}{n+1})')
    first_result=sumrow+3;ss.cell(first_result-1,1,'Firma / senaryo puanları ve sıraları')
    for index,oid in enumerate(offers):
        r=first_result+index*2;ss.cell(r,1,labels.get(oid,oid));ss.cell(r+1,1,'Senaryo sırası')
        sourcecol=letter(index+5);score_range=f"'{SHEET}'!{sourcecol}{normalized_first}:{sourcecol}{normalized_first+n-1}"
        for col,case in enumerate([{'id':'Baz'},*cases],start=2):
            let=letter(col);weights=f'{let}2:{let}{n+1}'
            formula(ss,let+str(r),f'=IF(AND(COUNT({weights})={n},COUNTIF({weights},"<0")=0,COUNTIF({weights},">100")=0,ABS(SUM({weights})-100)<0.000001,COUNT({score_range})={n}),SUMPRODUCT({weights},{score_range})/100,"")')
            # Noncontiguous score cells exclude intervening rank rows.
            peers=','.join(f'{let}{first_result+i*2}' for i in range(len(offers)))
            higher='+'.join(f'IF({let}{first_result+i*2}>{let}{r},1,0)' for i in range(len(offers)))
            formula(ss,let+str(r+1),f'=IF(COUNT({peers})={len(offers)},1+{higher},"")')
            if case['id'] in results:
                contract['reference'].setdefault(str(r),{})[SENSITIVITY+'!'+let]=results[case['id']]['scores'][oid]
                contract['reference'].setdefault(str(r+1),{})[SENSITIVITY+'!'+let]=results[case['id']]['ranks'][oid]
    ratio_row=first_result+len(offers)*2+1;ss.cell(ratio_row,1,'Ayırt edici ağırlık oranı')
    for col,case in enumerate([{'id':'Baz'},*cases],start=2):
        let=letter(col);weights=f'{let}2:{let}{n+1}'
        terms=[]
        for index in range(n):
            source=index+2;refs=f"'{SHEET}'!E{source}:{letter(4+len(offers))}{source}"
            terms.append(f'IF(MAX({refs})<>MIN({refs}),{let}{source},0)')
        scores=f"'{SHEET}'!E{normalized_first}:{letter(4+len(offers))}{normalized_first+n-1}"
        formula(ss,let+str(ratio_row),f'=IF(AND(COUNT({scores})={n*len(offers)},COUNT({weights})={n},ABS(SUM({weights})-100)<0.000001),SUM({",".join(terms)})/100,"")')
        ss[let+str(ratio_row)].number_format='0.0%'
        if case['id'] in results:contract['reference'].setdefault(str(ratio_row),{})[SENSITIVITY+'!'+let]=results[case['id']]['distinguishing_weight_ratio']
    contract['scoring']={'configuration':config,'normalized_first':normalized_first,'total_row':totalrow,'rank_row':rankrow}
    contract['scoring_results']=results
    for sheet in (ws,ss):
        sheet.freeze_panes='E2' if sheet==ws else 'C2';sheet.auto_filter.ref=sheet.dimensions;sheet.protection.sheet=True
        sheet.row_dimensions[1].height=60
        for cell in sheet[1]:cell.font=Font(bold=True)
        for cells in sheet:
            for cell in cells:cell.alignment=Alignment(wrap_text=True,vertical='top')
        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=28
        sheet.column_dimensions['A'].width=45

def read_scoring(wb,contract):
    entry=contract.get('scoring')
    if entry is None:return None
    config=deepcopy(entry['configuration']);ws=wb[SHEET];ss=wb[SENSITIVITY]
    for r,c in enumerate(config['criteria'],start=2):
        c.update(weight=ws.cell(r,2).value,direction=ws.cell(r,3).value,method=ws.cell(r,4).value,
                 values={oid:ws.cell(r,col).value for col,oid in enumerate(config['offers'],start=5)})
    for col,case in enumerate(config.get('sensitivity',[]),start=3):
        case['weights']={c['id']:ss.cell(r,col).value for r,c in enumerate(config['criteria'],start=2)}
    recompute(config)
    return config
