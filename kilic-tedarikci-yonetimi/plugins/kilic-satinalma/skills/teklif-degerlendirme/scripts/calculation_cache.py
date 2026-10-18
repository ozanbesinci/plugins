"""Validate a calculated copy and transfer only caches into original OOXML."""
import argparse
from copy import deepcopy
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import posixpath
import re
import zipfile

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PNS='http://schemas.openxmlformats.org/package/2006/relationships'
ns={'m':NS}
q=lambda name:'{'+NS+'}'+name
sha=lambda path:hashlib.sha256(Path(path).read_bytes()).hexdigest()

# This engine consumes generated report formulas, never arbitrary workbooks.
FUNCTIONS=set('ABS AND CODE COUNT COUNTA COUNTIF COUNTIFS DATE DAY EXACT FALSE IF IFERROR INT ISNUMBER ISTEXT LEFT LEN MAX MID MIN MONTH OR RIGHT SUM SUMPRODUCT TRIM TRUE VALUE YEAR ROUND'.split())

def preflight(source, cold):
    """Remove stale caches and refuse executable/external workbook features."""
    count=0
    with zipfile.ZipFile(source) as src:
        names=src.namelist()
        if len(names)!=len(set(names)) or sum(i.file_size for i in src.infolist())>128*1024*1024:
            raise ValueError('Invalid or oversized workbook')
        for name in names:
            low=name.lower()
            if any(x in low for x in ('vbaproject','externallinks/','connections.xml','querytables/','embeddings/','customxml/')):
                raise ValueError('External or executable workbook content')
            if name.endswith('.rels'):
                if any(r.get('TargetMode')=='External' for r in tree(src.read(name))):
                    raise ValueError('External workbook relationship')
        # Token checks before spawning the engine, including defined names.
        wb=load_workbook(BytesIO(Path(source).read_bytes()),data_only=False)
        try:
            if wb.defined_names:raise ValueError('Defined names outside generated report contract')
            for ws in wb:
                if ws.max_row*ws.max_column>500000:raise ValueError('Oversized sheet')
                for row in ws:
                    for c in row:
                        if c.data_type=='f':
                            formula_tokens(c.value);count+=1
        finally:wb.close()
        if not count:raise ValueError('No generated formulas to calculate')
        parts=set(sheet_parts(src).values())
        with zipfile.ZipFile(cold,'x',zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                data=src.read(info.filename)
                if info.filename in parts:
                    root=tree(data)
                    for c in root.findall('.//m:sheetData/m:row/m:c',ns):
                        if c.find(q('f')) is not None:
                            c.attrib.pop('t',None)
                            for old in list(c):
                                if old.tag in (q('v'),q('is')):c.remove(old)
                    data=ET.tostring(root,encoding='UTF-8',xml_declaration=True)
                dst.writestr(info,data)
    return count

def formula_tokens(expression):
    items=Tokenizer(expression).items; result=[]; i=0
    for token in items:
        if token.type=='FUNC' and token.subtype=='OPEN' and token.value[:-1].upper() not in FUNCTIONS:
            raise ValueError('Formula outside generated report contract: '+token.value)
        if token.type=='OPERAND' and token.subtype=='RANGE' and any(x in token.value for x in ('[',']','|','\\')):
            raise ValueError('External formula reference')
    while i<len(items):
        t=items[i]; value=t.value
        if (t.type=='FUNC' and t.subtype=='OPEN' and value in ('FALSE(','TRUE(')
                and i+1<len(items) and items[i+1].type=='FUNC' and items[i+1].subtype=='CLOSE'):
            result.append(('OPERAND','LOGICAL',value[:-1])); i+=2; continue
        if t.type=='OPERAND' and t.subtype=='RANGE':
            prefix,sep,ref=value.rpartition('!')
            found=re.fullmatch(r'(\$?[A-Z]{1,3}\$?[1-9][0-9]*):\1',ref)
            if found:value=(prefix+sep if sep else '')+found.group(1)
        result.append((t.type,t.subtype,value));i+=1
    return result

def same(a,b):
    if isinstance(a,bool) or isinstance(b,bool):return type(a)==type(b) and a==b
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):
        return math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-7)
    return type(a)==type(b) and a==b

def tree(data):
    return ET.fromstring(data,parser=ET.XMLParser(resolve_entities=False,no_network=True))

def sheet_parts(archive):
    rels=tree(archive.read('xl/_rels/workbook.xml.rels'))
    mapping={r.get('Id'):r.get('Target') for r in rels}
    result={}
    for sheet in tree(archive.read('xl/workbook.xml')).findall('m:sheets/m:sheet',ns):
        target=mapping[sheet.get('{'+RNS+'}id')]
        part=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        if not part.startswith('xl/worksheets/') or part not in archive.namelist():raise ValueError('Unexpected sheet part')
        result[sheet.get('name')]=part
    return result

def stripped(xml):
    root=tree(xml)
    for c in root.findall('.//m:sheetData/m:row/m:c',ns):
        if c.find(q('f')) is not None:
            c.attrib.pop('t',None)
            for old in list(c):
                if old.tag in (q('v'),q('is')):c.remove(old)
    return ET.tostring(root,method='c14n')

def validate_caches(path,contract):
    """Require all formula caches and compare known scenario values independently."""
    count=0
    with zipfile.ZipFile(path) as book:
        for part in sheet_parts(book).values():
            for cell in tree(book.read(part)).findall('.//m:sheetData/m:row/m:c',ns):
                if cell.find(q('f')) is None:continue
                value=cell.find(q('v'));typ=cell.get('t','n')
                if value is None or typ not in ('n','b','str'):raise ValueError('Invalid/missing formula cache')
                if typ=='n' and (value.text is None or not math.isfinite(float(value.text))):raise ValueError('Invalid numeric cache')
                if typ=='b' and value.text not in ('0','1'):raise ValueError('Invalid boolean cache')
                count+=1
    if count!=len(contract['formulas']):raise ValueError('Calculation cache count differs from contract')
    wb=load_workbook(BytesIO(Path(path).read_bytes()),data_only=True)
    compared=0
    try:
        for row,values in contract.get('reference',{}).items():
            for key,wanted in values.items():
                sheet,column=key.rsplit('!',1)
                actual=wb[sheet][column+str(row)].value
                if not same(actual,wanted):raise ValueError('Independent scenario cache mismatch: '+key+str(row))
                compared+=1
    finally:wb.close()
    return {'status':'PASS','formula_count':count,'independent_values_compared':compared}

def transfer(source,calculated,output):
    source,calculated,output=map(Path,(source,calculated,output))
    if output.exists() or output.resolve() in (source.resolve(),calculated.resolve()):raise ValueError('Output must be new')
    before={str(p):sha(p) for p in (source,calculated)}
    # Malformed cell values can make openpyxl fail while loading. Byte buffers
    # prevent such a rejected workbook from leaving an open Windows file handle.
    src=load_workbook(BytesIO(source.read_bytes()),data_only=False)
    calc=load_workbook(BytesIO(calculated.read_bytes()),data_only=False)
    vals=load_workbook(BytesIO(calculated.read_bytes()),data_only=True)
    expected_formulas=set();changes=0;boolean_cells=0
    try:
        if src.sheetnames!=calc.sheetnames:raise ValueError('Sheet names changed')
        for ws in src:
            cs,vs=calc[ws.title],vals[ws.title]
            for row in ws:
                for c in row:
                    other=cs[c.coordinate]
                    if c.data_type=='f':
                        if other.data_type!='f' or formula_tokens(c.value)!=formula_tokens(other.value):
                            raise ValueError('Formula changed: '+ws.title+'!'+c.coordinate)
                        if vs[c.coordinate].data_type=='e':raise ValueError('Formula error')
                        expected_formulas.add((ws.title,c.coordinate));changes+=int(c.value!=other.value)
                    elif c.data_type=='b' and other.data_type=='f' and other.value==('=TRUE()' if c.value else '=FALSE()'):
                        if not same(c.value,vs[c.coordinate].value):raise ValueError('Boolean value changed')
                        boolean_cells+=1
                    elif not same(c.value,other.value):
                        # Empty inline strings become empty cells in round trip.
                        if c.value in (None,'') and other.value in (None,''):continue
                        raise ValueError('Literal changed: '+ws.title+'!'+c.coordinate)
            for row in cs:
                for c in row:
                    if c.value is not None and ws[c.coordinate].value is None:
                        raise ValueError('New nonempty cell')
    finally:
        src.close();calc.close();vals.close()
    modified={};copied=0
    with zipfile.ZipFile(source) as original,zipfile.ZipFile(calculated) as computed:
        source_parts=sheet_parts(original);calc_parts=sheet_parts(computed)
        for title,part in source_parts.items():
            root=tree(original.read(part))
            target=tree(computed.read(calc_parts[title]))
            caches={c.get('r'):c for c in target.findall('.//m:sheetData/m:row/m:c',ns)}
            count=0
            for c in root.findall('.//m:sheetData/m:row/m:c',ns):
                if c.find(q('f')) is None:continue
                key=(title,c.get('r'))
                if key not in expected_formulas:raise ValueError('Unexpected formula')
                cc=caches.get(c.get('r'))
                if cc is None or cc.find(q('v')) is None:raise ValueError('Missing calculation cache')
                typ=cc.get('t','n')
                if typ not in ('n','str','b'):raise ValueError('Unsupported cache type')
                text=cc.find(q('v')).text
                if typ=='n' and (text is None or not math.isfinite(float(text))):raise ValueError('Invalid numeric cache')
                if typ=='b' and text not in ('0','1'):raise ValueError('Invalid boolean cache')
                for old in list(c):
                    if old.tag in (q('v'),q('is')):c.remove(old)
                c.set('t',typ)
                c.insert(list(c).index(c.find(q('f')))+1,deepcopy(cc.find(q('v'))))
                count+=1
            if count:
                raw=ET.tostring(root,encoding='UTF-8',xml_declaration=True)
                if stripped(raw)!=stripped(original.read(part)):raise ValueError('Non-cache XML changed')
                modified[part]=raw;copied+=count
        if copied!=len(expected_formulas):raise ValueError('Formula cache count mismatch')
        if any(sha(Path(p))!=h for p,h in before.items()):raise ValueError('Input changed')
        # Exclusive creation: neither the source nor an existing output can be replaced.
        with output.open('xb') as handle,zipfile.ZipFile(handle,'w') as dest:
            for info in original.infolist():dest.writestr(info,modified.get(info.filename,original.read(info.filename)))
    return {'formulas':copied,'formula_normalizations_observed':changes,'boolean_roundtrip_changes_observed':boolean_cells,
            'non_cache_xml_preserved':True,'source_hashes':before,'sources_unchanged':all(sha(Path(p))==h for p,h in before.items()),
            'output_sha256':sha(output),'modified_sheet_parts':list(modified),
            'engine':'LibreOffice','native_excel_validation':False}

if __name__=='__main__':
    ap=argparse.ArgumentParser()
    for name in ('source','calculated','output'):ap.add_argument('--'+name,required=True)
    args=ap.parse_args()
    print(json.dumps(transfer(args.source,args.calculated,args.output),ensure_ascii=False))
