from __future__ import annotations

import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from adapters.base import ParseResult, Unit


def _tracked_changes(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml")
        document = ET.fromstring(xml)
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        revision_tags = {
            "ins", "del", "moveFrom", "moveTo", "moveFromRangeStart", "moveFromRangeEnd",
            "moveToRangeStart", "moveToRangeEnd", "customXmlInsRangeStart", "customXmlInsRangeEnd",
            "customXmlDelRangeStart", "customXmlDelRangeEnd", "customXmlMoveFromRangeStart",
            "customXmlMoveFromRangeEnd", "customXmlMoveToRangeStart", "customXmlMoveToRangeEnd",
            "pPrChange", "rPrChange", "sectPrChange", "tblPrChange", "tblPrExChange",
            "tblGridChange", "trPrChange", "tcPrChange", "numberingChange", "cellIns", "cellDel", "cellMerge",
        }
        # Exact expanded names: table borders insideH/insideV are not insertions.
        expanded_tags = {namespace + tag for tag in revision_tags}
        return any(element.tag in expanded_tags for element in document.iter())
    except Exception:
        return False


def parse(path: Path) -> ParseResult:
    try:
        from docx import Document
        doc = Document(str(path))
    except Exception as exc:
        return ParseResult("python-docx", "unreadable", [], [str(exc)])
    units: list[Unit] = []
    idx = 1
    for pno, p in enumerate(doc.paragraphs, start=1):
        if p.text.strip():
            units.append(Unit("docx_paragraph", idx, {"paragraph": pno}, p.text))
            idx += 1
    for tno, table in enumerate(doc.tables, start=1):
        for rno, row in enumerate(table.rows, start=1):
            values = [c.text for c in row.cells]
            units.append(Unit("docx_table_row", idx, {"table": tno, "row": rno}, " | ".join(values), {"cells": values}))
            idx += 1
    tracked = _tracked_changes(path)
    notes = ["Tracked changes detected; inspect OOXML inserted/deleted runs before final use."] if tracked else []
    return ParseResult("python-docx+ooxml-check", "partial" if tracked else "read", units, notes, {"tracked_changes": tracked})
