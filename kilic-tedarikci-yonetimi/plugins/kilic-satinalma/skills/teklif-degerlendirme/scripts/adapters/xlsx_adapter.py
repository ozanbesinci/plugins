from __future__ import annotations

from pathlib import Path
import colorsys
from xml.etree import ElementTree

from adapters.base import ParseResult, Unit
from common import json_safe


_THEME_ORDER = ('lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3',
                'accent4', 'accent5', 'accent6', 'hlink', 'folHlink')


def _theme_colors(workbook) -> dict[int, str]:
    raw = getattr(workbook, 'loaded_theme', None)
    if not raw:
        return {}
    try:
        root = ElementTree.fromstring(raw)
        scheme = next(node for node in root.iter() if node.tag.rsplit('}', 1)[-1] == 'clrScheme')
    except (ElementTree.ParseError, StopIteration, TypeError):
        return {}
    colors = {}
    for child in scheme:
        name = child.tag.rsplit('}', 1)[-1]
        if name not in _THEME_ORDER:
            continue
        source = next(iter(child), None)
        if source is None:
            continue
        rgb = source.get('lastClr') or source.get('val')
        if rgb and len(rgb) in (6, 8):
            colors[_THEME_ORDER.index(name)] = rgb[-6:].upper()
    return colors


def _tinted_rgb(rgb: str, tint: float) -> str:
    channels = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    hue, lightness, saturation = colorsys.rgb_to_hls(*channels)
    lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
    adjusted = colorsys.hls_to_rgb(hue, min(1, max(0, lightness)), saturation)
    return ''.join(f'{round(value * 255):02X}' for value in adjusted)


def _fill_payload(cell, theme_colors: dict[int, str]) -> dict | None:
    fill = cell.fill
    if not fill or not fill.patternType:
        return None
    color = fill.fgColor
    raw = {'type': color.type, 'tint': color.tint}
    if color.type in {'rgb', 'indexed', 'theme', 'auto'}:
        value = getattr(color, color.type)
        raw[color.type] = value if isinstance(value, (str, int, float, bool)) else None
    result = {'pattern_type': fill.patternType, 'foreground': raw}
    rgb = None
    if color.type == 'rgb' and isinstance(color.rgb, str) and len(color.rgb) in (6, 8):
        rgb = color.rgb[-6:].upper()
    elif color.type == 'theme' and isinstance(color.theme, int):
        rgb = theme_colors.get(color.theme)
    elif color.type == 'indexed' and isinstance(color.indexed, int) and color.indexed < 64:
        from openpyxl.styles.colors import COLOR_INDEXED
        rgb = COLOR_INDEXED[color.indexed][-6:].upper()
    if fill.patternType == 'solid' and rgb:
        try:
            result['effective_rgb'] = _tinted_rgb(rgb, float(color.tint or 0))
        except (TypeError, ValueError):
            pass
    return result


def _sheet_payload(ws_formula, ws_values):
    cells = []
    theme_colors = _theme_colors(ws_formula.parent)
    hidden_rows = [i for i, d in ws_formula.row_dimensions.items() if d.hidden]
    hidden_cols = [k for k, d in ws_formula.column_dimensions.items() if d.hidden]
    max_row = ws_formula.max_row or 0
    max_col = ws_formula.max_column or 0
    for row in ws_formula.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            vcell = ws_values[cell.coordinate]
            is_formula = cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))
            record = {
                "cell": cell.coordinate,
                "formula_or_value": json_safe(cell.value),
                "calculated_value": json_safe(vcell.value),
                "cell_type": cell.data_type,
                "is_formula": bool(is_formula),
                "calculated_available": bool(not is_formula or vcell.value is not None),
                "number_format": cell.number_format,
            }
            fill = _fill_payload(cell, theme_colors)
            if fill:
                record['fill'] = fill
            cells.append(record)
    return {
        "title": ws_formula.title,
        "sheet_state": ws_formula.sheet_state,
        "max_row": max_row,
        "max_column": max_col,
        "hidden_rows": hidden_rows,
        "hidden_columns": hidden_cols,
        "merged_ranges": [str(r) for r in ws_formula.merged_cells.ranges],
        "cells": cells,
    }


def parse(path: Path) -> ParseResult:
    try:
        from openpyxl import load_workbook
        formula_wb = load_workbook(path, data_only=False, read_only=False, keep_vba=path.suffix.lower()==".xlsm")
        values_wb = load_workbook(path, data_only=True, read_only=False, keep_vba=path.suffix.lower()==".xlsm")
    except Exception as exc:
        return ParseResult("openpyxl", "unreadable", [], [str(exc)])

    units: list[Unit] = []
    notes: list[str] = []
    for i, ws in enumerate(formula_wb.worksheets, start=1):
        payload = _sheet_payload(ws, values_wb[ws.title])
        text_lines = []
        for c in payload["cells"]:
            if c.get("is_formula") and not c.get("calculated_available"):
                value = "<UNCALCULATED_FORMULA>"
            else:
                value = c["calculated_value"] if c["calculated_value"] is not None else c["formula_or_value"]
            text_lines.append(f"{c['cell']}={value}")
        units.append(Unit("xlsx_sheet", i, {"sheet": ws.title}, "\n".join(text_lines), payload))
        if payload["sheet_state"] != "visible" or payload["hidden_rows"] or payload["hidden_columns"]:
            notes.append(f"Hidden content detected in sheet {ws.title}")
        uncalculated = sum(1 for c in payload["cells"] if c.get("is_formula") and not c.get("calculated_available"))
        if uncalculated:
            notes.append(f"Uncalculated formula cells detected in sheet {ws.title}: {uncalculated}")
    return ParseResult("openpyxl", "read", units, notes, {"sheet_count": len(units)})
