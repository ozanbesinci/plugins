"""WP-046: trusted, location-bound formula templates and explicit scenario inputs."""
from __future__ import annotations
import math
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from procurement_math import line_gross, sequential_discount, fx_to_base, vat_split, present_value_days

INPUTS = ["quantity","unit_price","discount","fx_rate","vat_rate","vat_deductible",
          "freight","insurance","duty_rate","annual_rate","days","energy_annual",
          "maintenance_annual","years","residual","technical_score","technical_weight",
          "cost_score","cost_weight","base_total"]
INPUT_LABELS = ["Miktar","Birim fiyat","İskonto oranı","Hedef/kaynak kur","KDV oranı","İndirilebilir KDV payı",
                "Navlun","Sigorta","Gümrük oranı","Yıllık efektif oran","Ödeme günü","Yıllık enerji",
                "Yıllık bakım","Yıl sayısı","Kalıntı değer","Teknik puan","Teknik ağırlık (0–1)",
                "Maliyet puanı","Maliyet ağırlık (0–1)","Karşılaştırma baz tutarı"]
SHEETS = ["Senaryo Kontrol","Senaryo Girdileri","Miktar ve Kapsam","Kur Vergi Ithalat",
          "Odeme NBD Nakit","TCO Yasam Dongusu","Puanlama Duyarlilik","Senaryo Karsilastirma"]

def literal(ws, row):
    ws.append(row)
    for c in ws[ws.max_row]:
        if isinstance(c.value,str): c.data_type = "s"


def source_cell_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, date):
        if not isinstance(value, datetime):
            value = datetime.combine(value, time())
        return value.isoformat()
    return value

def _valid(key, value):
    if value is None: return
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
        raise ValueError("Senaryo girdisi sonlu sayı olmalı: "+key)
    if value < 0 or (key in {"quantity","fx_rate","years"} and value <= 0):
        raise ValueError("Senaryo girdisi aralığı geçersiz: "+key)
    if key in {"discount","vat_deductible","technical_weight","cost_weight"} and value > 1:
        raise ValueError("Senaryo oranı 0..1 olmalı: "+key)
    if key in {"technical_score","cost_score"} and value > 10:
        raise ValueError("Senaryo puanı 0..10 olmalı: "+key)
    if key == "days" and value != int(value):
        raise ValueError("Ödeme günü tam sayı olmalı")
    if key == "discount" and value >= 1:
        raise ValueError("İskonto 1'den küçük olmalı")

def add_scenarios(wb, scenarios=None):
    scenarios = scenarios or [{"id":"Teyit bekleyen senaryo","inputs":{}}]
    if len({x["id"] for x in scenarios}) != len(scenarios): raise ValueError("Duplicate scenario")
    sheets = {n:wb.create_sheet(n) for n in SHEETS}
    contract = {"version":"teklif-formula-contract/v1","formulas":{},"inputs":{},
                "source_cells":{},"reference":{},"recalculation":"NOT_EVALUATED_BY_STRUCTURAL_CHECK"}
    def formula(ws,cell,value):
        ws[cell] = value
        contract["formulas"][ws.title+"!"+cell] = value
    literal(sheets["Senaryo Girdileri"],["Senaryo"]+INPUT_LABELS+["Varsayım ve kur dayanağı"])
    for name, headers in {
        "Miktar ve Kapsam":["Senaryo","Net kalem tutarı","Durum"],
        "Kur Vergi Ithalat":["Senaryo","Hedef para net","İthalat matrahı","İndirilemeyen vergi","Nakit brüt","TCO baz tutarı","Durum"],
        "Odeme NBD Nakit":["Senaryo","Nominal ödeme","NBD","Durum"],
        "TCO Yasam Dongusu":["Senaryo","Yaşam döngüsü nominal toplam","Durum"],
        "Puanlama Duyarlilik":["Senaryo: iki puan / ağırlık toplamı 1","Ağırlıklı puan","Teknik ağırlık +0,05 duyarlılığı","Durum"],
        "Senaryo Karsilastirma":["Senaryo","Baz","Senaryo TCO","Fark","Oransal fark","Durum"],
        "Senaryo Kontrol":["Senaryo","Net bedel","NBD","TCO","Simülasyon durumu"],
    }.items(): literal(sheets[name],headers)
    for row, scenario in enumerate(scenarios,start=2):
        values = scenario.get("inputs",{})
        if set(values)-set(INPUTS): raise ValueError("Unknown scenario field")
        for key,value in values.items(): _valid(key,value)
        literal(sheets["Senaryo Girdileri"],[scenario["id"]]+[values.get(k) for k in INPUTS]+
                [scenario.get("basis") or "Kullanıcı varsayımı; kur yönü ve tarih/kaynak teyidi gerekli."])
        for i,key in enumerate(INPUTS,start=2):
            cell = sheets["Senaryo Girdileri"].cell(row,i)
            contract["inputs"][cell.coordinate] = key
            cell.protection = Protection(locked=False)
            cell.fill = PatternFill("solid",fgColor="DDEBFF")
        refs = {key:f"'Senaryo Girdileri'!{get_column_letter(i)}{row}"
                for i,key in enumerate(INPUTS,start=2)}
        def guard(keys): return "AND("+",".join("ISNUMBER("+refs[k]+")" for k in keys)+")"
        def f(sheet,col,expr,condition):
            formula(sheets[sheet],f"{col}{row}",f'=IF({condition},{expr},"")')
        def status(sheet,col,condition,why):
            formula(sheets[sheet],f"{col}{row}",f'=IF({condition},"SİMÜLASYON HESAPLANDI","HESAPLANAMADI — {why}")')
        for name in SHEETS:
            if name!="Senaryo Girdileri": literal(sheets[name],[scenario["id"]])
        net_guard = guard(["quantity","unit_price","discount"])
        net_guard = f"AND({net_guard},{refs['quantity']}>0,{refs['unit_price']}>=0,{refs['discount']}>=0,{refs['discount']}<1)"
        f("Miktar ve Kapsam","B",f"ROUND({refs['quantity']}*{refs['unit_price']}*(1-{refs['discount']}),2)",net_guard)
        status("Miktar ve Kapsam","C",net_guard,"miktar/fiyat/iskonto")
        net = f"'Miktar ve Kapsam'!B{row}"
        fx_guard = f"AND(ISNUMBER({net}),{guard(['fx_rate'])},{refs['fx_rate']}>0)"
        f("Kur Vergi Ithalat","B",f"{net}*{refs['fx_rate']}",fx_guard)
        converted = f"'Kur Vergi Ithalat'!B{row}"
        import_guard = f"AND(ISNUMBER({converted}),{guard(['freight','insurance','duty_rate','vat_rate','vat_deductible'])},{refs['freight']}>=0,{refs['insurance']}>=0,{refs['duty_rate']}>=0,{refs['vat_rate']}>=0,{refs['vat_deductible']}>=0,{refs['vat_deductible']}<=1)"
        base = f"C{row}"; duty = refs["duty_rate"]; vat=refs["vat_rate"]
        f("Kur Vergi Ithalat","C",f"{converted}+{refs['freight']}+{refs['insurance']}",import_guard)
        f("Kur Vergi Ithalat","D",f"{base}*{duty}+{base}*(1+{duty})*{vat}*(1-{refs['vat_deductible']})",import_guard)
        f("Kur Vergi Ithalat","E",f"{base}*(1+{duty})*(1+{vat})",import_guard)
        f("Kur Vergi Ithalat","F",f"{base}+D{row}",import_guard)
        status("Kur Vergi Ithalat","G",import_guard,"kur/vergi/navlun/sigorta/gümrük")
        cash=f"'Kur Vergi Ithalat'!E{row}"; cost=f"'Kur Vergi Ithalat'!F{row}"
        pv_guard=f"AND(ISNUMBER({cash}),{guard(['annual_rate','days'])},{refs['annual_rate']}>=0,{refs['days']}>=0)"
        f("Odeme NBD Nakit","B",cash,f"ISNUMBER({cash})")
        f("Odeme NBD Nakit","C",f"{cash}/(1+{refs['annual_rate']})^({refs['days']}/365)",pv_guard)
        status("Odeme NBD Nakit","D",pv_guard,"efektif oran/gün; tek ödeme 365 gün bazı")
        tco_guard=f"AND(ISNUMBER({cost}),{guard(['energy_annual','maintenance_annual','years','residual'])},{refs['years']}>0,{refs['energy_annual']}>=0,{refs['maintenance_annual']}>=0,{refs['residual']}>=0)"
        f("TCO Yasam Dongusu","B",f"{cost}+({refs['energy_annual']}+{refs['maintenance_annual']})*{refs['years']}-{refs['residual']}",tco_guard)
        status("TCO Yasam Dongusu","C",tco_guard,"enerji/bakım/dönem/kalıntı; nominal senaryo")
        score_guard=guard(["technical_score","technical_weight","cost_score","cost_weight"])
        score_guard=f"AND({score_guard},ABS({refs['technical_weight']}+{refs['cost_weight']}-1)<0.000001,{refs['technical_score']}>=0,{refs['technical_score']}<=10,{refs['cost_score']}>=0,{refs['cost_score']}<=10,{refs['technical_weight']}>=0,{refs['technical_weight']}<=1,{refs['cost_weight']}>=0,{refs['cost_weight']}<=1)"
        f("Puanlama Duyarlilik","B",f"{refs['technical_score']}*{refs['technical_weight']}+{refs['cost_score']}*{refs['cost_weight']}",score_guard)
        sensitivity=f"AND({score_guard},{refs['technical_weight']}<=0.95)"
        f("Puanlama Duyarlilik","C",f"{refs['technical_score']}*({refs['technical_weight']}+0.05)+{refs['cost_score']}*({refs['cost_weight']}-0.05)",sensitivity)
        status("Puanlama Duyarlilik","D",score_guard,"puan/ağırlık; rekabetçi sıralama değildir")
        tco=f"'TCO Yasam Dongusu'!B{row}"; comp=f"AND(ISNUMBER({tco}),{guard(['base_total'])})"
        f("Senaryo Karsilastirma","B",refs["base_total"],guard(["base_total"]))
        f("Senaryo Karsilastirma","C",tco,f"ISNUMBER({tco})")
        f("Senaryo Karsilastirma","D",f"C{row}-B{row}",comp)
        f("Senaryo Karsilastirma","E",f"D{row}/B{row}",f"AND({comp},B{row}<>0)")
        status("Senaryo Karsilastirma","F",comp,"aynı kapsamdaki baz/TCO")
        for col,ref in [("B",net),("C",f"'Odeme NBD Nakit'!C{row}"),("D",tco)]:
            f("Senaryo Kontrol",col,ref,f"ISNUMBER({ref})")
        sheets["Senaryo Kontrol"].cell(row,5,"Kullanıcı senaryosu; kaynak ve resmî baz değişmez; sıralama yok.")
        sheets["Senaryo Kontrol"].cell(row,5).data_type="s"
        contract["reference"][str(row)] = scenario_reference(values, scenario.get('pricing'))
    for ws in sheets.values():
        ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.protection.sheet=True
        ws.column_dimensions["A"].width=30
        for row in ws:
            for cell in row:
                if cell.row==1: cell.font=Font(bold=True)
                ws.column_dimensions[cell.column_letter].width=24
        ws.sheet_properties.pageSetUpPr.fitToPage=True
    input_sheet = sheets["Senaryo Girdileri"]
    for column, key in enumerate(INPUTS, start=2):
        upper = 1 if key in {"vat_deductible", "technical_weight", "cost_weight"} else 10 if key in {"technical_score", "cost_score"} else None
        positive = key in {"quantity", "fx_rate", "years"}
        validation = DataValidation(type="whole" if key == "days" else "decimal",
                                    operator="between" if upper is not None else "greaterThan" if positive else "greaterThanOrEqual",
                                    formula1="0", formula2=str(upper) if upper is not None else None,
                                    allow_blank=True)
        if key == "discount":
            validation = DataValidation(type="decimal", operator="between", formula1="0", formula2="0.999999999999", allow_blank=True)
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        validation.errorTitle = "Geçersiz senaryo girdisi"
        validation.error = "Alan için izinli aralıkta sayı girin; bilinmiyorsa boş bırakın."
        input_sheet.add_data_validation(validation)
        validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{len(scenarios)+1}")
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.forceFullCalc=True
    from scenario_pricing import add_pricing
    add_pricing(wb,scenarios,contract)
    from finance_workbook import add_finance
    add_finance(wb,scenarios,contract)
    from dated_workbook import add_dated_ledgers
    add_dated_ledgers(wb,scenarios,contract)
    contract['assumptions']={str(r):scenario for r,scenario in enumerate(scenarios,start=2)}
    return contract

def scenario_reference(v, pricing=None):
    """Reuse existing deterministic arithmetic; do not invent missing inputs."""
    result = {}
    def enough(*keys): return all(v.get(k) is not None for k in keys)
    from scenario_pricing import net_price
    if enough("quantity","unit_price","discount") and (net := net_price(v,pricing)) is not None:
        result["Miktar ve Kapsam!B"]=float(net)
        if enough("fx_rate"):
            fx=fx_to_base(net,v["fx_rate"])
            result["Kur Vergi Ithalat!B"]=float(fx)
            if enough("freight","insurance","duty_rate","vat_rate","vat_deductible"):
                from decimal import Decimal as D
                base=fx+D(str(v["freight"]))+D(str(v["insurance"]))
                duties=base*D(str(v["duty_rate"]))
                vat=vat_split(base+duties,v["vat_rate"],v["vat_deductible"])
                cash=vat["cash_gross"]; cost=vat["cost_basis"]
                result.update({"Kur Vergi Ithalat!C":float(base),"Kur Vergi Ithalat!D":float(duties+vat["non_deductible_vat"]),
                               "Kur Vergi Ithalat!E":float(cash),"Kur Vergi Ithalat!F":float(cost),"Odeme NBD Nakit!B":float(cash)})
                if enough("annual_rate","days"):
                    result["Odeme NBD Nakit!C"]=float(present_value_days(cash,int(v["days"]),v["annual_rate"]))
                if enough("energy_annual","maintenance_annual","years","residual"):
                    tco=cost+(D(str(v["energy_annual"]))+D(str(v["maintenance_annual"])))*D(str(v["years"]))-D(str(v["residual"]))
                    result["TCO Yasam Dongusu!B"]=float(tco)
                    result["Senaryo Karsilastirma!C"]=float(tco)
                    if enough("base_total"):
                        diff=tco-D(str(v["base_total"]))
                        result["Senaryo Karsilastirma!B"]=v["base_total"]
                        result["Senaryo Karsilastirma!D"]=float(diff)
                        if v["base_total"]!=0: result["Senaryo Karsilastirma!E"]=float(diff/D(str(v["base_total"])))
    if enough("technical_score","technical_weight","cost_score","cost_weight"):
        if abs(v["technical_weight"]+v["cost_weight"]-1)<0.000001:
            result["Puanlama Duyarlilik!B"]=v["technical_score"]*v["technical_weight"]+v["cost_score"]*v["cost_weight"]
            if v["technical_weight"]<=0.95:
                result["Puanlama Duyarlilik!C"]=v["technical_score"]*(v["technical_weight"]+0.05)+v["cost_score"]*(v["cost_weight"]-0.05)
    for target,source in [("Senaryo Kontrol!B","Miktar ve Kapsam!B"),
                          ("Senaryo Kontrol!C","Odeme NBD Nakit!C"),("Senaryo Kontrol!D","TCO Yasam Dongusu!B")]:
        if source in result: result[target]=result[source]
    return result

def validate_formulas(path: Path, contract):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False)
    errors=[]
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if any(token in name.casefold() for token in ["vbaproject","externallinks","macrosheets"]):
                errors.append("Forbidden workbook part: "+name)
    if wb.defined_names or getattr(wb,"_external_links",None): errors.append("Defined names/external links forbidden")
    observed={}
    for ws in wb:
        if ws.sheet_state!="visible": errors.append("Unexpected hidden sheet: "+ws.title)
        for row in ws:
            for cell in row:
                if cell.data_type=="f": observed[ws.title+"!"+cell.coordinate]=cell.value
    if observed!=contract["formulas"]: errors.append("Formula location/expression contract mismatch")
    for locator,expected in contract.get("source_cells",{}).items():
        sheet,cell=locator.rsplit("!",1)
        if sheet not in wb or source_cell_value(wb[sheet][cell].value)!=expected: errors.append("Immutable source changed: "+locator)
    for cell,key in contract["inputs"].items():
        try: _valid(key,wb["Senaryo Girdileri"][cell].value)
        except ValueError as exc: errors.append(str(exc))
    from finance_workbook import validate_extra_inputs
    errors.extend(validate_extra_inputs(wb,contract))
    if contract.get('scoring'):
        from scenario_scoring import read_scoring
        try: read_scoring(wb,contract)
        except (ValueError,KeyError,TypeError) as exc: errors.append(str(exc))
    wb.close()
    return {"status":"FAIL" if errors else "PASS","errors":errors,"formula_count":len(observed),
            "recalculation":"NOT_EVALUATED_BY_STRUCTURAL_CHECK"}
