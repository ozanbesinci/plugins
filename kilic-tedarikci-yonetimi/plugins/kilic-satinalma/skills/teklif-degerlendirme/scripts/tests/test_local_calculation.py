"""Real process faults and adversarial OOXML; these are not Excel equivalence QA."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile
from openpyxl import Workbook,load_workbook
from calculation_cache import transfer,preflight,tree,q,ns,stripped,validate_caches
from calculation_runtime import digest,validate,machine_id
from calculation_process import run_owned
import local_calculation

class CacheTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.source=self.root/'özgün.xlsx';self.calculated=self.root/'computed.xlsx';self.output=self.root/'result.xlsx'
        wb=Workbook();ws=wb.active;ws.title='Türkçe';ws['A1']=2;ws['A2']=True;ws['B1']='=IF(A2,A1*3,0)';ws['B2']='=IF(A2,"çalışır","")';ws['B1'].number_format='#,##0.00';wb.save(self.source);wb.close()
        self.produce()
    def tearDown(self):self.temp.cleanup()
    def produce(self,mutation=None):
        with zipfile.ZipFile(self.source) as src,zipfile.ZipFile(self.calculated,'w') as dst:
            for info in src.infolist():
                data=src.read(info.filename)
                if info.filename=='xl/worksheets/sheet1.xml':
                    root=tree(data)
                    for cell in root.findall('.//m:c',ns):
                        f=cell.find(q('f'))
                        if f is not None:
                            v=cell.find(q('v'));v.text='6' if cell.get('r')=='B1' else 'çalışır'
                            cell.set('t','n' if cell.get('r')=='B1' else 'str')
                            if mutation=='formula':f.text='42'
                            if mutation=='missing':cell.remove(v)
                            if mutation=='error':cell.set('t','e');v.text='#VALUE!'
                            if mutation=='nan':cell.set('t','n');v.text='NaN'
                        elif mutation=='literal' and cell.get('r')=='A1':cell.find(q('v')).text='99'
                    data=__import__('lxml').etree.tostring(root)
                dst.writestr(info,data)
    def test_only_caches_change_with_types_styles_and_original_formulas_preserved(self):
        before=digest(self.source);r=transfer(self.source,self.calculated,self.output)
        self.assertEqual(r['formulas'],2);self.assertEqual(before,digest(self.source))
        with zipfile.ZipFile(self.source) as a,zipfile.ZipFile(self.output) as b:
            for name in a.namelist():
                if name.startswith('xl/worksheets/'):self.assertEqual(stripped(a.read(name)),stripped(b.read(name)))
                else:self.assertEqual(a.read(name),b.read(name))
        wb=load_workbook(self.output,data_only=True)
        self.assertEqual(wb.active['B1'].value,6);self.assertEqual(wb.active['B2'].value,'çalışır');wb.close()
    def test_corrupt_results_never_create_output(self):
        for mutation in ('formula','missing','error','nan','literal'):
            with self.subTest(mutation=mutation):
                self.produce(mutation)
                with self.assertRaises((ValueError,TypeError)):transfer(self.source,self.calculated,self.output)
                self.assertFalse(self.output.exists())
    def test_preflight_rejects_external_and_unvalidated_functions_before_process(self):
        for formula in ('=WEBSERVICE("https://example.invalid")','=FORMULATEXT(B1)',"='[remote.xlsx]Sheet1'!A1"):
            wb=load_workbook(self.source);wb.active['B1']=formula;wb.save(self.source);wb.close()
            with self.assertRaises(ValueError):preflight(self.source,self.root/'cold.xlsx')
            self.assertFalse((self.root/'cold.xlsx').exists())
    def test_cold_copy_removes_all_old_results(self):
        cold=self.root/'cold.xlsx';self.assertEqual(preflight(self.calculated,cold),2)
        wb=load_workbook(cold,data_only=True);self.assertIsNone(wb.active['B1'].value);self.assertIsNone(wb.active['B2'].value);wb.close()
    def test_independent_values_reject_plausible_but_wrong_numeric_cache(self):
        contract={'formulas':{'Türkçe!B1':'=IF(A2,A1*3,0)','Türkçe!B2':'=IF(A2,"çalışır","")'},
                  'reference':{'1':{'Türkçe!B':6},'2':{'Türkçe!B':'çalışır'}}}
        self.assertEqual(validate_caches(self.calculated,contract)['independent_values_compared'],2)
        contract['reference']['1']['Türkçe!B']=7
        with self.assertRaisesRegex(ValueError,'Independent scenario cache mismatch'):
            validate_caches(self.calculated,contract)
    def test_delivery_rejects_missing_and_nonfinite_caches(self):
        contract={'formulas':{'Türkçe!B1':'formula','Türkçe!B2':'formula'}}
        for mutation in ('missing','nan','error'):
            with self.subTest(mutation=mutation):
                self.produce(mutation)
                with self.assertRaises(ValueError):validate_caches(self.calculated,contract)
    def test_failed_or_timed_out_engine_preserves_published_workbook(self):
        identity={'executable':sys.executable}
        before=digest(self.source)
        for outcome in ({'timed_out':True,'exit_code':3},{'timed_out':False,'exit_code':2}):
            with patch.object(local_calculation,'current',return_value=identity),patch.object(local_calculation,'run_owned',return_value=outcome):
                result=local_calculation.recalculate(self.source)
            self.assertFalse(result['available']);self.assertEqual(digest(self.source),before)
            self.assertEqual(list(self.root.glob('*.calc-*')),[])

class RegistryTests(unittest.TestCase):
    def test_changed_added_missing_and_foreign_engine_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            user=Path(temp);root=user/'KilicPython/calculation/test';root.mkdir(parents=True)
            file=root/'engine.bin';file.write_bytes(b'engine')
            manifest=root/'engine-manifest.json';manifest.write_text(json.dumps({'engine.bin':digest(file)}))
            expected={'id':'test','manifest_sha256':digest(manifest),'version':'test','sha256':'package'}
            record={'schema':'kilic-calculation-runtime/v1','profile_id':'test','machine_id':machine_id(),'root':str(root)}
            self.assertEqual(validate(record,expected,user_profile=user)['engine'],'LibreOffice')
            for change in ('changed','added','missing','foreign'):
                with self.subTest(change=change):
                    file.write_bytes(b'engine');extra=root/'extra.dll'
                    if extra.exists():extra.unlink()
                    current=dict(record)
                    if change=='changed':file.write_bytes(b'tampered')
                    if change=='added':extra.write_bytes(b'extra')
                    if change=='missing':file.unlink()
                    if change=='foreign':current['machine_id']='other'
                    with self.assertRaises((OSError,ValueError)):validate(current,expected,user_profile=user)

@unittest.skipUnless(os.name=='nt','Windows process ownership integration')
class OwnedProcessTests(unittest.TestCase):
    def test_timeout_kills_only_own_parent_and_child(self):
        import win32api,win32con,win32event,pywintypes
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);child=root/'child.py';parent=root/'parent.py';pid_file=root/'pid.txt'
            child.write_text('import time;time.sleep(60)')
            parent.write_text('import subprocess,sys,time\nfrom pathlib import Path\np=subprocess.Popen([sys.executable,sys.argv[1]])\nPath(sys.argv[2]).write_text(str(p.pid))\ntime.sleep(60)')
            unrelated=subprocess.Popen([sys.executable,str(child)],creationflags=subprocess.CREATE_NO_WINDOW)
            try:
                r=run_owned([sys.executable,str(parent),str(child),str(pid_file)],cwd=root,timeout_seconds=2)
                self.assertTrue(r['timed_out']);self.assertEqual(r['owned_process_cleanup'],'JOB_EMPTY');self.assertIsNone(unrelated.poll())
                pid=int(pid_file.read_text())
                try:handle=win32api.OpenProcess(win32con.SYNCHRONIZE,False,pid)
                except pywintypes.error:handle=None
                if handle:
                    try:self.assertEqual(win32event.WaitForSingleObject(handle,0),win32event.WAIT_OBJECT_0)
                    finally:handle.Close()
            finally:unrelated.kill();unrelated.wait(timeout=5)
    def test_invalid_deadline_never_starts_process(self):
        for value in (0,-1,181,True,float('nan')):
            with self.assertRaises(ValueError):run_owned([sys.executable],cwd=Path.cwd(),timeout_seconds=value)
