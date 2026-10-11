from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault('TEKLIF_TEST_SQLITE', '1')
SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from db import init_db, connect, rows_as_dict
from ingest import ingest_source


class MultiformatIngestTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix='teklif-v5-mf-'))
        self.src = self.root/'src'; self.run = self.root/'run'; self.src.mkdir()
        init_db(self.run)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _fixtures(self):
        from openpyxl import Workbook
        from docx import Document
        from PIL import Image, ImageDraw

        wb=Workbook(); ws=wb.active; ws.title='Teklif'; ws.append(['Fiyat Teklifi','Firma A']); ws.append(['Toplam','100000','TRY']); wb.save(self.src/'FirmaA_Teklif.xlsx')
        (self.src/'FirmaB_Teklif.csv').write_text('Fiyat Teklifi;Firma B\nToplam;110000\nPara Birimi;TRY\n',encoding='utf-8-sig')
        doc=Document(); doc.add_heading('Teknik Şartname',0); doc.add_paragraph('1. Kapasite en az 10 birim olacaktır.'); doc.save(self.src/'Teknik_Sartname.docx')
        img=Image.new('RGB',(500,220),'white'); draw=ImageDraw.Draw(img); draw.text((20,20),'Fiyat Teklifi - Firma C',fill='black'); draw.text((20,70),'Toplam: 120000 TRY',fill='black'); img.save(self.src/'FirmaC_Teklif.jpg')

    def test_formats_route_to_native_adapters(self):
        self._fixtures()
        inventory=ingest_source(self.src,self.run)
        families={x['format_family'] for x in inventory}
        self.assertTrue({'xlsx','csv','docx','image'}.issubset(families))
        by_name={x['filename']:x for x in inventory}
        self.assertEqual(by_name['FirmaA_Teklif.xlsx']['document_role'],'offer')
        self.assertEqual(by_name['Teknik_Sartname.docx']['document_role'],'specification')
        self.assertEqual(by_name['FirmaC_Teklif.jpg']['parse_status'],'partial')
        con=connect(self.run)
        try:
            self.assertGreater(con.execute("SELECT count(*) FROM document_units WHERE document_id=?",[by_name['FirmaA_Teklif.xlsx']['document_id']]).fetchone()[0],0)
            self.assertGreater(con.execute("SELECT count(*) FROM document_units WHERE document_id=?",[by_name['Teknik_Sartname.docx']['document_id']]).fetchone()[0],0)
        finally:
            con.close()


if __name__ == '__main__':
    unittest.main()
