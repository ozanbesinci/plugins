import copy
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from db import init_db, connect, get_meta
from scenario_finance import finance_reference, lifecycle_reference
from scenario_formulas import add_scenarios, validate_formulas
from matrix_workbook import seal_source_cells
from test_finance_workbook import VALUES
from test_scenario_finance import FINANCE

COSTS = [{'year':1,'energy':10,'maintenance':5,'other':0},
         {'year':2,'energy':20,'maintenance':6,'other':4},
         {'year':3,'energy':12,'maintenance':8,'other':5}]
YEARLY = {**FINANCE,'cost_mode':'YEARLY_SCHEDULE','annual_costs':COSTS}
YEAR_VALUES = {**VALUES,'years':3,'base_total':100}


class AnnualCostsTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);init_db(self.root)
    def tearDown(self):self.tmp.cleanup()
    def make(self,financial=YEARLY):
        wb=Workbook();wb.remove(wb.active)
        contract=add_scenarios(wb,[{'id':'Y1','inputs':copy.deepcopy(YEAR_VALUES),'financial':copy.deepcopy(financial)}])
        seal_source_cells(wb,contract);path=self.root/'annual.xlsx';wb.save(path);wb.close()
        return path,contract
    def test_yearly_costs_match_independent_golden_and_do_not_mutate_inputs(self):
        values=copy.deepcopy(YEAR_VALUES);financial=copy.deepcopy(YEARLY)
        result=finance_reference(values,financial,cash=120,cost=100)
        self.assertEqual(result['status'],'SCENARIO_CALCULATED')
        self.assertAlmostEqual(result['opex_present_value'],57.21262208865515)
        self.assertAlmostEqual(result['total_present_value'],149.8196844477836)
        self.assertEqual(result['lifecycle']['nominal'],70)
        self.assertEqual(values,YEAR_VALUES);self.assertEqual(financial,YEARLY)
    def test_missing_year_category_and_mode_never_become_zero(self):
        missing=copy.deepcopy(COSTS);missing[1]['other']=None
        for financial in ({**YEARLY,'annual_costs':COSTS[:2]},
                          {**YEARLY,'annual_costs':missing},{**YEARLY,'cost_mode':None}):
            result=finance_reference(YEAR_VALUES,financial,cash=120,cost=100)
            self.assertIsNone(result['opex_present_value']);self.assertIsNone(result['total_present_value'])
    def test_known_zero_costs_and_zero_rate_are_valid(self):
        zeros=[{'year':y,'energy':0,'maintenance':0,'other':0} for y in range(1,4)]
        result=finance_reference({**YEAR_VALUES,'annual_rate':0},{**YEARLY,'annual_costs':zeros},cash=120,cost=100)
        self.assertEqual(result['opex_present_value'],0);self.assertEqual(result['total_present_value'],95)
        self.assertEqual(finance_reference({**YEAR_VALUES,'annual_rate':0},YEARLY,cash=120,cost=100)['total_present_value'],165)
    def test_duplicate_outside_horizon_bad_type_and_ambiguous_mode_reject(self):
        for costs in (COSTS+[COSTS[0]],COSTS+[{'year':4,'energy':0,'maintenance':0,'other':0}],
                      [{'year':True,'energy':0,'maintenance':0,'other':0}],
                      [{'year':1,'energy':float('inf'),'maintenance':0,'other':0}],
                      [{'year':1,'energy':0,'maintenance':0,'other':-1}],
                      [{'year':1,'energy':0,'maintenance':0,'other':0,'forged':1}],False):
            with self.subTest(costs=costs),self.assertRaises(ValueError):
                lifecycle_reference(YEAR_VALUES,{**YEARLY,'annual_costs':costs})
        for mode in ('bad',True,[]):
            with self.subTest(mode=mode),self.assertRaises(ValueError):
                lifecycle_reference(YEAR_VALUES,{**YEARLY,'cost_mode':mode})
        with self.assertRaises(ValueError):lifecycle_reference(YEAR_VALUES,{**YEARLY,'cost_mode':'CONSTANT_ANNUAL'})
    def test_nominal_yearly_costs_do_not_require_discount_rate(self):
        result=lifecycle_reference({**YEAR_VALUES,'annual_rate':None},YEARLY)
        self.assertEqual(result['nominal'],70);self.assertIsNone(result['present_value'])
    def test_workbook_yearly_model_propagates_to_nominal_and_financial_scenarios(self):
        path,contract=self.make();ref=contract['reference']['2']
        self.assertEqual(validate_formulas(path,contract)['status'],'PASS')
        self.assertAlmostEqual(ref['Finansal TCO!F'],149.8196844477836)
        self.assertEqual(ref['TCO Yasam Dongusu!B'],165)
        self.assertEqual(ref['Senaryo Karsilastirma!D'],65)
        self.assertEqual(ref['Senaryo Kontrol!D'],165)
        wb=load_workbook(path)
        self.assertFalse(wb['Yillik Yasam Giderleri']['E2'].protection.locked)
        self.assertTrue(wb['Yillik Yasam Giderleri']['B2'].protection.locked)
        self.assertIn("'Finansal TCO'!H2",wb['TCO Yasam Dongusu']['B2'].value)
        self.assertTrue(all(len(formula)<8192 for formula in contract['formulas'].values()))
        wb.close()
    def test_user_yearly_cost_change_allowed_but_year_formula_and_outside_cost_rejected(self):
        path,contract=self.make();original=path.read_bytes()
        for sheet,cell,value,expected in [('Yillik Yasam Giderleri','E3',None,'PASS'),
                                          ('Yillik Yasam Giderleri','B2',2,'FAIL'),
                                          ('Yillik Yasam Giderleri','F2','=SUM(1,2)','FAIL'),
                                          ('Yillik Yasam Giderleri','C5',1,'FAIL'),
                                          ('Finansal Varsayimlar','K2','CONSTANT_ANNUAL','FAIL'),
                                          ('Finansal Varsayimlar','K2','bad','FAIL')]:
            path.write_bytes(original);wb=load_workbook(path);wb[sheet][cell]=value;wb.save(path);wb.close()
            self.assertEqual(validate_formulas(path,contract)['status'],expected,(sheet,cell,value))
    def test_multiple_scenarios_have_separate_cost_rows(self):
        wb=Workbook();wb.remove(wb.active)
        scenarios=[{'id':sid,'inputs':YEAR_VALUES,'financial':YEARLY} for sid in ('ONE','TWO')]
        contract=add_scenarios(wb,scenarios)
        self.assertEqual(wb['Yillik Yasam Giderleri']['A2'].value,'ONE')
        self.assertEqual(wb['Yillik Yasam Giderleri']['A102'].value,'TWO')
        self.assertIn('G102:G201',wb['Finansal TCO']['D3'].value)
        self.assertEqual(contract['reference']['2']['Finansal TCO!F'],contract['reference']['3']['Finansal TCO!F'])
        wb.close()
    def test_import_preserves_source_and_pdf_shows_unknown_year_cost(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary
        from scenario_import import import_scenarios
        run=self.root/'run';seed(run)
        before=connect(run)
        try:facts=before.execute('SELECT * FROM facts ORDER BY fact_id').fetchall()
        finally:before.close()
        initial=export_preliminary(run,'Synthetic yearly scenario',[{'id':'Y1','inputs':YEAR_VALUES,'financial':YEARLY}])
        self.assertEqual(initial['delivery_status'],'READY',initial)
        user=self.root/'user.xlsx';shutil.copy2(initial['xlsx'],user)
        wb=load_workbook(user);wb['Yillik Yasam Giderleri']['E3']=None;wb.save(user);wb.close()
        report=import_scenarios(run,user);self.assertEqual(report['scenario_import'],'IMPORTED',report)
        model=json.loads((Path(report['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
        self.assertIsNone(model['scenario_assumptions']['2']['financial']['annual_costs'][1]['other'])
        self.assertNotIn('Finansal TCO!F',model['scenario_reference']['2'])
        self.assertNotIn('TCO Yasam Dongusu!B',model['scenario_reference']['2'])
        after=connect(run)
        try:self.assertEqual(facts,after.execute('SELECT * FROM facts ORDER BY fact_id').fetchall())
        finally:after.close()
        import fitz
        with fitz.open(report['pdf']) as pdf:
            text=' '.join(page.get_text() for page in pdf)
        self.assertIn('Yıl 2 giderleri',text);self.assertIn('diğer: Bilinmiyor',text)
        self.assertIn('Hesaplanamadı',text);self.assertEqual(report['decision_assurance'],'NO_DECISION')
    def test_import_rejects_outside_horizon_without_replacing_previous_pair(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary
        from scenario_import import import_scenarios
        run=self.root/'run';seed(run)
        initial=export_preliminary(run,'Synthetic yearly scenario',[{'id':'Y1','inputs':YEAR_VALUES,'financial':YEARLY}])
        user=self.root/'user.xlsx';shutil.copy2(initial['xlsx'],user)
        wb=load_workbook(user);wb['Yillik Yasam Giderleri']['C5']=1;wb.save(user);wb.close()
        with self.assertRaises(ValueError):import_scenarios(run,user)
        self.assertEqual(get_meta(run,'preliminary_report_delivery',parse_json=True)['revision'],initial['revision'])
    def test_independent_pdf_fallback_preserves_missing_annual_category(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary
        from unittest.mock import patch
        import fitz
        costs=copy.deepcopy(COSTS);costs[1]['other']=None
        run=self.root/'run';seed(run)
        with patch('report_delivery.render_pdf',side_effect=RuntimeError('Synthetic PDF renderer failure')):
            report=export_preliminary(run,'Synthetic fallback',[{'id':'Y1','inputs':YEAR_VALUES,'financial':{**YEARLY,'annual_costs':costs}}])
        self.assertEqual(report['delivery_status'],'READY',report)
        with fitz.open(report['pdf']) as pdf:text=' '.join(page.get_text() for page in pdf)
        self.assertIn('Yıl 2 giderleri',text);self.assertIn('diğer: Bilinmiyor',text)
        self.assertIn('Hesaplanamadı',text)
    def test_clearing_model_preserves_assumptions_without_any_lifecycle_total(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary
        from scenario_import import import_scenarios
        run=self.root/'run';seed(run)
        initial=export_preliminary(run,'Synthetic yearly scenario',[{'id':'Y1','inputs':YEAR_VALUES,'financial':YEARLY}])
        user=self.root/'user.xlsx';shutil.copy2(initial['xlsx'],user)
        wb=load_workbook(user);wb['Finansal Varsayimlar']['K2']=None;wb.save(user);wb.close()
        report=import_scenarios(run,user);self.assertEqual(report['scenario_import'],'IMPORTED',report)
        model=json.loads((Path(report['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
        self.assertIsNone(model['scenario_assumptions']['2']['financial']['cost_mode'])
        self.assertEqual(model['scenario_assumptions']['2']['financial']['annual_costs'],COSTS)
        for key in ('Finansal TCO!F','Finansal TCO!H','TCO Yasam Dongusu!B','Senaryo Kontrol!D'):
            self.assertNotIn(key,model['scenario_reference']['2'])
