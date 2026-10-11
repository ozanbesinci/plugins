import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from db import init_db
from dated_cashflows import CATEGORIES,METHOD,SCOPE,ledger_reference
from dated_workbook import SUMMARY,DETAIL,read_ledger
from scenario_formulas import add_scenarios,validate_formulas
from matrix_workbook import seal_source_cells
from test_finance_workbook import VALUES
from test_scenario_finance import FINANCE

FLOWS=[{'cost_id':kind,'category':kind,'amount':0,'currency':'TRY','cash_date':'2026-10-11',
        'rate_date':'2026-10-11','fx_rate':1,'rate_source':'Explicit synthetic zero/exclusion'} for kind in CATEGORIES]
FLOWS[0].update(amount=100,currency='EUR',cash_date='2027-10-11',rate_date='2027-10-11',fx_rate=40,rate_source='User future cash-date assumption')
FLOWS[1].update(amount=100)
FLOWS[4].update(amount=200)
FLOWS[5].update(amount=200,cash_date='2027-10-11',rate_date='2027-10-11')
FLOWS[-1].update(amount=50,cash_date='2027-10-11',rate_date='2027-10-11')
LEDGER={'valuation_date':'2026-10-11','base_currency':'TRY','annual_rate':0.1,'baseline':3000,'scope':SCOPE,'method':METHOD,'flows':FLOWS}

class DatedCashflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);init_db(self.root)
    def tearDown(self):self.tmp.cleanup()
    def make(self,ledger=LEDGER):
        wb=Workbook();wb.remove(wb.active)
        contract=add_scenarios(wb,[{'id':'D','inputs':copy.deepcopy(VALUES),'financial':copy.deepcopy(FINANCE),'dated_costs':copy.deepcopy(ledger)}]);seal_source_cells(wb,contract)
        path=self.root/'ledger.xlsx';wb.save(path);wb.close();return path,contract
    def test_independent_golden_signed_cash_projected_fx_and_no_automatic_double_count(self):
        snapshot=copy.deepcopy(LEDGER);result=ledger_reference(LEDGER)
        self.assertEqual(result['nominal'],4050)
        self.assertAlmostEqual(result['present_value'],3709.090909090909)
        self.assertAlmostEqual(result['difference'],709.090909090909)
        self.assertEqual(result['rows'][0]['days'],365);self.assertEqual(LEDGER,snapshot)
        path,contract=self.make();self.assertAlmostEqual(contract['reference']['2'][SUMMARY+'!H'],3709.090909090909)
        self.assertAlmostEqual(contract['reference']['2']['Finansal TCO!F'],126.94214876033058)
        self.assertEqual(validate_formulas(path,contract)['status'],'PASS')
    def test_missing_kind_amount_date_rate_and_model_never_make_complete_total(self):
        for key in ('amount','cash_date','rate_date','fx_rate','rate_source','cost_id','category','currency'):
            ledger=copy.deepcopy(LEDGER);ledger['flows'][0][key]=None
            with self.subTest(key=key):self.assertIsNone(ledger_reference(ledger)['present_value'])
        for key in ('valuation_date','base_currency','annual_rate','scope','method'):
            ledger=copy.deepcopy(LEDGER);ledger[key]=None
            self.assertIsNone(ledger_reference(ledger)['present_value'],key)
        ledger=copy.deepcopy(LEDGER);ledger['flows'].pop(1);self.assertIsNone(ledger_reference(ledger)['nominal'])
    def test_zero_cash_zero_discount_and_zero_baseline_preserved(self):
        ledger=copy.deepcopy(LEDGER);ledger.update(annual_rate=0,baseline=0)
        for item in ledger['flows']:item['amount']=0
        result=ledger_reference(ledger);self.assertEqual(result['present_value'],0);self.assertIsNone(result['relative_difference'])
    def test_duplicate_identity_wrong_date_direction_and_unsupported_inputs_reject(self):
        for key,value in (('cost_id','FREIGHT'),('cash_date','2026-10-10'),('rate_date','2026-10-11'),('fx_rate',-1),('amount',True),('currency','eur'),('amount',float('inf')),('category',[])):
            ledger=copy.deepcopy(LEDGER);ledger['flows'][0][key]=value
            with self.subTest(key=key,value=value),self.assertRaises(ValueError):ledger_reference(ledger)
        ledger=copy.deepcopy(LEDGER);ledger['flows'][1]['fx_rate']=2
        with self.assertRaises(ValueError):ledger_reference(ledger)
        with self.assertRaises(ValueError):ledger_reference({**LEDGER,'flows':FLOWS*5})
    def test_actual_calendar_dates_ignore_fictitious_excel_leap_day(self):
        ledger=copy.deepcopy(LEDGER);ledger.update(valuation_date='1900-01-01')
        for item in ledger['flows']:item.update(cash_date='1900-03-01',rate_date='1900-03-01')
        self.assertEqual(ledger_reference(ledger)['rows'][0]['days'],59)
    def test_cost_ids_are_literal_case_sensitive_and_empty_typed_ledgers_reject(self):
        ledger=copy.deepcopy(LEDGER)
        ledger['flows'][0]['cost_id']='Cost*?~'
        ledger['flows'][1]['cost_id']='cost*?~'
        self.assertEqual(ledger_reference(ledger)['nominal'],4050)
        path,contract=self.make(ledger)
        self.assertEqual(validate_formulas(path,contract)['status'],'PASS')
        for invalid in (False,[],0,''):
            with self.subTest(value=invalid),self.assertRaises(ValueError):self.make(invalid)
    def test_editable_inputs_round_trip_but_cost_identity_dup_and_formulas_reject(self):
        path,contract=self.make();wb=load_workbook(path)
        self.assertEqual(read_ledger(wb,2),LEDGER);self.assertFalse(wb[DETAIL]['I2'].protection.locked)
        wb[DETAIL]['I2']=41;wb.save(path);wb.close()
        self.assertEqual(validate_formulas(path,contract)['status'],'PASS')
        for cell,value in (('C2','FREIGHT'),('C14','EXTRA'),('M2','=1'),('B2',9)):
            path,contract=self.make();wb=load_workbook(path);wb[DETAIL][cell]=value;wb.save(path);wb.close()
            self.assertEqual(validate_formulas(path,contract)['status'],'FAIL',(cell,value))
    def test_formulas_stay_bounded_and_missing_ledger_never_invents_totals(self):
        path,contract=self.make(None)
        self.assertNotIn(SUMMARY+'!H',contract['reference']['2'])
        self.assertTrue(all(len(expr)<=8192 for expr in contract['formulas'].values()))
        self.assertEqual(validate_formulas(path,contract)['status'],'PASS')
    def test_two_scenario_ledgers_have_disjoint_row_ranges(self):
        wb=Workbook();wb.remove(wb.active)
        contract=add_scenarios(wb,[{'id':key,'dated_costs':copy.deepcopy(LEDGER)} for key in ('A','B')]);seal_source_cells(wb,contract)
        self.assertEqual(wb[DETAIL]['A2'].value,'A');self.assertEqual(wb[DETAIL]['A50'].value,'B')
        self.assertEqual(read_ledger(wb,2),LEDGER);self.assertEqual(read_ledger(wb,3),LEDGER);wb.close()
    def test_pdf_shared_assumptions_show_dates_fx_missing_categories_and_detail(self):
        from pdf_renderer import assumption_rows,scenario_rows
        text=str(assumption_rows({'dated_costs':LEDGER}));self.assertIn('2027-10-11',text);self.assertIn('Yenileme',text)
        self.assertIn('40',text);self.assertIn('Bağımsız',str(scenario_rows({})))
    def test_import_updates_ledger_in_new_pair_and_preserves_source_and_decision(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary,inspect_preliminary
        from scenario_import import import_scenarios
        from db import connect,rows_as_dict
        seed(self.root)
        scenario={'id':'D','inputs':copy.deepcopy(VALUES),'dated_costs':copy.deepcopy(LEDGER)}
        first=export_preliminary(self.root,'Synthetic dated ledger',[scenario]);self.assertEqual(first['delivery_status'],'READY')
        con=connect(self.root)
        try:before=rows_as_dict(con,'SELECT * FROM facts ORDER BY fact_id')
        finally:con.close()
        path=self.root/'edited.xlsx';wb=load_workbook(first['xlsx']);wb[DETAIL]['I2']=41;wb.save(path);wb.close()
        result=import_scenarios(self.root,path);self.assertEqual(result['scenario_import'],'IMPORTED');self.assertNotEqual(result['xlsx'],first['xlsx'])
        con=connect(self.root)
        try:after=rows_as_dict(con,'SELECT * FROM facts ORDER BY fact_id')
        finally:con.close()
        self.assertEqual(before,after)
        import json
        model=json.loads((Path(result['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
        self.assertEqual(model['scenario_assumptions']['2']['dated_costs']['flows'][0]['fx_rate'],41)
        self.assertAlmostEqual(model['scenario_reference']['2'][SUMMARY+'!H'],3800)
        self.assertEqual(result['decision_assurance'],'NO_DECISION')
        wb=load_workbook(path);wb[DETAIL]['C2']='FREIGHT';wb.save(path);wb.close()
        with self.assertRaises(ValueError):import_scenarios(self.root,path)
        self.assertEqual(inspect_preliminary(self.root)['xlsx'],result['xlsx'])
    def test_real_pdf_primary_and_alternative_show_dated_ledger_and_unknowns(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary
        from unittest.mock import patch
        import fitz
        seed(self.root);ledger=copy.deepcopy(LEDGER);ledger['flows'][0]['fx_rate']=None
        scenario={'id':'D','dated_costs':ledger}
        for alternative in (False,True):
            context=patch('report_delivery.render_pdf',side_effect=RuntimeError('Synthetic PDF fault')) if alternative else patch('report_delivery.render_pdf',wraps=__import__('pdf_renderer').render_pdf)
            with context:result=export_preliminary(self.root,'Synthetic dated unknown',[scenario])
            self.assertEqual(result['delivery_status'],'READY')
            with fitz.open(result['pdf']) as pdf:
                text=' '.join(page.get_text() for page in pdf)
                self.assertGreaterEqual(len(pdf),2);self.assertIn('2027-10-11',text);self.assertIn('Yenileme',text);self.assertIn('Hesaplanamadı',text)
