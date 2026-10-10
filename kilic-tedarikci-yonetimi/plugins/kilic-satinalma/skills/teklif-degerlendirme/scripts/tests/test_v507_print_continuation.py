from pathlib import Path
import shutil, sys, tempfile, unittest

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook
from display_values import structured
from print_layout import MAX_LINES, build_print_workbook, pack_groups


class ReadableNumberTests(unittest.TestCase):
    def test_precision_artefacts_are_rounded_with_explicit_marker(self):
        self.assertEqual(structured('{"delta":"-1632.2597716800010316681073"}'), "Düzeltme tutarı: ≈\u00a0-1632.26")
        self.assertEqual(structured('{"ratio":"0.94123456789"}'), "Pay: ≈\u00a00.941235")
        self.assertEqual(structured('{"amount":103367.74022831999}'), "Tutar: ≈\u00a0103367.74")

    def test_short_numbers_and_text_are_unchanged(self):
        self.assertEqual(structured('{"amount":"105000.0","ratio":"1.0","days":60}'), "Tutar: 105000.0\nPay: 1.0\nGün: 60")
        self.assertEqual(structured('{"source_id":"off_3d0b4f5791489aac4c/field/total"}'), "Kaynak kimliği: off_3d0b4f5791489aac4c/field/total")
        self.assertEqual(structured('{"value_number":115000}'), "Sayısal değer: 115000")


class ContinuationTests(unittest.TestCase):
    def test_pack_groups_never_splits_a_logical_line_that_fits(self):
        groups = [["x"]] * (MAX_LINES - 1) + [["Kaynak kimliği:", "off_dd309138a5f3b4757c/field/total"]] + [["y"]]
        chunks = pack_groups(groups)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[1][:2], ["Kaynak kimliği:", "off_dd309138a5f3b4757c/field/total"])
        self.assertTrue(all(len(c) <= MAX_LINES for c in chunks))
        long = pack_groups([["z"] * (MAX_LINES * 2 + 3)])
        self.assertEqual([len(c) for c in long], [MAX_LINES, MAX_LINES, 3])

    def test_print_keeps_label_with_value_and_marks_continuation(self):
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook(); ws = wb.active; ws.title = "Iddia Kaydi"
            ws.append(["Claim ID", "Assertions"])
            parts = []
            for i in range(14):
                parts += ["Kaynak türü: Veri", f"Kaynak kimliği: off_{i:02d}3d0b4f5791489aac4c/field/total", "Sayısal değer: 120000.0"]
            ws.append(["claim_ff1147e0093e1d541d5514ed", "\n".join(parts)])
            ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 32
            wb.save(root / "in.xlsx"); wb.close()
            build_print_workbook(root / "in.xlsx", root / "print.xlsx")
            p = load_workbook(root / "print.xlsx"); sh = p.active
            rows = [(sh.cell(r, 1).value, sh.cell(r, 2).value) for r in range(2, sh.max_row + 1)]
            self.assertGreater(len(rows), 1)
            self.assertNotIn("(devam)", rows[0][0])
            for ident, _ in rows[1:]:
                self.assertTrue(ident.endswith("(devam)"), ident)
            for _, text in rows:
                self.assertFalse(text.split("\n")[-1].endswith(":"), "label separated from its value")
            joined = "\n".join(t for _, t in rows)
            self.assertEqual(joined.count("Sayısal değer: 120000.0"), 14)
            p.close()
        finally:
            shutil.rmtree(root)


if __name__ == "__main__":
    unittest.main()
