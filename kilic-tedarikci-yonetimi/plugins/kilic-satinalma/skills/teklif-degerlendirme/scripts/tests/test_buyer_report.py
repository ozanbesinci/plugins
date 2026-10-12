"""Synthetic counterparties and amounts only. The reviewed buyer declaration must remain visible and calculation-bound in exports."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from print_layout import build_print_workbook
from report_export import export_report


class BuyerReportTests(unittest.TestCase):
    def test_xlsx_and_pdf_print_source_show_buyer_inputs_without_rewriting_facts(self):
        declarations = {
            "currency": [{"kind": "currency", "offer_id": "off-synthetic-alpha", "value": "TRY",
                          "answer_text": "Evet, üçü de TL", "issue_id": "issue-currency"}],
            "reference_quantity": [{"kind": "reference_quantity", "offer_id": "off-synthetic-alpha",
                                    "quantity_item_id": "item-concrete", "quantity": "25", "unit": "m3",
                                    "answer_text": "Esas miktar 25 m3", "issue_id": "issue-quantity"}],
            "scope_cost": [{"kind": "scope_cost", "offer_id": "off-synthetic-alpha", "scope_id": "scope-water",
                            "amount": "0", "currency": "TRY",
                            "answer_text": "Mevcut imkânlarla, ilave maliyet yok", "issue_id": "issue-scope"}],
            "payment_terms": [{"kind": "payment_terms", "offer_id": "off-synthetic-alpha",
                               "milestones": [{"trigger": "work_start", "percent": 50},
                                              {"trigger": "work_delivery", "percent": 50}],
                               "answer_text": "Yarısı başlangıçta, yarısı teslimde", "issue_id": "issue-payment"}],
            "warranty_terms": [{"kind": "warranty_terms", "offer_id": "off-synthetic-alpha",
                                "duration_years": 5, "scope": "whole_work",
                                "answer_text": "İşin tamamına beş yıl", "issue_id": "issue-warranty"}],
        }
        result = {"currency": "TRY", "currency_basis": "buyer_declaration",
                  "buyer_quantity_adjustment": "-100", "buyer_quantity_rows": [],
                  "buyer_inputs": declarations, "known_scope_adjustment": "0",
                  "acquisition_cost_basis": "900", "comparable_total": "900"}
        snapshot = {"status": "PRELIMINARY", "meta": {"profile": "standart", "decision_status": "SUFFICIENT"},
                    "suppliers": [{"supplier_id": "supplier-synthetic-alpha", "display_name": "Sentetik Tedarikçi Alfa"}],
                    "offers": [{"offer_id": "off-synthetic-alpha", "supplier_id": "supplier-synthetic-alpha", "status": "active"}],
                    "facts": [{"fact_id": "fact-total", "offer_id": "off-synthetic-alpha", "field_key": "total",
                               "value_number": 1000, "evidence_status": "verified"},
                              {"fact_id": "fact-currency", "offer_id": "off-synthetic-alpha", "field_key": "currency",
                               "value_text": None, "evidence_status": "missing"}],
                    "calculations": [{"calculation_id": "calc-synthetic-alpha", "offer_id": "off-synthetic-alpha",
                                      "calculation_type": "cost_summary", "result_json": json.dumps(result)}]}

        from test_engine_support import ReportFixtureConnection
        class Connection(ReportFixtureConnection):
            def __init__(self): super().__init__("snapshot-sha")

        bindings = []
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp)
            with mock.patch("reviewed_snapshot.load_latest", return_value=(snapshot, {"sha256": "snapshot-sha", "relative_path": "reviewed/test.json"})), \
                 mock.patch("report_export.connect", return_value=Connection()), \
                 mock.patch("report_export._artifact", return_value="artifact-test"), \
                 mock.patch("report_export._insert_bindings", side_effect=lambda con, aid, sha, rows: bindings.extend(rows)), \
                 mock.patch("report_export._recalc_with_excel", return_value={"available": False}):
                exported = export_report(run)
            book = load_workbook(exported["xlsx"], data_only=True)
            try:
                summary = {row[0]: row[1] for row in book["Ozet"].values}
                self.assertEqual(summary["Karar İçin Bilgi Yeterliliği"], "Karar için yeterli")
                self.assertEqual(summary["İncelenen Kayıt Durumu"], "Ön değerlendirme")
                comparison = book["Teklif Karsilastirma"]
                self.assertIsNone(comparison["D2"].value)
                self.assertIn("fact-currency", comparison["D2"].comment.text)
                self.assertIn("snapshot-sha", comparison["D2"].comment.text)
                self.assertIn("Kaynak kanıtı yok", comparison["D2"].comment.text)
                self.assertEqual(comparison["T2"].value, "TRY")
                self.assertEqual(comparison["U2"].value, "Alıcı beyanı")
                self.assertEqual(comparison["V2"].value, -100)
                rows = list(book["Alici Beyanlari"].values)
                self.assertEqual(len(rows), 6)
                self.assertEqual([r[1] for r in rows[1:]], ["Para birimi", "Esas alınan miktar", "Kapsam maliyeti", "Ödeme koşulları", "Garanti koşulları"])
                self.assertEqual([r[3] for r in rows[1:3]], ["TRY", 25])
                self.assertEqual(rows[3][3], 0)
                self.assertIn("%50", rows[4][3])
                self.assertEqual(rows[5][3], "5 yıl, işin tamamı")
                self.assertEqual([r[6] for r in rows[1:]], ["issue-currency", "issue-quantity", "issue-scope", "issue-payment", "issue-warranty"])
                self.assertTrue(all(r[7] == "calc-synthetic-alpha" for r in rows[1:]))
                self.assertEqual(rows[3][5], "Mevcut imkânlarla, ilave maliyet yok")
                self.assertEqual(snapshot["facts"][1]["value_text"], None)
            finally:
                book.close()
            self.assertTrue(any(b["component_id"] == "xlsx:Teklif Karsilastirma:V2" and
                                b["calculation_ids"] == ["calc-synthetic-alpha"] for b in bindings))
            self.assertTrue(any(b["component_id"] == "xlsx:Alici Beyanlari:D3" and
                                b["locator"]["value_path"] == "buyer_inputs.reference_quantity.0.quantity" and
                                b["calculation_ids"] == ["calc-synthetic-alpha"] for b in bindings))
            printed = run / "output" / "print.xlsx"
            info = build_print_workbook(Path(exported["xlsx"]), printed)
            print_book = load_workbook(printed, data_only=True)
            try:
                self.assertTrue(any(m["source_sheet"] == "Alici Beyanlari" for m in info["mapping"]))
                values = "\n".join(str(cell.value or "") for sheet in print_book for row in sheet for cell in row)
                self.assertIn("Mevcutimkânlarla,ilavemaliyetyok", "".join(values.split()))
                self.assertIn("issue-scope", values)
            finally:
                print_book.close()


if __name__ == "__main__":
    unittest.main()
