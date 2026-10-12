"""Confirmed exclusions determine competition; source inventory remains intact."""
import json, tempfile, unittest
from pathlib import Path
import fitz
from openpyxl import load_workbook
from db import connect, get_meta
from report_model import build_preliminary
from report_delivery import export_preliminary
from pdf_renderer import label
from test_report_slice import seed

class EligibleSupplierComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def add_other_supplier(self,run,status='confirmed',add_alternative=False):
        con=connect(run)
        try:
            con.execute("INSERT INTO suppliers(supplier_id,display_name) VALUES ('T','Elenecek sentetik firma')")
            con.execute("INSERT INTO offers(offer_id,supplier_id,status) VALUES ('T1','T','active')")
            if add_alternative:
                con.execute("INSERT INTO offers(offer_id,supplier_id,status) VALUES ('O2','S','active')")
            con.execute("INSERT INTO exclusions(exclusion_id,offer_id,decision,reason,status) VALUES (?,?,?,?,?)",['X','T1','exclude','Teyitli kapsam dışı',status])
            con.commit()
        finally: con.close()
    def test_confirmed_exclusion_leaves_one_supplier_in_both_outputs_all_profiles(self):
        for profile in ['hizli','standart','yuksek_guvence']:
            with self.subTest(profile=profile):
                run=self.root/profile;seed(run,profile);self.add_other_supplier(run)
                model=build_preliminary(run,'Teyitli eleme sonrası tek tedarikçi kaldı')
                self.assertEqual(model['comparison_status'],'NOT_APPLICABLE_SINGLE_SUPPLIER')
                self.assertEqual(len(model['offers']),2);self.assertEqual(len(model['suppliers']),2)
                self.assertEqual(model['ranking'],[]);self.assertIsNone(model['recommendation'])
                result=export_preliminary(run,'Teyitli eleme sonrası tek tedarikçi kaldı')
                self.assertEqual(result['delivery_status'],'READY');self.assertEqual(result['decision_assurance'],'NO_DECISION')
                text=label('NOT_APPLICABLE_SINGLE_SUPPLIER')
                with fitz.open(result['pdf'])as pdf:
                    self.assertGreaterEqual(len(pdf),2);self.assertIn(text,'\n'.join(p.get_text()for p in pdf))
                wb=load_workbook(result['xlsx']);self.assertGreaterEqual(len(wb.sheetnames),20)
                self.assertIn(text,[cell.value for row in wb['Ozet']for cell in row]);wb.close()
                frozen=json.loads((Path(result['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
                self.assertEqual(len(frozen['offers']),2);self.assertEqual(frozen['comparison_status'],'NOT_APPLICABLE_SINGLE_SUPPLIER')
                self.assertIsNone(get_meta(run,'reviewed_snapshot'))
    def test_two_eligible_alternatives_are_one_supplier_after_other_supplier_excluded(self):
        run=self.root/'alternatives';seed(run);self.add_other_supplier(run,add_alternative=True)
        model=build_preliminary(run,'Tek firma iki alternatif')
        self.assertEqual(model['comparison_status'],'NOT_APPLICABLE_SINGLE_SUPPLIER')
        self.assertEqual(len(model['offers']),3);self.assertEqual(model['ranking'],[])
    def test_provisional_exclusion_does_not_remove_second_supplier(self):
        run=self.root/'pending';seed(run);self.add_other_supplier(run,status='proposed')
        model=build_preliminary(run,'Teyitsiz eleme')
        self.assertEqual(model['comparison_status'],'BLOCKED')
        self.assertEqual(len(model['offers']),2);self.assertEqual(model['decision_assurance'],'NO_DECISION')
