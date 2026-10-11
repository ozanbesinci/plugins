import copy
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook,load_workbook
from db import init_db
from scenario_formulas import add_scenarios,validate_formulas
from matrix_workbook import seal_source_cells
from test_scenario_finance import FINANCE

VALUES={"quantity":1,"unit_price":100,"discount":0,"fx_rate":1,"vat_rate":0.2,"vat_deductible":1,
        "freight":0,"insurance":0,"duty_rate":0,"annual_rate":0.1,"days":365,"energy_annual":10,
        "maintenance_annual":10,"years":2,"residual":5}
class FinanceWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);init_db(self.root);self.path=self.root/'finance.xlsx'
    def tearDown(self):self.tmp.cleanup()
    def make(self,financial=None):
        wb=Workbook();wb.remove(wb.active)
        scenario={'id':'S1','inputs':copy.deepcopy(VALUES)}
        if financial is not None:scenario['financial']=copy.deepcopy(financial)
        contract=add_scenarios(wb,[scenario]);seal_source_cells(wb,contract)
        wb.save(self.path);wb.close();return contract
    def test_four_financial_groups_and_independent_reference(self):
        contract=self.make(FINANCE)
        self.assertEqual(validate_formulas(self.path,contract)['status'],'PASS')
        self.assertAlmostEqual(contract['reference']['2']['Finansal TCO!F'],126.94214876033058)
        wb=load_workbook(self.path)
        self.assertEqual(wb['Finansal Varsayimlar']['B2'].value,2)
        self.assertFalse(wb['Odeme Plani']['C2'].protection.locked)
        self.assertTrue(wb['Odeme Plani']['E2'].protection.locked)
        self.assertIn('COUNT(B2:E2)=4',wb['Finansal TCO']['F2'].value)
        wb.close()
    def test_unknown_policy_has_no_financial_total_reference(self):
        contract=self.make()
        self.assertNotIn('Finansal TCO!F',contract['reference']['2'])
        self.assertEqual(validate_formulas(self.path,contract)['status'],'PASS')
    def test_bad_percentage_outside_count_and_formula_reject(self):
        for cell,value in [('C2',0.4),('C4',0.1),('F2','=SUM(1,2)')]:
            contract=self.make(FINANCE);wb=load_workbook(self.path);wb['Odeme Plani'][cell]=value;wb.save(self.path);wb.close()
            self.assertEqual(validate_formulas(self.path,contract)['status'],'FAIL',(cell,value))
    def test_unknown_date_reverse_rate_and_boolean_input_reject(self):
        for cell,value in [('G2','2026-02-31'),('I2','SOURCE_PER_TARGET'),('B2',True)]:
            contract=self.make(FINANCE);wb=load_workbook(self.path);wb['Finansal Varsayimlar'][cell]=value;wb.save(self.path);wb.close()
            self.assertEqual(validate_formulas(self.path,contract)['status'],'FAIL',(cell,value))
    def test_payment_or_tax_days_cleared_are_valid_unknown_inputs(self):
        for sheet,cell in [('Odeme Plani','D2'),('Finansal Varsayimlar','C2')]:
            contract=self.make(FINANCE);wb=load_workbook(self.path);wb[sheet][cell]=None;wb.save(self.path);wb.close()
            self.assertEqual(validate_formulas(self.path,contract)['status'],'PASS')
