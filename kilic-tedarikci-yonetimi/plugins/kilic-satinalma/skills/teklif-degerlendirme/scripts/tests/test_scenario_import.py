import copy,json,shutil
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from db import connect,get_meta,set_state
from report_delivery import export_preliminary
from scenario_import import import_scenarios
from test_report_slice import seed
from test_finance_workbook import VALUES
from test_scenario_finance import FINANCE

class ScenarioImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.run=self.root/'run';seed(self.run)
        self.initial=export_preliminary(self.run,'Initial synthetic report',[{'id':'S1','inputs':copy.deepcopy(VALUES),'financial':copy.deepcopy(FINANCE)}])
        self.book=self.root/'user-copy.xlsx';shutil.copy2(self.initial['xlsx'],self.book)
    def tearDown(self):self.tmp.cleanup()
    def facts(self):
        con=connect(self.run)
        try:return con.execute('SELECT * FROM facts ORDER BY fact_id').fetchall()
        finally:con.close()
    def test_user_cells_create_new_pair_without_rewriting_facts_or_decision(self):
        before=self.facts();wb=load_workbook(self.book);wb['Senaryo Girdileri']['B2']=2;wb.save(self.book);wb.close()
        result=import_scenarios(self.run,self.book)
        self.assertEqual(result['scenario_import'],'IMPORTED',result)
        self.assertEqual(result['decision_assurance'],'NO_DECISION')
        self.assertNotEqual(result['revision'],self.initial['revision']);self.assertEqual(self.facts(),before)
        self.assertIsNone(get_meta(self.run,'decision_summary'))
        model=json.loads((Path(result['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
        self.assertEqual(model['scenario_assumptions']['2']['inputs']['quantity'],2)
        self.assertTrue(Path(self.initial['xlsx']).is_file());self.assertTrue(Path(result['pdf']).is_file())
    def test_formula_and_source_mutation_are_rejected_without_new_pointer(self):
        for sheet,cell,value in [('Kaynak Olgulari','D2','forged source'),('Finansal TCO','F2','=SUM(1,2)')]:
            shutil.copy2(self.initial['xlsx'],self.book);wb=load_workbook(self.book);wb[sheet][cell]=value;wb.save(self.book);wb.close()
            with self.assertRaises(ValueError):import_scenarios(self.run,self.book)
            self.assertEqual(get_meta(self.run,'preliminary_report_delivery',parse_json=True)['revision'],self.initial['revision'])
    def test_closed_run_rejects_import(self):
        set_state(self.run,'COMPLETE','PASS',{'scope':'synthetic closed fixture'})
        with self.assertRaisesRegex(ValueError,'Completed run preserved'):import_scenarios(self.run,self.book)
    def test_changed_source_revision_rejects_stale_user_workbook(self):
        con=connect(self.run)
        try:
            con.execute("UPDATE facts SET value_text='Changed source basis' WHERE fact_id=(SELECT fact_id FROM facts LIMIT 1)")
            con.commit()
        finally:con.close()
        with self.assertRaisesRegex(ValueError,'Current preliminary source revision is not trusted'):
            import_scenarios(self.run,self.book)
        self.assertEqual(get_meta(self.run,'preliminary_report_delivery',parse_json=True)['revision'],self.initial['revision'])
    def test_missing_payment_input_remains_unknown_in_new_report(self):
        wb=load_workbook(self.book);wb['Odeme Plani']['D2']=None;wb.save(self.book);wb.close()
        result=import_scenarios(self.run,self.book)
        self.assertEqual(result['scenario_import'],'IMPORTED',result)
        model=json.loads((Path(result['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
        self.assertNotIn('Finansal TCO!F',model['scenario_reference']['2'])
