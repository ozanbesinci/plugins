import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from db import init_db
from scenario_finance import scenario_date,policy_valid
from scenario_formulas import add_scenarios,validate_formulas
from matrix_workbook import seal_source_cells
from test_finance_workbook import VALUES
from test_scenario_finance import FINANCE,POLICY

BAD_DATES=('20261011','2026-W41-7','0001-01-01','1899-12-31','1900-02-29','2026-02-31','2026- 1-01','2026-1-01','２０２６-10-11','2026-10-11T00:00:00')

class ScenarioDateTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);init_db(self.root)
    def tearDown(self):self.tmp.cleanup()
    def test_noncanonical_impossible_and_unsupported_dates_reject(self):
        for value in BAD_DATES+(True,123,None):
            with self.subTest(value=value),self.assertRaises(ValueError):scenario_date(value)
    def test_gregorian_boundaries_leap_day_and_policy_are_supported(self):
        for value in ('1900-01-01','1900-02-28','1900-03-01','2024-02-29','2026-10-11','9999-12-31'):
            self.assertEqual(scenario_date(value).isoformat(),value)
            self.assertTrue(policy_valid({**POLICY,'rate_date':value},1))
    def make(self):
        wb=Workbook();wb.remove(wb.active)
        contract=add_scenarios(wb,[{'id':'D','inputs':copy.deepcopy(VALUES),'financial':copy.deepcopy(FINANCE)}]);seal_source_cells(wb,contract)
        path=self.root/'dates.xlsx';wb.save(path);wb.close();return path,contract
    def test_workbook_pasted_bad_dates_are_rejected_by_input_contract(self):
        path,contract=self.make()
        for value in BAD_DATES:
            wb=load_workbook(path);wb['Finansal Varsayimlar']['G2']=value;wb.save(path);wb.close()
            self.assertEqual(validate_formulas(path,contract)['status'],'FAIL',value)
    def test_missing_date_is_valid_unknown_and_has_no_total_reference(self):
        financial=copy.deepcopy(FINANCE);financial['policy']['rate_date']=None
        wb=Workbook();wb.remove(wb.active)
        contract=add_scenarios(wb,[{'id':'D','inputs':copy.deepcopy(VALUES),'financial':financial}]);seal_source_cells(wb,contract)
        path=self.root/'unknown.xlsx';wb.save(path);wb.close()
        self.assertEqual(validate_formulas(path,contract)['status'],'PASS')
        self.assertNotIn('Finansal TCO!F',contract['reference']['2'])
    def test_formulas_use_component_dates_and_preserve_expression_seal(self):
        path,contract=self.make()
        self.assertFalse(any('DATEVALUE(' in expr for expr in contract['formulas'].values()))
        self.assertTrue(all(len(expr)<=8192 for expr in contract['formulas'].values()))
        wb=load_workbook(path);wb['Finansal TCO']['B2']='=1';wb.save(path);wb.close()
        self.assertEqual(validate_formulas(path,contract)['status'],'FAIL')
