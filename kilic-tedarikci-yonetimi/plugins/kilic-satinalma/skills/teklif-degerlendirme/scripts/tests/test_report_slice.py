"""Negative/golden slice for WP-047/048/046, using synthetic sources."""
import copy, json
from pathlib import Path
import tempfile, unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
import fitz
from common import sha256_file
from db import connect, init_db, set_meta, get_meta
from report_model import build_preliminary, matrix
from report_delivery import export_preliminary
from scenario_formulas import add_scenarios, validate_formulas, scenario_reference
GOLDEN={"quantity":10,"unit_price":100,"discount":0.1,"fx_rate":40,"vat_rate":0.2,
"vat_deductible":1,"freight":1000,"insurance":500,"duty_rate":0.1,"annual_rate":0.1,
"days":365,"energy_annual":100,"maintenance_annual":200,"years":5,"residual":1000,
"technical_score":8,"technical_weight":0.6,"cost_score":6,"cost_weight":0.4,"base_total":40000}

def seed(run,profile="standart",statuses=("noncompliant","not_stated","requires_rfi","compliant","conflict")):
    init_db(run); set_meta(run,"profile",profile)
    archive=run/"source_archive"; archive.mkdir(exist_ok=True)
    con=connect(run)
    try:
        con.execute("INSERT INTO suppliers(supplier_id,display_name) VALUES ('S','Örnek Firma')")
        con.execute("INSERT INTO offers(offer_id,supplier_id,status) VALUES ('O','S','active')")
        for did,offer,supplier,role in [("SPEC",None,None,"specification"),("OFFER","O","S","offer")]:
            source=archive/(did+".txt"); source.write_text("Sentetik kaynak "+did,encoding="utf-8"); sha=sha256_file(source)
            con.execute("INSERT INTO documents(document_id,relative_path,filename,format_family,document_role,supplier_id,offer_id,sha256,size_bytes,parse_status,metadata_json) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            [did,did+".txt",did+".txt","text",role,supplier,offer,sha,source.stat().st_size,"read",json.dumps({"archived_source_path":"source_archive/"+did+".txt"})])
            con.execute("INSERT INTO evidence(evidence_id,document_id,locator_json,quote,evidence_status,extraction_path,source_sha256) VALUES (?,?,?,?,?,?,?)",
            ["E"+did,did,'{"line":1}',"Sentetik belge konumu","verified","test",sha])
        for i,state in enumerate(statuses,start=1):
            rid="R"+str(i)
            con.execute("INSERT INTO requirements(requirement_id,source_document_id,text,critical,evidence_id) VALUES (?,?,?,?,?)",[rid,"SPEC","Şartname maddesi "+str(i),i==1,"ESPEC"])
            con.execute("INSERT INTO facts(fact_id,supplier_id,offer_id,requirement_id,field_key,value_text,evidence_status,evidence_id,extraction_path) VALUES (?,?,?,?,?,?,?,?,?)",
            ["F"+str(i),"S","O",rid,"requirement_status",state,"verified","EOFFER","test"])
        con.execute("INSERT INTO open_issues(issue_id,owner,issue_type,severity,question,status) VALUES ('I','IDARE','missing_price','critical','Eksik fiyat ve ödeme teyidi gerekli','open')"); con.commit()
    finally: con.close()

class ReportSliceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.run=Path(self.tmp.name)/"run"
    def tearDown(self): self.tmp.cleanup()
    def test_profiles_and_business_barriers_produce_a_pair(self):
        for profile in ["hizli","standart","yuksek_guvence"]:
            run=self.run/profile; seed(run,profile)
            result=export_preliminary(run,"Şartname uygunsuzluğu; eksik fiyat/belge.")
            self.assertEqual(result["delivery_status"],"READY",result)
            self.assertEqual(result["decision_assurance"],"NO_DECISION")
            self.assertEqual(result["comparison_status"],"NOT_APPLICABLE_SINGLE_SUPPLIER")
            wb=load_workbook(result["xlsx"]); self.assertGreaterEqual(len(wb.sheetnames),10)
            self.assertEqual(wb["Sartname Matrisi"]["E2"].value,"Uygunsuz"); wb.close()
            with fitz.open(result["pdf"]) as pdf:
                self.assertGreaterEqual(len(pdf),2); self.assertIn("Karar verilemedi",pdf[0].get_text())
                self.assertIn("Eksik fiyat","\n".join(p.get_text() for p in pdf))
            self.assertIsNone(get_meta(run,"reviewed_snapshot"))
    def test_matrix_distinguishes_statuses_without_fake_decision(self):
        seed(self.run); m=build_preliminary(self.run,"Eksik belge")
        self.assertEqual({r["status"] for r in m["matrix"]},{"NONCOMPLIANT","NOT_STATED","REQUIRES_CONFIRMATION","COMPLIANT","CONFLICTING"})
        self.assertEqual(m["matrix"][0]["decision_effect"],"BLOCK_RECOMMENDATION")
        self.assertIsNone(m["recommendation"]); self.assertEqual(m["ranking"],[])
    def test_same_supplier_alternatives_and_all_excluded(self):
        seed(self.run); con=connect(self.run)
        con.execute("INSERT INTO offers(offer_id,supplier_id,status) VALUES ('O2','S','active')"); con.commit(); con.close()
        self.assertEqual(build_preliminary(self.run,"Test")["comparison_status"],"NOT_APPLICABLE_SINGLE_SUPPLIER")
        con=connect(self.run)
        for i,oid in enumerate(["O","O2"]):
            con.execute("INSERT INTO exclusions(exclusion_id,offer_id,decision,reason,status) VALUES (?,?,?,?,?)",["X"+str(i),oid,"exclude","Kritik uygunsuzluk","confirmed"])
        con.commit(); con.close()
        result=export_preliminary(self.run,"Bütün teklifler elendi")
        self.assertEqual(result["delivery_status"],"READY",result); self.assertEqual(result["comparison_status"],"NO_ELIGIBLE_OFFERS")
    def test_hash_tamper_emits_safety_notice_without_source_values(self):
        seed(self.run); (self.run/"source_archive/OFFER.txt").write_text("Değişti",encoding="utf-8")
        result=export_preliminary(self.run,"Bütünlük")
        self.assertEqual(result["delivery_status"],"READY",result); self.assertEqual(result["decision_assurance"],"BLOCKED_INTEGRITY")
        m=json.loads((Path(result["xlsx"]).parent/"report_model.json").read_text(encoding="utf-8"))
        self.assertEqual(m["facts"],[]); self.assertEqual(m["matrix"],[])
    def test_foreign_supplier_evidence_rejected(self):
        seed(self.run); m=copy.deepcopy(build_preliminary(self.run,"Test"))
        next(d for d in m["documents"] if d["document_id"]=="OFFER")["supplier_id"]="OTHER"
        with self.assertRaisesRegex(ValueError,"another supplier"): matrix(m)
    def test_provisional_noncompliance_is_not_confirmed_exclusion(self):
        seed(self.run,statuses=("noncompliant",)); con=connect(self.run)
        con.execute("UPDATE facts SET evidence_status='requires_confirmation'"); con.commit(); con.close()
        row=build_preliminary(self.run,"Test")["matrix"][0]
        self.assertEqual(row["status"],"REQUIRES_CONFIRMATION"); self.assertNotEqual(row["decision_effect"],"EXCLUDE_OFFER")
    def test_primary_pdf_failure_uses_independent_fallback(self):
        seed(self.run)
        with patch("report_delivery.render_pdf",side_effect=RuntimeError("Primary PDF failed")):
            result=export_preliminary(self.run,"Eksik fiyat")
        self.assertEqual(result["delivery_status"],"READY",result)
        self.assertEqual(result["renderer"]["renderer"],"pymupdf_fallback")
    def test_two_pdf_failures_cannot_claim_ready(self):
        seed(self.run)
        with patch("report_delivery.render_pdf",side_effect=OSError("Disk")), patch("pdf_fallback.render_safety_pdf",side_effect=OSError("Disk")):
            result=export_preliminary(self.run,"Failure")
        self.assertEqual(result["delivery_status"],"FAILED"); self.assertFalse(result["deliverable_artifacts"])
        self.assertIsNone(get_meta(self.run,"preliminary_report_delivery"))
    def test_metadata_failure_preserves_previous_ready_revision(self):
        seed(self.run); first=export_preliminary(self.run,"İlk rapor"); before=get_meta(self.run,"preliminary_report_delivery")
        with patch("report_delivery._register",side_effect=OSError("DB failed")):
            failed=export_preliminary(self.run,"İkinci rapor")
        self.assertEqual(failed["delivery_status"],"FAILED"); self.assertEqual(before,get_meta(self.run,"preliminary_report_delivery"))
        self.assertEqual(sha256_file(Path(first["xlsx"])),first["files"]["Teklif_Degerlendirme.xlsx"])
    def test_reentry_and_new_scenario_keep_old_revision(self):
        seed(self.run); first=export_preliminary(self.run,"Ön rapor"); second=export_preliminary(self.run,"Ön rapor")
        self.assertTrue(second["reused"]); self.assertEqual(first["xlsx"],second["xlsx"])
        third=export_preliminary(self.run,"Yeni rapor",[{"id":"S1","inputs":GOLDEN}])
        self.assertNotEqual(first["revision"],third["revision"])
        self.assertEqual(sha256_file(Path(first["xlsx"])),first["files"]["Teklif_Degerlendirme.xlsx"])
    def test_broken_db_emits_two_safety_notice_files(self):
        self.run.mkdir(); (self.run/"analysis.duckdb").write_bytes(b"invalid duckdb")
        result=export_preliminary(self.run,"Veri deposu okunamadı")
        self.assertEqual(result["delivery_status"],"READY",result)
        self.assertEqual(result["metadata_mode"],"SAFETY_FILESYSTEM_ONLY")

    def test_unwritable_primary_directory_tries_safe_alternate(self):
        seed(self.run)
        (self.run / "reports").write_text("File blocks directory", encoding="utf-8")
        result = export_preliminary(self.run, "Eksik fiyat")
        self.assertEqual(result["delivery_status"], "READY", result)
        self.assertTrue(result["alternate_report_path_attempted"])
        self.assertIn("reports_fallback", Path(result["xlsx"]).parts)
        self.assertTrue(Path(result["pdf"]).is_file())

    def test_repaired_manifest_cannot_authorize_changed_formula_contract(self):
        seed(self.run)
        first = export_preliminary(self.run, "Ön rapor")
        directory = Path(first["xlsx"]).parent
        path = directory / "formula_contract.json"
        contract = json.loads(path.read_text(encoding="utf-8"))
        contract["formulas"]["Miktar ve Kapsam!B2"] = "=SUM(1,2)"
        path.write_text(json.dumps(contract), encoding="utf-8")
        manifest_path = directory / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["files"][path.name] = sha256_file(path)
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        second = export_preliminary(self.run, "Ön rapor")
        self.assertEqual(second["delivery_status"], "FAILED")
        self.assertFalse(second["deliverable_artifacts"])

    def test_source_text_is_literal_and_source_mutation_rejected(self):
        seed(self.run)
        con = connect(self.run)
        con.execute("UPDATE suppliers SET display_name='=HYPERLINK(\"https://invalid\",\"x\")'")
        con.commit(); con.close()
        result = export_preliminary(self.run, "Kaynak metni testi")
        wb = load_workbook(result["xlsx"])
        self.assertEqual(wb["Teklifler"]["A2"].data_type, "s")
        wb["Kaynak Olgulari"]["D2"] = "Changed source fact"
        wb.save(result["xlsx"]); wb.close()
        contract = json.loads((Path(result["xlsx"]).parent / "formula_contract.json").read_text(encoding="utf-8"))
        checked = validate_formulas(Path(result["xlsx"]), contract)
        self.assertEqual(checked["status"], "FAIL")
        self.assertTrue(any("Immutable source" in error for error in checked["errors"]))

    def test_official_export_contains_matrix_scenarios_and_pdf_for_every_profile(self):
        from reviewed_snapshot import freeze
        from report_export import export_report
        from qa import _workbook_checks
        for profile in ("hizli", "standart", "yuksek_guvence"):
            run = self.run / profile
            seed(run, profile)
            frozen = freeze(run)
            self.assertTrue(frozen["frozen"], frozen)
            with patch("report_export._recalc_with_excel", return_value={"available": False, "reason": "explicit_test_without_com"}):
                result = export_report(run)
            contract = get_meta(run, "formula_contract_" + result["xlsx_artifact_id"], parse_json=True)
            self.assertEqual(_workbook_checks(Path(result["xlsx"]), profile, contract)[0], [])
            wb = load_workbook(result["xlsx"])
            self.assertIn("TCO Yasam Dongusu", wb.sheetnames)
            self.assertIn("Sartname Sayim", wb.sheetnames)
            self.assertIsNone(wb["Senaryo Girdileri"]["B2"].value)
            wb.close()
            with fitz.open(result["pdf"]["path"]) as pdf:
                self.assertGreaterEqual(len(pdf), 2)
            self.assertFalse(result["recalc"]["available"])

    def test_duplicate_requirement_grain_and_unreadable_status(self):
        seed(self.run)
        m = build_preliminary(self.run, "Test")
        duplicate = copy.deepcopy(m)
        duplicate["facts"].append(copy.deepcopy(duplicate["facts"][0]))
        with self.assertRaisesRegex(ValueError, "Duplicate requirement"):
            matrix(duplicate)
        for fact in m["facts"]:
            fact["evidence_status"] = "unreadable"
        self.assertEqual({row["status"] for row in matrix(m)}, {"UNREADABLE"})

    def test_delivery_status_rechecks_both_files_and_current_sources(self):
        from report_delivery import inspect_preliminary
        seed(self.run)
        first = export_preliminary(self.run, "Ön rapor")
        self.assertEqual(inspect_preliminary(self.run)["delivery_status"], "READY")
        pdf = Path(first["pdf"])
        original = pdf.read_bytes()
        pdf.write_bytes(b"broken PDF")
        self.assertFalse(inspect_preliminary(self.run)["deliverable_artifacts"])
        pdf.write_bytes(original)
        con = connect(self.run)
        con.execute("UPDATE facts SET value_text='not_stated' WHERE fact_id='F1'")
        con.commit(); con.close()
        self.assertEqual(inspect_preliminary(self.run)["delivery_status"], "FAILED")

    def test_long_requirement_and_readable_scenario_pdf(self):
        seed(self.run)
        con = connect(self.run)
        con.execute("UPDATE requirements SET text=? WHERE requirement_id='R1'", ["Uzun şartname açıklaması. " * 200])
        con.commit(); con.close()
        result = export_preliminary(self.run, "Uzun şartname", [{"id":"S1", "inputs":GOLDEN}])
        self.assertEqual(result["delivery_status"], "READY", result)
        with fitz.open(result["pdf"]) as doc:
            text = "\n".join(page.get_text() for page in doc)
            self.assertIn("Net kalem tutarı", text)
            self.assertIn("900", text)
            self.assertNotIn('"Miktar ve Kapsam!B"', text)
            self.assertNotIn("PENDING_NATIVE_RECALC", text)

    def test_workflow_physical_failure_attempts_pair_without_business_success(self):
        from report_hooks import terminal_reports
        seed(self.run)
        @terminal_reports()
        def failing(run_root):
            raise OSError("Sentetik fiziksel analiz arızası")
        result = failing(self.run)
        self.assertEqual(result["status"], "BLOCKED")
        self.assertEqual(result["report_delivery_status"], "READY")
        self.assertNotEqual(result["decision_assurance"], "VERIFIED")
        self.assertIn("Fiziksel", result["reason"])

    def test_real_corrupt_db_through_workflow_emits_safety_pair(self):
        import os
        import workflow
        self.run.mkdir()
        filename = "analysis.test.sqlite" if os.environ.get("TEKLIF_TEST_SQLITE") == "1" else "analysis.duckdb"
        (self.run / filename).write_bytes(b"invalid database bytes")
        with patch("run_compatibility.assert_compatible"):
            result = workflow.advance(self.run)
        self.assertEqual(result["status"], "BLOCKED", result)
        self.assertEqual(result["report_delivery_status"], "READY", result)
        self.assertEqual(result["decision_assurance"], "BLOCKED_INTEGRITY")

    def test_programming_error_is_not_hidden_as_physical_success(self):
        from report_hooks import terminal_reports
        @terminal_reports()
        def invalid(run_root):
            raise ValueError("Logical bug")
        with self.assertRaisesRegex(ValueError, "Logical bug"):
            invalid(self.run)

class FormulaSliceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.path=Path(self.tmp.name)/"scenario.xlsx"
    def tearDown(self): self.tmp.cleanup()
    def make(self,inputs=None):
        wb=Workbook(); wb.remove(wb.active)
        contract=add_scenarios(wb,[{"id":"S1","inputs":inputs if inputs is not None else GOLDEN}])
        wb.save(self.path); wb.close(); return contract
    def test_eight_groups_and_independent_golden_values(self):
        contract=self.make(); self.assertEqual(validate_formulas(self.path,contract)["status"],"PASS")
        self.assertGreater(len(contract["formulas"]),20)
        expected={"Miktar ve Kapsam!B":900,"Kur Vergi Ithalat!B":36000,"Kur Vergi Ithalat!C":37500,
        "Kur Vergi Ithalat!D":3750,"Kur Vergi Ithalat!E":49500,"Kur Vergi Ithalat!F":41250,
        "Odeme NBD Nakit!C":45000,"TCO Yasam Dongusu!B":41750,"Puanlama Duyarlilik!B":7.2,
        "Puanlama Duyarlilik!C":7.3,"Senaryo Karsilastirma!D":1750,"Senaryo Karsilastirma!E":0.04375}
        ref=scenario_reference(GOLDEN)
        for key,value in expected.items(): self.assertAlmostEqual(ref[key],value,places=8)
    def test_missing_inputs_are_not_zero_or_cached_pass(self):
        contract=self.make({}); self.assertEqual(contract["reference"]["2"],{})
        wb=load_workbook(self.path,data_only=True); self.assertIsNone(wb["Miktar ve Kapsam"]["B2"].value); wb.close()
        self.assertEqual(validate_formulas(self.path,contract)["recalculation"],"NOT_EVALUATED_BY_STRUCTURAL_CHECK")
    def test_changed_allowed_function_other_row_and_deleted_formula_rejected(self):
        contract=self.make()
        for change in ["=SUM(1,2)","='Senaryo Girdileri'!B3",None]:
            wb=load_workbook(self.path); wb["Miktar ve Kapsam"]["B2"]=change; wb.save(self.path); wb.close()
            self.assertEqual(validate_formulas(self.path,contract)["status"],"FAIL")
    def test_external_link_macro_and_hidden_sheet_rejected(self):
        contract=self.make(); wb=load_workbook(self.path); wb["Miktar ve Kapsam"]["B2"]="='[other.xlsx]A'!B2"
        wb["Senaryo Girdileri"].sheet_state="hidden"; wb.save(self.path); wb.close()
        self.assertEqual(validate_formulas(self.path,contract)["status"],"FAIL")
        import zipfile
        with zipfile.ZipFile(self.path,"a") as z: z.writestr("xl/vbaProject.bin",b"not executed")
        self.assertEqual(validate_formulas(self.path,contract)["status"],"FAIL")
    def test_input_mutations_change_independent_reference(self):
        ref=scenario_reference(GOLDEN)
        for key,value,target in [("quantity",20,"Miktar ve Kapsam!B"),("fx_rate",45,"Kur Vergi Ithalat!B"),("annual_rate",0.2,"Odeme NBD Nakit!C"),("technical_weight",0.7,"Puanlama Duyarlilik!B")]:
            modified={**GOLDEN,key:value}
            if key=="technical_weight": modified["cost_weight"]=0.3
            self.assertNotEqual(ref[target],scenario_reference(modified)[target])
        self.assertEqual(GOLDEN["quantity"],10)
    def test_invalid_and_unknown_input_rejected(self):
        for bad in [True,-1,float("nan"),float("inf"),"10"]:
            with self.subTest(bad=bad), self.assertRaises(ValueError): self.make({"quantity":bad})
        with self.assertRaises(ValueError): self.make({"unknown":1})

    def test_official_qa_accepts_only_location_bound_scenario_formulas(self):
        from qa import _workbook_checks
        contract = self.make()
        errors, _ = _workbook_checks(self.path, "hizli", contract)
        self.assertFalse(any("formula" in e.lower() for e in errors), errors)
        wb = load_workbook(self.path)
        wb["Miktar ve Kapsam"]["B2"] = "=SUM(1,2)"
        wb.save(self.path); wb.close()
        errors, _ = _workbook_checks(self.path, "hizli", contract)
        self.assertTrue(any("Formula location/expression" in e for e in errors))
        self.assertTrue(any("Unverified formula" in e for e in errors))
