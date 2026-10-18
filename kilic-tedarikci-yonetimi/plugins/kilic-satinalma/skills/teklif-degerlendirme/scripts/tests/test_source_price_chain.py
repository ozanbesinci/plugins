"""Source-price contract regressions. All recorder metadata is a unit fixture;
no model is called and these receipts are never actual Work acceptance.
"""
from __future__ import annotations

import copy
import csv
import io
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from common import json_text, sha256_file
from db import connect, init_db
from record_results import record_task_result
from source_price_chain import calculate_item_pricing

SCRIPTS = Path(__file__).resolve().parents[1]


class SourcePriceChainTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='source-price-contract-')
        self.root = Path(self.tmp.name)
        self.counter = 0

    def tearDown(self):
        self.tmp.cleanup()

    def fixture(self, *, declared_net=8550, second_row=False, existing=False):
        self.counter += 1
        run = self.root / str(self.counter); init_db(run, SCRIPTS.parent)
        header = ['firma','miktar','birim_fiyat','para_birimi','iskonto_1_yuzde','iskonto_2_yuzde',
                  'kdv_yuzde','fiyat_kdv_durumu','net_toplam_kdv_haric','kdv_tutari','toplam_kdv_dahil']
        values = ['Supplier A','10','1000','TRY','10','5','20','KDV HARIC',str(declared_net),
                  str(declared_net*.2),str(declared_net*1.2)]
        rows = [header, values] + ([list(values)] if second_row else [])
        stream=io.StringIO();csv.writer(stream,lineterminator='\n').writerows(rows)
        text=stream.getvalue();source=run/'offer.csv';source.write_text(text,encoding='utf-8');digest=sha256_file(source)
        con=connect(run)
        try:
            con.execute("INSERT INTO suppliers(supplier_id,display_name,identification_status) VALUES ('A','Supplier A','confirmed')")
            con.execute("INSERT INTO offers(offer_id,supplier_id,status,source_document_ids_json) VALUES ('A','A','active','[\"doc-A\"]')")
            con.execute("INSERT INTO documents(document_id,relative_path,filename,format_family,document_role,supplier_id,offer_id,sha256,size_bytes,parse_status,active) VALUES ('doc-A','offer.csv','offer.csv','csv','offer','A','A',?,?,'read',true)",[digest,len(text.encode('utf-8'))])
            con.execute("INSERT INTO document_units(unit_id,document_id,unit_type,unit_index,locator_json,text_content,structured_json,content_sha256) VALUES ('unit-A','doc-A','csv_chunk',1,?,?,?,?)",[json_text({'start_row':1,'end_row':len(rows)}),text,json_text({'rows':rows}),digest])
            from task_queue import _add_task
            task_id=_add_task(con,'EXTRACT','semantic_extraction','offer_semantic_fields',{'offer_id':'A','fields':['quantity'],'document_ids':['doc-A']},offer_id='A',critical=True)
            con.commit()
        finally:con.close()
        names={'q':2,'price':3,'currency':4,'d1':5,'d2':6,'rate':7,'basis':8,'net':9,'vat':10,'gross':11}
        evidence=[{'client_id':'ev-'+name,'document_id':'doc-A','unit_id':'unit-A','locator':{'row':2,'column':column},'quote':values[column-1],'evidence_status':'verified'} for name,column in names.items()]
        item={'item_id':'STP-01','amount':1000,'quantity':10,'amount_basis':'unit_price','currency':'TRY','unit':'ADET','evidence_status':'verified',
              'evidence_refs':['ev-q','ev-price','ev-currency'],
              'pricing':{'discounts':{'percentages':[10,5],'evidence_refs':['ev-d1','ev-d2']},
                         'vat':{'percent':20,'basis':'excluded','evidence_refs':['ev-rate','ev-basis']},
                         'stated_amounts':{'scope':'offer','net':{'amount':declared_net,'evidence_refs':['ev-net']},
                                           'vat':{'amount':declared_net*.2,'evidence_refs':['ev-vat']},
                                           'gross':{'amount':declared_net*1.2,'evidence_refs':['ev-gross']}}}}
        fact={'offer_id':'A','field_key':'quantity','value_number':10,'evidence_status':'verified',
              'evidence_client_id':'ev-q','normalized':{'items':[item]}}
        result={'facts':[fact],'evidence':evidence}
        if existing:
            con=connect(run)
            try:
                for ev in evidence:
                    con.execute("INSERT INTO evidence(evidence_id,document_id,unit_id,locator_json,quote,evidence_status,extraction_path,confidence,source_sha256) VALUES (?,'doc-A','unit-A',?,?,'verified','unit_source_fixture',1.0,?)",[ev['client_id'],json_text(ev['locator']),ev['quote'],digest])
                con.commit()
            finally:con.close()
            result['evidence']=[];fact.pop('evidence_client_id');fact['evidence_id']='ev-q'
        return run,task_id,result

    @staticmethod
    def record(run,tid,result):
        return record_task_result(run,tid,result,model_name='unit-source-price-fixture-no-model-call',reasoning_effort='fixture',input_tokens=0,output_tokens=0)

    @staticmethod
    def counts(run,tid):
        con=connect(run)
        try:return ([con.execute('SELECT count(*) FROM '+t).fetchone()[0] for t in ['facts','evidence','model_receipts']],con.execute('SELECT status FROM task_queue WHERE task_id=?',[tid]).fetchone()[0])
        finally:con.close()

    def rejects_atomically(self,run,tid,result):
        before,_=self.counts(run,tid)
        with self.assertRaises(ValueError):self.record(run,tid,result)
        after,status=self.counts(run,tid)
        self.assertEqual(after,before);self.assertEqual(status,'pending')

    def seed_source_totals(self,run,declared_net=8550):
        con=connect(run)
        try:
            refs=con.execute('SELECT evidence_id,quote FROM evidence').fetchall()
            by_quote={str(quote):eid for eid,quote in refs}
            for field,num,text,eid in [('total',declared_net,None,by_quote[str(declared_net)]),('currency',None,'TRY',by_quote['TRY']),('vat',None,'KDV HARIC',by_quote['KDV HARIC'])]:
                con.execute("INSERT INTO facts(fact_id,offer_id,supplier_id,field_key,value_number,value_text,evidence_status,critical,evidence_id,extraction_path,normalized_json,updated_at) VALUES (?,'A','A',?,?,?,'verified',true,?,'unit_source_fixture','{}','now')",['source/'+field,field,num,text,eid])
            con.commit()
        finally:con.close()

    def test_real_recorder_accepts_complete_source_groups_without_double_adding_stated_total(self):
        run,tid,result=self.fixture();self.assertEqual(self.record(run,tid,result)['status'],'completed')
        con=connect(run)
        try:
            normalized=json.loads(con.execute("SELECT normalized_json FROM facts WHERE field_key='quantity'").fetchone()[0])
            self.assertEqual(Decimal(normalized['item_calculation']['amount']),Decimal(10000))
            calculated=calculate_item_pricing(normalized['items'][0])
            self.assertEqual([Decimal(calculated[k]) for k in ['net','vat','gross']],[Decimal(8550),Decimal(1710),Decimal(10260)])
        finally:con.close()

    def test_discount_order_is_verified_even_when_swapping_rates_has_the_same_product(self):
        run,tid,result=self.fixture();result['facts'][0]['normalized']['items'][0]['pricing']['discounts']['percentages']=[5,10]
        self.rejects_atomically(run,tid,result)

    def test_same_numeric_quantity_cell_cannot_prove_discount_percent(self):
        run,tid,result=self.fixture();result['facts'][0]['normalized']['items'][0]['pricing']['discounts']['evidence_refs']=['ev-q','ev-d2']
        self.rejects_atomically(run,tid,result)

    def test_discount_on_another_row_is_rejected(self):
        run,tid,result=self.fixture(second_row=True)
        next(e for e in result['evidence'] if e['client_id']=='ev-d1')['locator']['row']=3
        self.rejects_atomically(run,tid,result)

    def test_all_source_discount_stages_must_be_covered(self):
        run,tid,result=self.fixture();group=result['facts'][0]['normalized']['items'][0]['pricing']['discounts'];group['percentages']=[10];group['evidence_refs']=['ev-d1']
        self.rejects_atomically(run,tid,result)

    def test_percent_cell_is_not_a_stated_vat_amount(self):
        run,tid,result=self.fixture();amount=result['facts'][0]['normalized']['items'][0]['pricing']['stated_amounts']['vat'];amount.update(amount=20,evidence_refs=['ev-rate'])
        self.rejects_atomically(run,tid,result)

    def test_source_vat_exclusion_cannot_be_changed_to_inclusion(self):
        run,tid,result=self.fixture();result['facts'][0]['normalized']['items'][0]['pricing']['vat']['basis']='included'
        self.rejects_atomically(run,tid,result)

    def test_stale_or_pixel_source_rate_evidence_cannot_be_promoted(self):
        for field,value in [('source_sha256','stale'),('extraction_path','vision')]:
            run,tid,result=self.fixture(existing=True);con=connect(run)
            try:con.execute('UPDATE evidence SET '+field+'=? WHERE evidence_id=\'ev-d1\'',[value]);con.commit()
            finally:con.close()
            self.rejects_atomically(run,tid,result)

    def test_unknown_discount_or_vat_terms_are_not_zero(self):
        for field in ['discounts','vat']:
            run,tid,result=self.fixture();result['facts'][0]['normalized']['items'][0]['pricing'].pop(field)
            self.rejects_atomically(run,tid,result)

    def test_included_vat_and_price_unit_reference_math(self):
        _,_,result=self.fixture();item=copy.deepcopy(result['facts'][0]['normalized']['items'][0]);item.update(amount=1200,quantity=200);item['pricing'].update(discounts={'percentages':[10],'evidence_refs':['math-fixture']},vat={'percent':20,'basis':'included','evidence_refs':['math-fixture']},price_unit={'value':100,'evidence_refs':['math-fixture']})
        row=calculate_item_pricing(item)
        self.assertEqual([Decimal(row[k]) for k in ['net','vat','gross']],[Decimal(1800),Decimal(360),Decimal(2160)])

    def test_nonfinite_negative_and_boolean_rates_cannot_reach_canonical_facts(self):
        for bad in [True,-1,100,float('nan'),float('inf')]:
            run,tid,result=self.fixture();result['facts'][0]['normalized']['items'][0]['pricing']['discounts']['percentages'][0]=bad
            self.rejects_atomically(run,tid,result)

    def test_normal_cost_path_registers_reconciled_source_chain_separately(self):
        from procurement_engine import calculate_costs
        run,tid,result=self.fixture();self.record(run,tid,result);self.seed_source_totals(run)
        con=connect(run)
        try:
            calculate_costs(con,{'base_currency':'TRY','base_date':'2026-10-17','annual_discount_rate':0,'annual_discount_rate_currency':'TRY'})
            row=con.execute("SELECT result_json,status FROM calculations WHERE calculation_type='source_price_chain'").fetchone()
            self.assertIsNotNone(row);self.assertEqual(row[1],'complete');obj=json.loads(row[0]);self.assertTrue(obj['complete']);self.assertEqual(len(obj['reconciliations']),4)
            self.assertEqual(Decimal(obj['net_total']),Decimal(8550));self.assertEqual(Decimal(obj['vat_total']),Decimal(1710));self.assertEqual(Decimal(obj['gross_total']),Decimal(10260))
        finally:con.close()

    def test_known_stated_source_mismatch_creates_a_critical_issue(self):
        from procurement_engine import calculate_costs
        run,tid,result=self.fixture(declared_net=8600);self.record(run,tid,result);self.seed_source_totals(run,8600)
        con=connect(run)
        try:
            calculate_costs(con,{'base_currency':'TRY','base_date':'2026-10-17'})
            state=con.execute("SELECT status FROM calculations WHERE calculation_type='source_price_chain'").fetchone()[0]
            self.assertEqual(state,'blocked')
            self.assertEqual(con.execute("SELECT count(*) FROM open_issues WHERE issue_type='source_price_chain_reconciliation' AND severity='critical' AND status='open'").fetchone()[0],1)
        finally:con.close()

    def test_independent_path_detects_tampered_main_source_net(self):
        from procurement_engine import calculate_costs
        from independent_recompute import run as independent_run
        from unittest.mock import patch
        run,tid,result=self.fixture();self.record(run,tid,result);self.seed_source_totals(run)
        con=connect(run)
        try:
            calculate_costs(con,{'base_currency':'TRY','base_date':'2026-10-17'})
            con.commit()
        finally:con.close()
        with patch('source_price_chain.calculate_item_pricing',side_effect=AssertionError('independent path called primary source arithmetic')):
            baseline=independent_run(run)
        source_checks=[row for row in baseline['checks'] if row['calculation_type']=='source_price_chain']
        self.assertEqual(len(source_checks),1);self.assertEqual(source_checks[0]['status'],'pass',source_checks)
        con=connect(run)
        try:
            row=con.execute("SELECT calculation_id,result_json FROM calculations WHERE calculation_type='source_price_chain'").fetchone();obj=json.loads(row[1]);obj['net_total']='8551'
            con.execute('UPDATE calculations SET result_json=? WHERE calculation_id=?',[json_text(obj),row[0]]);con.commit()
        finally:con.close()
        report=independent_run(run)
        rows=[row for row in report['checks'] if row['calculation_type']=='source_price_chain']
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['status'],'fail')

    def test_frozen_price_chain_reaches_excel_and_both_pdf_renderers(self):
        from unittest import mock
        from db import rows_as_dict
        from procurement_engine import calculate_costs
        from report_export import export_report
        from pdf_fallback import render_safety_pdf
        from source_scenarios import from_sources
        from openpyxl import load_workbook
        import fitz
        from test_engine_support import ReportFixtureConnection
        run,tid,result=self.fixture();self.record(run,tid,result);self.seed_source_totals(run)
        con=connect(run)
        try:
            calculate_costs(con,{'base_currency':'TRY','base_date':'2026-10-17'})
            model={name:rows_as_dict(con,'SELECT * FROM '+name) for name in ['suppliers','offers','documents','evidence','facts','calculations']}
        finally:con.close()
        model.update(status='PRELIMINARY',meta={'profile':'hizli','decision_status':'INSUFFICIENT'})
        original=copy.deepcopy(model);bindings=[]
        class Con(ReportFixtureConnection):
            def __init__(self):super().__init__('snapshot')
        with mock.patch('reviewed_snapshot.load_latest',return_value=(model,{'sha256':'snapshot','relative_path':'reviewed/snapshot.json'})),mock.patch('report_export.connect',return_value=Con()),mock.patch('report_export._artifact',return_value='A'),mock.patch('report_export._insert_bindings',side_effect=lambda c,a,s,r:bindings.extend(r)),mock.patch('report_export._recalc_with_excel',return_value={'available':False}):
            report=export_report(run)
        wb=load_workbook(report['xlsx'],data_only=True)
        try:
            values={row[2]:row[3] for row in wb['Kaynak Fiyat Zinciri'].iter_rows(min_row=2,values_only=True) if row[1]=='STP-01'}
            for field,expected in [('KDV hariç net',8550),('KDV tutarı',1710),('KDV dahil brüt',10260)]:self.assertEqual(Decimal(values[field]),Decimal(expected))
            self.assertEqual(values['Birim fiyat KDV esası'],'KDV hariç')
            by_title={row[2]:row for row in wb['Kaynak Fiyat Zinciri'].iter_rows(min_row=2,values_only=True) if row[1]=='STP-01'}
            for field in ['Kaynak miktarı','Fiyatın miktar böleni','Kaynak iskonto 1 (%)','Kaynak iskonto 2 (%)','Kaynak KDV (%)','Birim fiyat KDV esası']:
                self.assertIsNone(by_title[field][4],field+' is not a monetary value')
            for field in ['Kaynak birim fiyatı','İskonto öncesi kaynak tutarı','İskonto sonrası kaynak tutarı','KDV hariç net','KDV tutarı','KDV dahil brüt']:
                self.assertEqual(by_title[field][4],'TRY',field)
        finally:wb.close()
        source_bindings=[b for b in bindings if b['locator'].get('sheet')=='Kaynak Fiyat Zinciri']
        self.assertTrue(any(b['locator']['value_path']=='items.0.net' and b['calculation_ids'] for b in source_bindings))
        line=next(s for s in from_sources(model) if s.get('item_id')=='STP-01')
        self.assertEqual(line['pricing']['discounts'],[.1,.05]);self.assertEqual(line['inputs']['vat_rate'],.2);self.assertEqual(line['pricing']['price_includes_vat'],0)
        primary=Path(report['xlsx']).with_suffix('.pdf')
        fallback=run/'fallback-source-prices.pdf'
        frozen=json.loads((primary.parent/'report_model.json').read_text(encoding='utf-8'))
        render_safety_pdf(frozen,fallback,'fixture-forced-renderer-failure')
        for path in [primary,fallback]:
            with fitz.open(path) as pdf:text=' '.join(page.get_text() for page in pdf)
            compact=' '.join(text.split())
            self.assertIn('Kaynak fiyat zinciri',compact)
            for expected in ['8550','1710','10260','Hesap eksi beyan: 0']:self.assertIn(expected,compact)
        self.assertEqual(model,original)

    def test_integrity_notice_never_displays_computed_source_prices(self):
        from source_price_display import sheet_entries,pdf_rows
        model={'offers':[{'offer_id':'A','status':'active'}],'calculations':[{'offer_id':'A','calculation_id':'C','calculation_type':'source_price_chain','result':{'net_total':'8550'}}]}
        for guard in [{'content_status':'SAFETY_NOTICE'},{'decision_assurance':'BLOCKED_INTEGRITY'}]:
            guarded={**model,**guard};self.assertEqual(sheet_entries(guarded),[]);self.assertEqual(pdf_rows(guarded),[])
