from __future__ import annotations
import json, shutil, tempfile, unittest
from pathlib import Path
from unittest import mock
from db import connect
from task_queue import _add_task, generate_review_tasks
from record_results import record_task_result
import test_v504_audit_regressions as audit_tests

class AcceptanceFixTests(unittest.TestCase):
    setUp = audit_tests.V504AuditRegressionTests.setUp
    tearDown = audit_tests.V504AuditRegressionTests.tearDown
    offer = audit_tests.V504AuditRegressionTests.offer
    valid_pair = audit_tests.V504AuditRegressionTests.valid_pair
    finish_valid = audit_tests.V504AuditRegressionTests.finish_valid

    def critical_roundtrip(self, explicit_false):
        self.valid_pair()
        con=connect(self.run)
        row=con.execute("SELECT fact_id,offer_id,evidence_id FROM facts WHERE field_key='currency' ORDER BY fact_id LIMIT 1").fetchone()
        fid,oid,eid=row
        did=con.execute("SELECT document_id FROM documents WHERE offer_id=?",[oid]).fetchone()[0]
        tid=_add_task(con,'EXTRACT','semantic_extraction','offer_semantic_fields',{'offer_id':oid,'fields':['currency']},offer_id=oid,critical=True)
        con.commit(); con.close()
        from workflow import _task_contexts, _policy
        from db import rows_as_dict
        con=connect(self.run); task=rows_as_dict(con,'SELECT * FROM task_queue WHERE task_id=?',[tid])[0]; con.close()
        _task_contexts(self.run,_policy(Path(__file__).resolve().parents[2]),[task])
        fact={'fact_id':fid,'offer_id':oid,'field_key':'currency','value_text':'EUR','evidence_status':'verified','evidence_client_id':'currency-new'}
        if explicit_false: fact['critical']=False
        record_task_result(self.run,tid,{'facts':[fact],'evidence':[{'client_id':'currency-new','document_id':did,'locator':{'sheet':'Teklif','cell':'B5'},'quote':'EUR','evidence_status':'verified'}]},model_name='unit-test-fixture',reasoning_effort='medium',input_tokens=1,output_tokens=1)
        con=connect(self.run)
        self.assertTrue(con.execute('SELECT critical FROM facts WHERE fact_id=?',[fid]).fetchone()[0])
        con.commit(); con.close()
        generate_review_tasks(self.run,'yuksek_guvence')
        con=connect(self.run)
        tasks=con.execute("SELECT payload_json FROM task_queue WHERE task_kind='independent_fact_review'").fetchall()
        con.commit(); con.close()
        self.assertTrue(any(fid in json.loads(x[0])['fact_ids'] for x in tasks))

    def test_optional_critical_cannot_narrow_independent_review(self): self.critical_roundtrip(False)
    def test_explicit_false_cannot_downgrade_critical_field(self): self.critical_roundtrip(True)

    def test_delivery_dual_representation_reconciles_but_tampering_fails(self):
        from cross_output_consistency import check
        from openpyxl import load_workbook
        self.valid_pair()
        con=connect(self.run)
        con.execute("UPDATE facts SET value_number=30 WHERE field_key='delivery'")
        con.commit(); con.close()
        with mock.patch('report_export._recalc_with_excel',return_value={'available':True}):
            self.assertEqual(self.finish_valid()['status'],'VERIFIED')
        self.assertEqual(check(self.run)['status'],'PASS')
        con=connect(self.run)
        locators=[json.loads(r[0]) for r in con.execute("SELECT locator_json FROM artifact_bindings WHERE component_id LIKE 'xlsx:Teklif Karsilastirma:N%'").fetchall()]
        con.commit(); con.close()
        self.assertTrue(locators)
        self.assertTrue(all(x['value_path']=='value_text' for x in locators))
        p=self.run/'output/Teklif_Degerlendirme.xlsx'
        wb=load_workbook(p); wb['Teklif Karsilastirma']['N2']='31 gun'; wb.save(p); wb.close()
        result=check(self.run)
        self.assertEqual(result['status'],'FAIL')
        self.assertTrue(any(x['locator']=='N2' for x in result['findings']))

    def test_numeric_only_delivery_is_not_blank(self):
        from cross_output_consistency import check
        from openpyxl import load_workbook
        from reviewed_snapshot import freeze
        from report_export import export_report
        self.valid_pair()
        with mock.patch('report_export._recalc_with_excel',return_value={'available':True}):
            self.assertEqual(self.finish_valid()['status'],'VERIFIED')
        con=connect(self.run)
        con.execute("UPDATE facts SET value_number=30,value_text=NULL WHERE field_key='delivery'")
        con.commit(); con.close()
        freeze(self.run)
        with mock.patch('report_export._recalc_with_excel',return_value={'available':True}):
            export_report(self.run)
        wb=load_workbook(self.run/'output/Teklif_Degerlendirme.xlsx')
        self.assertEqual(wb['Teklif Karsilastirma']['N2'].value,30)
        comparison=wb['Teklif Karsilastirma']
        self.assertEqual(comparison.page_setup.scale,100)
        self.assertEqual(comparison.print_title_cols,'$A:$B')
        # Native Excel repeats identities reliably with automatic horizontal
        # pagination; manually forced bands can omit A:B on intermediate pages.
        self.assertEqual(len(comparison.col_breaks.brk),0)
        wb.close()
        self.assertEqual(check(self.run)['status'],'PASS')

    def vat_result(self, first, second):
        from comparison_contract import assess
        self.valid_pair()
        con=connect(self.run)
        ids=[r[0] for r in con.execute("SELECT fact_id FROM facts WHERE field_key='vat' ORDER BY fact_id").fetchall()]
        for fid,value in zip(ids,(first,second)):
            con.execute('UPDATE facts SET value_text=? WHERE fact_id=?',[value,fid])
        con.commit(); con.close()
        assess(self.run)
        con=connect(self.run)
        row=con.execute("SELECT status FROM comparability_checks WHERE dimension='vat_basis'").fetchone()
        con.commit(); con.close()
        return row[0]

    def test_vat_haric_and_excluded_share_basis(self):
        self.assertEqual(self.vat_result('hariç','excluded'),'pass')

    def test_vat_included_and_excluded_still_require_normalization(self):
        self.assertEqual(self.vat_result('dahil','excluded'),'fail')

    def test_unknown_same_vat_does_not_become_zero_or_known_basis(self):
        from normalize import vat_basis
        self.assertIsNone(vat_basis('sonradan belirtilecek'))
        self.assertEqual(self.vat_result('sonradan belirtilecek','sonradan belirtilecek'),'fail')

    def test_docx_targeted_packet_keeps_quantity_unit_and_turkish_terms(self):
        from datetime import date
        from docx import Document
        from workflow import prepare
        from context_packet import task_packet
        self.offer('Alpha_teklif.xlsx',100000)
        doc=Document()
        doc.add_paragraph('Fiyat Teklifi')
        for i in range(35): doc.add_paragraph('İlgisiz açıklama '+str(i))
        for text in ['Miktar: 10','Birim: adet','Ürün: Pompa','Ödeme: %100 60 gün']:
            doc.add_paragraph(text)
        doc.save(self.src/'Beta_teklif.docx')
        prepare(self.src,self.run,'hizli',date(2026,10,3))
        con=connect(self.run)
        did,oid=con.execute("SELECT document_id,offer_id FROM documents WHERE filename='Beta_teklif.docx'").fetchone()
        tid=_add_task(con,'EXTRACT','semantic_extraction','offer_semantic_fields',{'document_ids':[did],'offer_id':oid,'fields':['quantity','scope_signature','payment']},offer_id=oid)
        con.commit(); con.close()
        packet=task_packet(self.run,tid,8,12000)
        text=' '.join(x.get('text_content','') for d in packet['documents'] for x in d['units'])
        for expected in ['Miktar: 10','Birim: adet','Ürün: Pompa','Ödeme: %100 60 gün']:
            self.assertIn(expected,text)
        self.assertNotIn('İlgisiz açıklama 0',text)
        self.assertLessEqual(len(json.dumps(packet,ensure_ascii=False,separators=(',',':'),default=str)),12000)
        self.assertLessEqual(sum(len(d['units']) for d in packet['documents']),8)

    def qa_visual_gate(self, decision_status, mechanical_errors, expected):
        from db import set_meta
        from qa import run_qa
        self.valid_pair()
        set_meta(self.run,'profile','yuksek_guvence')
        set_meta(self.run,'decision_status',decision_status)
        art={'artifact_id':'test-artifact','relative_path':'unit-fixture.pdf','sha256':'test-hash'}
        with mock.patch('cross_output_consistency.check',return_value={'status':'PASS'}), mock.patch('provenance_survival.check',return_value={'status':'PASS'}), mock.patch('qa._latest_artifact',return_value=art), mock.patch('qa.sha256_file',return_value='test-hash'), mock.patch('qa._workbook_checks',return_value=([],[])), mock.patch('xlsx_formula_audit.audit',return_value={'findings':[]}), mock.patch('qa._mechanical_db_checks',return_value=list(mechanical_errors)), mock.patch('qa._pdf_content_checks',return_value=[]), mock.patch('qa._render_pdf',return_value=[]) as render, mock.patch('qa._visual_tasks',return_value=1) as tasks:
            result=run_qa(self.run)
            self.assertEqual(render.called,expected)
            self.assertEqual(tasks.called,expected)
            if mechanical_errors: self.assertEqual(result['status'],'FAIL')

    def test_failed_mechanical_qa_does_not_queue_visual_acceptance(self):
        self.qa_visual_gate('SUFFICIENT',['numeric mismatch'],False)

    def test_preliminary_does_not_queue_final_visual_acceptance(self):
        self.qa_visual_gate('PRELIMINARY_ONLY',[],False)

    def test_sufficient_clean_artifact_still_requires_visual_acceptance(self):
        self.qa_visual_gate('SUFFICIENT',[],True)

    def local_id_roundtrip(self, foreign=False):
        self.valid_pair()
        con=connect(self.run)
        rows=con.execute("SELECT f.fact_id,f.offer_id,d.document_id FROM facts f JOIN documents d ON d.offer_id=f.offer_id WHERE field_key='currency' ORDER BY f.fact_id").fetchall()
        fid,oid,did=rows[0]
        alias=rows[1][0] if foreign else 'model_currency_alias'
        before=con.execute('SELECT count(*) FROM facts').fetchone()[0]
        tid=_add_task(con,'EXTRACT','semantic_extraction','offer_semantic_fields',{'offer_id':oid,'document_ids':[did],'fields':['currency']},offer_id=oid,critical=True)
        con.commit(); con.close()
        from workflow import _task_contexts, _policy
        from db import rows_as_dict
        con=connect(self.run); task=rows_as_dict(con,'SELECT * FROM task_queue WHERE task_id=?',[tid])[0]; con.close()
        _task_contexts(self.run,_policy(Path(__file__).resolve().parents[2]),[task])
        result={'facts':[{'fact_id':alias,'offer_id':oid,'field_key':'currency','value_text':'EUR','evidence_status':'verified','evidence_client_id':'client-currency'}],'evidence':[{'client_id':'client-currency','document_id':did,'locator':{'sheet':'Teklif','cell':'B5'},'quote':'EUR','evidence_status':'verified'}],'open_issues':[{'owner':'FIRMA:'+oid,'question':'Testte kanıt referansını doğrula','fact_ids':[alias],'evidence_ids':['client-currency']}]}
        if foreign:
            with self.assertRaisesRegex(ValueError,'Fact identity'):
                record_task_result(self.run,tid,result,model_name='unit-test-fixture',reasoning_effort='medium',input_tokens=1,output_tokens=1)
        else:
            record_task_result(self.run,tid,result,model_name='unit-test-fixture',reasoning_effort='medium',input_tokens=1,output_tokens=1)
        con=connect(self.run)
        self.assertEqual(con.execute('SELECT count(*) FROM facts').fetchone()[0],before)
        self.assertEqual(con.execute('SELECT offer_id FROM facts WHERE fact_id=?',[rows[1][0]]).fetchone()[0],rows[1][1])
        if not foreign:
            eid=con.execute('SELECT evidence_id FROM facts WHERE fact_id=?',[fid]).fetchone()[0]
            fids,eids=con.execute("SELECT fact_ids_json,evidence_ids_json FROM open_issues WHERE question='Testte kanıt referansını doğrula'").fetchone()
            self.assertEqual(json.loads(fids),[fid]); self.assertEqual(json.loads(eids),[eid])
        con.commit(); con.close()

    def test_model_local_ids_and_rfi_refs_use_one_canonical_grain(self):
        self.local_id_roundtrip()

    def test_model_id_cannot_overwrite_foreign_fact_identity(self):
        self.local_id_roundtrip(True)

    def test_visual_coverage_requires_all_actual_pdf_pages(self):
        from common import sha256_file
        from coverage import build
        from db import set_meta
        from pypdf import PdfWriter
        self.valid_pair()
        set_meta(self.run,'profile','yuksek_guvence')
        pdf=self.run/'coverage-fixture.pdf'
        w=PdfWriter();w.add_blank_page(width=200,height=200);w.add_blank_page(width=200,height=200)
        with pdf.open('wb') as h:w.write(h)
        con=connect(self.run)
        con.execute("INSERT INTO artifacts(artifact_id,artifact_type,relative_path,sha256,parent_hashes_json,status,created_at) VALUES (?,?,?,?,?,?,?)",['coverage-pdf','pdf',pdf.name,sha256_file(pdf),'[]','ready','2026-10-04'])
        con.commit();con.close()
        def state():
            report=build(self.run)
            unit=next(u for u in report['units'] if u['id']=='coverage-pdf')
            return unit['categories']['visual-qa']['state'],report
        st,report=state();self.assertEqual(st,'unverified')
        self.assertTrue(any(x['unit_id']=='coverage-pdf' and x['category']=='visual-qa' for x in report['blocking_gaps']))
        con=connect(self.run)
        con.execute("INSERT INTO render_hashes VALUES (?,?,?,?,?)",['coverage-pdf','page:1','hash1','pass','controlled unit fixture'])
        con.commit();con.close();self.assertEqual(state()[0],'unverified')
        con=connect(self.run)
        con.execute("INSERT INTO render_hashes VALUES (?,?,?,?,?)",['coverage-pdf','page:2','hash2','fail','controlled unit fixture'])
        con.commit();con.close();self.assertEqual(state()[0],'failed')
        con=connect(self.run);con.execute("UPDATE render_hashes SET qa_status='pass' WHERE artifact_id='coverage-pdf' AND render_key='page:2'")
        con.commit();con.close();self.assertEqual(state()[0],'passed')
