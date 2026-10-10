from __future__ import annotations
import json,unittest
from datetime import date
from unittest import mock
from openpyxl import load_workbook
import test_v504_audit_regressions as audit
from db import connect,get_meta,set_meta
from common import json_text
from comparison_contract import assess
from decision_gate import run_gate
from reviewed_snapshot import freeze,load_latest
from report_export import export_report
from provenance_survival import check as provenance
from context_packet import decision_packet_from_connection
from cross_output_consistency import check as reconcile

class ReportFixTests(unittest.TestCase):
    setUp=audit.V504AuditRegressionTests.setUp
    tearDown=audit.V504AuditRegressionTests.tearDown
    offer=audit.V504AuditRegressionTests.offer
    valid_pair=audit.V504AuditRegressionTests.valid_pair
    def finish(self):
        with mock.patch('report_export._recalc_with_excel',return_value={'available':True,'status':'unit_mock'}):
            return audit.V504AuditRegressionTests.finish_valid(self)
    def export(self):
        with mock.patch('report_export._recalc_with_excel',return_value={'available':True,'status':'unit_mock'}):
            return export_report(self.run)
    def add_old_scoring(self):
        c=connect(self.run);row=c.execute("SELECT offer_id FROM offers ORDER BY offer_id LIMIT 1").fetchone();oid=row[0]
        result={'status':'cost_only','price_scores':[{'offer_id':oid,'price_score_0_10':'999'}]}
        c.execute("INSERT INTO calculations VALUES (?,?,?,?,?,?,?,?,?)",['zz-old-score',None,'scoring','{}',json_text(result),'fixture','old-hash','complete','2000-01-01'])
        c.commit();c.close();return oid
    def test_latest_score_is_suppressed_and_old_score_never_exported(self):
        self.valid_pair();self.finish();self.add_old_scoring()
        c=connect(self.run);cid=c.execute("SELECT calculation_id FROM calculations WHERE calculation_type='scoring' AND calculation_id<>'zz-old-score' ORDER BY updated_at DESC LIMIT 1").fetchone()[0]
        c.execute("UPDATE facts SET evidence_status='requires_confirmation' WHERE field_key='total' AND offer_id=(SELECT min(offer_id) FROM offers)");c.commit();c.close()
        assess(self.run);set_meta(self.run,'decision_status','PRELIMINARY_ONLY')
        c=connect(self.run);status,result=c.execute('SELECT status,result_json FROM calculations WHERE calculation_id=?',[cid]).fetchone();c.close()
        self.assertEqual(status,'incomplete');self.assertEqual(json.loads(result)['status'],'suppressed_by_comparability')
        freeze(self.run);self.export();w=load_workbook(self.run/'output/Teklif_Degerlendirme.xlsx',data_only=True)
        try:self.assertTrue(all(w['Teklif Karsilastirma'].cell(r,k).value is None for r in range(2,4) for k in [12,13]))
        finally:w.close()
        self.assertEqual(reconcile(self.run)['status'],'PASS')
        w=load_workbook(self.run/'output/Teklif_Degerlendirme.xlsx');w['Teklif Karsilastirma']['L2']=9;w.save(self.run/'output/Teklif_Degerlendirme.xlsx');w.close()
        self.assertEqual(reconcile(self.run)['status'],'FAIL')
    def test_nonverified_total_is_visible_as_confirmation_without_value_loss(self):
        self.valid_pair();self.finish();c=connect(self.run);oid=c.execute('SELECT min(offer_id) FROM offers').fetchone()[0]
        c.execute("UPDATE facts SET evidence_status='requires_confirmation' WHERE offer_id=? AND field_key='total'",[oid]);c.commit();c.close();assess(self.run);set_meta(self.run,'decision_status','PRELIMINARY_ONLY');freeze(self.run);self.export()
        w=load_workbook(self.run/'output/Teklif_Degerlendirme.xlsx',data_only=True);ws=w['Teklif Karsilastirma'];r=next(r for r in range(2,ws.max_row+1) if ws.cell(r,2).value==oid)
        try:
            self.assertEqual(ws.cell(r,17).value,'Teyit bekleniyor');self.assertEqual(ws.cell(r,18).value,'Teyit bekleniyor');self.assertIsInstance(ws.cell(r,3).value,(int,float));self.assertIn('Teyit bekleniyor',ws.cell(r,3).number_format)
        finally:w.close()
    def test_provenance_uses_current_artifact_and_still_rejects_bad_parent(self):
        self.valid_pair();self.finish();set_meta(self.run,'decision_summary','Güncellenmiş ön değerlendirme.');freeze(self.run);self.export()
        self.assertEqual(provenance(self.run)['status'],'PASS')
        c=connect(self.run);aid=c.execute("SELECT artifact_id FROM artifacts WHERE artifact_type='xlsx' ORDER BY created_at DESC LIMIT 1").fetchone()[0];c.execute("UPDATE artifacts SET parent_hashes_json='[]' WHERE artifact_id=?",[aid]);c.commit();c.close()
        self.assertEqual(provenance(self.run)['status'],'FAIL')
    def test_canonical_packet_and_snapshot_exclude_old_calculations(self):
        self.valid_pair();self.finish();self.add_old_scoring();c=connect(self.run);packet=decision_packet_from_connection(c,40000);c.close();self.assertNotIn('zz-old-score',[x['calculation_id'] for x in packet['decision_context']['calculations']]);freeze(self.run);snap,_=load_latest(self.run);self.assertNotIn('zz-old-score',[x['calculation_id'] for x in snap['calculations']])
    def test_resolved_dimension_closes_only_its_generated_issue(self):
        self.valid_pair();c=connect(self.run);oid=c.execute('SELECT min(offer_id) FROM offers').fetchone()[0];c.execute("UPDATE facts SET evidence_status='requires_confirmation' WHERE offer_id=? AND field_key='currency'",[oid]);c.execute("UPDATE facts SET value_number=20 WHERE offer_id=? AND field_key='quantity'",[oid]);c.commit();c.close();assess(self.run)
        c=connect(self.run);iid=c.execute("SELECT issue_id FROM open_issues WHERE question LIKE 'Resolve comparison basis for currency before%' AND status='open'").fetchone()[0];c.execute("UPDATE facts SET evidence_status='verified' WHERE offer_id=? AND field_key='currency'",[oid]);c.commit();c.close();assess(self.run)
        c=connect(self.run);self.assertEqual(c.execute('SELECT status FROM open_issues WHERE issue_id=?',[iid]).fetchone()[0],'resolved');self.assertGreater(c.execute("SELECT count(*) FROM open_issues WHERE question LIKE 'Resolve comparison basis for quantity before%' AND status='open'").fetchone()[0],0);c.close()
    def test_passing_comparison_never_closes_manual_same_category_question(self):
        self.valid_pair();c=connect(self.run);c.execute("INSERT INTO open_issues VALUES (?,?,?,?,?,?,?,?,?,?)",['manual-comparison','IDARE','critical','comparability_contract','Kullanıcının kapsam sorusu','[]','[]','open','2026-10-04',None]);c.commit();c.close();assess(self.run)
        c=connect(self.run);self.assertEqual(c.execute("SELECT status FROM open_issues WHERE issue_id='manual-comparison'").fetchone()[0],'open');c.close()
    def test_verified_field_closes_generated_rfi_but_preserves_manual_and_reopens(self):
        self.valid_pair();c=connect(self.run);oid,sid=c.execute('SELECT offer_id,supplier_id FROM offers ORDER BY offer_id LIMIT 1').fetchone();c.execute("UPDATE facts SET evidence_status='requires_confirmation' WHERE offer_id=? AND field_key='total'",[oid]);c.execute("INSERT INTO open_issues VALUES (?,?,?,?,?,?,?,?,?,?)",['manual-field','FIRMA:'+sid,'critical','critical_field','Elle açılmış teyit sorusu','[]','[]','open','2026-10-04',None]);c.commit();c.close();run_gate(self.run,date(2026,10,3))
        c=connect(self.run);iid=c.execute("SELECT issue_id FROM open_issues WHERE question=?",['Confirm the total field for offer '+oid+'.']).fetchone()[0];c.execute("UPDATE facts SET evidence_status='verified' WHERE offer_id=? AND field_key='total'",[oid]);c.commit();c.close();run_gate(self.run,date(2026,10,3))
        c=connect(self.run);self.assertEqual(c.execute('SELECT status FROM open_issues WHERE issue_id=?',[iid]).fetchone()[0],'resolved');self.assertEqual(c.execute("SELECT status FROM open_issues WHERE issue_id='manual-field'").fetchone()[0],'open');c.execute("UPDATE facts SET evidence_status='requires_confirmation' WHERE offer_id=? AND field_key='total'",[oid]);c.commit();c.close();run_gate(self.run,date(2026,10,3));c=connect(self.run);self.assertEqual(c.execute('SELECT status,resolved_at FROM open_issues WHERE issue_id=?',[iid]).fetchone(),('open',None));c.close()
    def test_reused_snapshot_is_the_explicit_active_revision(self):
        self.valid_pair();self.finish();original=get_meta(self.run,'decision_summary');a=freeze(self.run);set_meta(self.run,'decision_summary','Başka revizyon.');b=freeze(self.run);self.assertNotEqual(a['sha256'],b['sha256']);set_meta(self.run,'decision_summary',original);again=freeze(self.run);self.assertEqual(again['sha256'],a['sha256']);self.assertEqual(load_latest(self.run)[1]['sha256'],a['sha256'])

    def test_manually_edited_generated_question_is_not_closed(self):
        from generated_issues import GeneratedIssues
        self.valid_pair();c=connect(self.run);g=GeneratedIssues(c,'fixture');iid=g.emit('basis','IDARE','critical','comparability_contract','Otomatik soru',[]);g.finish();c.execute('UPDATE open_issues SET question=? WHERE issue_id=?',['Kullanıcının değiştirip açık tuttuğu soru',iid]);c.commit();c.close()
        c=connect(self.run);GeneratedIssues(c,'fixture').finish();self.assertEqual(c.execute('SELECT status,question FROM open_issues WHERE issue_id=?',[iid]).fetchone(),('open','Kullanıcının değiştirip açık tuttuğu soru'));c.commit();c.close()

    def test_existing_unowned_issue_id_collision_is_preserved(self):
        from generated_issues import GeneratedIssues
        from common import stable_id
        self.valid_pair();c=connect(self.run);iid=stable_id('issue','generated','fixture','basis',length=22);c.execute('INSERT INTO open_issues VALUES (?,?,?,?,?,?,?,?,?,?)',[iid,'IDARE','critical','comparability_contract','Elle açılan soru','[]','[]','open','2026-10-04',None]);g=GeneratedIssues(c,'fixture');new=g.emit('basis','IDARE','critical','comparability_contract','Otomatik soru',[]);g.finish();self.assertNotEqual(iid,new);GeneratedIssues(c,'fixture').finish();self.assertEqual(c.execute('SELECT status,question FROM open_issues WHERE issue_id=?',[iid]).fetchone(),('open','Elle açılan soru'));c.commit();c.close()

    def test_model_issue_cannot_take_over_generated_ownership(self):
        from generated_issues import GeneratedIssues
        from record_results import _record_facts
        self.valid_pair();c=connect(self.run);g=GeneratedIssues(c,'fixture');iid=g.emit('basis','IDARE','critical','comparability_contract','Otomatik soru',[]);g.finish()
        _record_facts(c,{'task_id':'model-fixture'}, {}, {'open_issues':[{'issue_id':iid,'owner':'IDARE','severity':'critical','issue_type':'comparability_contract','question':'Modelin teyit sorusu'}]},'semantic_llm')
        GeneratedIssues(c,'fixture').finish();self.assertEqual(c.execute('SELECT status,question FROM open_issues WHERE issue_id=?',[iid]).fetchone(),('resolved','Otomatik soru'));self.assertEqual(c.execute("SELECT count(*) FROM open_issues WHERE question='Modelin teyit sorusu' AND status='open'").fetchone()[0],1);c.commit();c.close()

    def test_assumed_critical_field_still_blocks_decision(self):
        self.valid_pair();c=connect(self.run);c.execute("UPDATE facts SET evidence_status='assumption' WHERE field_key='total'");c.commit();c.close();result=run_gate(self.run,date(2026,10,3));self.assertEqual(result['status'],'PRELIMINARY_ONLY');self.assertTrue(any(r.get('field')=='total' and r.get('status')=='assumption' for r in result['reasons']))

    def test_bad_active_snapshot_pointer_does_not_fall_back(self):
        self.valid_pair();self.finish();set_meta(self.run,'reviewed_snapshot',{'snapshot_id':'missing','path':'reviewed/missing.json','sha256':'bad'})
        with self.assertRaisesRegex(RuntimeError,'pointer mismatch'):load_latest(self.run)

    def test_reassessment_timestamp_does_not_change_stable_snapshot(self):
        self.valid_pair();self.finish();a=freeze(self.run);assess(self.run);b=freeze(self.run);self.assertEqual(a['sha256'],b['sha256']);self.assertTrue(b['reused'])

    def test_historical_calculation_cannot_support_current_claim(self):
        from record_results import _record_decision_summary
        from claim_registry import rebuild
        self.valid_pair();self.finish();self.add_old_scoring();c=connect(self.run)
        _record_decision_summary(c,{'task_id':'history-claim-fixture'},{'summary':'Geçmiş puan denemesi.','recommendation':None,'claims':[{'type':'calculated','text':'Fiyat puanı 999.','calculation_ids':['zz-old-score'],'assertions':[{'source_type':'calculation','source_id':'zz-old-score','path':'price_scores.0.price_score_0_10','operator':'approx','value_number':999,'tolerance':0}]}]});c.commit();c.close()
        self.assertEqual(rebuild(self.run)['status'],'FAIL');c=connect(self.run);self.assertEqual(c.execute("SELECT support_status FROM claims WHERE source_task_id='history-claim-fixture'").fetchone()[0],'unsupported');c.close()

    def test_final_validator_uses_active_snapshot_and_current_calculation_count(self):
        from validate_run import validate
        self.valid_pair();self.finish();self.add_old_scoring();original=get_meta(self.run,'decision_summary');a=freeze(self.run);set_meta(self.run,'decision_summary','Başka revizyon.');freeze(self.run);set_meta(self.run,'decision_summary',original);freeze(self.run);self.export()
        result=validate(self.run);self.assertFalse(any('Independent recomputation coverage is incomplete' in r for r in result['blocking_reasons']));self.assertFalse(any('snapshot' in r.casefold() and ('hash' in r.casefold() or 'parent' in r.casefold()) for r in result['blocking_reasons']),result['blocking_reasons'])

    def test_coverage_uses_explicit_reused_snapshot(self):
        from coverage import build
        self.valid_pair();self.finish();original=get_meta(self.run,'decision_summary');a=freeze(self.run);set_meta(self.run,'decision_summary','Başka revizyon.');freeze(self.run);set_meta(self.run,'decision_summary',original);freeze(self.run)
        units=build(self.run)['units'];control=next(u for u in units if 'reviewed-snapshot' in u.get('categories',{}));self.assertEqual(control['id'],a['snapshot_id']);self.assertEqual(control['revision'],a['source_revision'])

    def test_workflow_does_not_fall_back_to_old_artifact(self):
        from workflow import _artifact_bound_to_snapshot
        from active_records import current_artifacts
        self.valid_pair();self.finish();snap,row=load_latest(self.run);c=connect(self.run);a=current_artifacts(c)['xlsx'];c.execute('INSERT INTO artifacts VALUES (?,?,?,?,?,?,?)',['new-unbound-xlsx','xlsx',a['relative_path'],a['sha256'],'[]','ready','2099-01-01']);c.commit();c.close()
        self.assertFalse(_artifact_bound_to_snapshot(self.run,'xlsx',row['sha256']))
