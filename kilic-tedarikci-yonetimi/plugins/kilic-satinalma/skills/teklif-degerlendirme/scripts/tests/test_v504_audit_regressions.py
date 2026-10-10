from __future__ import annotations

import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import types
import unittest
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path
from unittest import mock

os.environ['TEKLIF_TEST_SQLITE'] = '1'
SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook

from adapters.archive_adapter import extract_zip
from common import json_text
from context_packet import task_packet
from cost_bridge import build as build_cost_bridge
from data_quality import assess as assess_data_quality
from db import connect, get_meta, init_db, set_meta
from independent_recompute import run as independent_recompute
from procurement_engine import run_engine
from procurement_math import present_value_schedule
from record_results import record_task_result
from report_export import _write_table
from task_queue import _add_task
from workflow import _token_budget, advance, prepare


class V504AuditRegressionTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix='teklif-v504-audit-'))
        self.src = self.root / 'source'
        self.run = self.root / 'run'
        self.src.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def offer(self, name: str, total, *, currency='EUR', quantity=10, unit='adet', product='Pompa',
              payment='%100 60 gün', vat='hariç', native_date=False, line_totals=False):
        wb = Workbook(); ws = wb.active; ws.title = 'Teklif'
        qdate = datetime(2026, 10, 3) if native_date else '03.10.2026'
        rows = [
            ('Fiyat Teklifi',''), ('Teklif Tarihi',qdate), ('Geçerlilik','90 gün'),
            ('Genel Toplam',total), ('Para Birimi',currency), ('KDV',vat),
            ('Miktar',quantity), ('Birim',unit), ('Ürün',product), ('Teslim','30 gün'),
            ('Ödeme',payment), ('Incoterm','DDP'), ('Garanti','24 ay'),
        ]
        for r,(k,v) in enumerate(rows,1): ws.cell(r,1,k); ws.cell(r,2,v)
        if native_date: ws['B2'].number_format='dd.mm.yyyy'
        if line_totals:
            ws.append(['Malzeme','Miktar','Birim Fiyat','Tutar'])
            ws.append(['Kalem 1',10,100,1000])
            ws.append(['Kalem 2',10,100,1000])
        ws.append(['Tedarikçi\nFirma: ' + Path(name).stem.split('_')[0] + ' Ltd.'])
        wb.save(self.src / name)

    def valid_pair(self, **kwargs):
        self.offer('Alpha_teklif.xlsx',100000,**kwargs)
        self.offer('Beta_teklif.xlsx',105000,**kwargs)
        return prepare(self.src,self.run,'hizli',date(2026,10,3))

    def finish_valid(self):
        step=advance(self.run)
        self.assertEqual(step['status'],'NEEDS_LLM')
        self.assertTrue(all(t['task_kind']=='decision_summary' for t in step['tasks']))
        tid=step['tasks'][0]['task_id']
        con=connect(self.run)
        try:
            row=con.execute("SELECT calculation_id FROM calculations WHERE calculation_type='cost_summary' AND status='complete' ORDER BY calculation_id LIMIT 1").fetchone()
            self.assertIsNotNone(row); cid=row[0]
        finally: con.close()
        result={
            'summary':'Maliyet karşılaştırması tamamlandı.', 'recommendation':None,
            'claims':[{
                'type':'calculated','text':'Karşılaştırılabilir maliyet hesaplanmıştır.',
                'calculation_ids':[cid],
                'assertions':[{'source_type':'calculation','source_id':cid,'path':'comparable_total','operator':'exists'}],
            }]
        }
        record_task_result(self.run,tid,result,model_name='test-model',reasoning_effort='medium',input_tokens=100,cached_input_tokens=50,output_tokens=20)
        return advance(self.run)

    def test_uncalculated_formula_never_becomes_verified_total(self):
        self.offer('Alpha_teklif.xlsx','=100000+5000')
        self.offer('Beta_teklif.xlsx',110000)
        prepare(self.src,self.run,'hizli',date(2026,10,3))
        con=connect(self.run)
        try:
            row=con.execute("""SELECT f.value_number,f.evidence_status FROM facts f JOIN documents d ON d.offer_id=f.offer_id
                               WHERE d.filename='Alpha_teklif.xlsx' AND f.field_key='total' LIMIT 1""").fetchone()
        finally: con.close()
        self.assertIsNotNone(row)
        self.assertNotEqual(row[1],'verified')
        self.assertNotEqual(row[0],1000005000)

    def test_payment_ratio_amount_principal_mismatch_is_rejected(self):
        with self.assertRaises(ValueError):
            present_value_schedule(100000,[{'ratio':1,'amount':1,'days':0}],0.10)

    def test_different_product_and_unit_block_ranking(self):
        self.offer('Alpha_teklif.xlsx',100000,unit='kg',product='Balık yemi')
        self.offer('Beta_teklif.xlsx',105000,unit='adet',product='Su pompası')
        prepare(self.src,self.run,'hizli',date(2026,10,3))
        comp=get_meta(self.run,'comparability_contract',{},parse_json=True)
        self.assertFalse(comp['ranking_allowed'])
        self.assertIn('quantity',comp['failed_dimensions'])
        self.assertIn('scope_identity',comp['failed_dimensions'])

    def test_line_total_grand_total_mismatch_is_critical(self):
        self.offer('Alpha_teklif.xlsx',100000,line_totals=True)
        self.offer('Beta_teklif.xlsx',105000,line_totals=True)
        result=prepare(self.src,self.run,'hizli',date(2026,10,3))
        self.assertEqual(result['data_quality']['status'],'FAIL')
        con=connect(self.run)
        try:
            n=con.execute("SELECT count(*) FROM data_quality_findings WHERE check_id='line_total_grand_total_mismatch' AND severity='critical' AND status='open'").fetchone()[0]
        finally: con.close()
        self.assertGreaterEqual(n,2)

    def test_false_numeric_claim_with_unrelated_calculation_is_blocked(self):
        self.valid_pair(); step=advance(self.run); self.assertEqual(step['status'],'NEEDS_LLM')
        tid=step['tasks'][0]['task_id']
        con=connect(self.run)
        try:
            cid=con.execute("SELECT calculation_id FROM calculations WHERE calculation_type='cost_summary' AND status='complete' LIMIT 1").fetchone()[0]
        finally: con.close()
        text='Tüm tekliflerin toplam bedeli 1 EUR olarak hesaplanmıştır.'
        result={'summary':text,'recommendation':None,'claims':[{
            'type':'calculated','text':text,'calculation_ids':[cid],
            'assertions':[{'source_type':'calculation','source_id':cid,'path':'comparable_total','operator':'exists'}]
        }]}
        record_task_result(self.run,tid,result,model_name='test-model',reasoning_effort='medium',input_tokens=10,output_tokens=5)
        step=advance(self.run)
        self.assertEqual(step['status'],'NEEDS_LLM')
        self.assertEqual(step['summary_retry']['status'],'RETRY_PENDING')
        self.assertNotEqual(step['tasks'][0]['task_id'],tid)
        self.assertFalse((self.run/'output'/'Teklif_Degerlendirme.xlsx').exists())
        from workflow import _finish_outputs
        self.assertEqual(_finish_outputs(self.run,preliminary=False)['status'],'BLOCKED')
        con=connect(self.run)
        try: status=con.execute("SELECT support_status FROM claims LIMIT 1").fetchone()[0]
        finally: con.close()
        self.assertEqual(status,'unsupported')

    def _offer_docs(self):
        con=connect(self.run)
        try:
            rows=con.execute("SELECT o.offer_id,d.document_id FROM offers o JOIN documents d ON d.offer_id=o.offer_id ORDER BY o.offer_id").fetchall()
            return rows
        finally: con.close()

    def test_cross_offer_evidence_reference_is_rejected_atomically(self):
        self.valid_pair(); a,b=self._offer_docs()[:2]
        con=connect(self.run)
        try:
            foreign=con.execute("SELECT evidence_id FROM facts WHERE offer_id=? AND field_key='total'",[b[0]]).fetchone()[0]
            before=con.execute("SELECT value_number,evidence_id FROM facts WHERE offer_id=? AND field_key='total'",[a[0]]).fetchone()
            tid=_add_task(con,'EXTRACT','semantic_extraction','offer_semantic_fields',{'document_ids':[a[1]]},document_id=a[1],offer_id=a[0]); con.commit()
        finally: con.close()
        with self.assertRaises(ValueError):
            record_task_result(self.run,tid,{'facts':[{'field_key':'total','value_number':1,'evidence_status':'verified','evidence_id':foreign}],'evidence':[]},model_name='test-model',reasoning_effort='medium',input_tokens=1,output_tokens=1)
        con=connect(self.run)
        try:
            after=con.execute("SELECT value_number,evidence_id FROM facts WHERE offer_id=? AND field_key='total'",[a[0]]).fetchone()
            task=con.execute("SELECT status FROM task_queue WHERE task_id=?",[tid]).fetchone()[0]
        finally: con.close()
        self.assertEqual(tuple(before),tuple(after)); self.assertNotEqual(task,'completed')

    def test_nonexistent_evidence_locator_is_rejected(self):
        self.valid_pair(); a=self._offer_docs()[0]
        con=connect(self.run)
        try:
            tid=_add_task(con,'EXTRACT','semantic_extraction','offer_semantic_fields',{'document_ids':[a[1]]},document_id=a[1],offer_id=a[0]); con.commit()
        finally: con.close()
        payload={'facts':[{'field_key':'total','value_number':1,'evidence_status':'verified','evidence_client_id':'e'}],
                 'evidence':[{'client_id':'e','document_id':a[1],'locator':{'sheet':'DOES_NOT_EXIST','cell':'Z999999'},'quote':'TOTAL = 1 EUR','evidence_status':'verified'}]}
        with self.assertRaises(ValueError):
            record_task_result(self.run,tid,payload,model_name='test-model',reasoning_effort='medium',input_tokens=1,output_tokens=1)

    def test_missing_token_receipt_cannot_partially_write_decision(self):
        self.valid_pair(); step=advance(self.run); tid=step['tasks'][0]['task_id']
        with self.assertRaises(ValueError):
            record_task_result(self.run,tid,{'summary':'x','recommendation':None,'claims':[]},model_name='test-model',reasoning_effort='medium',input_tokens=1)
        con=connect(self.run)
        try:
            self.assertIsNone(con.execute("SELECT value FROM run_meta WHERE key='decision_summary'").fetchone())
            self.assertEqual(con.execute("SELECT count(*) FROM claims").fetchone()[0],0)
        finally: con.close()

    def test_context_packet_respects_hard_serialized_bound(self):
        init_db(self.run)
        con=connect(self.run)
        try:
            con.execute("INSERT INTO documents(document_id,relative_path,filename,format_family,document_role,sha256,size_bytes,parse_status,active) VALUES ('D1','x.txt','x.txt','text','offer','abc',1,'parsed',true)")
            huge='X'*600000
            con.execute("INSERT INTO document_units(unit_id,document_id,unit_type,unit_index,locator_json,text_content,structured_json,content_sha256) VALUES ('U1','D1','text',1,'{}',?,?,?)",[huge,json_text({'blob':huge}),'hash'])
            tid=_add_task(con,'EXTRACT','semantic_extraction','document_classification',{},document_id='D1'); con.commit()
        finally: con.close()
        packet=task_packet(self.run,tid,24,80000)
        size=len(json.dumps(packet,ensure_ascii=False,default=str,separators=(',',':')))
        self.assertLessEqual(size,80000)

    def test_unknown_historical_token_usage_fails_closed(self):
        init_db(self.run); set_meta(self.run,'profile','hizli')
        con=connect(self.run)
        try:
            con.execute("INSERT INTO task_queue(task_id,phase,role,task_kind,payload_json,priority,critical,status,attempts,created_at) VALUES ('T','X','r','decision_summary','{}',1,true,'completed',1,'now')")
            con.execute("INSERT INTO model_receipts(receipt_id,task_id,model_name,reasoning_effort,input_tokens,cached_input_tokens,output_tokens,recorded_at) VALUES ('R','T','m','medium',NULL,0,NULL,'now')")
            con.commit()
        finally: con.close()
        policy=json.loads((SCRIPTS.parent/'config'/'policy.json').read_text())
        budget=_token_budget(self.run,policy)
        self.assertTrue(budget['usage_unknown']); self.assertTrue(budget['hard_stop'])

    def test_native_excel_datetime_is_serializable(self):
        self.offer('Alpha_teklif.xlsx',100000,native_date=True)
        self.offer('Beta_teklif.xlsx',105000)
        result=prepare(self.src,self.run,'hizli',date(2026,10,3))
        self.assertEqual(result['status'],'PREPARED')

    def test_mixed_currency_cost_bridge_is_unverified_without_provenance(self):
        self.offer('Alpha_teklif.xlsx',100000,currency='EUR')
        self.offer('Beta_teklif.xlsx',105000,currency='USD')
        prepare(self.src,self.run,'hizli',date(2026,10,3),financial_params={'base_currency':'TRY','fx':{'EUR':40,'USD':36}})
        bridge=build_cost_bridge(self.run)
        self.assertEqual(bridge['status'],'UNVERIFIED')
        self.assertEqual(bridge['pairwise_differences'],[])

    def test_weighted_scoring_tamper_is_caught_independently(self):
        self.valid_pair()
        con=connect(self.run)
        try: offers=[r[0] for r in con.execute("SELECT offer_id FROM offers WHERE status='active' ORDER BY offer_id").fetchall()]
        finally: con.close()
        set_meta(self.run,'scoring_parameters',{'criteria':[{'id':'quality','weight':100,'scores':{offers[0]:8,offers[1]:6}}]})
        run_engine(self.run)
        con=connect(self.run)
        try:
            row=con.execute("SELECT calculation_id,result_json FROM calculations WHERE calculation_type='scoring' AND status='complete' ORDER BY updated_at DESC LIMIT 1").fetchone()
            obj=json.loads(row[1]); self.assertEqual(obj.get('status'),'weighted_complete')
            obj['weighted_scores'][0]['total_score_0_10']='999'
            con.execute("UPDATE calculations SET result_json=? WHERE calculation_id=?",[json.dumps(obj),row[0]]); con.commit()
        finally: con.close()
        result=independent_recompute(self.run)
        self.assertEqual(result['status'],'FAIL')
        self.assertTrue(any(x['calculation_type']=='scoring' for x in result['failed']))

    def test_duplicate_zip_member_is_rejected(self):
        z=self.root/'dup.zip'
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            with zipfile.ZipFile(z,'w') as zf:
                zf.writestr('a.txt','one'); zf.writestr('a.txt','two')
        with self.assertRaises(ValueError): extract_zip(z,self.root/'unzipped')

    def test_spreadsheet_formula_injection_is_escaped(self):
        wb=Workbook(); ws=wb.active; _write_table(ws,['Supplier text'],[['=1+1']])
        self.assertNotEqual(ws['A2'].data_type,'f')
        self.assertTrue(str(ws['A2'].value).startswith("'"))

    def test_completed_run_reentry_is_idempotent(self):
        self.valid_pair(); first=self.finish_valid(); self.assertEqual(first['status'],'VERIFIED')
        con=connect(self.run)
        try:
            before=con.execute("SELECT snapshot_id,sha256,relative_path FROM reviewed_snapshots ORDER BY created_at").fetchall()
        finally: con.close()
        second=advance(self.run); self.assertEqual(second['status'],'VERIFIED')
        con=connect(self.run)
        try: after=con.execute("SELECT snapshot_id,sha256,relative_path FROM reviewed_snapshots ORDER BY created_at").fetchall()
        finally: con.close()
        self.assertEqual(before,after)

    def test_resume_request_preserves_analysis_and_parameter_payloads(self):
        import environment_check
        finance={'discount_rates':{'EUR':0.1}}; scoring={'criteria':[]}
        with mock.patch.object(environment_check,'check',return_value={'status':'BLOCKED_ENVIRONMENT','blocking_reasons':['duckdb'],'missing_components':[{'component':'duckdb','required':'duckdb>=1.4,<2','installed':None,'reason':'Python paketi bulunamadı'}]}):
            old=os.environ.pop('TEKLIF_TEST_SQLITE',None)
            try: result=prepare(self.src,self.run,'standart',date(2026,9,1),finance,scoring)
            finally:
                if old is not None: os.environ['TEKLIF_TEST_SQLITE']=old
        rr=json.loads(Path(result['resume_request']).read_text())
        self.assertEqual(rr['analysis_date'],'2026-09-01')
        self.assertEqual(rr['financial_parameters'],finance)
        self.assertEqual(rr['scoring_parameters'],scoring)
        self.assertIn('--resume-request',result['resume_command'])

    def test_verified_output_is_formula_free_and_source_archive_is_hash_bound(self):
        import hashlib
        self.valid_pair(); final=self.finish_valid(); self.assertEqual(final['status'],'VERIFIED')
        out=self.run/'output'/'Teklif_Degerlendirme.xlsx'
        wb=load_workbook(out,data_only=False)
        formulas=[]
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type=='f' or (isinstance(cell.value,str) and cell.value.startswith('=')):
                        formulas.append(f"{ws.title}!{cell.coordinate}")
        self.assertEqual(formulas,[])
        con=connect(self.run)
        try:
            docs=con.execute("SELECT sha256, metadata_json FROM documents WHERE active=true").fetchall()
        finally: con.close()
        self.assertGreaterEqual(len(docs),2)
        for sha, meta_json in docs:
            rel=(json.loads(meta_json or '{}') or {}).get('archived_source_path')
            self.assertTrue(rel)
            fp=self.run/rel
            self.assertTrue(fp.exists())
            self.assertEqual(hashlib.sha256(fp.read_bytes()).hexdigest(),sha)

    def test_tampered_evidence_is_caught_by_data_quality_gate(self):
        self.valid_pair()
        con=connect(self.run)
        try:
            rows=con.execute("SELECT fact_id,offer_id,evidence_id FROM facts WHERE field_key='total' AND evidence_status='verified' ORDER BY offer_id").fetchall()
            self.assertGreaterEqual(len(rows),2)
            foreign=rows[1][2]
            con.execute("UPDATE facts SET evidence_id=? WHERE fact_id=?",[foreign,rows[0][0]])
            con.commit()
        finally: con.close()
        dq=assess_data_quality(self.run)
        self.assertEqual(dq['status'],'FAIL')
        self.assertTrue(any(f.get('check_id')=='evidence_offer_mismatch' and f.get('severity')=='critical' for f in dq['findings']))

    def test_legacy_excel_com_disables_macros_events_and_links(self):
        from adapters.office_legacy_adapter import parse_xls
        source=self.root/'legacy.xls'; source.write_bytes(b'fake')
        calls={}
        class MacroSheets: Count=0
        class FakeWB:
            Excel4MacroSheets=MacroSheets()
            def SaveAs(self,path,FileFormat=None):
                wb=Workbook(); ws=wb.active; ws['A1']='Genel Toplam'; ws['B1']=100; wb.save(path)
            def Close(self,*args): calls['closed']=True
        class Books:
            def Open(self,*args,**kwargs):
                calls['open_args']=args; calls['open_kwargs']=kwargs; calls['security_at_open']=excel.AutomationSecurity; calls['events_at_open']=excel.EnableEvents
                return FakeWB()
        class FakeExcel:
            def __init__(self):
                self.Visible=True; self.DisplayAlerts=True; self.EnableEvents=True; self.AutomationSecurity=1; self.Workbooks=Books()
            def Quit(self): calls['quit']=True
        excel=FakeExcel()
        client=types.ModuleType('win32com.client'); client.DispatchEx=lambda name: excel
        package=types.ModuleType('win32com'); package.client=client
        with mock.patch.dict(sys.modules,{'win32com':package,'win32com.client':client}), mock.patch('adapters.office_legacy_adapter.platform.system',return_value='Windows'):
            result=parse_xls(source)
        self.assertEqual(result.status,'read')
        self.assertEqual(calls['security_at_open'],3)
        self.assertFalse(calls['events_at_open'])
        self.assertEqual(calls['open_kwargs'].get('UpdateLinks'),0)
        self.assertTrue(calls['open_kwargs'].get('ReadOnly'))
        self.assertFalse(calls['open_kwargs'].get('AddToMru'))
        self.assertEqual(excel.AutomationSecurity,1)
        self.assertTrue(calls.get('quit'))

    def test_legacy_word_com_disables_macros_and_opens_read_only(self):
        from adapters.office_legacy_adapter import parse_doc
        source=self.root/'legacy.doc'; source.write_bytes(b'fake')
        calls={}
        class FakeDoc:
            def __init__(self): self.Content=types.SimpleNamespace(Text='Teklif metni')
            def Close(self,*args): calls['closed']=True
        class Docs:
            def Open(self,*args,**kwargs):
                calls['open_kwargs']=kwargs; calls['security_at_open']=word.AutomationSecurity
                return FakeDoc()
        class FakeWord:
            def __init__(self): self.Visible=True; self.AutomationSecurity=1; self.Documents=Docs()
            def Quit(self): calls['quit']=True
        word=FakeWord()
        client=types.ModuleType('win32com.client'); client.DispatchEx=lambda name: word
        package=types.ModuleType('win32com'); package.client=client
        with mock.patch.dict(sys.modules,{'win32com':package,'win32com.client':client}), mock.patch('adapters.office_legacy_adapter.platform.system',return_value='Windows'):
            result=parse_doc(source)
        self.assertEqual(result.status,'read')
        self.assertEqual(calls['security_at_open'],3)
        self.assertTrue(calls['open_kwargs'].get('ReadOnly'))
        self.assertFalse(calls['open_kwargs'].get('AddToRecentFiles'))
        self.assertEqual(word.AutomationSecurity,1)
        self.assertTrue(calls.get('quit'))

    def test_blocked_validate_run_cli_returns_nonzero(self):
        init_db(self.run)
        env=dict(os.environ)
        p=subprocess.run([sys.executable,str(SCRIPTS/'validate_run.py'),'--run',str(self.run)],capture_output=True,text=True,env=env)
        self.assertEqual(p.returncode,3)
        self.assertEqual(json.loads(p.stdout)['status'],'BLOCKED')


    def test_delivery_relative_payment_adds_delivery_lead_time(self):
        from normalize import _payment_schedule
        self.assertEqual(_payment_schedule('%100 teslimden 60 gün sonra',30),[{'ratio':1.0,'days':90}])
        self.assertIsNone(_payment_schedule('%100 teslimden 60 gün sonra',None))

    def test_invalid_payment_schedule_blocks_decision_end_to_end(self):
        self.valid_pair()
        con=connect(self.run)
        try:
            row=con.execute("SELECT fact_id,normalized_json FROM facts WHERE field_key='payment' ORDER BY offer_id LIMIT 1").fetchone()
            obj=json.loads(row[1] or '{}')
            obj['payment_schedule']=[{'ratio':1.0,'amount':1,'days':0}]
            con.execute("UPDATE facts SET normalized_json=? WHERE fact_id=?",[json.dumps(obj),row[0]])
            con.commit()
        finally: con.close()
        set_meta(self.run,'financial_parameters',{'discount_rates':{'EUR':0.10}})
        run_engine(self.run)
        result=independent_recompute(self.run)
        self.assertIn(result['status'],{'FAIL','UNVERIFIED'})
        from decision_gate import run_gate
        gate=run_gate(self.run,date(2026,10,3))
        self.assertEqual(gate['status'],'PRELIMINARY_ONLY')

    def test_mixed_currency_bridge_uses_one_sourced_base_currency_and_date(self):
        self.offer('Alpha_teklif.xlsx',100000,currency='EUR')
        self.offer('Beta_teklif.xlsx',105000,currency='USD')
        fp={
            'base_currency':'TRY','base_date':'2026-10-03',
            'fx':{
                'EUR':{'rate':40,'unit':1,'source':'TEST_SOURCE','date':'2026-10-03'},
                'USD':{'rate':36,'unit':1,'source':'TEST_SOURCE','date':'2026-10-03'},
            }
        }
        prepare(self.src,self.run,'hizli',date(2026,10,3),financial_params=fp)
        bridge=build_cost_bridge(self.run)
        self.assertEqual(bridge['status'],'PASS')
        self.assertEqual(len(bridge['pairwise_differences']),1)
        diff=bridge['pairwise_differences'][0]
        self.assertEqual(diff['currency'],'TRY')
        self.assertEqual(diff['base_date'],'2026-10-03')
        self.assertEqual(diff['basis'],'fx_normalized')
        self.assertEqual(abs(float(diff['difference'])),220000.0)

    def test_pdf_content_check_requires_critical_snapshot_values(self):
        import qa
        fake=types.SimpleNamespace(pages=[types.SimpleNamespace(extract_text=lambda:'Alpha teklif Beta teklif 100000 105000 sufficient')])
        with mock.patch.dict(sys.modules,{'pypdf':types.SimpleNamespace(PdfReader=lambda _:fake)}):
            self.valid_pair()
            from reviewed_snapshot import freeze
            # decision status is sufficient in this ordinary, comparable pair.
            set_meta(self.run,'decision_status','SUFFICIENT')
            snap=freeze(self.run)
            errors=qa._pdf_content_checks(self.root/'fake.pdf',self.run)
            self.assertFalse(any('critical numeric value' in e for e in errors),errors)
        fake_bad=types.SimpleNamespace(pages=[types.SimpleNamespace(extract_text=lambda:'Alpha teklif Beta teklif sufficient')])
        with mock.patch.dict(sys.modules,{'pypdf':types.SimpleNamespace(PdfReader=lambda _:fake_bad)}):
            errors=qa._pdf_content_checks(self.root/'fake.pdf',self.run)
            self.assertTrue(any('critical numeric value' in e for e in errors),errors)

    def test_blocked_environment_workflow_cli_returns_3_and_resume_file(self):
        import environment_check
        with mock.patch.object(environment_check,'check',return_value={'status':'BLOCKED_ENVIRONMENT','blocking_reasons':['duckdb'],'missing_components':[{'component':'duckdb','required':'duckdb>=1.4,<2','installed':None,'reason':'Python paketi bulunamadı'}]}):
            # Exercise main() directly so the patched preflight is visible in-process.
            argv=['workflow.py','prepare','--source',str(self.src),'--run',str(self.run),'--profile','standart','--analysis-date','2026-09-01']
            old_env=os.environ.pop('TEKLIF_TEST_SQLITE',None)
            try:
                with mock.patch.object(sys,'argv',argv):
                    from workflow import main as workflow_main
                    with self.assertRaises(SystemExit) as exc:
                        workflow_main()
                    self.assertEqual(exc.exception.code,3)
            finally:
                if old_env is not None: os.environ['TEKLIF_TEST_SQLITE']=old_env
        rr=json.loads((self.run/'resume_request.json').read_text())
        self.assertEqual(rr['analysis_date'],'2026-09-01')
        self.assertEqual(rr['profile'],'standart')


if __name__ == '__main__':
    unittest.main()
