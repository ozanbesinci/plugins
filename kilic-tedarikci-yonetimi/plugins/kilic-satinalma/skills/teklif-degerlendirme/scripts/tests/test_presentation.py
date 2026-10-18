from pathlib import Path
import unittest,tempfile,shutil,json
from openpyxl import Workbook,load_workbook
from display_values import label,structured,display_cell
from print_layout import build_print_workbook
class PresentationTests(unittest.TestCase):
 def test_pdf_uses_human_readable_recommendation_without_internal_record(self):
  import pymupdf
  from pdf_renderer import render_pdf
  with tempfile.TemporaryDirectory() as folder:
   path=Path(folder)/'recommendation.pdf'
   render_pdf({'recommendation':{'offer_id':'internal-id','action':'SHORTLIST'},
               'recommendation_text':'Alfa teklifini kısa listeye al.'},path)
   with pymupdf.open(path) as pdf:text='\n'.join(page.get_text() for page in pdf)
   self.assertIn('Önerilen aksiyon: Alfa teklifini kısa listeye al.',text)
   self.assertNotIn('offer_id',text)
   self.assertNotIn('internal-id',text)
   self.assertNotIn('SHORTLIST',text)
 def test_codes_are_turkish_but_source_quotes_are_verbatim(self):
  self.assertEqual(display_cell('Kanitlar',4,'requires_confirmation'),'Teyit bekleniyor')
  self.assertEqual(display_cell('Kanitlar',7,'Source says verified'),'Source says verified')
  self.assertEqual(label('SUFFICIENT'),'Karar için yeterli')
  self.assertEqual(structured('{\"note\":\"Cost score is not a total decision score.\"}'),'Not: Maliyet puanı toplam karar puanı değildir.')
 def test_locator_and_assertion_display_preserves_all_values(self):
  self.assertEqual(structured('{"sheet":"Teklif","cell":"B4"}'),'Sayfa: Teklif\nHücre: B4')
  txt=structured('[{"source_id":"id123","operator":"equals","value_number":115000}]')
  self.assertIn('id123',txt);self.assertIn('115000',txt);self.assertIn('Eşit',txt)
 def test_print_bands_keep_every_column_and_long_tail_without_shrinking(self):
  root=Path(tempfile.mkdtemp());w=Workbook();s=w.active;s.title='Genis Tablo'
  for c in range(1,20):s.cell(1,c,f'Alan {c}');s.cell(2,c,c);s.column_dimensions[s.cell(1,c).column_letter].width=32
  tail='son işaret';s.cell(2,19,'uzun metin '*250+tail);w.save(root/'input.xlsx');w.close()
  try:
   info=build_print_workbook(root/'input.xlsx',root/'print.xlsx');p=load_workbook(root/'print.xlsx')
   cols=set(c for r in info['mapping'] for c in r['source_columns']);self.assertEqual(cols,set(range(1,20)))
   text=''.join(str(c.value or '').replace('\n','') for sh in p for row in sh for c in row)
   self.assertIn(tail.replace(' ',''),text.replace(' ',''));self.assertGreater(len(p.worksheets),1)
   for sh in p:
    self.assertLessEqual(sum(sh.column_dimensions[sh.cell(1,c).column_letter].width for c in range(1,sh.max_column+1)),145)
    self.assertTrue(all(c.font.sz==11 for row in sh for c in row));self.assertTrue(all(x.height<=305 for k,x in sh.row_dimensions.items() if k>1))
   p.close()
  finally:shutil.rmtree(root)

 def test_grouped_excel_dimensions_keep_tail_column_width(self):
  root=Path(tempfile.mkdtemp());w=Workbook();s=w.active;s.title='Tablo'
  s.append(['Kimlik','Boyut','Detay','Gözlenen']);s.append(['x','y','z','uzun değer'])
  from openpyxl.worksheet.dimensions import ColumnDimension
  s.column_dimensions['C']=ColumnDimension(s,index='C',min=3,max=4,width=32)
  w.save(root/'input.xlsx');w.close()
  try:
   build_print_workbook(root/'input.xlsx',root/'print.xlsx')
   p=load_workbook(root/'print.xlsx')
   self.assertEqual(p.active.column_dimensions['D'].width,32)
   p.close()
  finally:shutil.rmtree(root)

 def test_empty_export_table_is_explicit_and_continuation_keeps_short_values(self):
  from report_export import _write_table
  root=Path(tempfile.mkdtemp());w=Workbook();s=w.active;s.title='Tablo'
  _write_table(s,['Kimlik','Durum','Detay'],[])
  self.assertEqual(s['A2'].value,'Kayıt bulunmuyor.')
  s.delete_rows(2);s.append(['kimlik','Doğrulandı','uzun metin '*300])
  for col in ('A','B','C'):s.column_dimensions[col].width=32
  w.save(root/'input.xlsx');w.close()
  try:
   info=build_print_workbook(root/'input.xlsx',root/'print.xlsx')
   p=load_workbook(root/'print.xlsx')
   self.assertGreater(p.active.max_row,2)
   for row in p.active.iter_rows(min_row=2):self.assertEqual(row[1].value,'Doğrulandı')
   p.close()
  finally:shutil.rmtree(root)
