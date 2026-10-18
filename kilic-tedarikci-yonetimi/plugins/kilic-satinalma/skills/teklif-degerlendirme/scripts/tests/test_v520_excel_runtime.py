import copy
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from scenario_formulas import add_scenarios,validate_formulas,scenario_reference
from matrix_workbook import seal_source_cells
from scenario_pricing import net_price,read_pricing
from scenario_scoring import add_scoring,configuration,recompute,read_scoring,SHEET,SENSITIVITY
from source_scenarios import from_sources

def source_model():
    model={'documents':[],'evidence':[],'facts':[],'suppliers':[],'offers':[],'requirements':[],'exclusions':[],
           'meta':{'scoring_parameters':{'criteria':[{'id':'Kalite','weight':50,'scores':{'A':8,'B':6}},
                      {'id':'Teslim','weight':50,'scores':{'A':7,'B':9}}]}}}
    for oid,amount in [('A',100),('B',120)]:
        model['suppliers'].append({'supplier_id':oid,'display_name':oid+' Firma'})
        model['offers'].append({'offer_id':oid,'supplier_id':oid,'status':'active'})
        model['documents'].append({'document_id':oid,'offer_id':oid,'supplier_id':oid,'sha256':'hash-'+oid})
        model['evidence'].append({'evidence_id':oid,'document_id':oid,'source_sha256':'hash-'+oid,'evidence_status':'verified'})
        model['facts'].extend([{'fact_id':'T'+oid,'offer_id':oid,'supplier_id':oid,'field_key':'total','value_number':amount,'evidence_id':oid,'evidence_status':'verified'},
                               {'fact_id':'C'+oid,'offer_id':oid,'supplier_id':oid,'field_key':'currency','value_text':'TRY','evidence_id':oid,'evidence_status':'verified'}])
    return model

class SourceScenarioTests(unittest.TestCase):
    def test_verified_totals_are_prefilled_without_inventing_tax_fx_or_costs(self):
        scenarios=from_sources(source_model());self.assertEqual(len(scenarios),2)
        self.assertEqual(scenarios[0]['inputs']['unit_price'],100)
        self.assertNotIn('vat_rate',scenarios[0]['inputs']);self.assertNotIn('fx_rate',scenarios[0]['inputs'])
        self.assertIsNone(scenarios[0]['pricing']['price_includes_vat'])
        self.assertNotIn('Miktar ve Kapsam!B',scenario_reference(scenarios[0]['inputs'],scenarios[0]['pricing']))
    def test_unverified_conflicting_and_cross_offer_evidence_do_not_seed_prices(self):
        for mutate in ('unverified','hash','owner','ambiguous'):
            model=source_model()
            if mutate=='unverified':model['facts'][0]['evidence_status']='requires_confirmation'
            if mutate=='hash':model['evidence'][0]['source_sha256']='different'
            if mutate=='owner':model['documents'][0]['offer_id']='B'
            if mutate=='ambiguous':model['facts'].append({**model['facts'][0],'fact_id':'other','value_number':999})
            self.assertFalse(any(s['offer_id']=='A' for s in from_sources(model)),mutate)
    def test_safety_notice_never_exposes_scenario_values(self):
        model=source_model();model['content_status']='SAFETY_NOTICE';self.assertEqual(from_sources(model),[])
    def test_recorded_effective_rate_and_analysis_day_payment_schedule_are_projected(self):
        model=source_model();model['meta']['financial_parameters']={'base_currency':'TRY','discount_rates':{'TRY':{'annual_effective_rate':'0.25'}},'offer_parameters':{'A':{'vat':{'price_basis':'net_ex_vat','rate':0.2,'deductible_share':1},'payment_schedule':[{'ratio':0.3,'days':0},{'ratio':0.7,'days':90}]}}}
        scenario=from_sources(model)[0]
        self.assertEqual(scenario['inputs']['annual_rate'],0.25)
        self.assertEqual(scenario['financial']['payments'][1],{'percent':0.7,'days':90})
        self.assertEqual(scenario['pricing']['price_includes_vat'],0)
    def test_foreign_currency_scenario_uses_target_rate_and_never_source_rate(self):
        model=source_model();model['facts'][1]['value_text']='EUR'
        model['meta']['financial_parameters']={'base_currency':'TRY','discount_rates':{'EUR':0.05,'TRY':0.4}}
        self.assertEqual(from_sources(model)[0]['inputs']['annual_rate'],0.4)
        del model['meta']['financial_parameters']['discount_rates']['TRY']
        self.assertNotIn('annual_rate',from_sources(model)[0]['inputs'])

class PricingScenarioTests(unittest.TestCase):
    def make(self,values,pricing=None):
        wb=Workbook();wb.remove(wb.active);scenario={'id':'S','inputs':values}
        if pricing is not None:scenario['pricing']=pricing
        contract=add_scenarios(wb,[scenario]);seal_source_cells(wb,contract);return wb,contract
    def test_price_unit_sequential_discounts_and_vat_included(self):
        values={'quantity':100,'unit_price':120,'discount':0.1,'vat_rate':0.2}
        pricing={'price_unit':10,'price_includes_vat':1,'discounts':[0.05,0.02]}
        wb,contract=self.make(values,pricing)
        self.assertAlmostEqual(contract['reference']['2']['Miktar ve Kapsam!B'],837.9)
        self.assertIn('/\'Kalem Fiyat Varsayimlari\'!B2',wb['Miktar ve Kapsam']['B2'].value)
        self.assertFalse(wb['Ardisik Iskontolar']['C2'].protection.locked);wb.close()
    def test_no_early_cent_rounding_and_unknown_vat_remains_unknown(self):
        values={'quantity':3,'unit_price':1.23456,'discount':0}
        self.assertAlmostEqual(scenario_reference(values)['Miktar ve Kapsam!B'],3.70368)
        self.assertIsNone(net_price(values,{'price_includes_vat':None}))
    def test_unused_discount_rows_do_not_accumulate_on_round_trip(self):
        wb,contract=self.make({'quantity':1,'unit_price':100,'discount':0})
        self.assertEqual(read_pricing(wb,2,contract)['discounts'],[]);wb.close()
    def test_invalid_price_unit_discount_and_formula_tampering_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'user.xlsx'
            for sheet,cell,value in [('Kalem Fiyat Varsayimlari','B2',0),('Ardisik Iskontolar','C2',1),('Miktar ve Kapsam','B2','=1')]:
                wb,contract=self.make({'quantity':1,'unit_price':100,'discount':0});wb[sheet][cell]=value;wb.save(path);wb.close()
                self.assertEqual(validate_formulas(path,contract)['status'],'FAIL')

class MultiCriterionTests(unittest.TestCase):
    def make(self):
        model=source_model();wb=Workbook();wb.remove(wb.active);contract=add_scenarios(wb)
        add_scoring(wb,model,contract);seal_source_cells(wb,contract);return wb,contract
    def test_base_plus_two_independent_alternatives_and_tied_ranks(self):
        config=configuration(source_model());self.assertEqual(len(config['sensitivity']),2)
        result=recompute(config);self.assertEqual(result['Baz']['scores'],{'A':7.5,'B':7.5})
        self.assertEqual(result['Baz']['ranks'],{'A':1,'B':1})
        self.assertNotEqual(result['Ornek 1']['scores'],result['Ornek 2']['scores'])
    def test_lower_higher_minmax_and_equal_values(self):
        config=configuration(source_model());c=config['criteria'][0]
        c.update(direction='lower',method='proportional',values={'A':100,'B':120})
        config['criteria'][1]['values']={'A':10,'B':10}
        self.assertAlmostEqual(recompute(config)['Baz']['scores']['B'],9.166666666666668)
        c.update(direction='higher',method='minmax',values={'A':1,'B':2})
        self.assertEqual(recompute(config)['Baz']['scores']['A'],5)
        c['values']={'A':0,'B':0};self.assertEqual(recompute(config)['Baz']['scores']['A'],10)
    def test_one_independent_supplier_never_ranked(self):
        model=source_model();model['offers'][1]['supplier_id']='A';self.assertIsNone(configuration(model))
    def test_workbook_weights_are_editable_and_formula_map_remains_immutable(self):
        wb,contract=self.make();self.assertFalse(wb[SHEET]['B2'].protection.locked)
        self.assertTrue(wb[SHEET].cell(contract['scoring']['total_row'],5).protection.locked)
        wb[SENSITIVITY]['C2']=35;wb[SENSITIVITY]['C3']=65
        config=read_scoring(wb,contract);self.assertEqual(config['sensitivity'][0]['weights']['Kalite'],35)
        wb.close()
    def test_duplicate_alternative_weight_sets_are_rejected(self):
        config=configuration(source_model());config['sensitivity'][1]['weights']=copy.deepcopy(config['sensitivity'][0]['weights'])
        with self.assertRaisesRegex(ValueError,'farklı'):recompute(config)

class ScenarioImportV520Tests(unittest.TestCase):
    def test_pricing_and_weight_changes_survive_new_pdf_pair_without_source_write(self):
        from test_report_slice import seed
        from report_delivery import export_preliminary
        from scenario_import import import_scenarios
        from db import set_meta,connect
        with tempfile.TemporaryDirectory() as directory:
            run=Path(directory)/'run';seed(run)
            con=connect(run)
            for oid in ('A','B'):
                con.execute('INSERT INTO suppliers(supplier_id,display_name) VALUES (?,?)',[oid,oid+' Firma'])
                con.execute('INSERT INTO offers(offer_id,supplier_id,status) VALUES (?,?,?)',[oid,oid,'active'])
            con.commit();con.close()
            config=configuration(source_model())
            first=export_preliminary(run,'Sentetik senaryo',[{'id':'S','inputs':{'quantity':10,'unit_price':120,'discount':0,'vat_rate':0.2}}],scoring=config)
            self.assertEqual(first['delivery_status'],'READY',first)
            user=Path(directory)/'user.xlsx';wb=load_workbook(first['xlsx']);wb['Kalem Fiyat Varsayimlari']['B2']=10;wb['Kalem Fiyat Varsayimlari']['C2']=1
            wb['Ardisik Iskontolar']['C2']=0.1;wb[SENSITIVITY]['C2']=35;wb[SENSITIVITY]['C3']=65;wb.save(user);wb.close()
            con=connect(run);before=con.execute('SELECT * FROM facts ORDER BY fact_id').fetchall();con.close()
            imported=import_scenarios(run,user);self.assertEqual(imported['scenario_import'],'IMPORTED',imported)
            model=json.loads((Path(imported['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
            self.assertAlmostEqual(model['scenario_reference']['2']['Miktar ve Kapsam!B'],90)
            self.assertEqual(model['scenario_scoring']['sensitivity'][0]['weights']['Kalite'],35)
            con=connect(run);self.assertEqual(con.execute('SELECT * FROM facts ORDER BY fact_id').fetchall(),before);con.close()

class ManagedRuntimeTests(unittest.TestCase):
    def test_registry_requires_exact_interpreter_lock_and_fresh_package_versions(self):
        path=Path(__file__).resolve().parents[4]/'scripts/runtime_registry.py'
        spec=importlib.util.spec_from_file_location('registry_test',path);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as directory:
            local=Path(directory);root=local/'KilicPython/env';root.mkdir(parents=True);python=root/'python.exe';python.write_bytes(b'fixture')
            expected={'id':'P','lock_sha256':'hash','python_series':'3.14','packages':{'pkg':'1'}}
            record={'schema':'kilic-offer-runtime/v1','runtime_id':'R','profile_id':'P','lock_sha256':'hash','environment_root':str(root),'python':str(python)}
            kwargs={'executable':python,'version':'3.14.8','package_version':lambda _: '1','user_profile':str(local)}
            identity=module.validate(record,expected,**kwargs);module.bind_run(local/'run',identity)
            with self.assertRaises(ValueError):module.validate({**record,'lock_sha256':'other'},expected,**kwargs)
            with self.assertRaises(ValueError):module.validate(record,expected,**{**kwargs,'package_version':lambda _: '2'})
            with self.assertRaises(ValueError):module.bind_run(local/'run',{**identity,'runtime_id':'changed'})

