"""A stored requirement exclusion cannot override the current mandatory facts."""
import json, tempfile, unittest
from pathlib import Path
import fitz
from openpyxl import load_workbook
from db import connect, get_meta, rows_as_dict
from report_model import build_preliminary
from report_delivery import export_preliminary
from test_report_slice import seed

class CurrentRequirementEligibilityTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def make_case(self, profile='standart', stale_both=False):
        run=self.root/profile; seed(run,profile,statuses=('compliant',))
        con=connect(run)
        try:
            con.execute("UPDATE requirements SET category='mandatory' WHERE requirement_id='R1'")
            con.execute("INSERT INTO suppliers(supplier_id,display_name) VALUES ('T','Ikinci sentetik firma')")
            con.execute("INSERT INTO offers(offer_id,supplier_id,status) VALUES ('T1','T','active')")
            for oid in (['O','T1'] if stale_both else ['T1']):
                con.execute("INSERT INTO exclusions(exclusion_id,offer_id,requirement_id,decision,reason,status) VALUES (?,?, 'R1','exclude','Historical mandatory failure','confirmed')",['old-'+oid,oid])
            con.commit()
        finally: con.close()
        return run
    def test_stale_requirement_exclusion_preserves_current_two_supplier_universe_and_reports(self):
        run=self.make_case(); con=connect(run)
        before={t:rows_as_dict(con,'SELECT * FROM '+t)for t in ['offers','facts','requirements','exclusions']}; con.close()
        model=build_preliminary(run,'Current facts override historical requirement exclusion')
        self.assertEqual(model['comparison_status'],'BLOCKED')
        result=export_preliminary(run,'Current facts override historical requirement exclusion')
        self.assertEqual(result['delivery_status'],'READY');self.assertEqual(result['comparison_status'],'BLOCKED')
        frozen=json.loads((Path(result['xlsx']).parent/'report_model.json').read_text(encoding='utf-8'))
        self.assertEqual(frozen['comparison_status'],'BLOCKED');self.assertEqual(frozen['ranking'],[])
        self.assertIsNone(frozen['recommendation']);self.assertEqual(len(frozen['offers']),2)
        wb=load_workbook(result['xlsx']);self.assertGreaterEqual(len(wb.sheetnames),20);wb.close()
        with fitz.open(result['pdf'])as pdf: self.assertGreaterEqual(len(pdf),2)
        con=connect(run)
        after={t:rows_as_dict(con,'SELECT * FROM '+t)for t in before};con.close()
        self.assertEqual(before,after);self.assertIsNone(get_meta(run,'reviewed_snapshot'))
    def test_all_stale_requirement_exclusions_do_not_claim_no_eligible_offers(self):
        run=self.make_case(stale_both=True)
        self.assertEqual(build_preliminary(run,'Stale records')['comparison_status'],'BLOCKED')
    def test_current_verified_mandatory_failure_excludes_without_persisted_exclusion(self):
        run=self.make_case();con=connect(run)
        con.execute("DELETE FROM exclusions")
        con.execute("UPDATE facts SET value_text='noncompliant' WHERE fact_id='F1'")
        con.commit();con.close()
        model=build_preliminary(run,'Current mandatory failure')
        self.assertEqual(model['comparison_status'],'NOT_APPLICABLE_SINGLE_SUPPLIER')
        self.assertEqual(len(model['offers']),2);self.assertEqual(model['ranking'],[])
    def test_unverified_current_failure_cannot_remove_supplier(self):
        run=self.make_case();con=connect(run)
        con.execute("DELETE FROM exclusions")
        con.execute("UPDATE facts SET value_text='noncompliant',evidence_status='requires_confirmation' WHERE fact_id='F1'")
        con.commit();con.close()
        self.assertEqual(build_preliminary(run,'Unverified failure')['comparison_status'],'BLOCKED')
    def test_current_verified_failures_remove_all_offers_without_mutating_history(self):
        run=self.make_case();con=connect(run)
        try:
            con.execute("DELETE FROM exclusions")
            con.execute("UPDATE facts SET value_text='noncompliant' WHERE fact_id='F1'")
            con.execute("INSERT INTO facts(fact_id,offer_id,requirement_id,field_key,value_text,evidence_status,extraction_path) VALUES ('T-failure','T1','R1','requirement_assessment','failed','verified','test')")
            con.commit()
        finally: con.close()
        model=build_preliminary(run,'Current failures')
        self.assertEqual(model['comparison_status'],'NO_ELIGIBLE_OFFERS')
        self.assertEqual(len(model['offers']),2);self.assertEqual(model['exclusions'],[])
