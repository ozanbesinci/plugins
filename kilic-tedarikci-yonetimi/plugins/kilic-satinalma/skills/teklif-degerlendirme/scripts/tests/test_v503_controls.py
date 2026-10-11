from __future__ import annotations
from test_engine_support import artifact_path

import json, os, shutil, sys, tempfile, unittest
from pathlib import Path
os.environ.setdefault('TEKLIF_TEST_SQLITE', '1')
SCRIPTS=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(SCRIPTS))
from openpyxl import Workbook, load_workbook
from db import connect
from workflow import prepare, advance
from record_results import record_task_result
from independent_recompute import run as independent_recompute
from provenance_survival import check as provenance_check
from grain_integrity import check as grain_check

class V503Controls(unittest.TestCase):
    def setUp(self):
        self.root=Path(tempfile.mkdtemp(prefix='teklif-v503-')); self.src=self.root/'source'; self.run=self.root/'run'; self.src.mkdir()
    def tearDown(self): shutil.rmtree(self.root,ignore_errors=True)
    def offer(self,name,total,qty=10,payment='%100 60 gun',vat='hariç'):
        wb=Workbook(); ws=wb.active; ws.title='Teklif'
        for r,(k,v) in enumerate([('Fiyat Teklifi',''),('Teklif Tarihi','03.10.2026'),('Geçerlilik','90 gün'),('Genel Toplam',total),('Para Birimi','EUR'),('KDV',vat),('Miktar',qty),('Birim','adet'),('Ürün','Pompa'),('Teslim','30 gün'),('Ödeme',payment),('Incoterm','DDP'),('Garanti','24 ay')],1): ws.cell(r,1,k); ws.cell(r,2,v)
        ws.append(['Tedarikçi\nFirma: ' + Path(name).stem.split('_')[0] + ' Ltd.'])
        wb.save(self.src/name)
    def finish_quick(self,bad_claim=False):
        step=advance(self.run); self.assertEqual(step['status'],'NEEDS_LLM'); tid=step['tasks'][0]['task_id']
        con=connect(self.run)
        try: cid=con.execute("SELECT calculation_id FROM calculations WHERE calculation_type='cost_summary' AND status='complete' ORDER BY calculation_id LIMIT 1").fetchone()[0]
        finally: con.close()
        ref='calc_does_not_exist' if bad_claim else cid
        record_task_result(self.run,tid,{'summary':'Maliyet karşılaştırması tamamlandı.','recommendation':{'type':'cost_observation','text':'Karar tablosuna bakınız.'},'claims':[{'type':'calculated','text':'Karşılaştırılabilir maliyet hesaplanmıştır.','calculation_ids':[ref],'assertions':[{'source_type':'calculation','source_id':ref,'path':'comparable_total','operator':'exists'}]},{'type':'recommendation','text':'Karar tablosuna bakınız.','calculation_ids':[ref]}]},model_name='test-model',reasoning_effort='medium',input_tokens=10,cached_input_tokens=5,output_tokens=5)
        return advance(self.run)
    def test_reviewed_snapshot_is_parent_of_outputs_and_bindings(self):
        self.offer('A_teklif.xlsx',100000); self.offer('B_teklif.xlsx',105000); prepare(self.src,self.run,'hizli'); final=self.finish_quick(); self.assertEqual(final['status'],'VERIFIED')
        con=connect(self.run)
        try:
            snap=con.execute("SELECT sha256 FROM reviewed_snapshots ORDER BY created_at DESC LIMIT 1").fetchone()[0]
            x=con.execute("SELECT artifact_id,parent_hashes_json FROM artifacts WHERE artifact_type='xlsx' ORDER BY created_at DESC LIMIT 1").fetchone()
            self.assertIn(snap,json.loads(x[1])); self.assertGreater(con.execute("SELECT count(*) FROM artifact_bindings WHERE artifact_id=? AND snapshot_sha256=?",[x[0],snap]).fetchone()[0],0)
        finally: con.close()
    def test_independent_recompute_detects_main_calculation_tamper(self):
        self.offer('A_teklif.xlsx',100000); self.offer('B_teklif.xlsx',105000); prepare(self.src,self.run,'hizli')
        con=connect(self.run)
        try:
            row=con.execute("SELECT calculation_id,result_json FROM calculations WHERE calculation_type='cost_summary' ORDER BY calculation_id LIMIT 1").fetchone(); obj=json.loads(row[1]); obj['comparable_total']=str(float(obj['comparable_total'])+1000); con.execute("UPDATE calculations SET result_json=? WHERE calculation_id=?",[json.dumps(obj),row[0]]); con.commit()
        finally: con.close()
        out=independent_recompute(self.run); self.assertEqual(out['status'],'FAIL'); self.assertTrue(out['failed'])
    def test_comparison_contract_blocks_ranking_on_quantity_mismatch(self):
        self.offer('A_teklif.xlsx',100000,qty=10); self.offer('B_teklif.xlsx',105000,qty=20); prepared=prepare(self.src,self.run,'hizli')
        con=connect(self.run)
        try:
            comp=json.loads(con.execute("SELECT value FROM run_meta WHERE key='comparability_contract'").fetchone()[0]); sc=json.loads(con.execute("SELECT result_json FROM calculations WHERE calculation_type='scoring' AND offer_id IS NULL ORDER BY updated_at DESC LIMIT 1").fetchone()[0])
        finally: con.close()
        self.assertFalse(comp['ranking_allowed']); self.assertIn('quantity',comp['failed_dimensions']); self.assertEqual(sc['status'],'suppressed_by_comparability')
    def test_comparison_contract_blocks_when_vat_basis_missing_for_one_offer(self):
        self.offer('A_teklif.xlsx',100000,vat='hariç'); self.offer('B_teklif.xlsx',105000,vat=''); prepare(self.src,self.run,'hizli')
        con=connect(self.run)
        try: comp=json.loads(con.execute("SELECT value FROM run_meta WHERE key='comparability_contract'").fetchone()[0])
        finally: con.close()
        self.assertFalse(comp['ranking_allowed']); self.assertIn('vat_basis',comp['failed_dimensions'])

    def test_unsupported_claim_blocks_pre_output_snapshot(self):
        self.offer('A_teklif.xlsx',100000); self.offer('B_teklif.xlsx',105000); prepare(self.src,self.run,'hizli'); step=self.finish_quick(bad_claim=True)
        self.assertEqual(step['status'],'NEEDS_LLM')
        self.assertEqual(step['summary_retry']['status'],'RETRY_PENDING')
        self.assertFalse((artifact_path(self.run)).exists())
        from workflow import _finish_outputs
        final=_finish_outputs(self.run,preliminary=False)
        self.assertEqual(final['status'],'BLOCKED'); self.assertEqual(final['reviewed_snapshot']['status'],'BLOCKED')
    def test_provenance_survival_detects_exported_evidence_tamper(self):
        self.offer('A_teklif.xlsx',100000); self.offer('B_teklif.xlsx',105000); prepare(self.src,self.run,'hizli'); final=self.finish_quick(); self.assertEqual(final['status'],'VERIFIED')
        xlsx=artifact_path(self.run); wb=load_workbook(xlsx); ws=wb['Kanitlar']; ws.cell(2,4).value='assumption'; wb.save(xlsx)
        out=provenance_check(self.run); self.assertEqual(out['status'],'FAIL')
    def test_grain_integrity_detects_duplicate_verified_offer_field(self):
        self.offer('A_teklif.xlsx',100000); self.offer('B_teklif.xlsx',105000); prepare(self.src,self.run,'hizli')
        con=connect(self.run)
        try:
            r=con.execute("SELECT offer_id,supplier_id,field_key,value_number,evidence_id FROM facts WHERE field_key='total' LIMIT 1").fetchone(); con.execute("INSERT INTO facts(fact_id,offer_id,supplier_id,field_key,value_number,evidence_status,evidence_id,extraction_path,updated_at) VALUES (?,?,?,?,?,'verified',?,'test','now')",['duplicate_fact',r[0],r[1],r[2],r[3],r[4]]); con.commit()
        finally: con.close()
        out=grain_check(self.run); self.assertEqual(out['status'],'FAIL'); self.assertIn('verified_offer_field_grain',out['failed'])

if __name__=='__main__': unittest.main()
