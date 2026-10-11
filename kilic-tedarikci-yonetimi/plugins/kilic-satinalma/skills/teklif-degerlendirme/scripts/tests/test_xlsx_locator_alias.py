"""A model's range alias must retain source binding across record and buyer cost."""
import json,unittest
from db import connect
from workflow import prepare
from user_question_gate import questions,record_response,open_correction
from record_results import _validate_locator_and_quote
import test_buyer_inputs as buyer_input_fixtures
import test_workflow_integration as workflow_fixtures
import test_buyer_costs as buyer_cost_fixtures

class LocatorGateTests(unittest.TestCase):
    setUp=workflow_fixtures.WorkflowIntegrationTests.setUp
    tearDown=workflow_fixtures.WorkflowIntegrationTests.tearDown
    _write_offer=workflow_fixtures.WorkflowIntegrationTests._write_offer
    _response=buyer_input_fixtures.BuyerInputsTests._response
    def fixture(self,locator):
        self._write_offer('Alpha_fiyat_teklifi.xlsx',100000)
        from openpyxl import load_workbook
        wb=load_workbook(self.src/'Alpha_fiyat_teklifi.xlsx');ws=wb.active
        ws['C7']=100;ws['D7']=1000;ws['E7']='adet';wb.save(self.src/'Alpha_fiyat_teklifi.xlsx');wb.close()
        prepare(self.src,self.run,'standart');c=connect(self.run)
        try:
            oid=c.execute("SELECT offer_id FROM offers WHERE status='active'").fetchone()[0]
            uid,did,sha=c.execute('SELECT u.unit_id,d.document_id,d.sha256 FROM document_units u JOIN documents d ON d.document_id=u.document_id WHERE d.offer_id=?',[oid]).fetchone()
            c.execute("INSERT INTO evidence(evidence_id,document_id,unit_id,locator_json,evidence_status,extraction_path,source_sha256) VALUES (?,?,?,?,?,?,?)",['ev-row',did,uid,json.dumps(locator),'verified','parser',sha])
            c.execute("UPDATE facts SET normalized_json=? WHERE offer_id=? AND field_key='quantity'",[json.dumps({'items':[{'item_id':'q','evidence_refs':['ev-row']}]}),oid])
            c.execute("INSERT INTO open_issues(issue_id,owner,severity,issue_type,question,fact_ids_json,evidence_ids_json,status) VALUES ('row-question','IDARE','critical','scope_comparability','Explicit source-row correction','[]','[]','open')")
            c.commit()
        finally:c.close()
        question=next(q for q in questions(self.run)['questions'] if q['issue_id']=='row-question')
        entry={'kind':'reference_quantity','offer_id':oid,'quantity_item_id':'q','quantity':20,'unit':'adet','source_row':{'unit_id':uid,'quantity_cell':'B7','unit_price_cell':'C7','line_total_cell':'D7','unit_cell':'E7'}}
        return question,entry,did,uid
    def test_cell_range_records_source_bound_quantity_correction(self):
        q,e,_,_=self.fixture({'cell_range':'A7:E7'})
        self.assertEqual(record_response(self.run,self._response(q,[e]),self.root/'next')['status'],'RECORDED_NEW_RUN_REQUIRED')
    def test_alias_cannot_expand_to_foreign_source_row(self):
        q,e,_,_=self.fixture({'cell_range':'A8:E8'})
        with self.assertRaisesRegex(ValueError,'locator kapsamına'):record_response(self.run,self._response(q,[e]),self.root/'next')
    def test_conflicting_aliases_cannot_authorize_quantity(self):
        q,e,_,_=self.fixture({'range':'A7:E7','cell_range':'A8:E8'})
        with self.assertRaisesRegex(ValueError,'locator kapsamına'):record_response(self.run,self._response(q,[e]),self.root/'next')
    def test_nonstring_alias_cannot_authorize_quantity(self):
        q,e,_,_=self.fixture({'cell_range':['A7:E7']})
        with self.assertRaisesRegex(ValueError,'locator kapsamına'):record_response(self.run,self._response(q,[e]),self.root/'next')
    def test_recorder_rejects_nonexistent_alias_even_when_quote_exists_in_sheet(self):
        _,_,did,uid=self.fixture({'cell_range':'A7:E7'});c=connect(self.run)
        try:
            with self.assertRaisesRegex(ValueError,'does not resolve'):
                _validate_locator_and_quote(c,did,uid,{'cell_range':'D88:G88'},'1000','semantic_llm','verified')
        finally:c.close()

class LocatorCostTests(unittest.TestCase):
    setUp=buyer_cost_fixtures.BuyerCostTests.setUp;tearDown=buyer_cost_fixtures.BuyerCostTests.tearDown
    _row=buyer_cost_fixtures.BuyerCostTests._row;_install=buyer_cost_fixtures.BuyerCostTests._install;_result=buyer_cost_fixtures.BuyerCostTests._result
    def test_alias_corrects_only_bound_cost_and_preserves_supplier_total(self):
        self.con.execute('DELETE FROM scope_items')
        for row in (26,27):self.con.execute('UPDATE evidence SET locator_json=? WHERE evidence_id=?',[json.dumps({'cell_range':f'D{row}:G{row}'}),f'ev{row}'])
        params=self._install([self._row(26,'q26'),self._row(27,'q27')])
        cost=self._result(params)['TEZ']
        self.assertEqual(float(cost['comparable_total']),2011640)
        self.assertEqual(self.con.execute("SELECT value_number FROM facts WHERE fact_id='TEZ/total'").fetchone()[0],2547640)

class ExplicitCorrectionTests(unittest.TestCase):
    setUp=workflow_fixtures.WorkflowIntegrationTests.setUp
    tearDown=workflow_fixtures.WorkflowIntegrationTests.tearDown
    _write_offer=workflow_fixtures.WorkflowIntegrationTests._write_offer
    def fixture(self):
        self._write_offer('Alpha_fiyat_teklifi.xlsx',100000);prepare(self.src,self.run,'standart');c=connect(self.run)
        try:
            oid=c.execute("SELECT offer_id FROM offers WHERE status='active'").fetchone()[0]
            fid=c.execute("SELECT fact_id FROM facts WHERE offer_id=? AND field_key='validity'",[oid]).fetchone()[0]
            c.execute("UPDATE facts SET value_text='Source one week',evidence_status='verified' WHERE fact_id=?",[fid]);c.commit()
        finally:c.close()
        return oid,fid
    def test_verified_source_can_receive_separate_explicit_date_correction(self):
        oid,fid=self.fixture();opened=open_correction(self.run,oid,'validity','Alıcı son geçerliliği 2026-10-30 olarak bildirdi.')
        q=opened['question'];response={'issue_id':q['issue_id'],'scope_sha256':q['scope_sha256'],'answer_text':'Son geçerlilik 2026-10-30.','answer_kind':'clarification','buyer_inputs':[{'kind':'validity_date','offer_id':oid,'date_status':'known','value':'2026-10-30'}]}
        self.assertEqual(record_response(self.run,response,self.root/'next')['status'],'RECORDED_NEW_RUN_REQUIRED')
        c=connect(self.run)
        try:self.assertEqual(c.execute('SELECT value_text,evidence_status FROM facts WHERE fact_id=?',[fid]).fetchone(),('Source one week','verified'))
        finally:c.close()
    def test_correction_cannot_target_foreign_offer_or_field(self):
        oid,_=self.fixture()
        for offer,field in [('FOREIGN','validity'),(oid,'total')]:
            with self.subTest(offer=offer,field=field),self.assertRaises(ValueError):open_correction(self.run,offer,field,'Gerçek kullanıcı düzeltmesi')
    def test_closed_run_cannot_be_reopened_by_correction(self):
        oid,_=self.fixture();c=connect(self.run)
        try:c.execute("INSERT INTO workflow_state(phase,status) VALUES ('COMPLETE','PASS')");c.commit()
        finally:c.close()
        with self.assertRaisesRegex(ValueError,'Tamamlanmış'):open_correction(self.run,oid,'validity','Yeni tarih')
    def test_correction_rejects_changed_source_before_question_write(self):
        oid,_=self.fixture();c=connect(self.run)
        try:metadata=c.execute('SELECT metadata_json FROM documents WHERE offer_id=?',[oid]).fetchone()[0]
        finally:c.close()
        (self.run/json.loads(metadata)['archived_source_path']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'integrity'):open_correction(self.run,oid,'validity','Yeni tarih')
    def test_correction_cannot_read_matching_source_outside_run_archive(self):
        oid,_=self.fixture();c=connect(self.run)
        try:
            meta=json.loads(c.execute('SELECT metadata_json FROM documents WHERE offer_id=?',[oid]).fetchone()[0])
            meta['archived_source_path']=str(self.src/'Alpha_fiyat_teklifi.xlsx')
            c.execute('UPDATE documents SET metadata_json=? WHERE offer_id=?',[json.dumps(meta),oid]);c.commit()
        finally:c.close()
        with self.assertRaisesRegex(ValueError,'integrity'):open_correction(self.run,oid,'validity','Yeni tarih')
