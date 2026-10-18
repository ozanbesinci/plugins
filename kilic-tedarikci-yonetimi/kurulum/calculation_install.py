"""Per-user file extraction only. Never execute MSI actions or modify Office."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import urllib.request
import uuid

HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE))
from msi_readonly import Database
PACKAGE=HERE.parent/'plugins/kilic-satinalma'
SCRIPTS=PACKAGE/'skills/teklif-degerlendirme/scripts'

def sha(path):
    with Path(path).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def atomic_json(path,value):
    pending=path.with_name(path.name+'.'+uuid.uuid4().hex+'.pending')
    pending.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    os.replace(pending,path)

def basename(value):
    value=value.split('|')[-1]
    if not value or Path(value).name!=value or value in ('.','..') or ':' in value:
        raise ValueError('MSI file path rejected')
    return value

def extract(package,stage,engine):
    db=Database(package)
    try:
        directories={r['Directory']:r for r in db.rows('Directory',['Directory','Directory_Parent','DefaultDir'])}
        components={r['Component']:r['Directory_'] for r in db.rows('Component',['Component','Directory_'])}
        files=db.rows('File',['File','Component_','FileName','FileSize'])
        media=db.rows('Media',['Cabinet'])
        if media!=[{'Cabinet':'#libreoffice1.cab'}]:raise ValueError('Unexpected MSI cabinet layout')
        cabinet=stage/'payload.cab';db.extract_stream('libreoffice1.cab',cabinet)
    finally:db.close()
    extracted=stage/'files';extracted.mkdir()
    with (stage/'extraction.log').open('wb') as log:
        subprocess.run([str(Path(os.environ['SystemRoot'])/'System32/expand.exe'),'-F:*',str(cabinet),str(extracted)],
                       stdout=log,stderr=log,timeout=300,check=True,creationflags=subprocess.CREATE_NO_WINDOW)
    def directory(key,seen=None):
        if key=='INSTALLLOCATION':return Path()
        if not key or key=='TARGETDIR' or key not in directories:return None
        seen=set() if seen is None else seen
        if key in seen:raise ValueError('MSI directory cycle')
        seen.add(key);entry=directories[key];parent=directory(entry['Directory_Parent'],seen)
        if parent is None:return None
        name=entry['DefaultDir'].split(':')[0]
        return parent if name=='.' else parent/basename(name)
    mapped=set()
    for entry in files:
        relative=directory(components[entry['Component_']])
        # The signed MSI includes x86, x64 and ARM runtimes. Extract only x64
        # as app-local DLLs; never write System32 or Windows Installer records.
        if relative is None and '_amd64.' in entry['File'] and entry['FileName'].endswith('.dll'):
            relative=Path('program')
        if relative is None:continue
        target=relative/basename(entry['FileName'])
        if target.as_posix().casefold() in mapped:raise ValueError('Duplicate engine destination')
        mapped.add(target.as_posix().casefold())
        source=extracted/basename(entry['File'])
        if source.stat().st_size!=int(entry['FileSize']):raise ValueError('MSI file size mismatch')
        destination=engine/target;destination.parent.mkdir(parents=True,exist_ok=True)
        # TemporaryDirectory is deliberately private on modern Windows Python.
        # A same-volume rename retains that private DACL. Create destination
        # files normally, inheriting the already-authorized installation folder;
        # do not copy security descriptors or change any existing permissions.
        with source.open('rb') as reader, destination.open('xb') as writer:
            shutil.copyfileobj(reader, writer, 1024 * 1024)

def install():
    sys.path.insert(0,str(SCRIPTS))
    from calculation_runtime import validate,machine_id,evidence,inventory,hash_files
    from local_calculation import calculate_with_engine
    from openpyxl import Workbook,load_workbook
    expected=json.loads((PACKAGE/'config/calculation-profile.json').read_text(encoding='utf-8'))
    root=(Path(os.environ['USERPROFILE'])/'KilicPython').resolve();root.mkdir(exist_ok=True)
    local=root/'calculation';local.mkdir(exist_ok=True)
    registry=root/'calculation-runtime.json'
    if registry.exists():
        try:
            installed=json.loads(registry.read_text())
            identity=validate(installed,expected)
            if installed.get('file_materialization')=='destination-inherited-v1':
                print('Hesap motoru zaten doğrulandı: '+identity['version'],flush=True)
                return
        except (OSError,ValueError,KeyError):pass
    engine=local/(expected['id']+'-'+uuid.uuid4().hex[:8]);engine.mkdir()
    with tempfile.TemporaryDirectory(prefix='calc-install-',dir=local) as temporary:
        stage=Path(temporary).resolve()
        if stage.parent!=local or not engine.resolve().is_relative_to(local):raise ValueError('Installation path escaped')
        package=stage/'LibreOffice.msi'
        override=os.environ.get('KILIC_CALCULATION_PACKAGE')
        if override:shutil.copyfile(Path(override).resolve(strict=True),package)
        else:
            print('Yerel hesap motoru indiriliyor (yaklaşık 358 MB)...',flush=True)
            with urllib.request.urlopen(expected['url'],timeout=120) as response,package.open('xb') as stream:
                if not response.url.startswith('https://'):raise ValueError('Engine download requires HTTPS')
                shutil.copyfileobj(response,stream,1024*1024)
        if sha(package)!=expected['sha256']:raise ValueError('Hesap motoru paket özeti uyuşmuyor.')
        print('Dosyalar kullanıcı klasörüne açılıyor; sistem kurulumu yapılmıyor...',flush=True)
        extract(package,stage,engine)
        manifest=hash_files(engine,inventory(engine))
        manifest_path=engine/'engine-manifest.json'
        manifest_path.write_text(json.dumps(manifest,sort_keys=True,separators=(',',':')),encoding='utf-8')
        record={'schema':'kilic-calculation-runtime/v1','profile_id':expected['id'],
                'machine_id':machine_id(),'root':str(engine),
                'file_materialization':'destination-inherited-v1'}
        identity=validate(record,expected)
        source=stage/'smoke.xlsx';output=stage/'calculated.xlsx'
        wb=Workbook();ws=wb.active;ws['A1']=2;ws['A2']=3;ws['A3']='=SUM(A1:A2)';ws['A4']='=IF(A3=5,"hazır","")';wb.save(source);wb.close()
        result=calculate_with_engine(source,output,identity)
        wb=load_workbook(output,data_only=True)
        try:
            if wb.active['A3'].value!=5 or wb.active['A4'].value!='hazır':raise ValueError('Hesap motoru gerçek hesap kontrolü başarısız.')
        finally:wb.close()
        # A real launch must also leave the pinned engine immutable.
        validate(record,expected)
        record['smoke']={'status':'PASS','formula_count':result['formulas'],'engine':evidence(identity)}
        atomic_json(registry,record)
    print('Yerel hesap motoru hazır: LibreOffice '+expected['version'],flush=True)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--prepared-environment',action='store_true');args=parser.parse_args()
    if not args.prepared_environment:
        # The separate upgrade entry uses the exact previously registered venv.
        record=json.loads((Path(os.environ['USERPROFILE'])/'KilicPython/teklif-runtime.json').read_text(encoding='utf-8'))
        python=Path(record['python']).resolve(strict=True)
        if not python.is_relative_to((Path(os.environ['USERPROFILE'])/'KilicPython/environments').resolve()):raise ValueError('Python kayıt yolu geçersiz.')
        return subprocess.call([str(python),'-X','utf8','-I','-B',str(Path(__file__).resolve()),'--prepared-environment'])
    install();return 0

if __name__=='__main__':
    try:raise SystemExit(main())
    except Exception as exc:
        print('Hesap motoru kurulamadı: '+str(exc),file=sys.stderr,flush=True)
        raise SystemExit(3)
