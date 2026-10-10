from pathlib import Path
import shutil, sys, tempfile, unittest

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook
from display_values import structured
from print_layout import MAX_LINES, build_print_workbook, paragraph_groups


class BandHeaderTests(unittest.TestCase):
    def test_comparison_bands_repeat_identity_and_retain_all_fields_at_readable_width(self):
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook(); ws = wb.active; ws.title = "Teklif Karsilastirma"
            ws.append([f"Field {c}" for c in range(1, 23)])
            ws.append([f"value {c}" for c in range(1, 23)])
            for c in range(1, 23):
                ws.column_dimensions[ws.cell(1, c).column_letter].width = 25
            wb.save(root / "in.xlsx"); wb.close()
            info = build_print_workbook(root / "in.xlsx", root / "print.xlsx")
            p = load_workbook(root / "print.xlsx")
            try:
                bands = [s for s in p if s.title.startswith("Teklif Karsilastirma")]
                self.assertGreater(len(bands), 3)
                covered = set()
                for band in bands:
                    headers = [cell.value for cell in band[1]]
                    self.assertEqual(headers[:2], ["Field 1", "Field 2"])
                    self.assertEqual([band.cell(2, 1).value, band.cell(2, 2).value], ["value 1", "value 2"])
                    self.assertLessEqual(sum(band.column_dimensions[band.cell(1, c).column_letter].width for c in range(1, band.max_column + 1)), 110)
                    covered.update(headers)
                self.assertEqual(covered, {f"Field {c}" for c in range(1, 23)})
                self.assertEqual(info["minimum_font_points"], 11)
            finally:
                p.close()
        finally:
            shutil.rmtree(root)

    def test_multi_band_header_counts_bands_even_with_multiline_cells(self):
        # RC5 regression: a per-cell variable shadowed the band list, printing "2/1", "3/1".
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook(); ws = wb.active; ws.title = "Kanitlar"
            ws.append([f"Alan {c}" for c in range(1, 11)])
            ws.append(["satir"] + ["kısa"] * 8 + ["bir\niki\nüç"])
            for c in range(1, 11):
                ws.column_dimensions[ws.cell(1, c).column_letter].width = 32
            wb.save(root / "in.xlsx"); wb.close()
            build_print_workbook(root / "in.xlsx", root / "print.xlsx")
            p = load_workbook(root / "print.xlsx")
            n = len(p.worksheets)
            self.assertGreater(n, 1)
            self.assertEqual(len({s.title for s in p.worksheets}), n)
            for i, sheet in enumerate(p.worksheets, 1):
                self.assertTrue(sheet.oddHeader.center.text.endswith(f"Sütun grubu {i}/{n}"), sheet.oddHeader.center.text)
            p.close()
        finally:
            shutil.rmtree(root)


class RecordGroupingTests(unittest.TestCase):
    def test_list_of_records_is_separated_by_blank_lines(self):
        text = structured('[{"source_id":"calc_a","path":"comparable_total","value_number":108290.01},{"source_id":"calc_b","value_number":1}]')
        self.assertEqual(text.split("\n\n")[0], "Kaynak kimliği: calc_a\nDeğer yolu: Karşılaştırılabilir toplam\nSayısal değer: 108290.01")
        self.assertEqual(structured('["a","b"]'), "a; b")

    def test_paragraph_kept_whole_and_long_paragraph_falls_back_to_lines(self):
        groups = paragraph_groups("Kaynak türü: Hesap\nKaynak kimliği: calc_c8d3612f1ce57f161ea3f94c\nSayısal değer: 108290.01\n\nKaynak türü: Veri", 24)
        self.assertEqual(groups[0][0], "Kaynak türü: Hesap")
        self.assertIn("Sayısal değer:", " ".join(groups[0]))
        self.assertEqual(groups[0][-1], "")
        long = paragraph_groups("\n".join(f"Etiket {i}: değer" for i in range(MAX_LINES + 5)), 40)
        self.assertEqual(len(long), MAX_LINES + 5)
        self.assertEqual(paragraph_groups("", 20), [[""]])

    def test_record_is_not_split_across_continuation_rows(self):
        root = Path(tempfile.mkdtemp())
        try:
            records = [{"source_type": "calculation", "source_id": f"calc_{i:02d}c8d3612f1ce57f161ea3f94c", "path": "comparable_total",
                        "operator": "approx", "value_number": 108290.01 + i, "tolerance": 0.01} for i in range(9)]
            import json
            wb = Workbook(); ws = wb.active; ws.title = "Iddia Kaydi"
            ws.append(["Claim ID", "Assertions"])
            ws.append(["claim_87561", structured(json.dumps(records))])
            ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 40
            wb.save(root / "in.xlsx"); wb.close()
            build_print_workbook(root / "in.xlsx", root / "print.xlsx")
            p = load_workbook(root / "print.xlsx"); sh = p.active
            cells = [sh.cell(r, 2).value for r in range(2, sh.max_row + 1)]
            self.assertGreater(len(cells), 1)
            for text in cells:
                blocks = [b for b in text.split("\n\n") if b.strip()]
                for b in blocks:
                    self.assertTrue(b.lstrip().startswith("Kaynak türü:"), b)
                    self.assertIn("Tolerans:", b)
            p.close()
        finally:
            shutil.rmtree(root)


if __name__ == "__main__":
    unittest.main()
