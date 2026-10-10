from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

os.environ['TEKLIF_TEST_SQLITE'] = '1'
SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook

from db import connect
from record_results import record_task_result
from workflow import advance, prepare, status


class WorkflowIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix='teklif-v5-int-'))
        self.src = self.root / 'source'
        self.run = self.root / 'run'
        self.src.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _write_offer(self, filename: str, total: float, quote_date: str = '03.10.2026', validity: str = '90 gün') -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = 'Teklif'
        rows = [
            ('Fiyat Teklifi', ''),
            ('Teklif Tarihi', quote_date),
            ('Geçerlilik', validity),
            ('Genel Toplam', total),
            ('Para Birimi', 'EUR'),
            ('KDV', 'hariç'),
            ('Miktar', 10),
            ('Birim', 'adet'),
            ('Ürün', 'Pompa'),
            ('Teslim', '30 gün'),
            ('Ödeme', '%100 60 gün'),
            ('Incoterm', 'DDP'),
            ('Garanti', '24 ay'),
        ]
        for r, (label, value) in enumerate(rows, 1):
            ws.cell(r, 1, label)
            ws.cell(r, 2, value)
        ws.append(['Tedarikçi\nFirma: ' + Path(filename).stem.split('_')[0] + ' Ltd.'])
        wb.save(self.src / filename)

    def _pending_decision_task(self) -> str:
        con = connect(self.run)
        try:
            row = con.execute("SELECT task_id FROM task_queue WHERE task_kind='decision_summary' AND status='pending'").fetchone()
            self.assertIsNotNone(row)
            return row[0]
        finally:
            con.close()

    def _support_refs(self):
        con = connect(self.run)
        try:
            fact = con.execute("SELECT fact_id FROM facts WHERE evidence_status='verified' ORDER BY fact_id LIMIT 1").fetchone()
            calc = con.execute("SELECT calculation_id FROM calculations WHERE calculation_type='cost_summary' AND status='complete' ORDER BY calculation_id LIMIT 1").fetchone()
            return fact[0] if fact else None, calc[0] if calc else None
        finally:
            con.close()

    def test_quick_profile_runs_to_verified_with_only_bounded_summary_llm(self):
        self._write_offer('Alpha_fiyat_teklifi.xlsx', 100000)
        self._write_offer('Beta_fiyat_teklifi.xlsx', 105000)

        prepared = prepare(self.src, self.run, 'hizli')
        self.assertEqual(prepared['status'], 'PREPARED')
        self.assertEqual(prepared['deterministic_fields']['fields_missing'], 0)

        step = advance(self.run)
        self.assertEqual(step['status'], 'NEEDS_LLM')
        self.assertEqual(step['pending_total'], 1)
        self.assertEqual(step['tasks'][0]['task_kind'], 'decision_summary')
        context_size = Path(step['tasks'][0]['context_path']).stat().st_size
        self.assertLess(context_size, 80000)

        task_id = self._pending_decision_task()
        fact_id, calc_id = self._support_refs()
        record_task_result(
            self.run,
            task_id,
            {
                'summary': 'Teklifler karşılaştırılabilir. Nihai seçim yetkili karar merciine aittir.',
                'recommendation': {'type': 'cost_observation', 'text': 'Daha düşük beyan edilen toplam maliyet ayrıca gösterilmiştir.'},
                'claims': [
                    {'type':'calculated','text':'Karşılaştırılabilir maliyetler hesaplanmıştır.','calculation_ids':[calc_id],'assertions':[{'source_type':'calculation','source_id':calc_id,'path':'comparable_total','operator':'exists'}]},
                    {'type':'recommendation','text':'Daha düşük maliyet ayrıca gösterilmiştir.','calculation_ids':[calc_id]}
                ],
            },
            model_name='test-model', reasoning_effort='medium', input_tokens=1200, cached_input_tokens=800, output_tokens=120,
        )
        final = advance(self.run)
        self.assertEqual(final['status'], 'VERIFIED')
        self.assertTrue((self.run / 'output' / 'Teklif_Degerlendirme.xlsx').exists())
        self.assertEqual(final['qa']['status'], 'PASS')
        self.assertEqual(final['qa']['provenance_survival']['status'], 'PASS')
        self.assertEqual(final['qa']['provenance_survival']['findings'], [])
        self.assertTrue(status(self.run)['deliverable_artifacts'])
        workbook = self.run / 'output' / 'Teklif_Degerlendirme.xlsx'
        with workbook.open('ab') as stream:
            stream.write(b'changed-after-closure')
        changed = status(self.run)
        self.assertFalse(changed['deliverable_artifacts'])
        self.assertEqual(changed['delivery_reason'], 'completed run integrity failed')

    def test_expired_offer_waits_for_user_before_summary_or_artifact(self):
        self._write_offer('Alpha_fiyat_teklifi.xlsx', 100000, quote_date='01.01.2026', validity='30 gün')
        self._write_offer('Beta_fiyat_teklifi.xlsx', 105000, quote_date='01.01.2026', validity='30 gün')
        prepare(self.src, self.run, 'standart')

        step = advance(self.run)
        # Contract change: a known critical business question is shown before
        # summary generation and never becomes an automatic preliminary file.
        self.assertEqual(step['status'], 'NEEDS_USER_INPUT')
        self.assertEqual(step.get('gate', {}).get('status'), 'PRELIMINARY_ONLY')
        self.assertTrue(step['questions'])
        self.assertFalse((self.run/'output'/'Teklif_Degerlendirme.xlsx').exists())
        con = connect(self.run)
        try:
            review_tasks = con.execute("SELECT count(*) FROM task_queue WHERE task_kind IN ('independent_fact_review','adjudicate_differences')").fetchone()[0]
            self.assertEqual(review_tasks, 0)
            self.assertEqual(con.execute("SELECT count(*) FROM task_queue WHERE task_kind='decision_summary'").fetchone()[0],0)
        finally:
            con.close()


if __name__ == '__main__':
    unittest.main()
