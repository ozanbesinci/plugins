import copy,json,tempfile,unittest
from pathlib import Path
from unittest import mock
from openpyxl import load_workbook
from effective_report_values import buyer_term,weighted_score_display
from report_export import export_report
from print_layout import build_print_workbook

class EffectiveReportTermTests(unittest.TestCase):
    def terms(self):
        date={'date_status':'known','value':'2026-10-30'}
        payment={'milestones':[{'trigger':'work_start','percent':50},{'trigger':'work_delivery','percent':50}]}
        return {'complete':True,'known':True,'validity_date_basis':'buyer_declaration','validity_date':date,
                'payment_basis':'buyer_declaration','payment_terms':payment,
                'buyer_inputs':{'validity_date':[date],'payment_terms':[payment]}}
    def test_current_terms_and_no_total_score_survive_pdf_print_with_original_source_retained(self):
        cost=self.terms()
        facts=[{'fact_id':'f-payment','offer_id':'O','field_key':'payment','value_text':'%50 nakit','evidence_status':'verified'},
               {'fact_id':'f-valid','offer_id':'O','field_key':'validity','value_date':'2026-10-14','evidence_status':'verified'}]
        snap={'status':'READY','meta':{'profile':'standart','decision_status':'SUFFICIENT','comparability_contract':{'ranking_allowed':True}},
              'suppliers':[{'supplier_id':'S','display_name':'Firma'}],
              'offers':[{'offer_id':'O','supplier_id':'S','status':'active','validity_end':'2026-10-14'}],
              'facts':facts,'calculations':[{'calculation_id':'C','offer_id':'O','calculation_type':'cost_summary','result_json':json.dumps(cost)},
                                          {'calculation_id':'SC','offer_id':None,'calculation_type':'scoring','result_json':json.dumps({'status':'cost_only','price_scores':[{'offer_id':'O','price_score_0_10':'10'}]})}]}
        original=copy.deepcopy(snap); bindings=[]
        class Con:
            def close(self):pass
        with tempfile.TemporaryDirectory() as temp:
            run=Path(temp)
            with mock.patch('reviewed_snapshot.load_latest',return_value=(snap,{'sha256':'snapshot','relative_path':'reviewed/snapshot.json'})),mock.patch('report_export.connect',return_value=Con()),mock.patch('report_export._artifact',return_value='A'),mock.patch('report_export._insert_bindings',side_effect=lambda c,a,s,r:bindings.extend(r)),mock.patch('report_export._recalc_with_excel',return_value={'available':False}):
                result=export_report(run)
            wb=load_workbook(result['xlsx'],data_only=True)
            try:
                comparison=wb['Teklif Karsilastirma']
                self.assertEqual(comparison['P2'].value,'2026-10-30')
                self.assertIn('başlama',comparison['O2'].value.lower())
                self.assertIn('%50',comparison['O2'].value)
                self.assertEqual(comparison['M2'].value,'Hesaplanmadı: yalnız maliyet puanı')
                self.assertEqual(comparison['Q1'].value,'Kaynak Belge Durumu')
                source=' '.join(str(c.value) for row in wb['Ticari Kosullar'] for c in row)
                self.assertIn('2026-10-14',source);self.assertIn('%50 nakit',source)
            finally:wb.close()
            self.assertEqual(snap,original)
            for cell,path in [('P2','validity_date.value'),('O2','payment_terms.milestones')]:
                self.assertTrue(any(b['component_id']=='xlsx:Teklif Karsilastirma:'+cell and b['calculation_ids']==['C'] and b['locator']['value_path']==path for b in bindings))
            printed=run/'output/print.xlsx';build_print_workbook(Path(result['xlsx']),printed)
            wb=load_workbook(printed,data_only=True)
            try:
                text=' '.join(str(c.value) for sheet in wb for row in sheet for c in row)
                self.assertIn('2026-10-30',text)
                self.assertIn('2026-10-14',text)
                self.assertIn('yalnızmaliyetpuanı',''.join(text.split()))
            finally:wb.close()
    def test_unapplied_incomplete_and_conflicting_terms_never_replace_source_or_invent_total_score(self):
        for mutation in [{'complete':False},{'known':False},{'validity_date_basis':'source'},{'buyer_inputs':{}},{'validity_date':{'date_status':'known','value':'2026-11-01'}}]:
            with self.subTest(mutation=mutation):
                value={**self.terms(),**mutation}
                self.assertIsNone(buyer_term(value,'validity'))
        self.assertIsNone(weighted_score_display({'status':'cost_only'},None,False))
        self.assertIsNone(weighted_score_display({'status':'incomplete'},None,True))
        self.assertEqual(weighted_score_display({'status':'weighted_complete'},'7.25',True),7.25)

if __name__=='__main__':unittest.main()
