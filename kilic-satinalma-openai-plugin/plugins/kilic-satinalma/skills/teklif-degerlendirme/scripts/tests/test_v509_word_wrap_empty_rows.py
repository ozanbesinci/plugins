from pathlib import Path
import shutil, sys, tempfile, unittest

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook
from display_values import present_workbook
from print_layout import build_print_workbook, wrap_text


class WordWrapTests(unittest.TestCase):
    def test_word_that_fits_the_column_is_not_cut(self):
        self.assertEqual(wrap_text("Belirtilmedi", 10, 13), ["Belirtilmedi"])
        # Without the real column capacity the conservative estimate would cut it.
        self.assertEqual(wrap_text("Belirtilmedi", 10), ["Belirtilme", "di"])
        # A token longer than the column is still cut, integer digits first.
        self.assertEqual(wrap_text("103367.7402283199989683318927", 10, 13)[0], "103367.740")

    def test_print_keeps_short_word_whole_in_narrow_column(self):
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook(); ws = wb.active; ws.title = "Degisim Kaydi"
            ws.append(["Event ID", "Asama", "Olay", "Detay"]); ws.append(["e1", "QA", "x", "Belirtilmedi"])
            for col in "ABCD":
                ws.column_dimensions[col].width = 14
            wb.save(root / "in.xlsx"); wb.close()
            build_print_workbook(root / "in.xlsx", root / "print.xlsx")
            p = load_workbook(root / "print.xlsx")
            self.assertEqual(p.active["D2"].value, "Belirtilmedi")
            p.close()
        finally:
            shutil.rmtree(root)


class EmptyTableRowTests(unittest.TestCase):
    def test_empty_table_row_is_not_filled_with_placeholders(self):
        wb = Workbook(); ws = wb.active; ws.title = "Degisim Kaydi"
        ws.append(["Event ID", "Asama", "Olay", "Detay"]); ws.append(["Kayıt bulunmuyor."])
        ws["D2"] = None
        present_workbook(wb)
        self.assertEqual(ws["A2"].value, "Kayıt bulunmuyor.")
        self.assertIsNone(ws["D2"].value)
        self.assertEqual(ws["A1"].value, "Olay kimliği")

    def test_rows_with_data_are_still_presented(self):
        wb = Workbook(); ws = wb.active; ws.title = "Degisim Kaydi"
        ws.append(["Event ID", "Asama", "Olay", "Detay"]); ws.append(["e1", "QA", "x", None])
        present_workbook(wb)
        self.assertEqual(ws["D2"].value, "Belirtilmedi")


if __name__ == "__main__":
    unittest.main()
