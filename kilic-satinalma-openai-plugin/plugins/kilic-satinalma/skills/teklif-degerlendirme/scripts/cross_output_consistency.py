from __future__ import annotations
from artifact_values import comparison_fact_value, evidence_label, offer_display_status, pending_offer_fields, ranking_visible
from effective_report_values import buyer_term, weighted_score_display
import argparse, json, math
from pathlib import Path
from typing import Any
from common import json_text, sha256_file, stable_id, utc_now
from db import connect, rows_as_dict
from active_records import current_artifacts
from display_values import display_cell,label
from reviewed_snapshot import load_latest

def _same(a:Any,b:Any)->bool:
    if a is None and b in (None,""): return True
    if b is None and a in (None,""): return True
    if isinstance(a,(int,float)) and isinstance(b,(int,float)): return math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-6)
    return " ".join(str(a).split()).casefold()==" ".join(str(b).split()).casefold()

def _j(v,default=None):
    try: return v if isinstance(v,(dict,list)) else json.loads(v) if v is not None else default
    except Exception: return default

def check(run_root:Path)->dict[str,Any]:
    from openpyxl import load_workbook
    snap,srow=load_latest(run_root); snapshot_sha=srow['sha256']; meta=snap.get('meta',{})
    con=connect(run_root)
    try:
        current=current_artifacts(con)
        arts=[current['xlsx']] if 'xlsx' in current else []
        if not arts: return {'status':'FAIL','findings':[{'severity':'critical','message':'XLSX artifact missing'}]}
        art=arts[0]; xlsx=run_root/art['relative_path']; wb=load_workbook(xlsx,data_only=True,read_only=False); findings=[]
        def add(sheet,locator,field,expected,observed): findings.append({'severity':'critical','sheet':sheet,'locator':locator,'field':field,'expected':expected,'observed':observed})
        facts=snap.get('facts',[]); fmap={(r.get('offer_id'),r['field_key']):r for r in facts if r.get('offer_id')}
        offers=[o for o in snap.get('offers',[]) if o.get('status')=='active']
        calcs=snap.get('calculations',[]); costs={}; scoring={}
        for c in calcs:
            obj=_j(c.get('result_json'),{}) or {}
            if c.get('calculation_type')=='cost_summary' and c.get('offer_id'): costs[c['offer_id']]=obj
            if c.get('calculation_type')=='scoring' and not c.get('offer_id'): scoring=obj
        price={x.get('offer_id'):x.get('price_score_0_10') for x in scoring.get('price_scores',[])}; weighted={x.get('offer_id'):x.get('total_score_0_10') for x in scoring.get('weighted_scores',[])}
        if 'Teklif Karsilastirma' in wb.sheetnames:
            ws=wb['Teklif Karsilastirma']; row_by={ws.cell(r,2).value:r for r in range(2,ws.max_row+1)}
            for off in offers:
                oid=off['offer_id']; r=row_by.get(oid)
                if r is None: add(ws.title,'B:B','offer_id',oid,None); continue
                for field,col in {'total':3,'currency':4,'delivery':14,'payment':15}.items():
                    f=fmap.get((oid,field));
                    if f:
                        term=buyer_term(costs.get(oid,{}),field)
                        exp=term['value'] if term else comparison_fact_value(f)
                        if not _same(exp,ws.cell(r,col).value): add(ws.title,ws.cell(r,col).coordinate,f'{oid}:{field}',exp,ws.cell(r,col).value)
                c=costs.get(oid,{})
                for field,col in [('known_scope_adjustment',5),('acquisition_present_value',7),('tco_present_value',8),('comparable_total',9),('base_currency',10),('base_amount',11)]:
                    exp=c.get(field); obs=ws.cell(r,col).value
                    if exp is not None and field not in {'base_currency'}:
                        try: exp=float(exp)
                        except Exception: pass
                    if not _same(exp,obs): add(ws.title,ws.cell(r,col).coordinate,f'{oid}:{field}',exp,obs)
                for score_map,col,field in [(price,12,'price_score'),(weighted,13,'weighted_score')]:
                    expected=(weighted_score_display(scoring,score_map.get(oid),ranking_visible(meta)) if col==13 else float(score_map[oid]) if ranking_visible(meta) and score_map.get(oid) is not None else None)
                    if not _same(expected,ws.cell(r,col).value): add(ws.title,ws.cell(r,col).coordinate,f'{oid}:{field}',expected,ws.cell(r,col).value)
                total=fmap.get((oid,'total'),{})
                for col,expected in [(17,offer_display_status(oid,facts,snap.get('exclusions',[]),meta)),(18,evidence_label(total)),(19,', '.join(pending_offer_fields(oid,facts)))]:
                    expected=display_cell(ws.title,col,expected)
                    if not _same(expected,ws.cell(r,col).value): add(ws.title,ws.cell(r,col).coordinate,f'{oid}:display_status',expected,ws.cell(r,col).value)
                if evidence_label(total)!='Doğrulandı' and evidence_label(total) not in ws.cell(r,3).number_format:
                    add(ws.title,ws.cell(r,3).coordinate,f'{oid}:total_confirmation_label',evidence_label(total),ws.cell(r,3).number_format)
                term=buyer_term(c,'validity'); expected=term['value'] if term else off.get('validity_end')
                if not _same(expected,ws.cell(r,16).value): add(ws.title,ws.cell(r,16).coordinate,f'{oid}:validity_end',expected,ws.cell(r,16).value)
                # A buyer payment can exist even when the source field is absent.
                term=buyer_term(c,'payment')
                if term and not _same(term['value'],ws.cell(r,15).value):add(ws.title,ws.cell(r,15).coordinate,f'{oid}:payment',term['value'],ws.cell(r,15).value)
        else: add('Teklif Karsilastirma','sheet','required_sheet','present','missing')
        if 'Ozet' in wb.sheetnames:
            ws=wb['Ozet']; sm={ws.cell(r,1).value:ws.cell(r,2).value for r in range(2,ws.max_row+1)}
            snapshot_status={'READY':'İnceleme hazır','PRELIMINARY':'Ön değerlendirme'}.get(snap.get('status'),snap.get('status') or 'Bilinmiyor')
            for key,val in [('Karar İçin Bilgi Yeterliliği',meta.get('decision_status','UNKNOWN')),('İncelenen Kayıt Durumu',snapshot_status),('Profil',meta.get('profile','standart')),('Yonetici Ozeti',meta.get('decision_summary') or ''),('İncelenen sürüm özeti',snapshot_sha)]:
                observed=sm.get(label(key))
                if not _same(display_cell('Ozet',2,val),observed): add(ws.title,key,key,val,observed)
        else: add('Ozet','sheet','required_sheet','present','missing')
        expected_items=[]
        for fact in facts:
            for item in (_j(fact.get('normalized_json'),{}) or {}).get('items',[]):
                expected_items.append([fact.get('offer_id'),fact.get('field_key'),item.get('item_id'),item.get('description'),item.get('quantity'),item.get('unit'),item.get('amount_basis'),item.get('amount'),item.get('currency'),item.get('evidence_status') or fact.get('evidence_status'),json_text(item.get('evidence_refs',[]))])
        if expected_items:
            if 'Kalem Detayi' not in wb.sheetnames:
                add('Kalem Detayi','sheet','required_item_detail','present','missing')
            else:
                ws=wb['Kalem Detayi']
                if ws.max_row != len(expected_items)+1:
                    add(ws.title,'rows','item_count',len(expected_items),ws.max_row-1)
                for rn,row in enumerate(expected_items,start=2):
                    for col,val in enumerate(row,start=1):
                        expected=display_cell(ws.title,col,val)
                        if not _same(expected,ws.cell(rn,col).value):
                            add(ws.title,ws.cell(rn,col).coordinate,'item_detail',expected,ws.cell(rn,col).value)
        expected_subtotals=[]
        for fact in facts:
            calculation=(_j(fact.get('normalized_json'),{}) or {}).get('item_calculation')
            if calculation:
                expected_subtotals.append([fact.get('offer_id'),fact.get('field_key'),calculation.get('basis'),calculation.get('status'),calculation.get('amount'),calculation.get('currency'),json_text(calculation.get('unresolved_items',[])),'Hayir'])
        if expected_subtotals:
            if 'Kalem Hesabi' not in wb.sheetnames:
                add('Kalem Hesabi','sheet','calculated_item_subtotal','present','missing')
            else:
                ws=wb['Kalem Hesabi']
                if ws.max_row!=len(expected_subtotals)+1:
                    add(ws.title,'rows','calculated_subtotal_count',len(expected_subtotals),ws.max_row-1)
                for rn,row in enumerate(expected_subtotals,start=2):
                    for col,val in enumerate(row,start=1):
                        expected=display_cell(ws.title,col,val)
                        if not _same(expected,ws.cell(rn,col).value):
                            add(ws.title,ws.cell(rn,col).coordinate,'calculated_item_subtotal',expected,ws.cell(rn,col).value)
        decision=run_root/'output'/'decision_snapshot.json'
        if decision.exists():
            obj=json.loads(decision.read_text(encoding='utf-8'))
            for key,val in [('reviewed_snapshot_sha256',snapshot_sha),('decision_status',meta.get('decision_status')),('summary',meta.get('decision_summary') or '')]:
                if not _same(val,obj.get(key)): add('decision_snapshot.json',key,key,val,obj.get(key))
        else: add('decision_snapshot.json','file','artifact','present','missing')
        pdfs=[current['pdf']] if 'pdf' in current else []
        if pdfs:
            parents=_j(pdfs[0].get('parent_hashes_json'),[]) or []
            if snapshot_sha not in parents: add('PDF','lineage','reviewed_snapshot_parent',snapshot_sha,parents)
            if art['sha256'] not in parents: add('PDF','lineage','xlsx_parent',art['sha256'],parents)
        out={'status':'PASS' if not findings else 'FAIL','reviewed_snapshot_sha256':snapshot_sha,'xlsx':str(xlsx),'xlsx_sha256':sha256_file(xlsx),'findings':findings,'checked_at':utc_now(),'note':'PASS means no observed mismatch in mapped decision outputs; independent arithmetic is checked separately.'}
        p=run_root/'output'/'cross_output_consistency.json'; p.write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding='utf-8'); sha=sha256_file(p); aid=stable_id('art','cross_output_consistency',sha,length=22)
        con.execute("INSERT OR REPLACE INTO artifacts(artifact_id,artifact_type,relative_path,sha256,parent_hashes_json,status,created_at) VALUES (?,?,?,?,?,?,?)",[aid,'cross_output_consistency',str(p.relative_to(run_root)),sha,json_text([snapshot_sha,art['sha256']]),'ready',utc_now()])
        con.execute("DELETE FROM qa_findings WHERE finding_type='cross_output_consistency'")
        for f in findings:
            fid=stable_id('qa','cross_output',f['sheet'],f['locator'],f['field'],length=22); msg=f"Expected {f['expected']!r}; observed {f['observed']!r}"
            con.execute("INSERT OR REPLACE INTO qa_findings(finding_id,artifact_id,finding_type,locator_json,severity,status,message,finding_hash,created_at) VALUES (?,?,?,?,?,?,?,?,?)",[fid,art['artifact_id'],'cross_output_consistency',json_text({'sheet':f['sheet'],'locator':f['locator'],'field':f['field']}),'critical','open',msg,stable_id('h',msg,length=20),utc_now()])
        if hasattr(con,'commit'): con.commit()
        return out
    finally: con.close()

def main():
    ap=argparse.ArgumentParser(description='Reconcile frozen reviewed procurement snapshot to generated outputs'); ap.add_argument('--run',required=True); a=ap.parse_args(); print(json.dumps(check(Path(a.run)),ensure_ascii=False,indent=2,default=str))
if __name__=='__main__': main()
