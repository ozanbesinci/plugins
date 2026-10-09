from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

ERROR_TOKENS = {"#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}
CELL_REF_RE = re.compile(r"(?<![A-Za-z0-9_])((?:'[^']+'|[A-Za-z0-9_ .]+)!)?\$?[A-Z]{1,3}\$?\d+")


def _norm_formula(value: str) -> str:
    return CELL_REF_RE.sub(lambda m: (m.group(1) or "") + "CELL", value.upper().replace(" ", ""))


def audit(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    findings: list[dict[str, Any]] = []

    def add(severity: str, category: str, sheet: str, locator: str, message: str) -> None:
        findings.append({"severity": severity, "category": category, "sheet": sheet, "locator": locator, "message": message})

    for ws in formulas.worksheets:
        vws = values[ws.title]
        if ws.sheet_state != "visible":
            add("warning", "hidden_sheet", ws.title, ws.title, f"Sheet state is {ws.sheet_state}.")
        hidden_rows = [i for i, dim in ws.row_dimensions.items() if dim.hidden]
        hidden_cols = [k for k, dim in ws.column_dimensions.items() if dim.hidden]
        if hidden_rows:
            add("warning", "hidden_rows", ws.title, ",".join(map(str, hidden_rows[:50])), f"{len(hidden_rows)} hidden row(s) detected.")
        if hidden_cols:
            add("warning", "hidden_columns", ws.title, ",".join(hidden_cols[:50]), f"{len(hidden_cols)} hidden column(s) detected.")

        by_col: dict[int, list[tuple[int, str, Any]]] = {}
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                cached = vws[cell.coordinate].value
                if isinstance(cached, str) and cached in ERROR_TOKENS:
                    add("critical", "formula_error", ws.title, cell.coordinate, f"Cached Excel result is {cached}.")
                if isinstance(value, str) and value.startswith("="):
                    if "[" in value and "]" in value:
                        add("critical", "external_link", ws.title, cell.coordinate, "Formula contains an external-workbook link.")
                    by_col.setdefault(cell.column, []).append((cell.row, value, cached))

        for col, rows in by_col.items():
            if len(rows) < 3:
                continue
            sigs = Counter(_norm_formula(v) for _, v, _ in rows)
            dominant, dominant_n = sigs.most_common(1)[0]
            if dominant_n < 2:
                continue
            for row_no, formula, _ in rows:
                sig = _norm_formula(formula)
                if sig != dominant and sigs[sig] == 1:
                    cell = ws.cell(row_no, col).coordinate
                    add("warning", "formula_pattern_anomaly", ws.title, cell, "Formula pattern differs from the dominant formula pattern in this column.")

        # Detect pasted-over values only in columns that are clearly formula-driven.
        for col in range(1, ws.max_column + 1):
            cells = [ws.cell(r, col) for r in range(2, ws.max_row + 1) if ws.cell(r, col).value not in (None, "")]
            if len(cells) < 4:
                continue
            formula_cells = [c for c in cells if isinstance(c.value, str) and c.value.startswith("=")]
            if len(formula_cells) / len(cells) < 0.75:
                continue
            for c in cells:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    add("warning", "pasted_over_formula", ws.title, c.coordinate, "Non-formula value appears inside a formula-dominant column.")

    critical = sum(1 for f in findings if f["severity"] == "critical")
    return {"file": str(path), "status": "FAIL" if critical else "PASS", "critical": critical, "warnings": sum(1 for f in findings if f["severity"] == "warning"), "findings": findings}


def main() -> None:
    ap = argparse.ArgumentParser(description="Deterministic XLSX formula and structure audit")
    ap.add_argument("--input", required=True)
    ap.add_argument("--output")
    args = ap.parse_args()
    result = audit(Path(args.input))
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
    else:
        print(text)
    raise SystemExit(1 if result["status"] == "FAIL" else 0)


if __name__ == "__main__":
    main()
