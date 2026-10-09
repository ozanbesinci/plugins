"""Production transitions with independently authored expected values, no model calls."""
import json
import unittest
from pathlib import Path
from unittest.mock import patch
from db import connect,rows_as_dict,init_db
from common import json_text,sha256_file
from workflow import advance,prepare,_task_contexts,_policy
from context_packet import task_packet
from task_queue import _add_task
from record_results import record_task_result
from llm_result_validate import validate_result,ResultValidationError
from item_totals import calculate
import test_v504_audit_regressions as baseline


class BoundaryIntegration(unittest.TestCase):
    setUp=baseline.V504AuditRegressionTests.setUp
    tearDown=baseline.V504AuditRegressionTests.tearDown
    offer=baseline.V504AuditRegressionTests.offer
    valid_pair=baseline.V504AuditRegressionTests.valid_pair

    def finish_saved(self):
        step=advance(self.run);self.assertEqual(step['status'],'NEEDS_LLM')
        task=step['tasks'][0]
        c=connect(self.run)
        calc=c.execute("SELECT calculation_id FROM calculations WHERE status='complete' AND calculation_type='cost_summary' ORDER BY calculation_id LIMIT 1").fetchone()[0];c.close()
        result={'summary':'Maliyet karşılaştırması tamamlandı.','recommendation':None,'claims':[{'type':'calculated','text':'Karşılaştırılabilir maliyet hesaplanmıştır.','calculation_ids':[calc],'assertions':[{'source_type':'calculation','source_id':calc,'path':'comparable_total','operator':'exists'}]}]}
        path=self.run/'saved-summary.json';path.write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
        record_task_result(self.run,task['task_id'],result,str(path),model_name='unit-test-fixture',reasoning_effort='high',input_tokens=0,output_tokens=0)
        final=advance(self.run)
        c=connect(self.run);details=rows_as_dict(c,"SELECT check_id,risk,remediation FROM data_quality_findings WHERE status='open'");c.close()
        self.assertEqual(final['status'],'VERIFIED',json.dumps({'final':final,'quality':details},ensure_ascii=False))
        return path

    def test_closed_run_rechecks_receipt_archive_result_context_and_output(self):
        self.valid_pair();result_path=self.finish_saved()
        c=connect(self.run)
        docs=rows_as_dict(c,'SELECT metadata_json FROM documents ORDER BY document_id')
        archive=self.run/json.loads(docs[0]['metadata_json'])['archived_source_path']
        context=self.run/c.execute('SELECT context_path FROM task_dispatches LIMIT 1').fetchone()[0]
        xlsx=self.run/c.execute("SELECT relative_path FROM artifacts WHERE artifact_type='xlsx'").fetchone()[0]
        c.close()
        before={p: p.read_bytes() for p in [result_path,archive,context,xlsx]}
        for path in before:
            with self.subTest(file=path.name):
                path.write_bytes(b'corrupted')
                self.assertEqual(advance(self.run)['status'],'BLOCKED')
                self.assertEqual(path.read_bytes(),b'corrupted')
                path.write_bytes(before[path])
                self.assertEqual(advance(self.run)['status'],'VERIFIED')
        c=connect(self.run)
        receipts=rows_as_dict(c,'SELECT * FROM model_receipts');c.execute('DELETE FROM model_receipts');c.commit();c.close()
        self.assertEqual(advance(self.run)['status'],'BLOCKED')
        c=connect(self.run)
        for row in receipts:
            keys=','.join(row);marks=','.join('?' for _ in row)
            c.execute(f'INSERT INTO model_receipts({keys}) VALUES ({marks})',list(row.values()))
        c.commit();c.close()
        self.assertEqual(advance(self.run)['status'],'VERIFIED')

    def test_closed_run_new_pending_task_does_not_reuse_old_approval(self):
        self.valid_pair();self.finish_saved()
        c=connect(self.run);tid=_add_task(c,'REVIEW','independent_review','independent_fact_review',{'fact_ids':[]});c.commit();c.close()
        self.assertEqual(advance(self.run)['status'],'BLOCKED')
        c=connect(self.run);self.assertEqual(c.execute('SELECT status FROM task_queue WHERE task_id=?',[tid]).fetchone()[0],'pending');c.close()

    def test_final_validation_is_read_only_and_checks_closed_evidence(self):
        from validate_run import validate
        from closure_integrity import _closed_state
        self.valid_pair();self.finish_saved()
        c=connect(self.run);before=_closed_state(c);context=self.run/c.execute('SELECT context_path FROM task_dispatches LIMIT 1').fetchone()[0];c.close()
        files={p:p.read_bytes() for p in (self.run/'output').glob('*') if p.is_file()}
        self.assertEqual(validate(self.run)['status'],'VERIFIED')
        self.assertEqual({p:p.read_bytes() for p in files},files)
        c=connect(self.run);self.assertEqual(_closed_state(c),before);c.close()
        saved=context.read_bytes();context.write_bytes(b'changed dispatched packet')
        self.assertEqual(validate(self.run)['status'],'BLOCKED')
        self.assertEqual({p:p.read_bytes() for p in files},files)
        context.write_bytes(saved)
        self.assertEqual(advance(self.run)['status'],'VERIFIED')

    def test_dispatch_change_is_rejected_without_overwriting_context(self):
        self.valid_pair();step=advance(self.run);task=step['tasks'][0]
        context=Path(task['context_path']);before=context.read_bytes()
        c=connect(self.run);pinned=c.execute('SELECT packet_sha256 FROM task_dispatches WHERE task_id=?',[task['task_id']]).fetchone()[0]
        c.execute("UPDATE facts SET value_number=120000 WHERE field_key='total' AND value_number=100000");c.commit()
        queue=rows_as_dict(c,'SELECT * FROM task_queue WHERE task_id=?',[task['task_id']]);c.close()
        with self.assertRaisesRegex(ValueError,'preserved original context'):
            _task_contexts(self.run,_policy(Path(__file__).resolve().parents[2]),queue)
        self.assertEqual(context.read_bytes(),before)
        c=connect(self.run);self.assertEqual(c.execute('SELECT packet_sha256 FROM task_dispatches WHERE task_id=?',[task['task_id']]).fetchone()[0],pinned);c.close()

    def test_invalid_response_then_retry_uses_identical_bounded_packet(self):
        self.valid_pair();step=advance(self.run);task=step['tasks'][0];path=Path(task['context_path']);before=path.read_bytes()
        with self.assertRaises(ResultValidationError):
            record_task_result(self.run,task['task_id'],{'unknown':'invalid'},model_name='unit-test-fixture',reasoning_effort='high',input_tokens=0,output_tokens=0)
        second=advance(self.run);self.assertEqual(second['status'],'NEEDS_LLM')
        self.assertEqual(Path(second['tasks'][0]['context_path']).read_bytes(),before)

    def semantic_task(self,oid,did,field):
        c=connect(self.run);tid=_add_task(c,'EXTRACT','semantic_extraction','offer_semantic_fields',{'offer_id':oid,'document_ids':[did],'fields':[field]},offer_id=oid);c.commit();tasks=rows_as_dict(c,'SELECT * FROM task_queue WHERE task_id=?',[tid]);c.close()
        _task_contexts(self.run,_policy(Path(__file__).resolve().parents[2]),tasks)
        return tid

    def test_foreign_scope_evidence_rejected_atomically(self):
        self.valid_pair();c=connect(self.run)
        rows=rows_as_dict(c,"SELECT o.offer_id,d.document_id,f.evidence_id FROM offers o JOIN documents d ON d.offer_id=o.offer_id JOIN facts f ON f.offer_id=o.offer_id AND f.field_key='total' ORDER BY o.offer_id")
        oid,did=rows[0]['offer_id'],rows[0]['document_id'];foreign=rows[1]['evidence_id'];c.close()
        tid=self.semantic_task(oid,did,'currency')
        result={'evidence':[],'facts':[{'field_key':'currency','value_text':'EUR','evidence_status':'missing'}],'scope_items':[{'classification':'freight','cost_role':'incremental_adjustment','description':'extra freight','status':'verified','amount':105000,'currency':'EUR','evidence_ids':[foreign]}]}
        c=connect(self.run);before=c.execute('SELECT count(*) FROM model_receipts').fetchone()[0];c.close()
        with self.assertRaisesRegex(ValueError,'outside task scope|does not belong'):
            record_task_result(self.run,tid,result,model_name='unit-test-fixture',reasoning_effort='high',input_tokens=0,output_tokens=0)
        c=connect(self.run);self.assertEqual(c.execute('SELECT count(*) FROM scope_items').fetchone()[0],0);self.assertEqual(c.execute('SELECT count(*) FROM model_receipts').fetchone()[0],before);c.close()

    def test_item_evidence_keeps_own_scale_through_record_refresh_and_xlsx(self):
        # Item table headers are distinct from the form's explicit product
        # label; this fixture isolates item/aggregate evidence transitions.
        from openpyxl import load_workbook
        # Both stated totals reconcile with two 10 x 100 commercial lines.
        # A contradictory 100000/105000 fixture correctly trips quality QA.
        for name,total in [('Alpha_teklif.xlsx',2000),('Beta_teklif.xlsx',2000)]:
            self.offer(name,total,quantity=20,line_totals=True)
            wb=load_workbook(self.src/name);wb.active['A14']='Kalem kodu';wb.save(self.src/name);wb.close()
        from datetime import date
        prepare(self.src,self.run,'hizli',date(2026,10,3))
        c=connect(self.run);oid,did=c.execute("SELECT o.offer_id,d.document_id FROM offers o JOIN documents d ON d.offer_id=o.offer_id WHERE d.filename='Alpha_teklif.xlsx'").fetchone();c.close()
        tid=self.semantic_task(oid,did,'quantity')
        result={'evidence':[
            {'client_id':'aggregate','document_id':did,'locator':{'sheet':'Teklif','cell':'B7'},'quote':'20','evidence_status':'verified'},
            {'client_id':'item-a-q','document_id':did,'locator':{'sheet':'Teklif','cell':'B15'},'quote':'10','evidence_status':'verified'},
            {'client_id':'item-a-p','document_id':did,'locator':{'sheet':'Teklif','cell':'C15'},'quote':'100','evidence_status':'verified'},
            {'client_id':'item-b-q','document_id':did,'locator':{'sheet':'Teklif','cell':'B16'},'quote':'10','evidence_status':'verified'},
            {'client_id':'item-b-p','document_id':did,'locator':{'sheet':'Teklif','cell':'C16'},'quote':'100','evidence_status':'verified'}],
            'facts':[{'field_key':'quantity','value_number':20,'value_text':'20','unit':'adet','evidence_status':'verified','evidence_client_id':'aggregate','normalized':{'items':[
                {'item_id':'a','description':'Kalem 1','quantity':10,'unit':'adet','amount':100,'amount_basis':'unit_price','currency':'EUR','evidence_status':'verified','evidence_refs':['item-a-q','item-a-p']},
                {'item_id':'b','description':'Kalem 2','quantity':10,'unit':'adet','amount':100,'amount_basis':'unit_price','currency':'EUR','evidence_status':'verified','evidence_refs':['item-b-q','item-b-p']} ]}}]}
        path=self.run/'saved-items.json';path.write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
        record_task_result(self.run,tid,result,str(path),model_name='unit-test-fixture',reasoning_effort='high',input_tokens=0,output_tokens=0)
        from field_extract import extract
        from normalize import normalize
        for _ in range(2):extract(self.run);normalize(self.run)
        c=connect(self.run);items=json.loads(c.execute('SELECT normalized_json FROM facts WHERE fact_id=?',[oid+'/field/quantity']).fetchone()[0])['items'];c.close()
        self.assertEqual([i['quantity'] for i in items],[10,10]);self.assertEqual([i['amount'] for i in items],[100,100])
        self.assertNotEqual(items[0]['evidence_refs'],['item-a'])
        self.finish_saved()
        from openpyxl import load_workbook
        wb=load_workbook(self.run/'output/Teklif_Degerlendirme.xlsx',data_only=True);ws=wb['Kalem Detayi']
        self.assertEqual([ws.cell(r,5).value for r in [2,3]],[10,10]);self.assertEqual([ws.cell(r,8).value for r in [2,3]],[100,100])
        self.assertEqual(wb['Kalem Hesabi']['D2'].value,'calculated');self.assertEqual(wb['Kalem Hesabi']['E2'].value,'2000');self.assertEqual(wb['Kalem Hesabi']['H2'].value,'Hayir');wb.close()
        from cross_output_consistency import check
        wb=load_workbook(self.run/'output/Teklif_Degerlendirme.xlsx');wb['Kalem Detayi']['H2']=999;wb.save(self.run/'output/Teklif_Degerlendirme.xlsx');wb.close()
        self.assertEqual(check(self.run)['status'],'FAIL')

    def test_alternative_item_keys_do_not_silently_export_blank_prices(self):
        root=Path(__file__).resolve().parents[2]
        result={'evidence':[],'facts':[{'field_key':'total','evidence_status':'requires_confirmation','normalized':{'items':[{'item_id':'a','evidence_refs':['e'],'evidence_status':'requires_confirmation','order_quantity':10,'discounted_unit_price':100}]}}]}
        with self.assertRaises(ResultValidationError):validate_result(root,'offer_semantic_fields',result)

    def test_pixel_evidence_cannot_verify_nested_numeric_item(self):
        self.valid_pair();c=connect(self.run)
        oid,did,eid=c.execute("SELECT o.offer_id,d.document_id,f.evidence_id FROM offers o JOIN documents d ON d.offer_id=o.offer_id JOIN facts f ON f.offer_id=o.offer_id AND f.field_key='total' WHERE d.filename='Alpha_teklif.xlsx'").fetchone()
        c.execute("UPDATE evidence SET extraction_path='vision' WHERE evidence_id=?",[eid]);c.commit();c.close()
        tid=self.semantic_task(oid,did,'total')
        result={'evidence':[],'facts':[{'field_key':'total','evidence_status':'requires_confirmation','evidence_id':eid,'normalized':{'items':[{'item_id':'pixel-price','amount':100000,'amount_basis':'line_total','currency':'EUR','evidence_status':'verified','evidence_refs':[eid]}]}}]}
        with self.assertRaisesRegex(ValueError,'non-pixel|independent confirmation'):
            record_task_result(self.run,tid,result,model_name='unit-test-fixture',reasoning_effort='high',input_tokens=0,output_tokens=0)
        c=connect(self.run);self.assertEqual(c.execute('SELECT status FROM task_queue WHERE task_id=?',[tid]).fetchone()[0],'pending');self.assertEqual(c.execute("SELECT value_number FROM facts WHERE fact_id=?",[oid+'/field/total']).fetchone()[0],100000);c.close()

    def test_unit_price_without_quantity_does_not_become_procurement_total(self):
        result=calculate([{'item_id':'a','amount':100,'amount_basis':'unit_price','currency':'EUR','evidence_status':'verified'}])
        self.assertEqual(result['status'],'incomplete');self.assertIsNone(result['amount']);self.assertFalse(result['is_stated_grand_total'])

    def test_line_subtotal_is_explicitly_calculated_not_stated_total(self):
        result=calculate([{'item_id':'a','amount':1000,'amount_basis':'line_total','currency':'EUR','evidence_status':'verified'},{'item_id':'b','amount':2000,'amount_basis':'line_total','currency':'EUR','evidence_status':'verified'}])
        self.assertEqual(result['amount'],'3000');self.assertEqual(result['status'],'calculated');self.assertFalse(result['is_stated_grand_total'])

    def test_mixed_views_of_same_commercial_line_are_not_added_twice(self):
        result=calculate([{'item_id':'unit','quantity':2,'amount':100,'amount_basis':'unit_price','currency':'EUR','evidence_status':'verified'},{'item_id':'line','amount':200,'amount_basis':'line_total','currency':'EUR','evidence_status':'verified'}])
        self.assertEqual(result['status'],'incomplete');self.assertIsNone(result['amount']);self.assertIn('mixed_price_bases_require_reconciliation',result['unresolved_items'])

    def test_old_database_missing_column_is_migrated_and_failed_alter_not_stamped(self):
        init_db(self.run);c=connect(self.run);c.execute('ALTER TABLE scope_items DROP COLUMN cost_role');c.execute("UPDATE run_meta SET value='5.0.5' WHERE key='schema_version'");c.commit();c.close()
        init_db(self.run);c=connect(self.run);self.assertIn('cost_role',{r[1] for r in c.execute("PRAGMA table_info('scope_items')").fetchall()});c.execute('ALTER TABLE scope_items DROP COLUMN cost_role');c.execute("UPDATE run_meta SET value='5.0.5' WHERE key='schema_version'");c.commit();c.close()
        real_connect=connect
        class BrokenAlter:
            def __init__(self,c):self.c=c
            def execute(self,sql,*args):
                if sql=='SHOW TABLES' and hasattr(self.c,'executescript'):sql="SELECT name FROM sqlite_master WHERE type='table'"
                if sql.startswith('ALTER TABLE scope_items ADD'):raise RuntimeError('injected ALTER failure')
                return self.c.execute(sql,*args)
            def executescript(self,sql):return self.c.executescript(sql)
            def commit(self):return self.c.commit()
            def close(self):return self.c.close()
        with patch('db.connect',side_effect=lambda run:BrokenAlter(real_connect(run))),patch('db.execute_script',side_effect=lambda con,sql:None):
            with self.assertRaisesRegex(RuntimeError,'injected ALTER failure'):init_db(self.run)
        c=connect(self.run);self.assertEqual(c.execute("SELECT value FROM run_meta WHERE key='schema_version'").fetchone()[0],'5.0.5');c.close()
