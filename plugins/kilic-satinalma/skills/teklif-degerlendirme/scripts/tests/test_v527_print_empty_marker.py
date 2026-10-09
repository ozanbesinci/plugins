from pathlib import Path
import shutil
import sys
import tempfile
import unittest

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from print_layout import build_print_workbook


class EmptyPrintMarkerTests(unittest.TestCase):
    def test_empty_notice_is_prominent_on_every_continuation_band(self):
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Kalem Hesabi"
            for col in range(1, 20):
                ws.cell(1, col, f"Alan {col}")
                ws.column_dimensions[ws.cell(1, col).column_letter].width = 32
            ws.cell(2, 1, "Kalem hesabı kaydı bulunmuyor.")
            wb.save(root / "input.xlsx")
            wb.close()

            info = build_print_workbook(root / "input.xlsx", root / "print.xlsx")
            printed = load_workbook(root / "print.xlsx")
            self.assertGreater(len(printed.worksheets), 1)
            self.assertEqual(len(info["mapping"]), len(printed.worksheets))
            for sheet in printed.worksheets:
                marker = sheet.cell(2, 1)
                self.assertEqual(marker.value, "Kayıt bulunmuyor")
                self.assertEqual(marker.font.sz, 18)
                self.assertTrue(marker.font.bold)
                self.assertEqual(marker.fill.fgColor.rgb, "00FFF2CC")
                self.assertEqual(sheet.row_dimensions[2].height, 54)
                self.assertEqual(sheet.print_title_rows, "$1:$1")
                self.assertIn("Sütun grubu", sheet.oddHeader.center.text)
                if sheet.max_column > 1:
                    self.assertEqual([str(rng) for rng in sheet.merged_cells.ranges], [f"A2:{get_column_letter(sheet.max_column)}2"])
            printed.close()
        finally:
            shutil.rmtree(root)

    def test_marker_detection_requires_the_row_to_contain_only_the_notice(self):
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Tablo"
            ws.append(["Kimlik", "Durum"])
            ws.append(["Kayıt bulunmuyor.", "Gerçek kayıt"])
            wb.save(root / "input.xlsx")
            wb.close()

            build_print_workbook(root / "input.xlsx", root / "print.xlsx")
            printed = load_workbook(root / "print.xlsx")
            self.assertEqual(" ".join(str(printed.active["A2"].value).split()), "Kayıt bulunmuyor.")
            self.assertEqual(" ".join(str(printed.active["B2"].value).split()), "Gerçek kayıt")
            self.assertFalse(printed.active.merged_cells.ranges)
            printed.close()
        finally:
            shutil.rmtree(root)


if __name__ == "__main__":
    unittest.main()
