from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
from adapters.docx_adapter import _tracked_changes, parse

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


class DocxRevisionTagTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def source(self, body, prefix='w', namespace=W):
        path = self.root / 'offer.docx'
        xml = f'<{prefix}:document xmlns:{prefix}="{namespace}"><{prefix}:body>{body}</{prefix}:body></{prefix}:document>'
        with zipfile.ZipFile(path, 'w') as z:
            z.writestr('word/document.xml', xml)
        return path

    def test_table_borders_are_not_revisions(self):
        p = self.source('<w:tbl><w:tblPr><w:tblBorders><w:insideH/><w:insideV/></w:tblBorders></w:tblPr></w:tbl>')
        self.assertFalse(_tracked_changes(p))

    def test_real_insert_delete_move_and_property_changes_are_detected(self):
        for tag in ['ins','del','moveFrom','moveTo','pPrChange','rPrChange','sectPrChange',
                    'tblPrChange','trPrChange','tcPrChange','cellIns','cellDel','cellMerge',
                    'customXmlInsRangeStart','customXmlDelRangeStart','numberingChange']:
            with self.subTest(tag=tag):
                self.assertTrue(_tracked_changes(self.source(f'<w:{tag}/>')))

    def test_namespace_not_prefix_controls_detection(self):
        self.assertTrue(_tracked_changes(self.source('<x:ins/>', prefix='x')))
        self.assertFalse(_tracked_changes(self.source('<w:ins/>', namespace='urn:not-word')))

    def test_revision_words_in_text_are_not_xml_revisions(self):
        self.assertFalse(_tracked_changes(self.source('<w:p><w:r><w:t>&lt;w:ins, &lt;w:del and insideH</w:t></w:r></w:p>')))

    def test_border_document_parse_is_read_and_real_revision_is_partial(self):
        from docx import Document
        from docx.oxml import OxmlElement
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0,0).text = 'Toplam 1250 EUR'
        borders = OxmlElement('w:tblBorders')
        borders.append(OxmlElement('w:insideH'))
        borders.append(OxmlElement('w:insideV'))
        table._tbl.tblPr.append(borders)
        path = self.root / 'full.docx'
        doc.save(path)
        result = parse(path)
        self.assertEqual(result.status, 'read')
        self.assertFalse(result.metadata['tracked_changes'])
        self.assertIn('1250 EUR', result.units[0].text)
        table._tbl.append(OxmlElement('w:ins'))
        doc.save(path)
        result = parse(path)
        self.assertEqual(result.status, 'partial')
        self.assertTrue(result.metadata['tracked_changes'])

if __name__ == '__main__':
    unittest.main()
