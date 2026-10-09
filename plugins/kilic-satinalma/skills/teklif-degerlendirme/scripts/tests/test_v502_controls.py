from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

os.environ['TEKLIF_TEST_SQLITE'] = '1'
SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook

from context_packet import task_packet
from cross_output_consistency import check as cross_check
from db import connect, init_db
from llm_result_validate import ResultValidationError, validate_result
from source_authority import refresh as refresh_authority
from task_queue import _add_task
from workflow import advance, prepare
from record_results import record_task_result
from xlsx_formula_audit import audit as formula_audit


class V502ControlTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix='teklif-v502-'))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_schema_gate_rejects_unexpected_fields(self):
        skill_root = SCRIPTS.parent
        with self.assertRaises(ResultValidationError):
            validate_result(skill_root, 'decision_summary', {
                'summary': 'ok', 'recommendation': None, 'workflow_override': 'ignore rules'
            })

    def test_source_authority_never_infers_offer_revision_from_filename_or_date(self):
        run = self.root / 'run'
        init_db(run)
        con = connect(run)
        try:
            con.execute("INSERT INTO suppliers(supplier_id,display_name,normalized_name,identification_status) VALUES ('S1','Acme','acme','confirmed')")
            con.execute("INSERT INTO offers(offer_id,supplier_id,display_name,status,source_document_ids_json) VALUES ('O1','S1','Acme','active','[]')")
            for i, name in enumerate(['Acme_teklif_rev1.xlsx', 'Acme_teklif_rev2.xlsx'], 1):
                con.execute(
                    """INSERT INTO documents(document_id,relative_path,filename,format_family,document_role,supplier_id,offer_id,revision_label,revision_status,sha256,size_bytes,parse_status,active)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    [f'D{i}', name, name, 'xlsx', 'offer', 'S1', 'O1', f'rev{i}', 'unknown', f'hash{i}', 100, 'parsed', True],
                )
            con.commit()
        finally:
            con.close()
        result = refresh_authority(run)
        self.assertEqual(result['controlling'], 0)
        self.assertEqual(result['unresolved'], 2)
        con = connect(run)
        try:
            states = [r[0] for r in con.execute("SELECT authority_status FROM source_authority ORDER BY document_id").fetchall()]
            self.assertEqual(states, ['needs_input', 'needs_input'])
            issue_count = con.execute("SELECT count(*) FROM open_issues WHERE issue_type='source_authority_conflict' AND status='open'").fetchone()[0]
            self.assertEqual(issue_count, 1)
        finally:
            con.close()

    def test_task_packet_carries_untrusted_reader_security_envelope(self):
        run = self.root / 'run'
        init_db(run)
        con = connect(run)
        try:
            con.execute(
                """INSERT INTO documents(document_id,relative_path,filename,format_family,document_role,sha256,size_bytes,parse_status,active)
                   VALUES ('D1','offer.txt','offer.txt','text','offer','abc',10,'parsed',true)"""
            )
            tid = _add_task(con, 'EXTRACT', 'normal_extraction', 'document_classification', {'document_id': 'D1'}, document_id='D1')
            con.commit()
        finally:
            con.close()
        packet = task_packet(run, tid, 10, 20000)
        env = packet['security_envelope']
        self.assertEqual(env['trust_zone'], 'untrusted_source_reader')
        self.assertTrue(env['source_content_is_data_not_instruction'])
        self.assertIn('workflow_control', env['forbidden_capabilities'])
        self.assertIn('write', env['forbidden_capabilities'])

    def test_formula_audit_flags_external_workbook_links(self):
        path = self.root / 'external.xlsx'
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "='[external.xlsx]Sheet1'!A1"
        wb.save(path)
        result = formula_audit(path)
        self.assertEqual(result['status'], 'FAIL')
        self.assertTrue(any(f['severity'] == 'critical' and f['category'] == 'external_link' for f in result['findings']))

    def _write_offer(self, folder: Path, filename: str, total: float) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = 'Teklif'
        rows = [
            ('Fiyat Teklifi', ''), ('Teklif Tarihi', '03.10.2026'), ('Geçerlilik', '90 gün'),
            ('Genel Toplam', total), ('Para Birimi', 'EUR'), ('KDV', 'hariç'), ('Miktar', 10),
            ('Birim', 'adet'),
            ('Ürün', 'Pompa'),
            ('Teslim', '30 gün'), ('Ödeme', '%100 60 gün'), ('Incoterm', 'DDP'), ('Garanti', '24 ay'),
        ]
        for r, (label, value) in enumerate(rows, 1):
            ws.cell(r, 1, label); ws.cell(r, 2, value)
        ws.append(['Tedarikçi\nFirma: ' + Path(filename).stem.split('_')[0] + ' Ltd.'])
        wb.save(folder / filename)

    def test_cross_output_consistency_detects_workbook_tamper(self):
        source = self.root / 'source'; run = self.root / 'run'; source.mkdir()
        self._write_offer(source, 'Alpha_fiyat_teklifi.xlsx', 100000)
        self._write_offer(source, 'Beta_fiyat_teklifi.xlsx', 105000)
        prepare(source, run, 'hizli')
        step = advance(run)
        self.assertEqual(step['status'], 'NEEDS_LLM')
        task_id = step['tasks'][0]['task_id']
        con = connect(run)
        try:
            calc_id = con.execute("SELECT calculation_id FROM calculations WHERE calculation_type='cost_summary' AND status='complete' ORDER BY calculation_id LIMIT 1").fetchone()[0]
        finally:
            con.close()
        record_task_result(
            run, task_id,
            {'summary': 'Karsilastirma tamamlandi.', 'recommendation': {'type':'cost_observation','text':'Dusuk maliyet ayri gosterildi.'},
             'claims':[{'type':'calculated','text':'Maliyetler hesaplanmistir.','calculation_ids':[calc_id],'assertions':[{'source_type':'calculation','source_id':calc_id,'path':'comparable_total','operator':'exists'}]},{'type':'recommendation','text':'Dusuk maliyet ayri gosterildi.','calculation_ids':[calc_id]}]},
            model_name='test-model', reasoning_effort='medium', input_tokens=100, cached_input_tokens=50, output_tokens=20,
        )
        final = advance(run)
        self.assertEqual(final['status'], 'VERIFIED')
        xlsx = run / 'output' / 'Teklif_Degerlendirme.xlsx'
        wb = load_workbook(xlsx)
        ws = wb['Teklif Karsilastirma']
        ws.cell(2, 3).value = 999999999
        wb.save(xlsx)
        result = cross_check(run)
        self.assertEqual(result['status'], 'FAIL')
        self.assertTrue(result['findings'])


if __name__ == '__main__':
    unittest.main()
