from pathlib import Path
import json, shutil, sys, tempfile, unittest
from decimal import Decimal

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

from openpyxl import Workbook, load_workbook
from calculation_detail import COLUMN_WIDTHS, HEADERS, MONEY_FORMAT, calculation_rows
from print_layout import build_print_workbook, wrap_text

COST = {"known": True, "complete": True, "stated_total": "105000.0", "known_scope_adjustment": "0",
        "acquisition_cost_basis": "105000.0", "acquisition_present_value": "103367.7402283199989683318927",
        "tco_present_value": None, "comparable_total": "103367.7402283199989683318927", "currency": "EUR",
        "unknown_scope_items": 0, "base_currency": "EUR", "base_date": None, "base_amount": "103367.7402283199989683318927",
        "fx_used": {"rate": "1", "unit": "1", "source": "same_currency", "date": None},
        "payment_present_value_detail": [{"days": 60, "amount": "105000.00", "ratio": "1.0", "present_value": "103367.7402283199989683318927"}],
        "warnings": [], "steps": [{"stage": "stated_total", "amount": "105000.0", "currency": "EUR"},
                                  {"stage": "scope_5A_known", "amount": "0", "currency": "EUR", "unknown_items": 0}]}
SCORING = {"status": "cost_only", "method": "10_x_min_cost_over_offer_cost",
           "price_scores": [{"offer_id": "off_b", "cost_metric": "103367.7402283199989683318927", "price_score_0_10": "9.52"},
                            {"offer_id": "off_a", "cost_metric": "98445.46688411428473174465970", "price_score_0_10": "10.00"}],
           "cost_ranking": ["off_a", "off_b"], "note": "Cost score is not a total decision score."}
CALCS = [{"calculation_id": "calc_cost_b", "offer_id": "off_b", "calculation_type": "cost_summary", "status": "complete", "result_json": json.dumps(COST)},
         {"calculation_id": "calc_score", "offer_id": None, "calculation_type": "scoring", "status": "complete", "result_json": json.dumps(SCORING)}]
NAMES = {"off_a": "Alpha teklif", "off_b": "Beta teklif"}


def _leaves(obj):
    if isinstance(obj, dict):
        for v in obj.values():
            yield from _leaves(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _leaves(v)
    else:
        yield obj


class CalculationDetailTests(unittest.TestCase):
    def test_amounts_are_numeric_cells_with_currency_and_full_double_precision(self):
        rows = calculation_rows(CALCS, NAMES, {"off_a": "EUR", "off_b": "EUR"})
        pv = [r for r in rows if r["values"][3] == "Alım bugünkü değeri"]
        self.assertEqual(len(pv), 1)
        value = pv[0]["values"][4]
        self.assertIsInstance(value, float)
        self.assertEqual(value, float(Decimal("103367.7402283199989683318927")))
        self.assertEqual(pv[0]["values"][5], "EUR")
        self.assertEqual(pv[0]["number_format"], MONEY_FORMAT)
        self.assertEqual(pv[0]["values"][:3], ["Beta teklif", "off_b", "cost_summary"])
        self.assertEqual(pv[0]["calculation_id"], "calc_cost_b")

    def test_no_leaf_value_is_dropped(self):
        rows = calculation_rows(CALCS, NAMES, {})
        cost_rows = [r for r in rows if r["calculation_id"] == "calc_cost_b"]
        numeric = sorted(r["values"][4] for r in cost_rows if isinstance(r["values"][4], float))
        expected = sorted(float(Decimal(str(v))) for v in _leaves(COST)
                          if v is not None and not isinstance(v, bool) and str(v).replace(".", "", 1).isdigit())
        self.assertEqual(numeric, expected)

    def test_scores_sit_with_their_offer_cost_and_rank_in_cost_order(self):
        rows = [r for r in calculation_rows(CALCS, NAMES, {"off_a": "EUR", "off_b": "EUR"}) if r["calculation_id"] == "calc_score"]
        per_offer = [r for r in rows if r["offer_id"]]
        self.assertEqual([r["offer_id"] for r in per_offer][:3], ["off_a"] * 3)
        a = {r["values"][3]: r for r in per_offer if r["offer_id"] == "off_a"}
        self.assertEqual(a["Maliyet ölçüsü"]["values"][5], "EUR")
        self.assertEqual(a["Fiyat puanı (0–10)"]["values"][4], 10.0)
        self.assertEqual(a["Maliyet sırası"]["values"][4], 1.0)
        self.assertEqual(a["Maliyet ölçüsü"]["values"][0], "Alpha teklif")
        b = {r["values"][3]: r for r in per_offer if r["offer_id"] == "off_b"}
        self.assertEqual(b["Maliyet sırası"]["values"][4], 2.0)
        self.assertEqual(b["Fiyat puanı (0–10)"]["values"][4], 9.52)
        general = [r for r in rows if not r["offer_id"]]
        self.assertTrue(all(r["values"][0] == "Tüm teklifler" for r in general))
        self.assertIn("Not", [r["values"][3] for r in general])

    def test_print_keeps_numbers_numeric_and_never_splits_integer_digits(self):
        root = Path(tempfile.mkdtemp())
        try:
            wb = Workbook(); ws = wb.active; ws.title = "Hesap Detayi"; ws.append(HEADERS)
            for r in calculation_rows(CALCS, NAMES, {"off_a": "EUR", "off_b": "EUR"}):
                ws.append(r["values"])
                if r["number_format"]:
                    ws.cell(ws.max_row, 5).number_format = r["number_format"]
            for col, width in COLUMN_WIDTHS.items():
                ws.column_dimensions[col].width = width
            wb.save(root / "in.xlsx"); wb.close()
            info = build_print_workbook(root / "in.xlsx", root / "print.xlsx")
            p = load_workbook(root / "print.xlsx")
            self.assertEqual(len(p.worksheets), 1, "detail table must stay in one column band")
            sheet = p.worksheets[0]
            numbers = [sheet.cell(r, 5) for r in range(2, sheet.max_row + 1) if isinstance(sheet.cell(r, 5).value, float)]
            self.assertTrue(numbers)
            self.assertTrue(any(c.number_format == MONEY_FORMAT for c in numbers))
            for r in range(2, sheet.max_row + 1):
                value = sheet.cell(r, 5).value
                if isinstance(value, str):
                    self.assertNotRegex(value.replace("\n", ""), r"^\s*-?[\d.,]+\s*$", f"amount/score left as text: E{r} {value!r}")
            self.assertTrue(info["all_columns_retained"])
            p.close()
        finally:
            shutil.rmtree(root)

    def test_wrap_text_splits_long_numbers_only_after_decimal_point(self):
        lines = wrap_text("Düzeltme tutarı: -1632.2597716800010316681073", 12)
        self.assertEqual(lines[:3], ["Düzeltme", "tutarı:", "-1632.259771"])
        self.assertEqual("".join(lines[2:]), "-1632.2597716800010316681073")
        # The former defect: label line ended with the first integer digits.
        lines = wrap_text("Alım bugünkü değeri: 103367.7402283199989683318927", 26)
        self.assertEqual(lines, ["Alım bugünkü değeri:", "103367.7402283199989683318", "927"])
        lines = wrap_text("toplam:103367.7402283199989683318927", 10)
        self.assertEqual(lines[0], "toplam:")
        self.assertTrue(lines[1].startswith("103367."))
        self.assertEqual(wrap_text("off_3d0b4f5791489aac4c", 16), ["off_3d0b4f579148", "9aac4c"])


if __name__ == "__main__":
    unittest.main()
