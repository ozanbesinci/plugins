"""Bounded source cell fill metadata for priced worksheet cells."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Color, PatternFill

from adapters.xlsx_adapter import _fill_payload, _theme_colors, _tinted_rgb, parse


class XlsxFillPayloadTests(unittest.TestCase):
    def test_rgb_and_theme_tint_survive_parse_without_blank_style_cells(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'priced.xlsx'
            wb = Workbook()
            ws = wb.active
            ws['F7'] = 100
            ws['F7'].fill = PatternFill(patternType='solid', fgColor='FF00FF00')
            ws['G7'] = 200
            ws['G7'].fill = PatternFill(patternType='solid', fgColor=Color(theme=4, tint=0.5))
            ws['H7'].fill = PatternFill(patternType='solid', fgColor='FFFF0000')
            wb.save(path)
            result = parse(path)
            self.assertEqual(result.status, 'read')
            cells = {cell['cell']: cell for cell in result.units[0].structured['cells']}
            self.assertNotIn('H7', cells)
            self.assertEqual(cells['F7']['fill']['foreground']['rgb'], 'FF00FF00')
            self.assertEqual(cells['F7']['fill']['effective_rgb'], '00FF00')
            self.assertEqual(cells['G7']['fill']['foreground']['theme'], 4)
            self.assertEqual(cells['G7']['fill']['foreground']['tint'], 0.5)
            source_theme = _theme_colors(load_workbook(path))
            self.assertIn(4, source_theme)
            self.assertEqual(cells['G7']['fill']['effective_rgb'],
                             _tinted_rgb(source_theme[4], 0.5))

    def test_theme_xml_last_clr_and_tint(self):
        wb = Workbook()
        wb.loaded_theme = (b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                           b'<a:themeElements><a:clrScheme name="test">'
                           b'<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
                           b'<a:accent1><a:srgbClr val="000000"/></a:accent1>'
                           b'</a:clrScheme></a:themeElements></a:theme>')
        colors = _theme_colors(wb)
        self.assertEqual(colors[1], '000000')
        self.assertEqual(colors[4], '000000')
        cell = wb.active['A1']
        cell.fill = PatternFill(patternType='solid', fgColor=Color(theme=4, tint=0.5))
        self.assertEqual(_fill_payload(cell, colors)['effective_rgb'], '808080')


if __name__ == '__main__':
    unittest.main()
