"""Sparse OOXML cell inventory; cached values are never called recalculated values."""
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS
from openpyxl.utils.cell import coordinate_to_tuple
from ..errors import CoreError
from .common import Package, feature, issue, locator, result_base, unit

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
q = lambda n: "{" + S + "}" + n


def read_xlsx(data):
    report = result_base(data, "XLSX", "ooxml-xlsx", ["lxml", "openpyxl"])
    package = Package(data)
    main = package.main("XLSX")
    root = package.roots[main]
    if root.tag != q("workbook"):
        raise CoreError("DOSYA_TURU", "Spreadsheet namespace desteklenmiyor")
    package.describe(report)
    shared, styles, custom_formats = [], [], {}
    for item in package.roots.values():
        if item.tag == q("sst"):
            shared = ["".join(t.text or "" for t in si.iter(q("t"))) for si in item]
        if item.tag == q("styleSheet"):
            custom_formats = {int(x.get("numFmtId")): x.get("formatCode", "")
                              for x in item.findall("./" + q("numFmts") + "/" + q("numFmt"))}
            styles = [int(x.get("numFmtId", "0")) for x in item.findall("./" + q("cellXfs") + "/" + q("xf"))]
    props = root.find(q("workbookPr"))
    date_system = "1904" if props is not None and props.get("date1904") in {"1", "true"} else "1900"
    feature(report, "date_system", 1, "TESTED", date_system)
    names = root.findall("./" + q("definedNames") + "/" + q("definedName"))
    feature(report, "defined_names", len(names), "PARTIAL", "Adlar envanterlendi; hesaplanmadı")
    sheet_names, total_cells = [], 0
    for sheet in root.findall("./" + q("sheets") + "/" + q("sheet")):
        title = sheet.get("name")
        if not title or title in sheet_names:
            raise CoreError("BELGE_BOZUK", "Sheet adı yok/yineleniyor")
        sheet_names.append(title)
        relation = next((x for x in package.relationships if x["source"] == main
                         and x["id"] == sheet.get("{" + R + "}id")), None)
        if not relation or relation["external"]:
            raise CoreError("ILISKI_EKSIK", "Sheet için iç ilişki gerekli")
        part = relation["resolved_part"]
        ws = package.roots.get(part)
        if ws is None or ws.tag != q("worksheet"):
            issue(report, "UNSUPPORTED_SHEET", "Worksheet dışındaki sheet içeriği okunmadı",
                  locator(part, sheet=title), "BLOCKING")
            continue
        state = sheet.get("state", "visible")
        feature(report, "sheet:" + title, 1, "TESTED", "state=" + state)
        if state != "visible":
            issue(report, "HIDDEN_SHEET", "Gizli sheet içeriği atılmadan okundu", locator(part, sheet=title))
        rows = ws.findall("./" + q("sheetData") + "/" + q("row"))
        hidden_rows = {x.get("r") for x in rows if x.get("hidden") in {"1", "true"}}
        hidden_cols = [(int(x.get("min")), int(x.get("max"))) for x in
                       ws.findall("./" + q("cols") + "/" + q("col")) if x.get("hidden") in {"1", "true"}]
        merged = [x.get("ref", "") for x in ws.findall("./" + q("mergeCells") + "/" + q("mergeCell"))]
        dimension = ws.find(q("dimension"))
        declared = "" if dimension is None else dimension.get("ref", "")
        feature(report, "declared_range:" + title, 1, "PARTIAL", declared)
        feature(report, "hidden_rows:" + title, len(hidden_rows), "TESTED", ",".join(sorted(hidden_rows)))
        feature(report, "hidden_column_spans:" + title, len(hidden_cols), "TESTED", repr(hidden_cols))
        feature(report, "merged_cells:" + title, len(merged), "TESTED", ",".join(merged))
        autofilter = ws.find(q("autoFilter"))
        feature(report, "autofilter:" + title, int(autofilter is not None), "PARTIAL",
                "" if autofilter is None else autofilter.get("ref", ""))
        seen, coordinates = set(), []
        for cell in ws.findall("./" + q("sheetData") + "/" + q("row") + "/" + q("c")):
            address = cell.get("r", "")
            if address in seen:
                raise CoreError("BELGE_BOZUK", "Yinelenen hücre adresi")
            seen.add(address)
            try:
                row_index, column_index = coordinate_to_tuple(address)
                if not (1 <= row_index <= 1048576 and 1 <= column_index <= 16384):
                    raise ValueError()
            except (ValueError, KeyError) as exc:
                raise CoreError("BELGE_BOZUK", "Hücre konumu geçersiz") from exc
            coordinates.append((row_index, column_index))
            total_cells += 1
            if total_cells > 20000:
                raise CoreError("BELGE_SINIRI", "P2 toplam hücre sınırı 20.000")
            cell_type = cell.get("t", "n")
            value = cell.find(q("v"))
            cached = "" if value is None or value.text is None else value.text
            if cell_type == "s":
                try:
                    index = int(cached)
                    if index < 0:
                        raise ValueError()
                    display = shared[index]
                except (ValueError, IndexError) as exc:
                    raise CoreError("BELGE_BOZUK", "Shared string referansı geçersiz") from exc
            elif cell_type == "inlineStr":
                display = "".join(t.text or "" for t in cell.iter(q("t")))
            else:
                display = cached  # Numeric serial/raw precision preserved, no float conversion.
            formula_node = cell.find(q("f"))
            formula = None if formula_node is None else formula_node.text or ""
            style_index = int(cell.get("s", "0"))
            if style_index < 0 or (styles and style_index >= len(styles)):
                raise CoreError("BELGE_BOZUK", "Hücre style referansı geçersiz")
            format_id = styles[style_index] if styles else 0
            attrs = {"cell_type": cell_type, "cached_value": cached, "date_system": date_system,
                     "number_format": custom_formats.get(format_id, BUILTIN_FORMATS.get(format_id, "")),
                     "formula": "" if formula is None else formula,
                     "calculation_status": "NOT_RECALCULATED" if formula_node is not None else "NOT_A_FORMULA",
                     "hidden_row": str(row_index) in hidden_rows,
                     "hidden_column": any(a <= column_index <= b for a, b in hidden_cols),
                     "sheet_state": state}
            location = locator(part, "cell:" + address, sheet=title, cell=address)
            unit(report, "CELL", location, display, display=display, attributes=attrs)
            if formula_node is not None:
                issue(report, "FORMULA_CACHE_UNVERIFIED", "Formül ve saklanan sonuç okundu; hesap doğrulanmadı",
                      location, "BLOCKING")
        feature(report, "stored_cells:" + title, len(seen), "TESTED",
                "Sparse XML; boş satırları/dimension aralığını enumerate etmez")
        if coordinates:
            extent = (min(x[0] for x in coordinates), min(x[1] for x in coordinates),
                      max(x[0] for x in coordinates), max(x[1] for x in coordinates))
            feature(report, "actual_cell_extent:" + title, len(coordinates), "TESTED", repr(extent))
    if len(sheet_names) > 100:
        raise CoreError("BELGE_SINIRI", "P2 sheet sınırı 100")
    # Same OOXML source, not independent visual/financial evidence.
    workbook = None
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if workbook.sheetnames != sheet_names:
            issue(report, "SECONDARY_READER_DISAGREEMENT", "Sheet envanteri okuyucular arasında farklı", severity="BLOCKING")
    except Exception as exc:
        issue(report, "SECONDARY_READER_FAILED", type(exc).__name__, severity="BLOCKING")
    finally:
        if workbook is not None:
            workbook.close()
    report["verification"]["methods"] = ["ZIP_CRC_AND_PART_HASHES", "SAFE_SPARSE_XML", "OPENPYXL_SHEET_CROSS_CHECK"]
    report["fallback"]["performed"].append("SECONDARY_SAME_SOURCE_INVENTORY")
    report["fallback"]["next"] = ["TARGETED_DISPLAY_CHECK", "AUTHORIZED_FORMULA_RECALCULATION_IF_NEEDED"]
    return report
