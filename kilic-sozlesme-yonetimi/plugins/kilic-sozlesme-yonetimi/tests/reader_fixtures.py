"""Synthetic source bytes and explicit goldens. No production data."""
import io
import zipfile
from datetime import datetime
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"


def package_bytes(parts):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            entry = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(entry, value.encode("utf-8") if isinstance(value, str) else value)
    return output.getvalue()


def unpack(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def docx_parts():
    reltype = R + "/"
    body = """<w:p><w:pPr><w:pStyle w:val="Title"/></w:pPr><w:r><w:t>Sentetik sözleşme okuma kontrolü</w:t></w:r></w:p>
    <w:p><w:r><w:t>Bu belge yalnız okuyucu testidir. İ ı ş ğ ç ö ü.</w:t></w:r></w:p>
    <w:p><w:r><w:t xml:space="preserve">Garanti süresi </w:t></w:r>
      <w:del w:id="10" w:author="Sentetik Tedarikçi" w:date="2026-10-01T09:00:00Z"><w:r><w:delText>24</w:delText></w:r></w:del>
      <w:ins w:id="11" w:author="Sentetik Hukuk" w:date="2026-10-02T09:00:00Z"><w:r><w:t>36</w:t></w:r></w:ins>
      <w:r><w:t xml:space="preserve"> aydır.</w:t></w:r></w:p>
    <w:p><w:commentRangeStart w:id="0"/><w:r><w:t>Avans oranı %20.</w:t></w:r>
      <w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r>
      <w:r><w:footnoteReference w:id="1"/></w:r><w:r><w:endnoteReference w:id="1"/></w:r></w:p>
    <w:tbl><w:tblPr><w:tblW w:w="8000" w:type="dxa"/></w:tblPr><w:tblGrid><w:gridCol w:w="3000"/><w:gridCol w:w="5000"/></w:tblGrid>
      <w:tr><w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:t>Bakım</w:t></w:r></w:p></w:tc>
        <w:tc><w:tcPr><w:tcW w:w="5000" w:type="dxa"/></w:tcPr><w:p><w:r><w:t>100.000 TL</w:t></w:r></w:p>
          <w:tbl><w:tblGrid><w:gridCol w:w="3000"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>İç tablo kapsamı</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p/></w:tc></w:tr></w:tbl>
    <w:sectPr><w:headerReference w:type="default" r:id="h1"/><w:footerReference w:type="default" r:id="f1"/>
      <w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720"/></w:sectPr>"""
    def wrap(tag, content):
        return '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:' + tag + ' xmlns:w="' + W + '" xmlns:r="' + R + '">' + content + "</w:" + tag + ">"
    parts = {
        "word/document.xml": wrap("document", "<w:body>" + body + "</w:body>"),
        "word/header1.xml": wrap("hdr", "<w:p><w:r><w:t>Sentetik üstbilgi</w:t></w:r></w:p>"),
        "word/footer1.xml": wrap("ftr", "<w:p><w:r><w:t>Sentetik altbilgi</w:t></w:r></w:p>"),
        "word/comments.xml": wrap("comments", '<w:comment w:id="0" w:author="Sentetik Hukuk" w:date="2026-10-02T09:00:00Z"><w:p><w:r><w:t>Oran teyidi gerekir.</w:t></w:r></w:p></w:comment>'),
        "word/footnotes.xml": wrap("footnotes", '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/><w:t>Dipnot kapsam açıklaması.</w:t></w:r></w:p></w:footnote>'),
        "word/endnotes.xml": wrap("endnotes", '<w:endnote w:id="1"><w:p><w:r><w:endnoteRef/><w:t>Sonnot teslim açıklaması.</w:t></w:r></w:p></w:endnote>'),
        "word/styles.xml": wrap("styles", '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="24"/></w:rPr></w:rPrDefault></w:docDefaults><w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:rPr><w:b/><w:color w:val="000000"/><w:sz w:val="32"/></w:rPr></w:style>'),
    }
    suffixes = {"document": "document.main", "header1": "header", "footer1": "footer", "comments": "comments",
                "footnotes": "footnotes", "endnotes": "endnotes", "styles": "styles"}
    parts["[Content_Types].xml"] = '<Types xmlns="' + CT + '"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>' + "".join(
        '<Override PartName="/word/' + name + '.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.' + typ + '+xml"/>' for name, typ in suffixes.items()) + "</Types>"
    parts["_rels/.rels"] = '<Relationships xmlns="' + REL + '"><Relationship Id="root1" Type="' + reltype + 'officeDocument" Target="word/document.xml"/></Relationships>'
    parts["word/_rels/document.xml.rels"] = '<Relationships xmlns="' + REL + '">' + "".join(
        '<Relationship Id="' + rid + '" Type="' + reltype + typ + '" Target="' + target + '"/>' for rid, typ, target in
        [("h1", "header", "header1.xml"), ("f1", "footer", "footer1.xml"), ("c1", "comments", "comments.xml"),
         ("n1", "footnotes", "footnotes.xml"), ("e1", "endnotes", "endnotes.xml"), ("s1", "styles", "styles.xml")]) + "</Relationships>"
    return parts


def docx_bytes():
    return package_bytes(docx_parts())


def xlsx_bytes():
    from openpyxl import Workbook
    wb = Workbook()
    wb.properties.created = datetime(2020, 1, 1)
    wb.properties.modified = datetime(2020, 1, 1)
    ws = wb.active
    ws.title = "Teklif İı"
    ws.append(["Kalem", "Tutar", "Hesap"])
    ws.append(["Bakım", 12, "=B2*2"])
    ws["B2"].number_format = '#,##0.00 "TL"'
    ws["B5"] = 45292
    ws["B5"].number_format = "dd.mm.yyyy"
    ws.merge_cells("A3:B3")
    ws["A3"] = "Birleşik kapsam"
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["B"].hidden = True
    ws.auto_filter.ref = "A1:C2"
    hidden = wb.create_sheet("Gizli")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "Gizli kritik teminat"
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    parts = unpack(output.getvalue())
    root = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root.find(".//s:c[@r='C2']/s:v", ns).text = "999"  # Deliberately wrong cached result.
    parts["xl/worksheets/sheet1.xml"] = etree.tostring(root)
    return package_bytes(parts)


def pdf_bytes(text="Sentetik garanti 24 ay; avans %20.", *, blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=300)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 12 Tf 30 250 Td (" + text + ") Tj ET").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    writer.close()
    return output.getvalue()
